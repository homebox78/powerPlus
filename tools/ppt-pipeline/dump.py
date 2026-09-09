# -*- coding: utf-8 -*-
"""슬라이드별 텍스트 추출(제목 후보 + 본문 키워드용)."""
import glob, json, os, re
from pptx import Presentation
SRC = r"d:\powerPlus\data\_incoming\ppt"
f = glob.glob(os.path.join(SRC, "*.pptx"))[0]
p = Presentation(f)
out = []
def walk(shapes, acc):
    for sh in shapes:
        if sh.shape_type == 6:  # GROUP
            walk(sh.shapes, acc); continue
        if sh.has_text_frame:
            t = sh.text_frame.text.strip()
            if t: acc.append(t)
        if getattr(sh, "has_table", False) and sh.has_table:
            for row in sh.table.rows:
                for c in row.cells:
                    t = c.text.strip()
                    if t: acc.append(t)
for i, s in enumerate(p.slides, start=1):
    acc = []
    walk(s.shapes, acc)
    txt = "\n".join(acc)
    txt = re.sub(r"\n{2,}", "\n", txt)
    out.append({"no": i, "text": txt})
json.dump(out, open(os.path.join(os.path.dirname(os.path.abspath(__file__)), "slides.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)
print("슬라이드", len(out), "· 총 문자", sum(len(o["text"]) for o in out))
