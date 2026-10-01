import sys, glob, io
from PIL import Image, ImageDraw
from pptx import Presentation
sys.stdout.reconfigure(encoding="utf-8")
F="제안서/청주시_발표자료(제안요약서)_v0.73_7쪽하단.pptx"
D=Image.open(sorted(glob.glob('제안서/design_ppt/*.png'))[1]).convert("RGB"); Z=D.size[0]/2000
def crop(x0,y0,x1,y1,cut,th=24):
    im=D.crop((int(x0*Z),int(y0*Z),int(x1*Z),int(y1*Z))); key=(255,0,255)
    ref=D.getpixel((int(cut[0]*Z),int(cut[1]*Z))); w,h=im.size
    near=lambda p: sum(abs(a-b) for a,b in zip(p,ref))<=th
    for pt in [(x,y) for x in range(0,w,6) for y in (0,h-1)]+[(x,y) for y in range(0,h,6) for x in (0,w-1)]:
        p=im.getpixel(pt)
        if p!=key and near(p): ImageDraw.floodfill(im,pt,key,thresh=th)
    im=im.convert("RGBA"); px=im.load()
    for y in range(h):
        for x in range(w):
            if px[x,y][:3]==key: px[x,y]=(255,255,255,0)
    b=io.BytesIO(); im.save(b,"PNG"); return b.getvalue(), im
pr=Presentation(F); s=pr.slides[5]; by={sh.shape_id:sh for sh in s.shapes}
# id, 옛 박스, 새 박스, 카드 x0
for sid,old,new,cx in ((122,(1150,598,1296,722),(1160,598,1296,722),1067),):
    sh=by[sid]; ox0,oy0,ox1,oy1=old; nx0,ny0,nx1,ny1=new
    sx=sh.width/(ox1-ox0); sy=sh.height/(oy1-oy0)
    data,im=crop(*new,cut=(cx+20,780))
    im.save(f"{sys.argv[1]}/kpi_{sid}.png")
    # 그림 데이터 교체
    rId=sh._element.blipFill.blip.rEmbed
    part=sh.part.related_part(rId); part._blob=data
    sh.left=int(sh.left+(nx0-ox0)*sx); sh.top=int(sh.top+(ny0-oy0)*sy)
    sh.width=int((nx1-nx0)*sx); sh.height=int((ny1-ny0)*sy)
pr.save(F); print("ok")
