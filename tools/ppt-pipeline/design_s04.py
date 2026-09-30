"""청주시 4쪽 — 시안(design_ppt 1번)과 같게 재조립.

그림(노트북·인물·건물 타일·도시 풍경·작은 아이콘)은 시안 이미지에서 그대로 오려 쓰고,
글자·알약·화살표·상자는 시안 픽셀 좌표·색을 읽어 PowerPoint 도형으로 그린다.
좌표는 시안을 2000×1125 로 본 값(원본 2560×1440 = ×1.28).
사용: py design_s04.py <src.pptx> <dst.pptx> <시안.png> <작업 폴더>
"""
import sys, os
from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_LINE
from pptx.oxml.ns import qn

src, dst, design, work = sys.argv[1:5]
os.makedirs(work, exist_ok=True)
D = Image.open(design).convert("RGB")
Z = D.size[0] / 2000.0            # 2000 기준 좌표 → 원본 픽셀
K = 10.83 / 2000.0                # 2000 기준 px → inch (가로 맞춤)
F2, F3 = "a시월구일2", "a시월구일3"
NAVY, BLUE, PINK, WHITE = RGBColor(0x10, 0x2A, 0x5C), RGBColor(0x2B, 0x6D, 0xE8), RGBColor(0xF2, 0x4F, 0x6E), RGBColor(255, 255, 255)


def rgb_at(x, y):
    return RGBColor(*D.getpixel((int(x * Z), int(y * Z))))


def crop(name, x0, y0, x1, y1, fade_top=0):
    im = D.crop((int(x0 * Z), int(y0 * Z), int(x1 * Z), int(y1 * Z)))
    if fade_top:
        im = im.convert("RGBA"); px = im.load(); n = int(fade_top * Z)
        for y in range(n):
            a = int(255 * y / n)
            for x in range(im.size[0]):
                r, g, b, _ = px[x, y]; px[x, y] = (r, g, b, a)
    p = os.path.join(work, name + ".png"); im.save(p)
    return p


# 세로: 시안 본문(2:1)을 장표 본문에 맞게 덩어리 중심만 옮긴다(크기는 가로 배율 그대로)
def cy(y):
    return 1.5 + (y - 160) * 0.00622


class Grp:
    """시안의 한 덩어리 — 중심 y 만 새 자리로, 안쪽은 같은 배율."""
    def __init__(self, y0, y1):
        self.c = (y0 + y1) / 2.0; self.t = cy(self.c)

    def X(self, x): return x * K
    def Y(self, y): return self.t + (y - self.c) * K


prs = Presentation(src)
s = prs.slides[3]
tree = s.shapes._spTree
title_font = None
for sh in list(s.shapes):
    if sh.name in ("그룹 90", "그림 89", "직사각형 18"):
        if sh.name == "직사각형 18":
            for r in sh.text_frame.paragraphs[0].runs:
                ea = r._r.find(".//" + qn("a:ea"))
                title_font = (ea.get("typeface") if ea is not None else None) or r.font.name
        tree.remove(sh._element)
title_font = title_font or F3


def shape(kind, l, t, w, h, fill=None, line=None, lw=0.75, adj=None, dash=False):
    b = s.shapes.add_shape(kind, Inches(l), Inches(t), Inches(w), Inches(h))
    b.shadow.inherit = False
    if adj is not None:
        b.adjustments[0] = adj
    if fill is None: b.fill.background()
    else: b.fill.solid(); b.fill.fore_color.rgb = fill
    if line is None: b.line.fill.background()
    else:
        b.line.color.rgb = line; b.line.width = Pt(lw)
        if dash: b.line.dash_style = MSO_LINE.DASH
    tf = b.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE; tf.word_wrap = False
    return b


def write(b, lines, align=PP_ALIGN.CENTER):
    """lines = [[(글, pt, 색, 서체), …], …]"""
    tf = b.text_frame
    for i, runs in enumerate(lines):
        pg = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        pg.alignment = align
        for tx, size, color, font in runs:
            r = pg.add_run(); r.text = tx
            r.font.size = Pt(size); r.font.color.rgb = color; r.font.name = font
            rPr = r._r.get_or_add_rPr()
            ea = rPr.find(qn("a:ea"))
            if ea is None: ea = etree.SubElement(rPr, qn("a:ea"))
            ea.set("typeface", font)
    return b


def pic(path, l, t, w):
    return s.shapes.add_picture(path, Inches(l), Inches(t), width=Inches(w))


def line(x0, y0, x1, y1, color, w=1.0):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x0), Inches(y0), Inches(x1), Inches(y1))
    c.line.color.rgb = color; c.line.width = Pt(w)
    return c


# 0) 본문 바탕(시안의 아주 옅은 파랑) + 하단 도시 풍경 — 맨 뒤
bg = shape(MSO_SHAPE.RECTANGLE, 0, 0.87, 10.83, 6.63, fill=rgb_at(1000, 270))
CITY_Y0 = 884
city = pic(crop("city", 0, CITY_Y0, 2000, 1125, fade_top=70), 0, 0, 10.83)
city.top = Inches(7.5) - city.height
for el in (city._element, bg._element):
    tree.remove(el); tree.insert(2, el)

# 1) 제목
t = shape(MSO_SHAPE.RECTANGLE, 0, cy(215) - 0.3, 10.83, 0.6)
write(t, [[("청주시 ", 27, NAVY, title_font), ("업무지원포털시스템", 27, BLUE, title_font)]])

# 2) 노트북 덩어리
g = Grp(280, 870)
pic(crop("laptop", 40, 280, 925, 872), g.X(40), g.Y(280), 885 * K)
shape(MSO_SHAPE.RECTANGLE, g.X(143), g.Y(312), 687 * K, 443 * K, fill=rgb_at(700, 440))
shape(MSO_SHAPE.ROUNDED_RECTANGLE, g.X(165), g.Y(405), 645 * K, 340 * K, fill=rgb_at(500, 470), adj=0.08)
b = shape(MSO_SHAPE.ROUNDED_RECTANGLE, g.X(343), g.Y(328), 310 * K, 67 * K, fill=rgb_at(360, 362), adj=0.5)
write(b, [[("핵심 시스템", 16, WHITE, F3)]])
for sx, dx in ((300, 1), (697, -1)):          # 알약 양옆 강조 선
    for dy in (-14, 0, 14):
        line(g.X(sx), g.Y(362 + dy * 1.4), g.X(sx + 30 * dx), g.Y(362 + dy * 0.6), PINK, 1.5)
b = shape(MSO_SHAPE.RECTANGLE, g.X(165), g.Y(415), 645 * K, 56 * K)
write(b, [[("행정포털", 16.5, NAVY, F3), ("(굿모닝)", 16.5, NAVY, F2)]])
pic(crop("woman", 165, 488, 562, 747), g.X(165), g.Y(488), 397 * K)
shape(MSO_SHAPE.ROUNDED_RECTANGLE, g.X(580), g.Y(493), 220 * K, 247 * K, fill=rgb_at(590, 600), line=rgb_at(580, 600), adj=0.1)
line(g.X(600), g.Y(620), g.X(782), g.Y(620), rgb_at(580, 600), 0.75)
for iy, lab, big, small in ((527, "분야", "5", "개 분야"), (645, "업무", "227", "종")):
    pic(crop("ico%d" % iy, 601, iy - 4, 650, iy + 50), g.X(601), g.Y(iy - 4), 49 * K)
    b = shape(MSO_SHAPE.RECTANGLE, g.X(668), g.Y(iy - 6), 120 * K, 28 * K)
    write(b, [[(lab, 8, NAVY, F2)]], PP_ALIGN.LEFT)
    b = shape(MSO_SHAPE.RECTANGLE, g.X(668), g.Y(iy + 24), 125 * K, 60 * K)
    write(b, [[(big, 18, BLUE, F3), (small, 10, BLUE, F3)]], PP_ALIGN.LEFT)

# 3) 가운데 화살표
g = Grp(462, 612)
b = shape(MSO_SHAPE.LEFT_RIGHT_ARROW, g.X(857), g.Y(462), 251 * K, 150 * K, fill=rgb_at(980, 500))
b.adjustments[0] = 0.62; b.adjustments[1] = 0.38
write(b, [[("시스템 연계", 10, WHITE, F3)], [("공동 이용", 9.5, WHITE, F2)]])

# 4) 사용자 조직
g = Grp(305, 820)
LN = RGBColor(0x8D, 0xB8, 0xF2)
shape(MSO_SHAPE.ROUNDED_RECTANGLE, g.X(1120), g.Y(305), 835 * K, 515 * K, line=LN, lw=1.0, adj=0.04, dash=True)
pill = rgb_at(1240, 390)
b = shape(MSO_SHAPE.ROUNDED_RECTANGLE, g.X(1210), g.Y(347), 660 * K, 86 * K, fill=pill, adj=0.5)
b = shape(MSO_SHAPE.RECTANGLE, g.X(1358), g.Y(347), 500 * K, 86 * K)
write(b, [[("청주시 공무원 ", 13, WHITE, F3), ("5,000명", 18, WHITE, F3), (" 대상", 13, WHITE, F3)]], PP_ALIGN.LEFT)
pic(crop("people", 1262, 356, 1345, 420), g.X(1262), g.Y(356), 83 * K)
b = shape(MSO_SHAPE.RECTANGLE, g.X(1120), g.Y(452), 835 * K, 56 * K)
write(b, [[("사용자 조직 구성", 15, NAVY, F3)]])
xs = [1146, 1308, 1470, 1632, 1794]; TW = 146
line(g.X(xs[0] + TW / 2), g.Y(533), g.X(xs[-1] + TW / 2), g.Y(533), LN, 0.75)
labels = ["시본청", "4개구청", "43개읍면동", "16개도서관", "사업소"]
lab_fill = rgb_at(1160, 738)
for x, lab in zip(xs, labels):
    line(g.X(x + TW / 2), g.Y(533), g.X(x + TW / 2), g.Y(567), LN, 0.75)
    pic(crop("tile%d" % x, x, 566, x + TW, 712), g.X(x), g.Y(566), TW * K)
    b = shape(MSO_SHAPE.ROUNDED_RECTANGLE, g.X(x + 5), g.Y(716), (TW - 10) * K, 46 * K, fill=lab_fill, adj=0.5)
    write(b, [[(lab, 8, NAVY, F3)]])

prs.save(dst)
print("saved")
