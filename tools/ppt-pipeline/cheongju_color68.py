"""v0.68 색·잔여 보정. 인자: <src> <dst>"""
import copy, sys
from pptx import Presentation
from pptx.dml.color import RGBColor
from lxml import etree
sys.stdout.reconfigure(encoding="utf-8")
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
IN = 914400
src, dst = sys.argv[1:3]
pr = Presentation(src)
W, H = pr.slide_width, pr.slide_height


def walk(shs):
    for sh in shs:
        if sh.shape_type == 6:
            yield from walk(sh.shapes)
        else:
            yield sh


def txt(sh):
    return sh.text_frame.text.strip() if sh.has_text_frame else ""


def runs(sh):
    for pg in sh.text_frame.paragraphs:
        for r in pg.runs:
            yield r


def set_runs(sh, hexc):
    for r in runs(sh):
        if r.text.strip():
            r.font.color.rgb = RGBColor.from_string(hexc)


def rgb(el):
    c = el.find(A + "srgbClr") if el is not None else None
    return c.get("val").upper() if c is not None else None


log = []
S = lambda i: pr.slides[i - 1]

# 8쪽: 네 번째 '고민' 글만 핑크 → 다른 셋과 같은 색
ref = [sh for sh in S(8).shapes if "특성이 다른 사용자" in txt(sh)][0]
refc = next(r.font.color.rgb for r in runs(ref) if r.text.strip())
for sh in S(8).shapes:
    if "장기계속계약" in txt(sh):
        set_runs(sh, str(refc)); log.append("08 고민 글 색 " + str(refc))

# 18쪽: 그림에 박힌 번호를 가리는 파란 사각형이 0.05in 짧아 '0' 왼쪽 가장자리가 '(' 로 보임 → 왼쪽으로 넓힘
from pptx.util import Emu
s18 = list(S(18).shapes)
nums = [sh for sh in s18 if txt(sh) in ("01", "02", "03", "04") and sh.top > 5.5 * IN]
for nb in nums:
    for sh in s18:
        if sh.has_text_frame and not txt(sh) and sh.shape_type == 1 and abs(sh.top - nb.top) < 0.1 * IN \
                and abs(sh.left - nb.left) < 0.15 * IN and sh.width < 0.4 * IN:
            d = int(0.08 * IN)
            sh.left, sh.width = Emu(sh.left - d), Emu(sh.width + d)
            log.append("18 번호 가리개 넓힘 %s" % sh.name)

# 20쪽: '문제 해결' 노드만 초록 → '변경 처리' 노드와 같은 면·선·글자
s20 = list(walk(S(20).shapes))
tpl = [sh for sh in s20 if txt(sh) == "변경 처리"][0]
for sh in s20:
    if txt(sh) == "문제 해결":
        spPr, tsp = sh._element.spPr, tpl._element.spPr
        for tag in ("solidFill", "gradFill", "noFill", "ln"):
            for e in spPr.findall(A + tag):
                spPr.remove(e)
        geom = spPr.find(A + "prstGeom")
        idx = list(spPr).index(geom) + 1 if geom is not None else len(spPr)
        for tag in ("solidFill", "gradFill", "ln"):
            e = tsp.find(A + tag)
            if e is not None:
                spPr.insert(idx, copy.deepcopy(e)); idx += 1
        tc = next(r.font.color.rgb for r in runs(tpl) if r.text.strip())
        set_runs(sh, str(tc)); log.append("20 문제 해결 노드 → 변경 처리 색")

# 22쪽: '100%' 보라 → 파랑
for sh in walk(S(22).shapes):
    if txt(sh) == "100%":
        set_runs(sh, "2070E8"); log.append("22 100% 파랑")

# 25쪽: 진한 빨강 채움 → 강조 핑크
for sh in walk(S(25).shapes):
    sp = getattr(sh._element, "spPr", None)
    if sp is None:
        continue
    f = sp.find(A + "solidFill")
    c = rgb(f)
    if c:
        r, g, b = int(c[:2], 16), int(c[2:4], 16), int(c[4:], 16)
        if r > 170 and g < 90 and b < 90:
            f.find(A + "srgbClr").set("val", "EC1C68"); log.append("25 %s 빨강 %s→EC1C68" % (sh.name, c))

# 28쪽: 기밀성·무결성·가용성 제목 세 색 → 남색 하나로, 슬라이드 밖 잔여 도형 삭제
for sh in list(walk(S(28).shapes)):
    if txt(sh) in ("기밀성", "무결성", "가용성"):
        set_runs(sh, "143A69"); log.append("28 %s 남색" % txt(sh))
for sh in list(S(28).shapes):
    if sh.left is not None and (sh.left > W or sh.left + sh.width < 0 or sh.top > H or sh.top + sh.height < 0):
        sh._element.getparent().remove(sh._element); log.append("28 슬라이드 밖 삭제 %s" % sh.name)

# 37쪽: 왼쪽 솔리데오 쪽 파랑 → 로고 버건디 C01348 계열
BURG = {"dark": "C01348", "deep": "9A0F3A", "line": "E8B4C6", "soft": "FBEEF2", "mid": "F6DCE5"}


def burgundy(c):
    r, g, b = int(c[:2], 16), int(c[2:4], 16), int(c[4:], 16)
    lum = (r * 3 + g * 6 + b) / 10
    if not (b - r > (20 if lum < 200 else 5)):    # 파랑 계열만(옅은 면은 기운만 있어도)
        return None
    if lum < 70:
        return BURG["deep"]
    if lum < 170:
        return BURG["dark"]
    if lum < 215:
        return BURG["line"]
    if lum < 238:
        return BURG["mid"]
    return BURG["soft"]


s37 = S(37)
for sh in walk(s37.shapes):
    if sh.left is None or sh.shape_type == 13 or sh.width > W * 0.6:   # 바탕·전폭 도형 제외
        continue
    cx = sh.left + sh.width / 2
    solideo_row = txt(sh) in ("솔리데오", "60%")
    if not (cx < W / 2 - 0.05 * IN or solideo_row):
        continue
    el = sh._element
    n = 0
    for c in el.iter(A + "srgbClr"):
        nv = burgundy(c.get("val"))
        if nv:
            c.set("val", nv); n += 1
    if n:
        log.append("37 %s 버건디 %d" % (sh.name, n))

# 41쪽: 오타
for sh in walk(S(41).shapes):
    if sh.has_table:
        for row in sh.table.rows:
            for cell in row.cells:
                for pg in cell.text_frame.paragraphs:
                    for r in pg.runs:
                        if "요지보수" in r.text:
                            r.text = r.text.replace("요지보수", "유지보수"); log.append("41 요지보수→유지보수")
    elif sh.has_text_frame:
        for r in runs(sh):
            if "요지보수" in r.text:
                r.text = r.text.replace("요지보수", "유지보수"); log.append("41 요지보수→유지보수")

pr.save(dst)
print("\n".join(log))
print("저장", dst)
