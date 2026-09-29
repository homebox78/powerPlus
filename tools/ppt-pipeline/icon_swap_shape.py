# -*- coding: utf-8 -*-
"""그림 도형 하나하나를 powerPlus 자산으로 바꾼다(미디어 공유 없이 도형 단위).
인자: <src> <dst> <map.json>   map = {"슬라이드|shape_id": {"asset": 경로, "white": false}}
- 원본 그림의 잉크 영역(알파/비흰색 바운딩박스)에 새 아이콘을 비율 유지·중앙으로 맞춘다
- white=true 면 아이콘을 흰색으로 칠한다(진한 칩 위 선 아이콘용)
- 그룹 안 도형도 그 자리(z순서·그룹 소속) 그대로 교체된다
"""
import sys, io, json
from pptx import Presentation
from PIL import Image
NS="{http://schemas.openxmlformats.org/drawingml/2006/main}"
RNS="{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
src,dst,mp=sys.argv[1],sys.argv[2],json.load(open(sys.argv[3],encoding="utf-8"))
p=Presentation(src); done=[]; miss=set(mp)
def ink_box(im):
    if im.mode in ("RGBA","LA") or "transparency" in im.info:
        a=im.convert("RGBA").getchannel("A").point(lambda v:255 if v>8 else 0)
        b=a.getbbox()
    else:
        g=im.convert("L").point(lambda v:255 if v<235 else 0)
        b=g.getbbox()
    return b or (0,0,im.width,im.height)
cache={}
def asset(path, white):
    k=(path,white)
    if k not in cache:
        im=Image.open(path).convert("RGBA")
        b=im.getchannel("A").point(lambda v:255 if v>8 else 0).getbbox()
        if b: im=im.crop(b)
        if max(im.size)>600: im.thumbnail((600,600), Image.LANCZOS)
        if white:
            a=im.getchannel("A"); im=Image.new("RGBA",im.size,(255,255,255,255)); im.putalpha(a)
        buf=io.BytesIO(); im.save(buf,"PNG",optimize=True); cache[k]=(buf.getvalue(), im.size)
    return cache[k]
def walk(shapes, sn):
    for sh in shapes:
        st=str(sh.shape_type).split()[0]
        if st=="GROUP": walk(sh.shapes, sn); continue
        key="%d|%d"%(sn, sh.shape_id)
        if st!="PICTURE" or key not in mp: continue
        spec=mp[key]
        old=Image.open(io.BytesIO(sh.image.blob))
        bf=sh._element.find(".//{*}blipFill")   # 그림은 p:blipFill(프레젠테이션 네임스페이스)
        sr=bf.find(NS+"srcRect") if bf is not None else None
        if sr is not None:   # 잘라 쓰던 그림이면 보이는 영역만 보고 잰다
            W,H=old.size
            l,t,r,b=[int(sr.get(k,"0"))/100000 for k in ("l","t","r","b")]
            old=old.crop((int(W*l),int(H*t),int(W*(1-r)),int(H*(1-b))))
        bx=ink_box(old); W,H=old.size
        L,T,SW,SH=sh.left,sh.top,sh.width,sh.height
        ix=L+SW*bx[0]/W; iy=T+SH*bx[1]/H; iw=SW*(bx[2]-bx[0])/W; ih=SH*(bx[3]-bx[1])/H
        blob,(aw,ah)=asset(spec["asset"], spec.get("white",False))
        s=min(iw/aw, ih/ah); nw,nh=aw*s, ah*s
        sh.left=int(ix+(iw-nw)/2); sh.top=int(iy+(ih-nh)/2); sh.width=int(nw); sh.height=int(nh)
        _, rId = sh.part.get_or_add_image_part(io.BytesIO(blob))
        blip=bf.find(NS+"blip"); blip.set(RNS+"embed", rId)
        # 옛 그림에 맞춰 둔 색 효과(듀오톤·투명도·밝기 등)는 새 그림을 흐리게 만든다 → 걷어낸다
        for c in list(blip):
            if c.tag.split("}")[1] in ("duotone","alphaModFix","lum","grayscl","clrChange","biLevel","tint","hsl"):
                blip.remove(c)
        if sr is not None: bf.remove(sr)
        done.append(key); miss.discard(key)
for i,s in enumerate(p.slides,1): walk(s.shapes,i)
p.save(dst)
print("교체", len(done), "/ 못 찾음", sorted(miss))
