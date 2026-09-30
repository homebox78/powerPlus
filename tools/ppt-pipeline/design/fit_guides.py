"""시안 반영본 마감 — 키메시지 통일 + 본문 안내선(좌우 13cm·상 6.9·하 8.3cm) 맞춤.

1) 키메시지: 모든 쪽을 한 규칙으로 다시 쓴다 — 서체 G마켓 산스 TTF Bold, 한 크기(가장 긴 문장이 한 줄에 드는 크기, 24pt 이하),
   남색 바탕글 + 파랑 강조, 가운데 정렬, 원래 장표 자리(L0.20 T1.03 W10.43; 4쪽만 T1.60).
   키메시지를 두르던 띠·알약('전략 N')·장식 선·곁 그림은 지운다(쪽마다 모양이 달라지는 원인).
2) 본문: 렌더에서 보이는 외곽을 재어, 좌·우·하 안내선과 키메시지 아래(GAP)에 오도록 한 번의 선형 변환(글자 크기 그대로).
사용: py fit_guides.py <src.pptx> <dst.pptx> <작업 폴더>
"""
import os, sys, time
import numpy as np
from PIL import Image
from lxml import etree
from pptx import Presentation
from pptx.util import Emu, Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

sys.stdout.reconfigure(encoding="utf-8")
IN, CM = 914400, 360000
PAGES = [4, 6, 7, 8, 10, 11, 12, 13, 14, 16, 17, 18, 19, 20, 21, 22]
KEY = {4: ["시안 4"], 6: ["시안 6", "시안 7"], 8: ["시안 3"], 10: ["시안 3"], 11: ["시안 7"], 12: ["시안 5"],
       13: ["시안 7"], 14: ["시안 4"], 16: ["시안 6"], 17: ["시안 4"], 18: ["시안 3"], 19: ["시안 1"],
       20: ["시안 6"], 21: ["시안 2"], 22: ["시안 6"]}
DEL = {6: ["시안 3", "시안 그림 4", "시안 그림 5"], 10: ["시안 선 4", "시안 선 5"],
       11: ["시안 3", "시안 4", "시안 그림 5", "시안 6"], 12: ["시안 3", "시안 4"],
       13: ["시안 3", "시안 4", "시안 5", "시안 6"], 14: ["시안 1", "시안 2", "시안 3"], 17: ["시안 3"],
       20: ["시안 그림 3", "시안 4", "시안 5"], 21: ["시안 그림 1"], 22: ["시안 4", "시안 그림 5"]}
FIXED = {4: ["그룹 7"], 16: ["시안 7"]}          # 키메시지와 한 몸(움직이지 않음)
KEY_TOP = {4: 1.60}
NO_Y = {4}     # 아래 전폭 풍경 띠와 한 구성이라 세로는 그대로
NO_X = {8}     # 왼쪽 가장자리에 붙은 큰 그림과 한 구성이라 가로는 그대로
FONT = "G마켓 산스 TTF Bold"
NAVY, BLUE = RGBColor(0x0B, 0x2E, 0x6B), RGBColor(0x20, 0x70, 0xE8)
GAP = 0.16 * IN
LINE_H = 0.46
src, dst, work = [os.path.abspath(a) for a in sys.argv[1:4]]
os.makedirs(work, exist_ok=True)
p = Presentation(src)
W, H = p.slide_width, p.slide_height
G = dict(L=W / 2 - 13.0 * CM, R=W / 2 + 13.0 * CM, T=H / 2 - 6.9 * CM, B=H / 2 + 8.3 * CM)


def by_name(s, n):
    for sh in s.shapes:
        if sh.name == n:
            return sh
    raise KeyError(n)


def box(sh):
    return sh.left, sh.top, sh.left + sh.width, sh.top + sh.height


def em(t):
    w = 0
    for ch in t:
        w += 0.3 if ch == " " else (0.62 if ord(ch) < 128 else 1.0)
    return w


def is_blue(c):
    return c is not None and c[2] - c[0] > 70 and c[2] > 150


def run_color(r):
    try:
        return tuple(r.font.color.rgb)
    except Exception:
        return None


def render(path, out, pages):
    import win32com.client
    lock = os.path.join(os.environ["TEMP"], "pp_render.lock")
    for _ in range(600):
        try:
            os.mkdir(lock); break
        except FileExistsError:
            time.sleep(1)
    try:
        app = win32com.client.Dispatch("PowerPoint.Application")
        pr = app.Presentations.Open(path, True, False, False)
        for i in pages:
            pr.Slides(i).Export(os.path.join(out, "s%02d.png" % i), "PNG", 2000, 1385)
        pr.Close()
    finally:
        os.rmdir(lock)


# ── 1) 키메시지 통일 ──
msgs = {}
for i, names in KEY.items():
    s = p.slides[i - 1]
    lines = []
    for n in names:
        sh = by_name(s, n)
        for pg in sh.text_frame.paragraphs:
            runs = [(r.text, is_blue(run_color(r))) for r in pg.runs if r.text]
            if "".join(t for t, _ in runs).strip():
                lines.append(runs)
    msgs[i] = lines
widest = max(em("".join(t for t, _ in ln).strip()) for ls in msgs.values() for ln in ls)
SIZE = min(24.0, int(10.1 * 72 / widest * 2) / 2)
print("키메시지 크기 %.1fpt (가장 긴 줄 %.1fem)" % (SIZE, widest))

keybox = {}
for i in PAGES:
    s = p.slides[i - 1]
    tree = s.shapes._spTree
    for n in DEL.get(i, []) + KEY.get(i, []):
        tree.remove(by_name(s, n)._element)
    if i not in KEY:
        continue
    lines = msgs[i]
    top = KEY_TOP.get(i, 1.03)
    b = s.shapes.add_textbox(Inches(0.20), Inches(top), Inches(10.43), Inches(LINE_H * len(lines)))
    b.name = "키메시지"
    tf = b.text_frame; tf.word_wrap = False; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for k, runs in enumerate(lines):
        pg = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        pg.alignment = PP_ALIGN.CENTER
        for j, (t, blue) in enumerate(runs):
            if j == 0: t = t.lstrip()
            if j == len(runs) - 1: t = t.rstrip()
            r = pg.add_run(); r.text = t
            r.font.size = Pt(SIZE); r.font.bold = False; r.font.name = FONT
            r.font.color.rgb = BLUE if blue else NAVY
            rPr = r._r.get_or_add_rPr()
            ea = etree.SubElement(rPr, qn("a:ea")); ea.set("typeface", FONT)
    keybox[i] = (b.shape_id, b.top + b.height)
    kb = b.top + b.height
    if i == 16:   # 부제는 키메시지 바로 아래
        sub = by_name(s, "시안 7"); sub.top = Emu(kb + int(0.02 * IN))
        keybox[i] = (b.shape_id, sub.top + sub.height)
    if i == 14:   # 오른쪽에 있던 ACT.5 띠를 키메시지 아래 왼쪽으로, 본문은 그만큼 아래로
        act = [by_name(s, n) for n in ("시안 5", "시안 6", "시안 7")]
        dx = int(0.17 * IN) - act[0].left
        at = kb + int(0.12 * IN)
        dy_act = at - act[0].top
        rest = [sh for sh in s.shapes if sh.top is not None and sh.top >= 1.6 * IN and sh not in act]
        shift = at + act[0].height + int(0.12 * IN) - min(sh.top for sh in rest)
        for sh in act:
            sh.left, sh.top = Emu(sh.left + dx), Emu(sh.top + dy_act)
        for sh in rest:
            sh.top = Emu(sh.top + shift)
    print("%02d 키메시지 %d줄 T%.2f" % (i, len(lines), top))

mid = os.path.join(work, "_mid.pptx"); p.save(mid)
rdir = os.path.join(work, "r1"); os.makedirs(rdir, exist_ok=True)
render(mid, rdir, PAGES)

# ── 2) 본문 안내선 맞춤 ──
p = Presentation(mid)


def content(s, i):
    kid, kb = keybox.get(i, (None, None))
    fixed = set(FIXED.get(i, []))
    keep, mask = [], []
    for sh in s.shapes:
        if sh.left is None or sh.width is None:
            continue
        l, t, r, b = box(sh)
        if b < 0.87 * IN or t > 7.12 * IN:
            continue
        edge = l < 0.2 * CM or r > W - 0.2 * CM
        texty = sh.has_text_frame and sh.text_frame.text.strip()
        if sh.shape_id == kid or sh.name in fixed:
            mask.append((l, t, r, b)); continue
        if edge and not texty:      # 바탕·전폭 장식: 움직이지 않고, 측정도 가리지 않는다
            continue
        keep.append(sh)
    return keep, mask, kb


def visual(png, keep, mask):
    a = np.asarray(Image.open(png).convert("RGB")).astype(int)
    k = a.shape[1] / W
    ink = (a.max(axis=2) - a.min(axis=2) > 12) | (a.sum(axis=2) < 735)
    bg = np.median(a[int(3.0 * IN * k):int(7.0 * IN * k), 2:8].reshape(-1, 3), axis=0)
    ink &= np.abs(a - bg).sum(axis=2) > 24          # 옅은 파랑 바탕은 잉크 아님
    lim = np.zeros(ink.shape, bool)
    for sh in keep:
        l, t, r, b = box(sh)
        lim[max(0, int(t * k) - 3):int(b * k) + 3, max(0, int(l * k) - 3):int(r * k) + 3] = True
    for l, t, r, b in mask:
        lim[max(0, int(t * k)):int(b * k) + 1, max(0, int(l * k)):int(r * k) + 1] = False
    lim[:int(0.9 * IN * k), :] = False
    lim[int(7.12 * IN * k):, :] = False
    ys, xs = np.nonzero(ink & lim)
    if len(xs) == 0:
        return None
    return xs.min() / k, ys.min() / k, (xs.max() + 1) / k, (ys.max() + 1) / k


for i in PAGES:
    s = p.slides[i - 1]
    keep, mask, kb = content(s, i)
    vb = visual(os.path.join(rdir, "s%02d.png" % i), keep, mask)
    if not vb:
        continue
    L, T, R, B = vb
    GT = max(G["T"], kb + GAP) if kb is not None else G["T"]
    sx = (G["R"] - G["L"]) / (R - L)
    sy = (G["B"] - GT) / (B - T)
    note = ""
    if not 0.85 <= sx <= 1.15: note += " 가로배율 %.3f 제한" % sx; sx = max(0.85, min(1.15, sx))
    if not 0.85 <= sy <= 1.15: note += " 세로배율 %.3f 제한" % sy; sy = max(0.85, min(1.15, sy))
    cxv, cyv = (L + R) / 2, (T + B) / 2
    tcx, tcy = (G["L"] + G["R"]) / 2, (GT + G["B"]) / 2
    if i in NO_X: sx, tcx, note = 1.0, cxv, note + " 가로 고정"
    if i in NO_Y: sy, tcy, note = 1.0, cyv, note + " 세로 고정"
    fx = lambda x: tcx + (x - cxv) * sx
    fy = lambda y: tcy + (y - cyv) * sy
    for sh in keep:
        l, t, r, b = box(sh)
        if sh.shape_type == 13:
            k = (sx + sy) / 2
            cx, cy = fx((l + r) / 2), fy((t + b) / 2)
            w, h = sh.width * k, sh.height * k
            sh.left, sh.top = Emu(int(cx - w / 2)), Emu(int(cy - h / 2))
            sh.width, sh.height = Emu(int(w)), Emu(int(h))
            continue
        nl, nt, nr, nb = fx(l), fy(t), fx(r), fy(b)
        if sx < 1 and sh.has_text_frame and sh.text_frame.text.strip():
            c = (nl + nr) / 2; nl, nr = c - (r - l) / 2, c + (r - l) / 2   # 글자는 그대로라 폭도 유지
        sh.left, sh.top = Emu(int(nl)), Emu(int(nt))
        sh.width, sh.height = Emu(max(1, int(nr - nl))), Emu(max(1, int(nb - nt)))
        if sh.shape_type == 19:
            for c in sh.table.columns: c.width = Emu(int(c.width * sx))
            for r_ in sh.table.rows: r_.height = Emu(int(r_.height * sy))
    print("%02d 좌%+.2f 우%+.2f 상%+.2f 하%+.2f in → 가로 %.3f 세로 %.3f%s" % (
        i, (L - G["L"]) / IN, (R - G["R"]) / IN, (T - GT) / IN, (B - G["B"]) / IN, sx, sy, note))
# 쪽별 보정: 16쪽 알약 글이 알약보다 길다 → 글만 10% 줄임
for sh in p.slides[15].shapes:
    if sh.has_text_frame and sh.text_frame.text.strip().startswith("환경변화와"):
        for pg in sh.text_frame.paragraphs:
            for r in pg.runs:
                if r.font.size: r.font.size = Pt(round(r.font.size.pt * 0.9 * 2) / 2)
        sh.left = Emu(sh.left + int(0.07 * IN))      # 앞 아이콘과 띄움
p.save(dst)
r2 =os.path.join(work, "r2"); os.makedirs(r2, exist_ok=True)
render(dst, r2, PAGES)
print("저장", dst)
