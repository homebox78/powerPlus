"""시안 재조립 실행기.

  py run.py build <src.pptx> <dst.pptx> 4,6,7      쪽별 모듈(sNN.py)의 build(c) 를 적용해 저장
  py run.py test <쪽> <out_dir>                    그 쪽만 적용 → 렌더 → 시안과 나란히(cmp_sNN.png)
  py run.py dump <쪽> <out.txt>                    원본 그 쪽의 도형·글 목록

PowerPoint 렌더는 한 번에 하나만(잠금 폴더) 돈다.
"""
import sys, os, glob, time, importlib, shutil
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
ROOT = "D:/powerPlus/제안서"
SRC = ROOT + "/청주시_발표자료(제안요약서)_v0.61_속카드2.pptx"
MAP = {4: 1, 6: 2, 7: 3, 8: 4, 10: 5, 11: 6, 12: 7, 13: 8, 14: 9, 16: 10, 17: 11, 18: 12, 19: 13, 20: 14, 21: 15, 22: 16}
LOCK = os.path.join(os.environ.get("TEMP", "C:/Temp"), "pp_render.lock")


# 2차(v0.66~): 시안이 없던 쪽은 GPT Image 로 만든 시안(design_ppt2/sNN.png)을 쓰고 v0.65 위에서 작업한다
SRC2 = ROOT + "/청주시_발표자료(제안요약서)_v0.65_가이드.pptx"
NEW = [5, 23, 24, 25, 26, 27, 28, 29, 31, 32, 33, 34, 35, 37, 38, 39]


def src_for(idx):
    return SRC2 if idx in NEW else SRC


def design_path(idx):
    if idx in NEW:
        return ROOT + "/design_ppt2/s%02d.png" % idx
    return sorted(glob.glob(ROOT + "/design_ppt/*.png"))[MAP[idx] - 1]


def apply(prs, idx, work):
    from lib import Ctx
    mod = importlib.import_module("s%02d" % idx)
    mod.build(Ctx(prs, idx, design_path(idx), work))


def render(pptx, idx, out):
    import win32com.client, pythoncom
    t0 = time.time()
    while True:
        try:
            os.mkdir(LOCK); break
        except FileExistsError:
            if time.time() - os.path.getmtime(LOCK) > 180:   # 죽은 잠금
                try: os.rmdir(LOCK)
                except OSError: pass
            if time.time() - t0 > 900: raise SystemExit("render lock timeout")
            time.sleep(2)
    try:
        pythoncom.CoInitialize()
        app = win32com.client.Dispatch("PowerPoint.Application")
        p = app.Presentations.Open(os.path.abspath(pptx), True, False, False)
        p.Slides(idx).Export(os.path.abspath(out), "PNG", 2000, 1385)
        p.Close()
    finally:
        try: os.rmdir(LOCK)
        except OSError: pass


def main():
    from pptx import Presentation
    cmd = sys.argv[1]
    if cmd == "build":
        src, dst, pages = sys.argv[2], sys.argv[3], [int(x) for x in sys.argv[4].split(",")]
        prs = Presentation(src)
        work = os.path.join(os.path.dirname(os.path.abspath(dst)), "_design_crops")
        for i in pages:
            apply(prs, i, work); print("built", i)
        prs.save(dst); print("saved")
    elif cmd == "test":
        idx, out = int(sys.argv[2]), sys.argv[3]
        os.makedirs(out, exist_ok=True)
        prs = Presentation(src_for(idx))
        apply(prs, idx, os.path.join(out, "crops"))
        one = os.path.join(out, "test_s%02d.pptx" % idx); prs.save(one)
        png = os.path.join(out, "render_s%02d.png" % idx); render(one, idx, png)
        from PIL import Image
        a = Image.open(design_path(idx)).convert("RGB"); b = Image.open(png).convert("RGB")
        a = a.resize((2000, 1125))
        S = Image.new("RGB", (2000, 1125 + 1385 + 10), "#ff0000")
        S.paste(a, (0, 0)); S.paste(b, (0, 1135))
        S.save(os.path.join(out, "cmp_s%02d.png" % idx))
        print("ok", png)
    elif cmd == "dump":
        idx, out = int(sys.argv[2]), sys.argv[3]
        prs = Presentation(src_for(idx)); s = prs.slides[idx - 1]
        f = open(out, "w", encoding="utf-8")

        def walk(sh, d=0):
            t = sh.text_frame.text.replace("\n", " / ").replace("\x0b", " / ") if sh.has_text_frame else ""
            if sh.shape_type == 19:
                t = " | ".join(" ; ".join(c.text.replace("\n", " ") for c in r.cells) for r in sh.table.rows)
            f.write("  " * d + "[%s] %s L%.2f T%.2f W%.2f H%.2f %s\n" % (
                sh.shape_type, sh.name, (sh.left or 0) / 914400, (sh.top or 0) / 914400, (sh.width or 0) / 914400, (sh.height or 0) / 914400, t))
            if sh.shape_type == 6:
                for c in sh.shapes: walk(c, d + 1)
        for sh in s.shapes: walk(sh)
        f.close(); print("ok")


if __name__ == "__main__":
    main()
