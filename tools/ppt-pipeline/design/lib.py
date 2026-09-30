"""시안 통이미지 → 장표 재조립 공용 도구.

그림은 시안에서 오려 쓰고(crop), 글·도형은 시안 좌표·색을 읽어 PowerPoint 도형으로 그린다.
좌표는 전부 시안을 2000×1125 로 본 값. 가로는 K(=10.83in/2000) 배율 그대로,
세로는 덩어리(Grp) 중심만 vmap 으로 옮기고 덩어리 안쪽은 같은 배율을 쓴다.
"""
import os
from lxml import etree
from PIL import Image
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_LINE
from pptx.oxml.ns import qn

F2, F3, F4 = "a시월구일2", "a시월구일3", "a시월구일4"
WHITE = RGBColor(255, 255, 255)
NAVY = RGBColor(0x10, 0x2A, 0x5C)
BLUE = RGBColor(0x2B, 0x6D, 0xE8)
PINK = RGBColor(0xF2, 0x4F, 0x6E)
LEFT, CENTER, RIGHT = PP_ALIGN.LEFT, PP_ALIGN.CENTER, PP_ALIGN.RIGHT
SLIDE_W, SLIDE_H = 10.83, 7.5
K = SLIDE_W / 2000.0


def hexrgb(h):
    h = h.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


class Grp:
    def __init__(self, ctx, y0, y1):
        self.c = (y0 + y1) / 2.0
        self.t = ctx.cy(self.c)

    def X(self, x): return x * K
    def Y(self, y): return self.t + (y - self.c) * K


class Ctx:
    def __init__(self, prs, idx, design, work):
        self.prs = prs; self.s = prs.slides[idx - 1]; self.tree = self.s.shapes._spTree
        self.D = Image.open(design).convert("RGB"); self.Z = self.D.size[0] / 2000.0
        self.work = work; os.makedirs(work, exist_ok=True)
        self.idx = idx; self._n = 0
        # 세로 기본 대응: 시안 y=160 → 1.5in, y=1125 → 7.5in
        self.vmap(160, 1.5, 1125, 7.5)

    # ── 좌표 ──
    def vmap(self, y0, in0, y1, in1):
        """시안 y0..y1 을 장표 in0..in1 inch 로 (덩어리 중심 배치용)."""
        self._v = (y0, in0, (in1 - in0) / float(y1 - y0))

    def cy(self, y):
        y0, in0, k = self._v
        return in0 + (y - y0) * k

    def grp(self, y0, y1):
        return Grp(self, y0, y1)

    # ── 시안 읽기 ──
    def rgb(self, x, y):
        return RGBColor(*self.D.getpixel((int(x * self.Z), int(y * self.Z))))

    def crop(self, name, x0, y0, x1, y1, fade_top=0, cut=None, cut_thresh=45):
        """시안에서 오려 PNG 로 저장. fade_top: 위쪽을 투명하게 흐림(px, 2000 기준).
        cut=(x,y): 그 점의 바탕색과 같은(가장자리에서 이어진) 영역을 투명하게."""
        im = self.D.crop((int(x0 * self.Z), int(y0 * self.Z), int(x1 * self.Z), int(y1 * self.Z)))
        if cut is not None:
            from PIL import ImageDraw
            key = (255, 0, 255)
            ref = self.D.getpixel((int(cut[0] * self.Z), int(cut[1] * self.Z)))
            w, h = im.size
            def near(p): return sum(abs(a - b) for a, b in zip(p, ref)) <= cut_thresh
            pts = [(x, y) for x in range(0, w, 6) for y in (0, h - 1)] + [(x, y) for y in range(0, h, 6) for x in (0, w - 1)]
            for pt in pts:
                p = im.getpixel(pt)
                if p != key and near(p):
                    ImageDraw.floodfill(im, pt, key, thresh=cut_thresh)
            im = im.convert("RGBA"); px = im.load()
            for y in range(h):
                for x in range(w):
                    if px[x, y][:3] == key:
                        px[x, y] = (255, 255, 255, 0)
        if fade_top:
            im = im.convert("RGBA"); px = im.load(); n = int(fade_top * self.Z)
            for y in range(min(n, im.size[1])):
                a = y / float(n)
                for x in range(im.size[0]):
                    r, g, b, al = px[x, y]; px[x, y] = (r, g, b, int(al * a))
        p = os.path.join(self.work, "s%02d_%s.png" % (self.idx, name)); im.save(p)
        return p

    # ── 기존 도형 ──
    def names(self):
        return [sh.name for sh in self.s.shapes]

    def remove(self, *names):
        for sh in list(self.s.shapes):
            if sh.name in names:
                self.tree.remove(sh._element)

    def keep_only(self, *names):
        """적은 이름만 남기고 나머지 최상위 도형을 지운다."""
        for sh in list(self.s.shapes):
            if sh.name not in names:
                self.tree.remove(sh._element)

    def to_back(self, *shapes):
        for sh in shapes:
            self.tree.remove(sh._element); self.tree.insert(2, sh._element)

    # ── 그리기 (단위 inch) ──
    def shape(self, kind, l, t, w, h, fill=None, line=None, lw=0.75, adj=None, dash=False, alpha=None, name=None):
        b = self.s.shapes.add_shape(kind, Inches(l), Inches(t), Inches(w), Inches(h))
        b.shadow.inherit = False
        self._n += 1; b.name = name or "시안 %d" % self._n
        if adj is not None:
            for i, a in enumerate(adj if isinstance(adj, (list, tuple)) else [adj]):
                b.adjustments[i] = a
        if fill is None: b.fill.background()
        else:
            b.fill.solid(); b.fill.fore_color.rgb = fill
            if alpha is not None:       # 0~1 불투명도
                sf = b.fill._xPr.find(qn("a:solidFill"))[0]
                a = etree.SubElement(sf, qn("a:alpha")); a.set("val", str(int(alpha * 100000)))
        if line is None: b.line.fill.background()
        else:
            b.line.color.rgb = line; b.line.width = Pt(lw)
            if dash: b.line.dash_style = MSO_LINE.DASH
        tf = b.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE; tf.word_wrap = False
        return b

    def rect(self, l, t, w, h, **kw): return self.shape(MSO_SHAPE.RECTANGLE, l, t, w, h, **kw)
    def rrect(self, l, t, w, h, adj=0.1, **kw): return self.shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h, adj=adj, **kw)
    def pill(self, l, t, w, h, **kw): return self.shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h, adj=0.5, **kw)
    def oval(self, l, t, w, h, **kw): return self.shape(MSO_SHAPE.OVAL, l, t, w, h, **kw)

    def write(self, b, lines, align=CENTER, anchor="middle", wrap=False, spacing=None, margin=None):
        """lines = [[(글, pt, 색, 서체), …], …] — 문단마다 런 목록. 글자는 7pt 이상."""
        tf = b.text_frame; tf.word_wrap = wrap
        tf.vertical_anchor = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}[anchor]
        if margin is not None:
            tf.margin_left = tf.margin_right = Inches(margin)
        for i, runs in enumerate(lines):
            pg = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            pg.alignment = align
            if spacing: pg.line_spacing = spacing
            for tx, size, color, font in runs:
                r = pg.add_run(); r.text = tx
                r.font.size = Pt(size); r.font.color.rgb = color; r.font.name = font
                rPr = r._r.get_or_add_rPr()
                ea = rPr.find(qn("a:ea"))
                if ea is None: ea = etree.SubElement(rPr, qn("a:ea"))
                ea.set("typeface", font)
        return b

    def text(self, l, t, w, h, lines, align=CENTER, **kw):
        """글만 있는 상자."""
        return self.write(self.rect(l, t, w, h), lines, align, **kw)

    def pic(self, path, l, t, w=None, h=None):
        kw = {}
        if w is not None: kw["width"] = Inches(w)
        if h is not None and w is None: kw["height"] = Inches(h)
        p = self.s.shapes.add_picture(path, Inches(l), Inches(t), **kw)
        self._n += 1; p.name = "시안 그림 %d" % self._n
        return p

    def line(self, x0, y0, x1, y1, color, w=1.0, dash=False, arrow=False):
        c = self.s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x0), Inches(y0), Inches(x1), Inches(y1))
        c.line.color.rgb = color; c.line.width = Pt(w)
        if dash: c.line.dash_style = MSO_LINE.DASH
        if arrow:
            ln = c.line._get_or_add_ln(); te = etree.SubElement(ln, qn("a:tailEnd")); te.set("type", "triangle")
        self._n += 1; c.name = "시안 선 %d" % self._n
        return c

    def background(self, x=1000, y=270, top=0.87):
        """본문 바탕을 시안 바탕색으로 (맨 뒤)."""
        # 오른쪽 아래 쪽 번호(레이아웃) 자리는 비운다
        col = self.rgb(x, y)
        a = self.rect(0, top, SLIDE_W, 7.12 - top, fill=col, name="시안 바탕")
        b = self.rect(0, 7.12, 9.3, SLIDE_H - 7.12, fill=col, name="시안 바탕 2")
        self.to_back(b); self.to_back(a)
        return a
