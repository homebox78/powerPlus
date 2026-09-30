"""v0.66 → v0.67: 안내선 넘침 보정 + 8쪽 그림 교체 + 34쪽 키메시지 크기.

본문은 렌더(audit 과 같은 잉크 측정)를 기준으로, 넘친 축만 줄여 안내선 안에 넣는다(늘리지는 않는다).
사용: py fix67.py <src.pptx> <dst.pptx> <렌더 폴더(sNN.png)>
"""
import os, sys
import numpy as np
from PIL import Image
from pptx import Presentation
from pptx.util import Emu, Pt, Inches

sys.stdout.reconfigure(encoding="utf-8")
IN, CM = 914400, 360000
src, dst, rdir = [os.path.abspath(a) for a in sys.argv[1:4]]
p = Presentation(src)
W, H = p.slide_width, p.slide_height
G = dict(L=W / 2 - 13.0 * CM, R=W / 2 + 13.0 * CM, T=H / 2 - 6.9 * CM, B=H / 2 + 8.3 * CM)
PAGES = [8, 10, 12, 14, 17, 23, 24, 26, 27, 28, 29, 32, 33, 34, 35, 38]
ONLY_X = {27}          # 27쪽 위쪽 넘침은 카드 모서리의 '추가 제안' 배지(의도)
NEW_IMG = r"D:/powerPlus/제안서/_design_crops/s08_office_a.png"


def box(sh):
    return sh.left, sh.top, sh.left + sh.width, sh.top + sh.height


# ── 8쪽: 서버실 그림 → 사무실 모니터 대시보드 ──
s = p.slides[7]
old = [sh for sh in s.shapes if sh.shape_type == 13 and sh.left < 0.2 * CM and sh.height > 3 * IN][0]
t, h = old.top, old.height
iw, ih = Image.open(NEW_IMG).size
w = int(h * iw / ih)
pic = s.shapes.add_picture(NEW_IMG, Emu(int(G["L"])), t, Emu(w), h)
pic.name = old.name
old._element.addprevious(pic._element)          # 원래 z 순서 자리
old._element.getparent().remove(old._element)
print("08 그림 교체 %.2fx%.2fin" % (w / IN, h / IN))

# ── 34쪽: 키메시지가 안내선보다 넓다 → 22pt ──
for sh in p.slides[33].shapes:
    if sh.name.startswith("키메시지"):
        for pg in sh.text_frame.paragraphs:
            for r in pg.runs:
                r.font.size = Pt(22)
        print("34 키메시지 22pt")


# ── 본문 안내선 맞춤 ──
def measure(png, s, ks, keep):
    a = np.asarray(Image.open(png).convert("RGB")).astype(int)
    k = a.shape[1] / W
    ink = (a.max(axis=2) - a.min(axis=2) > 12) | (a.sum(axis=2) < 735)
    bg = np.median(a[int(3.0 * IN * k):int(7.0 * IN * k), 2:8].reshape(-1, 3), axis=0)
    ink &= np.abs(a - bg).sum(axis=2) > 24
    lim = np.zeros(ink.shape, bool)
    for sh in keep:
        l, t, r, b = box(sh)
        lim[max(0, int(t * k) - 2):int(b * k) + 2, max(0, int(l * k) - 2):int(r * k) + 2] = True
    for sh in ks:
        lim[int(sh.top * k):int((sh.top + sh.height) * k) + 1, int(sh.left * k):int((sh.left + sh.width) * k) + 1] = False
    lim[:int(0.9 * IN * k), :] = False
    lim[int(7.02 * IN * k) + 40:, int(9.2 * IN * k):] = False
    ys, xs = np.nonzero(ink & lim)
    return xs.min() / k, ys.min() / k, (xs.max() + 1) / k, (ys.max() + 1) / k


for i in PAGES:
    s = p.slides[i - 1]
    ks = [sh for sh in s.shapes if sh.name.startswith("키메시지")]
    keep = []
    for sh in s.shapes:
        if sh.left is None or sh.width is None or sh in ks:
            continue
        l, t, r, b = box(sh)
        if b < 0.87 * IN or t > 7.3 * IN:
            continue
        edge = l < 0.2 * CM or r > W - 0.2 * CM
        texty = sh.has_text_frame and sh.text_frame.text.strip()
        if edge and not texty:
            continue
        keep.append(sh)
    png = os.path.join(rdir, "s%02d.png" % i)
    if i == 8:   # 그림을 바꿨으니 그림 자리는 새 좌표로 계산
        L, T, R, B = measure(png, s, ks, [sh for sh in keep if sh is not pic])
        L = min(L, pic.left)
    else:
        L, T, R, B = measure(png, s, ks, keep)
    kb = max((sh.top + sh.height for sh in ks), default=G["T"])
    GT = max(G["T"], kb + 0.1 * IN)
    sx = sy = 1.0
    cx0, cy0 = (L + R) / 2, (T + B) / 2
    tcx, tcy = cx0, cy0
    pad = 0.02 * IN
    if L < G["L"] or R > G["R"]:
        span = min(R - L, G["R"] - G["L"] - 2 * pad)
        sx = span / (R - L)
        tcx = min(max(cx0, G["L"] + pad + span / 2), G["R"] - pad - span / 2)
    if i not in ONLY_X and (T < GT or B > G["B"]):
        span = min(B - T, G["B"] - GT - 2 * pad)
        sy = span / (B - T)
        tcy = min(max(cy0, GT + pad + span / 2), G["B"] - pad - span / 2)
    if sx == 1 and sy == 1 and tcx == cx0 and tcy == cy0:
        continue
    fx = lambda x: tcx + (x - cx0) * sx
    fy = lambda y: tcy + (y - cy0) * sy
    for sh in keep:
        l, t, r, b = box(sh)
        if sh.shape_type == 13:            # 그림은 비율 유지
            kk = min(sx, sy)
            cx, cy = fx((l + r) / 2), fy((t + b) / 2)
            ww, hh = sh.width * kk, sh.height * kk
            sh.left, sh.top = Emu(int(cx - ww / 2)), Emu(int(cy - hh / 2))
            sh.width, sh.height = Emu(int(ww)), Emu(int(hh))
            continue
        nl, nt, nr, nb = fx(l), fy(t), fx(r), fy(b)
        if sh.has_text_frame and sh.text_frame.text.strip() and sh.shape_type == 17:
            # 글자는 그대로라 글상자 크기는 유지하고 자리만 옮긴다
            c = (nl + nr) / 2; nl, nr = c - (r - l) / 2, c + (r - l) / 2
            m = (nt + nb) / 2; nt, nb = m - (b - t) / 2, m + (b - t) / 2
        sh.left, sh.top = Emu(int(nl)), Emu(int(nt))
        sh.width, sh.height = Emu(max(1, int(nr - nl))), Emu(max(1, int(nb - nt)))
        if sh.shape_type == 19:
            for c in sh.table.columns: c.width = Emu(int(c.width * sx))
    print("%02d 좌%+.2f 우%+.2f 상%+.2f 하%+.2f → 가로 %.3f 세로 %.3f" % (
        i, (L - G["L"]) / IN, (G["R"] - R) / IN, (T - GT) / IN, (G["B"] - B) / IN, sx, sy))
p.save(dst)
print("저장", dst)
