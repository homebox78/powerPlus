"""27쪽 — 7. 기능 및 성능 개선 (2/2) (2차 시안 design_ppt2/s27.png).

좌: 대량의 동시접속 시 성능 향상 / 우: 변경부서 데이터 이관 기능 구축.
글은 원래 장표 문구, 배치·색·그림은 시안. 화면 캡처 3장과 '추가 제안' 배지는 원래 것을 남겨 자리만 옮긴다.
"""
from lib import *
from pptx.util import Emu

KM_FONT = "G마켓 산스 TTF Bold"
KM_NAVY, KM_BLUE = hexrgb("0B2E6B"), hexrgb("2070E8")
TXT = hexrgb("162E63")        # 본문 글
TTL = hexrgb("16357F")        # 소제목 글
SUB = hexrgb("1E3A70")        # 패널 부제
DOT = hexrgb("2466D8")
PANEL = hexrgb("EFF8FE")
PANEL_LN = hexrgb("D2E7FD")
HDR = hexrgb("3587EA")
BOX_HDR = hexrgb("2B75CD")
BOX_LN = hexrgb("D2E7FD")


def build(c):
    # ── 원래 키메시지 문구(색 다른 구절 = 강조) ──
    km_runs = []
    for sh in c.s.shapes:
        if sh.name == "직사각형 124":
            first = None
            for pg in sh.text_frame.paragraphs:
                for r in pg.runs:
                    if not r.text:
                        continue
                    try:
                        col = tuple(r.font.color.rgb)
                    except Exception:
                        col = None
                    if first is None:
                        first = col
                    km_runs.append((r.text, col != first))
    c.keep_only("TextBox 67", "Picture 246", "Picture 247", "Picture 248", "그룹 118")
    shots = {sh.name: sh for sh in c.s.shapes}

    c.background(1000, 245)
    c.vmap(252, 1.62, 1098, 7.02)

    # ── 키메시지 ──
    km = c.rect(0.20, 1.03, 10.43, 0.46, name="키메시지")
    lines = []
    for i, (t, acc) in enumerate(km_runs):
        if i == 0: t = t.lstrip()
        if i == len(km_runs) - 1: t = t.rstrip()
        lines.append((t, 23, KM_BLUE if acc else KM_NAVY, KM_FONT))
    c.write(km, [lines])

    def box(g, fn, x0, y0, x1, y1, **kw):
        return fn(g.X(x0), g.Y(y0), (x1 - x0) * K, (y1 - y0) * K, **kw)

    def txt(g, x0, y0, x1, y1, lines, align=CENTER, **kw):
        return c.text(g.X(x0), g.Y(y0), (x1 - x0) * K, (y1 - y0) * K, lines, align, **kw)

    def r2(l, t, w, h, **kw):
        return c.shape(MSO_SHAPE.ROUND_2_SAME_RECTANGLE, l, t, w, h, **kw)

    def section(g, x, y, label):
        """● 소제목 (y = 글 가운데)."""
        box(g, c.oval, x, y - 15, x + 30, y + 15, fill=DOT)
        box(g, c.oval, x + 9, y - 6, x + 21, y + 6, fill=WHITE)
        txt(g, x + 42, y - 16, x + 700, y + 16, [[(label, 10.5, TTL, F3)]], LEFT)

    # ── 두 패널 틀 + 머리 띠 + 부제 ──
    top, bot = c.cy(252), c.cy(1098)
    gh = c.grp(263, 338)
    gs = c.grp(354, 380)
    for x0, x1, head, sub in ((45, 992, "대량의 동시접속 시 성능 향상",
                               "시스템의 응답 대기 낭비 시간 제거 , 3초 이내 응답 가능토록 지원"),
                              (1012, 1956, "변경부서 데이터 이관 기능 구축",
                               "조직개편 시 부서·팀 단위 업무 데이터 이관으로 신속한 업무 연속성 확보")):
        c.rrect(x0 * K, top, (x1 - x0) * K, bot - top, adj=0.025, fill=PANEL, line=PANEL_LN, lw=1.0)
        b = box(gh, c.rrect, x0 + 10, 263, x1 - 10, 338, adj=0.22, fill=HDR)
        c.write(b, [[(head, 15, WHITE, F3)]])
        txt(gs, x0, 354, x1, 380, [[(sub, 8.5, SUB, F2)]])

    # ════ 왼쪽 ════
    g = c.grp(398, 425)
    section(g, 70, 411, "성능 최적화 절차")

    # 단계 화살표 + 아래 꼬리표
    g = c.grp(439, 718)      # 화살표·그림 칸은 한 덩어리(시안처럼 붙게)
    steps = [(70, 311, "상시 모니터링", hexrgb("3D8BF0"), MSO_SHAPE.PENTAGON),
             (289, 538, "분석 및 조치단계", hexrgb("67ABF6"), MSO_SHAPE.CHEVRON),
             (516, 765, "보고서 작성", hexrgb("4A9CF3"), MSO_SHAPE.CHEVRON),
             (743, 973, "최적화 단계", hexrgb("2E83EA"), MSO_SHAPE.CHEVRON)]
    for x0, x1, lab, col, kind in steps:
        box(g, lambda l, t, w, h, **kw: c.shape(kind, l, t, w, h, **kw), x0, 439, x1, 536, fill=col, adj=0.36)
    for x0, x1, lab, col, kind in steps:
        tx0 = x0 + (20 if kind == MSO_SHAPE.PENTAGON else 60)
        txt(g, tx0, 450, x1 - 30, 500, [[(lab, 8.5, WHITE, F3)]])
    for x0, x1, lab in ((185, 382, "기준설정 및 취합"), (453, 609, "보고서 작성"), (688, 837, "조치/개선")):
        b = box(g, lambda l, t, w, h, **kw: c.shape(MSO_SHAPE.PENTAGON, l, t, w, h, **kw),
                x0, 506, x1, 539, fill=WHITE, adj=0.4)
        c.write(b, [[(lab, 7.5, TXT, F3)]])

    # 단계별 그림 칸
    box(g, c.rrect, 70, 553, 970, 718, adj=0.05, fill=hexrgb("F9FCFF"), line=BOX_LN, lw=0.75)
    seps = [205, 357, 510, 665, 824]
    for x in seps:
        c.line(g.X(x), g.Y(560), g.X(x), g.Y(712), hexrgb("A9CCF0"), 0.75, dash=True)
    cells = [70] + seps + [970]
    icons = [(92, 564, 182, 652), (244, 562, 322, 650), (390, 562, 482, 650),
             (549, 562, 635, 650), (702, 562, 794, 650), (854, 562, 944, 650)]
    labels = [["상시 모니터링", "원인 도출"], ["성능 분석용", "데이터 취합"], ["원인분석 및", "해결방안 결정"],
              ["보고서 작성"], ["결정대로 조치"], ["조치결과 확인"]]
    for i, (ic, lab) in enumerate(zip(icons, labels)):
        x0, y0, x1, y1 = ic
        p = c.crop("ico%d" % i, x0, y0, x1, y1, cut=(x0 + 2, y0 + 2), cut_thresh=40)
        c.pic(p, g.X(x0), g.Y(y0), (x1 - x0) * K)
        cx0, cx1 = cells[i], cells[i + 1]
        ln = [[("• " + lab[0], 7.2, TXT, F2)]] + ([[(lab[1], 7.2, TXT, F2)]] if len(lab) > 1 else [])
        txt(g, cx0 + 1, 658, cx1 - 1, 710, ln, spacing=0.95)

    # 성능 관리 사이클 주요 활동
    g = c.grp(733, 888)
    g.t -= 0.18              # 위 그림 칸에 붙게(시안 간격)
    box(g, c.rrect, 68, 733, 972, 888, adj=0.07, fill=WHITE, line=BOX_LN, lw=1.0)
    b = box(g, r2, 68, 733, 972, 776, adj=[0.3, 0], fill=BOX_HDR)
    c.write(b, [[("성능 관리 사이클 주요 활동", 10, WHITE, F3)]])
    for i, t in enumerate(["상시 모니터링으로 성능저하 원인 도출",
                           "성능저하 원인 데이터를 취합하여 문제점을 분석하고 문제 해결방안 도출",
                           "도출된 결과를 바탕으로 성능 보고서 작성 및 개선방안 도출",
                           "변경내역에 대한 이력관리 실시"]):
        y = 793 + i * 26
        txt(g, 92, y - 12, 960, y + 12, [[("•  " + t, 8, TXT, F2)]], LEFT)

    g = c.grp(900, 930)
    g.t -= 0.06
    section(g, 70, 915, "주요 성능개선 대상")

    # 성능 도식 + 대상 목록
    g = c.grp(932, 1090)
    box(g, c.oval, 98, 936, 290, 1076, line=hexrgb("CBE3FB"), lw=5)
    b = box(g, c.oval, 157, 986, 246, 1052, fill=hexrgb("3D8CF0"))
    c.write(b, [[("성능", 12, WHITE, F3)]])
    for cx, cy, lab in ((200, 959, "CPU"), (131, 1047, "I/O"), (276, 1047, "Mem")):
        b = box(g, c.oval, cx - 27, cy - 27, cx + 27, cy + 27, fill=hexrgb("2E80EC"), line=WHITE, lw=1.5)
        c.write(b, [[(lab, 7.5, WHITE, F3)]])
    for i, t in enumerate(["DB 튜닝 (접근경로, 인덱스, 조인, 클러스터, 부분 범위 스캔)", "악성 모듈 추출 및 보완",
                           "WAS 성능 개선 안 도출 및 보고", "응용 프로그램 구조 개선"]):
        y = 966 + i * 31
        txt(g, 342, y - 14, 975, y + 14, [[("•  " + t, 8.5, TXT, F2)]], LEFT)

    # ════ 오른쪽 ════
    g = c.grp(398, 425)
    section(g, 1037, 411, "변경부서현황 등록(엑셀 업로드)·목록 관리")

    g = c.grp(436, 652)
    p1, p2 = shots["Picture 246"], shots["Picture 247"]
    a1, a2 = p1.width / p1.height, p2.width / p2.height
    gap = 0.07
    left = 1068 * K
    H = ((1932 - 1068) * K - gap) / (a1 + a2)     # 두 캡처가 시안 폭(1068~1932)을 채우게
    w1, w2 = H * a1, H * a2
    ty = g.Y(544) - H / 2
    p1.left, p1.top, p1.width, p1.height = Inches(left), Inches(ty), Inches(w1), Inches(H)
    p2.left, p2.top, p2.width, p2.height = Inches(left + w1 + gap), Inches(ty), Inches(w2), Inches(H)

    g = c.grp(675, 702)
    section(g, 1037, 688, "변경부서현황 상세(예시)")

    g = c.grp(710, 872)
    p3 = shots["Picture 248"]
    h3 = 160 * K
    w3 = h3 * p3.width / p3.height
    p3.left, p3.top, p3.width, p3.height = Inches(left), Inches(g.Y(711)), Inches(w3), Inches(h3)
    ex0 = left / K + w3 / K + 18
    box(g, c.rrect, ex0, 711, 1935, 871, adj=0.08, fill=WHITE, line=hexrgb("D7EAFC"), lw=1.0)
    b = box(g, r2, ex0, 711, 1935, 755, adj=[0.3, 0], fill=hexrgb("2973CD"))
    c.write(b, [[("기대효과", 10.5, WHITE, F3)]])
    for i, t in enumerate(["조직개편 당일 업무 공백 없는 연속성 확보", "수작업 이관 대비 오류·소요시간 감소",
                           "이관 이력 관리로 변경 추적성 확보"]):
        y = 784 + i * 29
        txt(g, ex0 + 30, y - 13, 1930, y + 13, [[("•  " + t, 8.5, TXT, F2)]], LEFT)

    g = c.grp(888, 916)
    section(g, 1037, 902, "데이터 이관 처리 절차")

    g = c.grp(926, 1078)
    xs = [1123, 1311, 1494, 1675, 1854]
    labs = [("단위업무 분석", "부서·팀 업무 식별"), ("변경부서 정보", "엑셀 일괄 등록"),
            ("이관 대상 테이블", "건수 산출·검토"), ("데이터이관 실행", "변경 이력 관리"),
            ("조직개편 즉시", "업무 연속성 확보")]
    for i, (x, lab) in enumerate(zip(xs, labs)):
        box(g, c.oval, x - 48, 927, x + 48, 1023, fill=WHITE, line=hexrgb("5E9FF0"), lw=1.25)
        p = c.crop("proc%d" % i, x - 34, 941, x + 34, 1009, cut=(x - 33, 942), cut_thresh=40)
        c.pic(p, g.X(x - 34), g.Y(941), 68 * K)
        txt(g, x - 92, 1028, x + 92, 1078, [[(lab[0], 8, TXT, F2)], [(lab[1], 8, TXT, F2)]], spacing=0.95)
        if i < 4:
            mx = (x + xs[i + 1]) / 2.0
            box(g, lambda l, t, w, h, **kw: c.shape(MSO_SHAPE.CHEVRON, l, t, w, h, **kw),
                mx - 7, 962, mx + 7, 986, fill=hexrgb("BFDDF9"), adj=0.55)

    # '추가 제안' 배지 — 원래 것, 오른쪽 머리 띠 오른쪽 위에 걸치게
    bd = shots["그룹 118"]
    s = 0.62 / (bd.width / 914400.0)
    bd.width, bd.height = Emu(int(bd.width * s)), Emu(int(bd.height * s))
    bd.left, bd.top = Inches(10.62) - bd.width, Inches(c.cy(263) - 0.30)
    # 원래 실물(캡처 3장·배지)은 새로 그린 패널 위로
    for sh in (p1, p2, p3, bd):
        c.tree.remove(sh._element); c.tree.append(sh._element)
