"""원래 장표와 시안 반영본에서 쪽마다 키메시지 후보(머리 띠 아래 2.6in 안, 글자 최대)를 찾아 보여 준다."""
import sys
from pptx import Presentation

sys.stdout.reconfigure(encoding="utf-8")
PAGES = [4, 6, 7, 8, 10, 11, 12, 13, 14, 16, 17, 18, 19, 20, 21, 22]
IN = 914400


def maxpt(sh):
    m = 0
    for pg in sh.text_frame.paragraphs:
        for r in pg.runs:
            if r.font.size and r.text.strip():
                m = max(m, r.font.size.pt)
    return m


def cands(s):
    out = []
    for sh in s.shapes:
        if not sh.has_text_frame or not sh.text_frame.text.strip() or sh.top is None:
            continue
        t = sh.top / IN
        if 0.87 <= t < 2.6:
            out.append((maxpt(sh), sh))
    out.sort(key=lambda x: -x[0])
    return out[:3]


a = Presentation(sys.argv[1]); b = Presentation(sys.argv[2])
for i in PAGES:
    for tag, p in (("원", a), ("새", b)):
        for pt, sh in cands(p.slides[i - 1]):
            print("%02d %s %5.1fpt %-14s L%.2f T%.2f W%.2f H%.2f | %s" % (
                i, tag, pt, sh.name[:14], sh.left / IN, sh.top / IN, sh.width / IN, sh.height / IN,
                sh.text_frame.text.replace("\n", " / ").replace("\x0b", " ")[:40]))
    print()
