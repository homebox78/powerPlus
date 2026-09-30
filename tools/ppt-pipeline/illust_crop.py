# -*- coding: utf-8 -*-
"""일러스트 키우기·높이 맞춰 자르기·방향 반전 — 작아서 허전한 인물 일러스트를 텍스트와 짝이 보이게.
   (청주 v0.17: 사용자가 5장을 직접 손본 방식 = 폭은 키우고 아래(다리)를 잘라 높이는 유지, 인물이 텍스트를 보게 좌우 반전)

   그림을 라이브러리 원본(고해상도)으로 다시 넣고, 잉크 영역만 남긴 뒤
   srcRect b(아래 자르기) + a:xfrm flipH 로 표현한다. 그룹 안 그림도 그룹 축척을 풀어 슬라이드 기준으로 계산한다.

   인자: <src> <dst> <spec.json> <lib_dir>
   spec = [{"slide":5, "name":"SW01", "asset":873, "flip":true, "H":0.47, "cut":0.3, "ax":"r", "ay":"c"}, ...]
     H   : 보이는 높이(in, 슬라이드 기준). 생략하면 지금 높이
     cut : 잉크 높이 중 아래에서 잘라 낼 비율(0이면 자르지 않음)
     ax  : 가로 기준 — r 오른쪽 변 고정(텍스트가 오른쪽) / l 왼쪽 변 / c 가운데
     ay  : 세로 기준 — c 가운데 / b 아래 변 / t 위 변
     maxW: 넘으면 cut 을 줄여 폭을 맞춘다(선택)"""
import io, json, os, sys
from PIL import Image, ImageChops
from pptx import Presentation
sys.stdout.reconfigure(encoding="utf-8")
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
E = 914400
src, dst, spec, lib = sys.argv[1:5]
jobs = json.load(open(spec, encoding="utf-8"))
p = Presentation(src)


def walk(shapes):
    for s in shapes:
        if s.shape_type == 6:
            yield from walk(s.shapes)
        else:
            yield s


def grp_chain(el):
    """그림에서 슬라이드까지 조상 그룹 목록(바깥 → 안쪽)."""
    out = []
    g = el.getparent()
    while g is not None and g.tag == P + "grpSp":
        out.append(g)
        g = g.getparent()
    return out[::-1]


def gx(g):
    x = g.find(P + "grpSpPr").find(A + "xfrm")
    off, ext = x.find(A + "off"), x.find(A + "ext")
    co, ce = x.find(A + "chOff"), x.find(A + "chExt")
    sx = int(ext.get("cx")) / max(1, int(ce.get("cx")))
    sy = int(ext.get("cy")) / max(1, int(ce.get("cy")))
    return int(off.get("x")), int(off.get("y")), int(co.get("x")), int(co.get("y")), sx, sy


def to_slide(chain, x, y, w, h):
    for g in reversed(chain):
        ox, oy, cx, cy, sx, sy = gx(g)
        x, y, w, h = ox + (x - cx) * sx, oy + (y - cy) * sy, w * sx, h * sy
    return x, y, w, h


def to_child(chain, x, y, w, h):
    for g in chain:
        ox, oy, cx, cy, sx, sy = gx(g)
        x, y, w, h = cx + (x - ox) / sx, cy + (y - oy) / sy, w / sx, h / sy
    return x, y, w, h


def ink(im):
    a = im.getchannel("A")
    if a.getextrema()[0] < 250:                      # 투명 배경
        return a.point(lambda v: 255 if v > 8 else 0).getbbox()
    bg = Image.new("RGB", im.size, (255, 255, 255))   # 불투명 흰 배경
    d = ImageChops.difference(im.convert("RGB"), bg).convert("L")
    return d.point(lambda v: 255 if v > 12 else 0).getbbox()


done = 0
for j in jobs:
    hits = [s for s in walk(p.slides[j["slide"] - 1].shapes) if s.shape_type == 13 and s.name == j["name"]]
    if not hits:
        print("없음", j["slide"], j["name"]); continue
    sh = hits[0]
    el = sh._element
    chain = grp_chain(el)
    xf = el.find(P + "spPr").find(A + "xfrm")
    off, ext = xf.find(A + "off"), xf.find(A + "ext")
    X, Y, W, H0 = to_slide(chain, int(off.get("x")), int(off.get("y")), int(ext.get("cx")), int(ext.get("cy")))

    im = Image.open(os.path.join(lib, f"illust_{j['asset']}.png")).convert("RGBA")
    bb = ink(im)
    im = im.crop(bb)
    iw, ih = im.size
    Ht = j.get("H", H0 / E) * E
    cut = j.get("cut", 0.0)
    Wt = iw / ih * Ht / (1 - cut)
    if j.get("maxW") and Wt > j["maxW"] * E:          # 너무 넓으면 덜 자른다
        Wt = j["maxW"] * E
        cut = max(0.0, 1 - iw / ih * Ht / Wt)
        Wt = iw / ih * Ht / (1 - cut)
    ax, ay = j.get("ax", "c"), j.get("ay", "c")
    nx = {"l": X, "r": X + W - Wt, "c": X + (W - Wt) / 2}[ax] + j.get("dx", 0) * E
    ny = {"t": Y, "b": Y + H0 - Ht, "c": Y + (H0 - Ht) / 2}[ay] + j.get("dy", 0) * E
    cx, cy, cw, chh = to_child(chain, nx, ny, Wt, Ht)
    off.set("x", str(round(cx))); off.set("y", str(round(cy)))
    ext.set("cx", str(round(cw))); ext.set("cy", str(round(chh)))
    if j.get("flip"):
        xf.set("flipH", "1")
    elif "flipH" in xf.attrib:
        del xf.attrib["flipH"]

    buf = io.BytesIO(); im.save(buf, "PNG", optimize=True); buf.seek(0)
    _, rid = sh.part.get_or_add_image_part(buf)
    blip = el.find(".//" + A + "blip")
    blip.set(R + "embed", rid)
    for ch in list(blip):                             # 다시 칠하기 제거, extLst 의 원본 레이어도 제거
        if ch.tag != A + "extLst":
            blip.remove(ch)
    ex = blip.find(A + "extLst")
    if ex is not None:
        for e in list(ex):
            if any(c.tag.endswith("}imgProps") for c in e):
                ex.remove(e)
    fill = blip.getparent()
    sr = fill.find(A + "srcRect")
    if sr is None:
        sr = fill.makeelement(A + "srcRect", {})
        blip.addnext(sr)
    for k in list(sr.attrib):
        del sr.attrib[k]
    if cut > 0:
        sr.set("b", str(round(cut * 100000)))
    print(f"{j['slide']:>2} {j['name']:<9} {W/E:.2f}x{H0/E:.2f} → {Wt/E:.2f}x{Ht/E:.2f} cut{cut:.2f} flip{int(bool(j.get('flip')))}")
    done += 1
p.save(dst)
print("처리", done, "/", len(jobs))
