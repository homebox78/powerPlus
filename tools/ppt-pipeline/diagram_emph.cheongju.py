"""v0.59 - diagram nodes that the earlier deck (v0.36) drew as strong blue + white text but are now light
   get the emphasis back: blue face 2F78E0 + white text. args: old src dst"""
import sys, collections
from lxml import etree
from pptx import Presentation
sys.stdout.reconfigure(encoding="utf-8")
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"; P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"; E = 914400
STRONG = {"0456B6", "1973D1", "0B2E6B"}; BLUE = "2F78E0"
SLIDES = {20, 21, 25, 27, 28}
def flat(shapes, tf=(0, 0, 1, 1)):
    ox, oy, sx, sy = tf
    for s in shapes:
        if s.width is None: continue
        if s.shape_type == 6:
            x = s._element.find(P + "grpSpPr").find(A + "xfrm"); co, ce = x.find(A + "chOff"), x.find(A + "chExt")
            kx = s.width / int(ce.get("cx")) if int(ce.get("cx")) else 1; ky = s.height / int(ce.get("cy")) if int(ce.get("cy")) else 1
            gx, gy = ox + s.left * sx, oy + s.top * sy
            yield from flat(s.shapes, (gx - int(co.get("x")) * kx * sx, gy - int(co.get("y")) * ky * sy, sx * kx, sy * ky))
        else: yield s, ox + s.left * sx, oy + s.top * sy, s.width * sx, s.height * sy
def fc(s):
    sp = s._element.find(P + "spPr")
    sf = sp.find(A + "solidFill") if sp is not None else None
    return sf.find(A + "srgbClr") if sf is not None else None
def txt(s): return s.text_frame.text.strip() if s.has_text_frame else ""
def white(el):
    for tag in ("r", "endParaRPr"):
        for r in el.iter(A + tag):
            rp = r if tag == "endParaRPr" else r.find(A + "rPr")
            if rp is None:
                rp = etree.Element(A + "rPr"); r.insert(0, rp)
            for t in ("solidFill", "gradFill", "noFill", "pattFill"):
                for e in rp.findall(A + t): rp.remove(e)
            sf = etree.Element(A + "solidFill"); etree.SubElement(sf, A + "srgbClr", val="FFFFFF")
            ln = rp.find(A + "ln"); rp.insert((list(rp).index(ln) + 1) if ln is not None else 0, sf)
old = Presentation(sys.argv[1]); p = Presentation(sys.argv[2]); cnt = collections.Counter()
for n, (a, b) in enumerate(zip(old.slides, p.slides), 1):
    if n not in SLIDES: continue
    om = collections.defaultdict(set)
    for s, *_ in flat(a.shapes):
        c = fc(s); om[(s.name, txt(s))].add(c.get("val").upper() if c is not None else None)
    fs = list(flat(b.shapes)); done = []
    for s, x, y, w, h in fs:
        c = fc(s); o = om.get((s.name, txt(s)))
        if c is None or not o or not o <= STRONG or w > 2.0 * E or h > 0.65 * E or w < 0.5 * E: continue
        v = c.get("val").upper()
        if sum(int(v[i:i + 2], 16) for i in (0, 2, 4)) < 600: continue
        for e in list(c): c.remove(e)
        c.set("val", BLUE)
        ln = s._element.find(P + "spPr").find(A + "ln")
        if ln is not None and ln.find(A + "solidFill") is not None:
            lf = ln.find(A + "solidFill")
            for e in list(lf): lf.remove(e)
            etree.SubElement(lf, A + "srgbClr", val=BLUE)
        if s.has_text_frame: white(s._element)
        done.append((x, y, w, h)); cnt[n] += 1
    for s, x, y, w, h in fs:
        if not getattr(s, "has_text_frame", False) or not txt(s) or fc(s) is not None: continue
        cx, cy = x + w / 2, y + h / 2
        if any(a_ <= cx <= a_ + c_ and b_ <= cy <= b_ + d and h < d * 1.6 for a_, b_, c_, d in done): white(s._element); cnt["txt"] += 1
print(dict(cnt)); p.save(sys.argv[3])
