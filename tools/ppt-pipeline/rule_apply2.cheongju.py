"""디자인 시스템 7절 적용 2단계 — 남은 글자색, 도형 안 글자 크기(내림만), 21·22쪽 라벨 단계.
   인자: 현재판 기준판(v0.36) 결과"""
import sys, collections
from pptx import Presentation
sys.stdout.reconfigure(encoding="utf-8")
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
SRC, BASE, DST = sys.argv[1:4]
p = Presentation(SRC); b = Presentation(BASE)
TXT_TAGS = {A + "rPr", A + "defRPr", A + "endParaRPr"}
SZ = {"900": "800", "1100": "1000", "1300": "1200"}
MEMO = {"FFFF00", "FFFFCC", "FFFF99"}
cnt = collections.Counter(); reds = []
def walk(sh):
    for s in sh:
        if s.shape_type == 6: yield from walk(s.shapes)
        else: yield s
def fill(s):
    sp = s._element.find(P + "spPr")
    if sp is None: return None, None
    sf = sp.find(A + "solidFill")
    c = sf.find(A + "srgbClr") if sf is not None else None
    ln = sp.find(A + "ln")
    has_ln = ln is not None and ln.find(A + "solidFill") is not None
    return c, (sf is not None or has_ln)
for n, (sl, bl) in enumerate(zip(p.slides, b.slides), 1):
    for c in sl._element.iter(A + "srgbClr"):
        v = c.get("val").upper(); anc = [a.tag for a in c.iterancestors()]
        if A + "gradFill" in anc: continue
        if v == "1F4E79": c.set("val", "0B2E6B"); cnt[v] += 1
        elif v == "000000" and any(t in TXT_TAGS for t in anc) and A + "ln" not in anc:
            c.set("val", "2F3B6F"); cnt[v] += 1
    cur = list(walk(sl.shapes)); old = list(walk(bl.shapes))
    same = len(cur) == len(old)
    for i, s in enumerate(cur):
        c, filled = fill(s)
        v = c.get("val").upper() if c is not None else None
        if v == "D23737": reds.append((n, s.text_frame.text.strip()[:12] if s.has_text_frame else ""))
        if n in (21, 22) and same and v == "0B2E6B":
            oc, _ = fill(old[i])
            if oc is not None and oc.get("val").upper() == "3A46A0":
                c.set("val", "0456B6"); cnt["label"] += 1
        if filled and v not in MEMO and getattr(s, "has_text_frame", False):
            for rp in s._element.iter(A + "rPr"):
                if rp.get("sz") in SZ: rp.set("sz", SZ[rp.get("sz")]); cnt["sz"] += 1
print(dict(cnt)); print(reds)
p.save(DST)
