# 시안 톤 2차: 팔레트 축 교체(남색 082850 · 밝은 파랑 2070E8), 머리 띠 남색, 큰 흰 카드 → 옅은 하늘색 면·선 없음
import sys,collections
from pptx import Presentation
sys.stdout.reconfigure(encoding="utf-8")
A="{http://schemas.openxmlformats.org/drawingml/2006/main}"
P="{http://schemas.openxmlformats.org/presentationml/2006/main}"
E=914400
MAP={"0B2E6B":"082850","002060":"082850","1F4E79":"082850","3A46A0":"0F3D91","0456B6":"2070E8","1973D1":"2070E8","133193":"2070E8","DEEBF7":"E8F0F8","222A35":"082850"}
p=Presentation(sys.argv[1]); n=collections.Counter()
def remap(el):
    for c in el.iter(A+"srgbClr"):
        v=c.get("val").upper()
        if v in MAP: c.set("val",MAP[v]); n["색"]+=1
for m in p.slide_masters:
    remap(m._element)
    for l in m.slide_layouts: remap(l._element)
    for s in m.shapes:
        if s.name=="자유형 20":
            sp=s._element.find(P+"spPr")
            for t in ("solidFill","gradFill","noFill","blipFill","pattFill"):
                for f in sp.findall(A+t): sp.remove(f)
            sf=sp.makeelement(A+"solidFill",{}); sf.append(sf.makeelement(A+"srgbClr",{"val":"082850"}))
            sp.find(A+"custGeom").addnext(sf); n["머리띠"]+=1
def walk(sh):
    for s in sh:
        if s.shape_type==6: yield from walk(s.shapes)
        else: yield s
for sd in p.slides:
    remap(sd._element)
    for s in walk(sd.shapes):
        if s.shape_type==13 or not (s.width>1.5*E and s.height>0.8*E): continue
        sp=s._element.find(P+"spPr")
        if sp is None: continue
        f=sp.find(A+"solidFill"); ln=sp.find(A+"ln")
        if f is None or ln is None or ln.find(A+"solidFill") is None: continue
        c=f.find(A+"srgbClr"); sc=f.find(A+"schemeClr")
        white=(c is not None and c.get("val").upper()=="FFFFFF") or (sc is not None and sc.get("val")=="bg1" and len(sc)==0)
        if not white: continue
        for x in list(f): f.remove(x)
        f.append(f.makeelement(A+"srgbClr",{"val":"F0F5FB"}))
        for x in list(ln):
            if x.tag in (A+"solidFill",A+"gradFill"): ln.remove(x)
        ln.insert(0,ln.makeelement(A+"noFill",{})); n["카드"]+=1
p.save(sys.argv[2]); print(dict(n))
