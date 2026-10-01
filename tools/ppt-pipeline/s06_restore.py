# -*- coding: utf-8 -*-
"""6쪽 원본 문구 복원: 흐름 글자 +20%, 5대 목표(배지 삭제·아이콘 -15%·제목↑·원본 설명), 4대 추진배경(원본 키메시지·제목·설명). 인자: <src> <dst>"""
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
sys.stdout.reconfigure(encoding="utf-8")
src, dst = sys.argv[1:3]
p = Presentation(src); s = p.slides[5]
by = {sh.shape_id: sh for sh in s.shapes}
I = lambda v: Inches(v)
INK, GRAY, NAVY = RGBColor(0x1F, 0x4E, 0x79), RGBColor(0x5F, 0x74, 0x96), RGBColor(0x14, 0x3A, 0x69)
F2, F3, F4 = "a시월구일2", "a시월구일3", "a시월구일4"


def face(r, f):
    rPr = r._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea"):
        e = rPr.find(qn(tag))
        if e is None:
            e = rPr.makeelement(qn(tag), {}); rPr.append(e)
        e.set("typeface", f)


def set_text(sh, lines, size, color, f, align=PP_ALIGN.CENTER, spacing=None, bullet=False):
    tf = sh.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.04); tf.margin_top = tf.margin_bottom = 0
    for pg in list(tf.paragraphs)[1:]:
        pg._p.getparent().remove(pg._p)
    p0 = tf.paragraphs[0]
    for r in list(p0.runs):
        r._r.getparent().remove(r._r)
    for i, ln in enumerate(lines):
        pg = p0 if i == 0 else tf.add_paragraph()
        pg.alignment = align
        if spacing: pg.line_spacing = spacing
        if bullet:
            pg.space_after = Pt(0.5)
        if bullet:
            pPr = pg._p.get_or_add_pPr(); pPr.set("marL", "82296"); pPr.set("indent", "-82296")
            bu = pPr.makeelement(qn("a:buChar"), {"char": "•"}); pPr.append(bu)
        r = pg.add_run(); r.text = ln
        r.font.size = Pt(size); r.font.color.rgb = color; r.font.bold = False; face(r, f)
    return sh


def textbox(x, y, w, h, lines, size, color, f, **kw):
    tb = s.shapes.add_textbox(I(x), I(y), I(w), I(h)); tb.text_frame.vertical_anchor = MSO_ANCHOR.TOP
    return set_text(tb, lines, size, color, f, **kw)


def kill(*ids):
    for i in ids:
        if i in by:
            by[i]._element.getparent().remove(by[i]._element)


# 1) 흐름 글자 +20%
for i in (108, 109, 110):
    for pg in by[i].text_frame.paragraphs:
        for r in pg.runs:
            r.font.size = Pt(round(r.font.size.pt * 1.2 * 2) / 2)
by[110].left = I(6.17); by[110].width = I(1.76); by[110].text_frame.word_wrap = False
by[108].width = I(1.5); by[109].width = I(1.3)

# 2) 5대 목표
kpi = [(113, 114, 116, ["시스템 안정성 확보"], ["사전 예방정비", "및 장애 시 신속한 조치"]),
       (117, 118, 120, ["업무 연속성 보장"], ["업무지원포털 시스템", "운영·구축 유경험자 투입"]),
       (121, 122, 124, ["보안을 강화하여", "철통보안 보장"], ["보안사고 예방", "지침 및 절차 준수"]),
       (125, 126, 128, ["체계적이고", "지속적인 지원"], ["지속적인 서비스", "확대 및 안정화"]),
       (129, 130, 132, ["시스템 기능개선"], ["이용자 중심의", "사용편의성 증대"])]
kill(115, 119, 123, 127, 131)
for card, icon, title, tl, desc in kpi:
    c = by[card]; ic = by[icon]
    cx = ic.left + ic.width // 2
    ic.width = int(ic.width * 0.85); ic.height = int(ic.height * 0.85)
    ic.left = cx - ic.width // 2; ic.top = I(3.73)
    t = by[title]; t.left = c.left; t.width = c.width; t.top = I(4.27); t.height = I(0.42)
    set_text(t, tl, 11, INK, F3, spacing=0.95); t.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    d = textbox(c.left / 914400, 4.72, c.width / 914400, 0.40, desc, 8.5, GRAY, F2, spacing=1.0)
    t._element.addnext(d._element)

# 3) 4대 추진배경: 원본 키메시지·제목·설명
lab = by[101]; lab.width = I(7.8)
set_text(lab, ["청주시 내부 업무 대표 관문으로 지속적이고 안정적인 운영이 필요"], 11, NAVY, F4, align=PP_ALIGN.LEFT)
lab.text_frame.margin_left = 0; lab.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
bg = [(133, 134, 135, 136, 137, ["안정적이고 신속한", "유지관리 체계 확립"],
       ["청주시 대표 관문인 업무지원 포털 시스템의 안정적인 서비스 제공 체계 마련",
        "신속하고 정확하게 사용자의 요구사항을 반영하기 위한 유지관리 체계 확립 필요"]),
      (138, 139, 140, 141, 142, ["신속한", "장애복구 체계 수립"],
       ["정기적인 예방·보수 점검으로 신속한 장애 식별 체계 마련",
        "식별된 장애를 신속하게 복구 할 수 있는 장애 복구 체계 수립 필요",
        "장애 이력을 분석하여 지속적인 근본 원인 해결 방안 마련"]),
      (143, 144, 145, 146, 147, ["무중단", "서비스 환경 구축"],
       ["체계적인 장애 관리를 통하여 무중단 서비스를 제공",
        "청주시 상주 운영 담당자를 통하여 AP의 상태 점검으로 무중단 서비스 환경 마련",
        "클라우드 전환에 따른 가상화 운영역량 확보 및 어플리케이션 모니터링"]),
      (148, 149, 150, 151, 152, ["업무추진 환경변화에", "신속·능동적 대처"],
       ["업무 관련 법제도·규정변경에 대하여 모니터링하여 신속하게 대처",
        "사용자의 업무 환경(OS, 웹 브라우저 등)의 변경에 능동적으로 대응하기 위한 지원 체계 마련 필요"])]
for card, icon, num, title, desc, tl, bl in bg:
    c = by[card]; L, W = c.left / 914400, c.width / 914400
    kill(icon)
    n = by[num]; n.left = I(L + 0.09); n.top = I(5.79)
    t = by[title]; t.left = I(L + 0.45); t.width = I(W - 0.5); t.top = I(5.73); t.height = I(0.38)
    set_text(t, tl, 10, INK, F3, align=PP_ALIGN.LEFT, spacing=0.95); t.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    c.height = int(p.slide_height / 2 + 8.3 * 360000) - int(0.04 * 914400) - c.top
    d = by[desc]; d.left = I(L + 0.08); d.width = I(W - 0.14); d.top = I(6.13); d.height = I(0.84)
    set_text(d, bl, 7, GRAY, F2, align=PP_ALIGN.LEFT, spacing=0.95, bullet=True)
    d.text_frame.vertical_anchor = MSO_ANCHOR.TOP
p.save(dst); print("ok")
