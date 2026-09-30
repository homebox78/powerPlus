import sys
from pptx import Presentation
from pptx.util import Emu,Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
E=914400
D=r"D:\powerPlus\제안서\청주시_발표자료(제안요약서)_v0.33_s7유지보수대상위로.pptx"
p=Presentation("l32.pptx"); sl=p.slides[6]
DY=2.09; junk=[]
for s in sl.shapes:
    y=s.top/E
    if s.name=="그룹 33": s.top=Emu(int(1.10*E)); s.height=Emu(int(1.46*E))
    elif s.name=="그룹 34": s.top=Emu(int(2.60*E))
    elif s.name in("직사각형 302","직사각형 319","직사각형 299"): junk.append(s)
    elif 1.0<=y<4.95: s.top=Emu(int((y+DY)*E))
c=sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Emu(int(0.30*E)),Emu(int(1.05*E)),Emu(int(10.23*E)),Emu(int(2.04*E)))
c.name="CARD_유지보수대상"; c.adjustments[0]=0.06
c.fill.solid(); c.fill.fore_color.rgb=RGBColor(0xFF,0xFF,0xFF)
c.line.color.rgb=RGBColor(0xC0,0xD2,0xE6); c.line.width=Pt(0.75); c.shadow.inherit=False
t=sl.shapes._spTree; t.remove(c._element); t.insert(2,c._element)
for j in junk: j._element.getparent().remove(j._element)
p.save(D)
