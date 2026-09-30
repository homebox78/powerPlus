import sys
from collections import Counter
from pptx import Presentation
from pptx.util import Emu,Pt
from pptx.enum.text import PP_ALIGN
sys.stdout.reconfigure(encoding="utf-8")
A="{http://schemas.openxmlformats.org/drawingml/2006/main}"
P="{http://schemas.openxmlformats.org/presentationml/2006/main}"
E=914400
D=r"D:\powerPlus\제안서\청주시_발표자료(제안요약서)_v0.32_전체검수.pptx"
p=Presentation("l31.pptx"); n=Counter(); ch=set()
SK=(1,2,3,9,14,23,30,36,40)
def walk(sh):
    for s in sh:
        yield s
        if s.shape_type==6: yield from walk(s.shapes)
def setfont(r,name):
    rp=r._r.get_or_add_rPr()
    for t in("latin","ea"):
        e=rp.find(A+t)
        if e is None:
            e=rp.makeelement(A+t,{}); 
            # 순서: latin, ea, cs, sym 은 fill·effect 뒤
            anchor=None
            for tag in("cs","sym","hlinkClick","hlinkMouseOver","rtl","extLst"):
                anchor=rp.find(A+tag)
                if anchor is not None: break
            if t=="latin" and rp.find(A+"ea") is not None: rp.find(A+"ea").addprevious(e)
            elif anchor is not None: anchor.addprevious(e)
            else: rp.append(e)
        e.set("typeface",name)
def has(r):
    rp=r._r.find(A+"rPr"); return rp is not None and rp.find(A+"ea") is not None
def white(r):
    rp=r._r.find(A+"rPr")
    if rp is None: return False
    x=rp.find(A+"solidFill")
    return x is not None and len(x) and x[0].get("val") in("FFFFFF","bg1")
for i,sl in enumerate(p.slides,1):
    for s in walk(sl.shapes):
        if s.left/E>=10.8: continue
        if s.shape_type==19:
            for rix,row in enumerate(s.table.rows):
                for c in row.cells:
                    for pg in c.text_frame.paragraphs:
                        for r in pg.runs:
                            if not r.text.strip(): continue
                            if not has(r): setfont(r,"a시월구일3" if rix==0 else "a시월구일2"); n["표 상속"]+=1; ch.add(i)
                            elif rix==0 and r._r.find(A+"rPr").find(A+"ea").get("typeface")=="a시월구일2": setfont(r,"a시월구일3"); n["표 머리"]+=1; ch.add(i)
            continue
        if not s.has_text_frame or not s.text_frame.text.strip(): continue
        x,y,w=s.left/E,s.top/E,s.width/E
        km=i not in SK and abs(y-1.03)<0.08 and w>8
        if km:
            if abs(x-0.20)>0.011: s.left=Emu(int(0.20*E)); n["키 위치"]+=1; ch.add(i)
        for pg in s.text_frame.paragraphs:
            if km and (pg.alignment is None or int(pg.alignment)!=2): pg.alignment=PP_ALIGN.CENTER; n["키 정렬"]+=1; ch.add(i)
            for r in pg.runs:
                if not r.text.strip(): continue
                if km:
                    if not has(r): setfont(r,"G마켓 산스 TTF Bold"); n["키 서체"]+=1; ch.add(i)
                    if r.font.size is None: r.font.size=Pt(24)
                elif not has(r):
                    sz=r.font.size.pt if r.font.size else 0
                    setfont(r,"a시월구일3" if (white(r) and sz>=12) else "a시월구일2"); n["상속"]+=1; ch.add(i)
for s in walk(p.slides[9].shapes):
    if s.name=="AutoShape 64":
        ef=s._element.find(P+"spPr").find(A+"effectLst")
        if ef is not None: s._element.find(P+"spPr").remove(ef); n["효과"]+=1; ch.add(10)
p.save(D); print(dict(n)); print(sorted(ch))
