# -*- coding: utf-8 -*-
"""이미 넣은 자산 A 를 자산 B 로 전 슬라이드 일괄 교체(그림 바이트 해시로 식별, 위치·크기 유지)."""
import hashlib, io, sys
from PIL import Image
from pptx import Presentation
from pptx.util import Emu
sys.stdout.reconfigure(encoding="utf-8")
src, dst, a_path, b_path = sys.argv[1:5]
def prep(p):
    im = Image.open(p).convert("RGBA")
    bb = im.getchannel("A").point(lambda v: 255 if v > 8 else 0).getbbox()
    if bb: im = im.crop(bb)
    buf = io.BytesIO(); im.save(buf, "PNG", optimize=True); return im, buf.getvalue()
_, ab = prep(a_path); HA = hashlib.sha1(ab).hexdigest()
bim, bb = prep(b_path)
pres = Presentation(src); n = 0
for sno, sl in enumerate(pres.slides, 1):
    for sh in list(sl.shapes):
        if sh.shape_type != 13 or hashlib.sha1(sh.image.blob).hexdigest() != HA: continue
        L, T, W, H = Emu(sh.left), Emu(sh.top), Emu(sh.width), Emu(sh.height)
        cx, cy = L + W / 2, T + H / 2
        s = min(W / bim.width, H / bim.height); nw, nh = int(bim.width * s), int(bim.height * s)
        sh._element.getparent().remove(sh._element)
        sl.shapes.add_picture(io.BytesIO(bb), int(cx - nw / 2), int(cy - nh / 2), nw, nh); n += 1
        print(f"  s{sno} 교체")
pres.save(dst); print(f"{n}개 → {dst}")
