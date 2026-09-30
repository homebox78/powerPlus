import sys, colorsys, collections, json
from pptx import Presentation
sys.stdout.reconfigure(encoding="utf-8")
A="{http://schemas.openxmlformats.org/drawingml/2006/main}"
P="{http://schemas.openxmlformats.org/presentationml/2006/main}"
E=914400
SRC=sys.argv[1]
p=Presentation(SRC)
def walk(sh):
    for s in sh:
        if s.shape_type==6: yield from walk(s.shapes)
        else: yield s
def hsv(h):
    r,g,b=[int(h[i:i+2],16)/255 for i in (0,2,4)]
    return colorsys.rgb_to_hsv(r,g,b)
def kind(h):
    H,S,V=hsv(h)
    if S<0.10: return "gray"
    d=H*360
    if 195<=d<=255: return "blue"
    return "other"
fillA=collections.Counter(); txtA=collections.Counter(); lineA=collections.Counter()
per=[]; sizeIn=collections.Counter(); ins=collections.Counter(); sizeBy=collections.defaultdict(collections.Counter)
for n,sl in enumerate(p.slides,1):
    fc=set(); tc=set(); blues=set(); others=set()
    for s in walk(sl.shapes):
        sp=s._element.find(P+"spPr")
        if sp is not None:
            sf=sp.find(A+"solidFill")
            if sf is not None and sf.find(A+"srgbClr") is not None:
                h=sf.find(A+"srgbClr").get("val").upper(); fc.add(h); fillA[h]+=1
            ln=sp.find(A+"ln")
            if ln is not None and ln.find(A+"solidFill") is not None and ln.find(A+"solidFill").find(A+"srgbClr") is not None:
                h=ln.find(A+"solidFill").find(A+"srgbClr").get("val").upper(); lineA[h]+=1; fc.add(h)
        if getattr(s,"has_text_frame",False) and s.text_frame.text.strip():
            filled = sp is not None and (sp.find(A+"solidFill") is not None or (sp.find(A+"ln") is not None and sp.find(A+"ln").find(A+"solidFill") is not None))
            bp=s._element.find(P+"txBody").find(A+"bodyPr")
            if filled:
                l=int(bp.get("lIns","91440"))/E; t=int(bp.get("tIns","45720"))/E
                ins[(round(l,2),round(t,2))]+=1
            for pg in s.text_frame.paragraphs:
                for r in pg.runs:
                    if not r.text.strip(): continue
                    rp=r._r.find(A+"rPr")
                    if rp is not None:
                        sf=rp.find(A+"solidFill")
                        if sf is not None and sf.find(A+"srgbClr") is not None:
                            h=sf.find(A+"srgbClr").get("val").upper(); tc.add(h); txtA[h]+=len(r.text)
                        sz=rp.get("sz")
                        if sz and filled:
                            sizeIn[int(sz)/100]+=len(r.text)
                            hgt=s.height/E
                            b="S(<0.4in)" if hgt<0.4 else ("M(<1in)" if hgt<1 else "L")
                            sizeBy[b][int(sz)/100]+=len(r.text)
    allc=fc|tc
    for h in allc:
        k=kind(h)
        if k=="blue": blues.add(h)
        elif k=="other": others.add(h)
    per.append((n,len(allc),len(blues),sorted(others)))
print("== 장당 색 수(채움+선+글자) / 블루 수 / 비블루 유채색")
for r in per: print(r)
import statistics
print("중앙값 색",statistics.median(x[1] for x in per),"블루",statistics.median(x[2] for x in per))
print("== 채움 상위"); print(fillA.most_common(25))
print("== 글자색 상위"); print(txtA.most_common(20))
print("== 선색 상위"); print(lineA.most_common(12))
print("== 도형 안 글자 크기(글자수)"); print(sorted(sizeIn.items(),key=lambda x:-x[1])[:20])
for b,c in sizeBy.items(): print(b,sorted(c.items(),key=lambda x:-x[1])[:8])
print("== 도형 내 여백 (좌,상) in"); print(ins.most_common(12))
