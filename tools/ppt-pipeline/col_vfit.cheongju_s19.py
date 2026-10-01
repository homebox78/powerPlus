"""19쪽 왼쪽 묶음을 오른쪽 묶음과 같은 위·아래로 세로 맞춤 + 11~14쪽 ACT 옆 제목 시작 통일. 인자: <pptx>(제자리)"""
import sys
from pptx import Presentation
from pptx.util import Emu
sys.stdout.reconfigure(encoding="utf-8")
IN = 914400
f = sys.argv[1]
p = Presentation(f)
s = p.slides[18]
sh = [x for x in s.shapes if x.top is not None and x.name not in ("키메시지",) and x.top > 1.5 * IN]
left = [x for x in sh if x.left < 6.65 * IN]
right = [x for x in sh if x.left >= 6.65 * IN]
l0 = min(x.top for x in left) / IN
l1 = max(x.top + x.height for x in left) / IN
r0 = min(x.top for x in right) / IN
r1 = max(x.top + x.height for x in right) / IN
T, B = r0, max(r1, 7.02)
k = (B - T) / (l1 - l0)
print("왼 %.2f~%.2f 오른 %.2f~%.2f → %.2f~%.2f 비율 %.3f" % (l0, l1, r0, r1, T, B, k))
for x in left:
    top, h = x.top / IN, x.height / IN
    if x.shape_type == 13:
        c = T + (top + h / 2 - l0) * k
        x.top = Emu(int((c - h / 2) * IN))
    else:
        x.top = Emu(int((T + (top - l0) * k) * IN))
        x.height = Emu(int(h * k * IN))

# 11~14쪽: ACT 배지(0.30~1.20) 옆 제목 시작 1.36in 로 통일(폭은 오른쪽 끝 유지)
X = 1.36
TITLE = {11: ["시안 14", "시안 57"], 12: ["시안 7"], 13: ["시안 9"], 14: ["시안 7"]}
for n, names in TITLE.items():
    for x in p.slides[n - 1].shapes:
        if x.name in names:
            right_edge = x.left + x.width
            x.left = Emu(int(X * IN))
            x.width = Emu(right_edge - x.left)
            print(n, x.name, "제목 시작", X)
p.save(f)
