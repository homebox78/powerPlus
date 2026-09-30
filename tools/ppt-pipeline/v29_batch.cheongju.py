import sys,io
from pptx import Presentation
from pptx.util import Emu,Pt
from PIL import Image
sys.stdout.reconfigure(encoding="utf-8")
A="{http://schemas.openxmlformats.org/drawingml/2006/main}"
E=914400
LIB=r"C:\Users\hbox7\AppData\Local\Temp\claude\d--powerPlus\818b2ce8-2260-4251-abfd-8c6a90609354\scratchpad\v16\lib"
SRC=r"D:\powerPlus\제안서\청주시_발표자료(제안요약서)_v0.28_s11아이콘.pptx"
DST=r"D:\powerPlus\제안서\청주시_발표자료(제안요약서)_v0.29_s32·s33·s34·s39·맺음.pptx"
p=Presentation(SRC)
def walk(sh):
    for s in sh:
        yield s
        if s.shape_type==6: yield from walk(s.shapes)
def ink(path):
    im=Image.open(path).convert("RGBA"); im=im.crop(im.getchannel("A").point(lambda a:255 if a>8 else 0).getbbox())
    b=io.BytesIO(); im.save(b,"PNG"); b.seek(0); return b,im.size
def put(sl,path,cx,cy,w=None,h=None,name="ICON"):
    b,(iw,ih)=ink(path)
    if w is None: w=h*iw/ih
    h=w*ih/iw
    pic=sl.shapes.add_picture(b,Emu(int((cx-w/2)*E)),Emu(int((cy-h/2)*E)),Emu(int(w*E)),Emu(int(h*E))); pic.name=name; return pic
def rm(s): s._element.getparent().remove(s._element)
# s39 일러스트 → 아이콘
sl=p.slides[38]
for s in list(sl.shapes):
    if s.name=="ADD_1046": rm(s); put(sl,"ic/icon_1434.png",9.72,3.17,h=1.35,name="IC_1434")
    if s.name=="ADD_792": rm(s); put(sl,"ic/icon_1262.png",1.02,6.45,w=1.25,name="IC_1262")
# s40 (+전 장) 글자 그림자 제거
n=0
for sl in p.slides:
    for s in walk(sl.shapes):
        if s.has_text_frame:
            for r in s._element.iter(A+"rPr",A+"defRPr",A+"endParaRPr"):
                ef=r.find(A+"effectLst")
                if ef is not None and len(ef): r.remove(ef); n+=1
print("그림자 제거",n)
# s32 로제트 → 월계관, 108 아이콘 하단 맞춤
sl=p.slides[31]
for s in list(sl.shapes):
    if s.name=="그룹 280":
        tb=[c for c in s.shapes if c.has_text_frame][0]
        import copy; el=copy.deepcopy(tb._element); rm(s)
        W=2.75; cx=8.83; top=1.62
        pic=put(sl,"ic/icon_1907.png",cx,top+W*354/404/2,w=W,name="IC_1907")
        sl.shapes._spTree.append(el)
        t=sl.shapes[-1]; t.width=Emu(int(1.1*E)); t.height=Emu(int(0.6*E))
        ccx=cx; ccy=top+W*354/404*0.455
        t.left=Emu(int((ccx-0.55)*E)); t.top=Emu(int((ccy-0.30)*E))
        for pg in t.text_frame.paragraphs:
            for r in pg.runs:
                if r.font.size and r.font.size.pt==10: r.font.size=Pt(11)
                elif r.font.size: r.font.size=Pt(20)
    if s.name=="그림 123": s.top=Emu(int((3.95+0.40)*E))-s.height
# s33 품질 일러스트: 원본으로, 틀 가득(10% 더 키워 자르기)
sl=p.slides[32]
for g in list(sl.shapes):
    if g.shape_type!=6: continue
    for c in list(g.shapes):
        if c.name in("SW19","SW20"):
            aid={"SW19":1001,"SW20":1114}[c.name]
            fx=g.left/E+0.03; fy=g.top/E+0.03; fw=g.width/E-0.06; fh=5.90-fy
            rm(c)
            b,(iw,ih)=ink(f"{LIB}\illust_{aid}.png")
            sc=max(fw/iw,fh/ih)*1.10
            vw=fw/(iw*sc); vh=fh/(ih*sc)
            pic=sl.shapes.add_picture(b,Emu(int(fx*E)),Emu(int(fy*E)),Emu(int(fw*E)),Emu(int(fh*E)))
            pic.crop_left=pic.crop_right=(1-vw)/2; pic.crop_top=0.02 if vh<0.98 else 0; pic.crop_bottom=1-vh-pic.crop_top; pic.name=c.name
            print(c.name,iw,ih,round(vw,2),round(vh,2))
# s33 알약 글자 9pt
k=0
for s in walk(sl.shapes):
    if s.has_text_frame and any(x in s.text_frame.text for x in("유지관리방법론 적용","품질보증팀 구성","분석/설계 도구")):
        for pg in s.text_frame.paragraphs:
            for r in pg.runs: r.font.size=Pt(9); k+=1
print("9pt",k)
# s34 표 '내용' 열 7pt
k=0
for s in walk(p.slides[33].shapes):
    if s.shape_type==19:
        t=s.table; hdr=[c.text.strip() for c in t.rows[0].cells]
        if "내용" in hdr:
            j=hdr.index("내용")
            for row in list(t.rows)[1:]:
                for pg in row.cells[j].text_frame.paragraphs:
                    for r in pg.runs: r.font.size=Pt(7); k+=1
                    e=pg._p.find(A+"endParaRPr")
                    if e is not None: e.set("sz","700")
print("7pt",k)
p.save(DST)
