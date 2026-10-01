# -*- coding: utf-8 -*-
"""'키메시지' 이름 도형의 세로 위치를 한 값으로 통일(같은 장의 '키메시지 2' 등도 같은 만큼). 인자: <src> <dst> [cm=2.62] [--skip 4,…]"""
import sys
from pptx import Presentation
sys.stdout.reconfigure(encoding="utf-8")
args = [a for a in sys.argv[1:] if not a.startswith("--")]
src, dst = args[:2]
CM = float(args[2]) if len(args) > 2 else 2.62
SKIP = set()
if "--skip" in sys.argv:
    SKIP = {int(x) for x in sys.argv[sys.argv.index("--skip") + 1].split(",")}
TOP = int(CM / 2.54 * 914400)
prs = Presentation(src)
for no, sl in enumerate(prs.slides, 1):
    if no in SKIP:
        continue
    ks = [s for s in sl.shapes if s.name.startswith("키메시지")]
    if not ks:
        continue
    d = TOP - min(s.top for s in ks)
    if abs(d) < 3000:
        continue
    for s in ks:
        s.top = s.top + d
    print(no, "%+.2f cm" % (d / 914400 * 2.54))
prs.save(dst)
