"""v0.72 구조 수정: 그림 교체·지도 핀·머리띠 모양/색·33쪽 분리·32쪽 아이콘. 인자: <src> <dst>"""
import copy, sys
from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
sys.stdout.reconfigure(encoding="utf-8")
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
I = 914400
OR = "C:/Users/hbox7/AppData/Local/Temp/claude/d--powerPlus/1b336add-4aaf-4831-b1b4-ba8dc70560ab/scratchpad/orig/"
NAVY = "1A5294"
BLUE = "2F78E6"
src, dst = sys.argv[1:3]
p = Presentation(src)


def S(n):
    return p.slides[n - 1]


def find(n, name, shapes=None):
    for s in (shapes if shapes is not None else S(n).shapes):
        if s.name == name:
            return s
    raise KeyError((n, name))


def put_pic(slide, path, x, y, w, h, before=None, align="c", valign="c"):
    """상자(x,y,w,h, in) 안에 비율 유지로 맞춰 넣기. before 요소 바로 앞(같은 z)으로."""
    im = Image.open(path)
    r = im.width / im.height
    if w / h > r:
        ph, pw = h, h * r
    else:
        pw, ph = w, w / r
    px = x + (w - pw) / 2 if align == "c" else x
    py = {"c": y + (h - ph) / 2, "b": y + h - ph, "t": y}[valign]
    pic = slide.shapes.add_picture(path, Emu(int(px * I)), Emu(int(py * I)), Emu(int(pw * I)), Emu(int(ph * I)))
    if before is not None:
        before.addprevious(pic._element)
    return pic


def replace_pic(n, name, path, box=None, shapes=None, valign="c"):
    s = find(n, name, shapes)
    x, y, w, h = (s.left / I, s.top / I, s.width / I, s.height / I) if box is None else box
    el = s._element
    put_pic(S(n), path, x, y, w, h, before=el, valign=valign)
    el.getparent().remove(el)
    print(n, name, "→", path.split("/")[-1])


def set_fill(s, val):
    sp = s._element.spPr
    for t in ("solidFill", "gradFill", "noFill"):
        for e in sp.findall(A + t):
            sp.remove(e)
    f = etree.Element(A + "solidFill")
    etree.SubElement(f, A + "srgbClr", val=val)
    g = sp.find(A + "prstGeom")
    if g is None:
        g = sp.find(A + "custGeom")
    g.addnext(f)


def to_band(s, x, w, color):
    """머리띠: 위 두 모서리만 둥근 사각형(오른쪽 띠와 같은 모양)."""
    sp = s._element.spPr
    for t in ("prstGeom", "custGeom"):
        for e in sp.findall(A + t):
            sp.remove(e)
    g = etree.Element(A + "prstGeom", prst="round2SameRect")
    av = etree.SubElement(g, A + "avLst")
    etree.SubElement(av, A + "gd", name="adj1", fmla="val 30000")
    etree.SubElement(av, A + "gd", name="adj2", fmla="val 0")
    sp.find(A + "xfrm").addnext(g)
    s.left, s.width = Emu(int(x * I)), Emu(int(w * I))
    set_fill(s, color)


def drop(n, name):
    s = find(n, name)
    s._element.getparent().remove(s._element)


# ---------- 1. 18쪽 여성 → 3D 인물(노트북) ----------
replace_pic(18, "시안 그림 4", OR + "illust_332.png", valign="b")

# ---------- 2. 22쪽 지도 속 인물 → 지도 핀 ----------
sl = S(22)


def grp_xf(g):
    x = g._element.find(P + "grpSpPr").find(A + "xfrm")
    co, ce = x.find(A + "chOff"), x.find(A + "chExt")
    kx = g.width / int(ce.get("cx"))
    ky = g.height / int(ce.get("cy"))
    return g.left - int(co.get("x")) * kx, g.top - int(co.get("y")) * ky, kx, ky


PIN_H = 0.30
pins = []
for g in list(sl.shapes):
    if g.shape_type != 6 or g.name not in ("그룹 2", "그룹 159", "그룹 5"):
        continue
    if g.name == "그룹 5":
        ox, oy, kx, ky = grp_xf(g)
        for c in list(g.shapes):
            if c.shape_type == 6:  # 사람 그림 묶음
                cx = (ox + (c.left + c.width / 2) * kx) / I
                cy = (oy + (c.top + c.height / 2) * ky) / I
                c._element.getparent().remove(c._element)
                pins.append((cx, cy, g._element))
    else:
        cx = (g.left + g.width / 2) / I
        cy = (g.top + g.height / 2) / I
        pins.append((cx, cy, g._element))
        g._element.getparent().remove(g._element) if False else None
pins = [(cx, cy + 0.22 if abs(cy - 4.36) < 0.05 else cy, el) for cx, cy, el in pins]
for cx, cy, el in pins:
    # 핀 끝(아래)이 그 자리를 가리키게
    pw = PIN_H * 326 / 442
    pic = sl.shapes.add_picture(OR + "icon_1068.png", Emu(int((cx - pw / 2) * I)), Emu(int((cy - PIN_H * 0.92) * I)),
                                Emu(int(pw * I)), Emu(int(PIN_H * I)))
    print("22 핀", round(cx, 2), round(cy, 2))
for g in list(sl.shapes):
    if g.shape_type == 6 and g.name in ("그룹 2", "그룹 159"):
        g._element.getparent().remove(g._element)

# 22쪽 분홍 원 100% 글자 흰색
for r in find(22, "시안 30").text_frame.paragraphs[0].runs:
    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
for pg in find(22, "시안 30").text_frame.paragraphs:
    for r in pg.runs:
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

# ---------- 3. 얼굴 없는 실루엣 → 얼굴 있는 3D 인물 ----------
def bust(n, name, fid, scale=1.15):
    s = find(n, name)
    cx, cy = (s.left + s.width / 2) / I, (s.top + s.height / 2) / I
    h = max(s.width, s.height) / I * scale
    el = s._element
    put_pic(S(n), OR + fid + ".png", cx - h / 2, cy - h / 2, h, h, before=el)
    el.getparent().remove(el)
    print(n, name, "→", fid)


bust(25, "시안 그림 17", "illust_743")
bust(28, "시안 그림 63", "illust_741")
bust(32, "시안 그림 38", "illust_740")
bust(23, "시안 그림 10", "illust_740", scale=1.0)

# ---------- 4. 32쪽 작은 아이콘 2배(가운데 x·아래 유지) / 월계관 원본 ----------
for nm in ("시안 그림 17", "시안 그림 23"):
    s = find(32, nm)
    cx, b = s.left + s.width / 2, s.top + s.height
    s.width, s.height = Emu(int(s.width * 2)), Emu(int(s.height * 2))
    s.left, s.top = Emu(int(cx - s.width / 2)), Emu(int(b - s.height))
    print("32", nm, "2배")
lau = find(32, "시안 그림 34")
ell = find(32, "시안 35")
ecx, ecy, ed = (ell.left + ell.width / 2) / I, (ell.top + ell.height / 2) / I, ell.width / I
# icon_1907: 가운데 짙은 원 지름 비율 측정
im = Image.open(OR + "icon_1907.png").convert("RGBA")
px = im.load()
dark = [(x, y) for x in range(im.width) for y in range(im.height)
        if px[x, y][3] > 200 and sum(px[x, y][:3]) < 200]
xs = [d[0] for d in dark]; ys = [d[1] for d in dark]
dx0, dx1, dy0, dy1 = min(xs), max(xs), min(ys), max(ys)
k = ed * 1.06 / ((dx1 - dx0) / im.width)  # 그림 폭(in): 짙은 원이 타원보다 살짝 크게
W = k; H = k * im.height / im.width
ccx = (dx0 + dx1) / 2 / im.width * W
ccy = (dy0 + dy1) / 2 / im.height * H
put_pic(S(32), OR + "icon_1907.png", ecx - ccx, ecy - ccy, W, H, before=lau._element)
lau._element.getparent().remove(lau._element)
print("32 월계관 → icon_1907", round(W, 2), round(H, 2))

# ---------- 5. 27쪽 아이콘 두 줄: 같은 크기·원본 ----------
ROW1 = [("시안 그림 30", "icon_1513"), ("시안 그림 32", "icon_1236"), ("시안 그림 34", "icon_1252"),
        ("시안 그림 36", "icon_1588"), ("시안 그림 38", "icon_1096"), ("시안 그림 40", "icon_1558")]
ROW2 = [("시안 그림 75", "icon_1450"), ("시안 그림 79", "icon_1345"), ("시안 그림 83", "icon_1454"),
        ("시안 그림 87", "icon_1187"), ("시안 그림 91", "icon_1461")]
for row, box in ((ROW1, 0.50), (ROW2, 0.36)):
    for nm, fid in row:
        s = find(27, nm)
        cx, cy = (s.left + s.width / 2) / I, (s.top + s.height / 2) / I
        el = s._element
        put_pic(S(27), OR + fid + ".png", cx - box / 2, cy - box / 2, box, box, before=el)
        el.getparent().remove(el)
    print("27 아이콘 줄", len(row), "상자", box)

# ---------- 6. 33쪽 통이미지 → 카드·인물·원·글자 분리 ----------
sl = S(33)
old = find(33, "시안 그림 55")
X, Y, W3, H3 = old.left / I, old.top / I, old.width / I, old.height / I
anchor = old._element
CW = 1.32  # 카드 폭
CH = H3 - 0.12
cy0 = Y + 0.10


def card(x, people, title, sub):
    c = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(int(x * I)), Emu(int(cy0 * I)), Emu(int(CW * I)), Emu(int(CH * I)))
    c.adjustments[0] = 0.08
    c.fill.solid(); c.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    c.line.color.rgb = RGBColor(0x3B, 0x82, 0xF0); c.line.width = Pt(1.25)
    c.shadow.inherit = False
    anchor.addprevious(c._element)
    # 아래 이름 띠
    bh = 0.48
    b = sl.shapes.add_shape(MSO_SHAPE.ROUND_2_SAME_RECTANGLE, Emu(int(x * I)), Emu(int((cy0 + CH - bh) * I)), Emu(int(CW * I)), Emu(int(bh * I)))
    b.rotation = 180
    b.adjustments[0] = 0.16; b.adjustments[1] = 0
    b.fill.solid(); b.fill.fore_color.rgb = RGBColor(0x14, 0x3A, 0x69)
    b.line.fill.background(); b.shadow.inherit = False
    anchor.addprevious(b._element)
    t = sl.shapes.add_textbox(Emu(int(x * I)), Emu(int((cy0 + CH - bh) * I)), Emu(int(CW * I)), Emu(int(bh * I)))
    tf = t.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, 0)
    for i, (txt, sz) in enumerate(((title, 11), (sub, 8))):
        pg = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        pg.alignment = PP_ALIGN.CENTER
        r = pg.add_run(); r.text = txt
        r.font.size = Pt(sz); r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        fn = "a시월구일4" if i == 0 else "a시월구일2"
        r.font.name = fn
        rp = r._r.get_or_add_rPr(); rp.append(rp.makeelement(A + "ea", {"typeface": fn}))
    anchor.addprevious(t._element)
    # 인물
    ah = CH - bh - 0.06
    n = len(people)
    pw = CW / n
    for i, fid in enumerate(people):
        put_pic(sl, OR + fid + ".png", x + i * pw + 0.04, cy0 + 0.06, pw - 0.08, ah, before=t._element, valign="b")


card(X, ["illust_745"], "품질통제", "(사업수행팀)")
card(X + W3 - CW, ["illust_740", "illust_749"], "품질보증", "(전사품질팀)")
# 가운데 원 + 순환 아이콘
D = 0.82
ccx, ccy = X + W3 / 2, cy0 + CH / 2
c = sl.shapes.add_shape(MSO_SHAPE.OVAL, Emu(int((ccx - D / 2) * I)), Emu(int((ccy - D / 2) * I)), Emu(int(D * I)), Emu(int(D * I)))
c.fill.solid(); c.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
c.line.color.rgb = RGBColor(0x3B, 0x82, 0xF0); c.line.width = Pt(1.5); c.shadow.inherit = False
anchor.addprevious(c._element)
put_pic(sl, OR + "icon_1452.png", ccx - 0.25, ccy - 0.25, 0.5, 0.5, before=anchor)
anchor.getparent().remove(anchor)
for nm in ("시안 56","시안 57","시안 58","시안 59"):
    drop(33, nm)
print("33 통이미지 분리")

# ---------- 7. 사선 탭 → 오른쪽과 같은 띠, 좌우 색 다르게 ----------
# 7쪽 시스템 구성도
tab = find(7, "시안 탭")
card7 = find(7, "시안 77")
to_band(tab, card7.left / I, card7.width / I, NAVY)
# 그림(굿모닝 건물)·글자는 띠 앞에 있어야 함 → 이미 뒤 순서
# 19쪽
h19 = find(19, "시안 3"); c19 = find(19, "시안 2")
to_band(h19, c19.left / I, c19.width / I, BLUE)
# 20쪽
h20 = find(20, "시안 8"); c20 = find(20, "시안 7")
to_band(h20, c20.left / I, c20.width / I, BLUE)
drop(20, "시안 9")
set_fill(find(20, "시안 81"), NAVY)
# 21쪽
h21 = find(21, "시안 5"); c21 = find(21, "시안 3")
to_band(h21, c21.left / I, c21.width / I, BLUE)
drop(21, "시안 6"); drop(21, "시안 4")
set_fill(find(21, "시안 137"), NAVY)
# 29쪽 청주시(옅은 사선) → 제안사 띠와 같은 모양, 파랑
h29 = find(29, "시안 7"); band29 = find(29, "시안 22"); c29 = find(29, "시안 6")
to_band(h29, c29.left / I + 0.01, c29.width / I - 0.02, BLUE)
h29.top, h29.height = band29.top, band29.height
for pg in h29.text_frame.paragraphs:
    for r in pg.runs:
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
print("머리띠 사선 탭 → 띠 5곳")

# ---------- 8. 같은 색 좌우 머리띠 → 오른쪽 남색 ----------
PAIRS = {23: ["시안 9"], 26: ["시안 8"], 27: ["시안 8"], 34: ["시안 7"], 35: ["시안 8"],
         19: ["시안 42"], 21: ["시안 103"]}
for n, names in PAIRS.items():
    for nm in names:
        set_fill(find(n, nm), NAVY)
        print(n, nm, "→", NAVY)

p.save(dst)
print("저장", dst)
