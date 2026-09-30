"""v0.54 slide 16 - match the reference composition: pink lead pill with side rules, pink key phrases,
   both labels navy (left bigger), one highlighted card, thin rule under card titles. args: src dst"""
import sys, copy, re
from lxml import etree
from pptx import Presentation
from pptx.util import Pt
sys.stdout.reconfigure(encoding="utf-8")
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
E = 914400; NAVY, PINK = "143A69", "EC1C68"
p = Presentation(sys.argv[1]); sl = p.slides[15]
def color(rp, hexv):
    for t in ("solidFill", "gradFill", "noFill", "pattFill"):
        for e in rp.findall(A + t): rp.remove(e)
    sf = etree.Element(A + "solidFill"); etree.SubElement(sf, A + "srgbClr", val=hexv)
    ln = rp.find(A + "ln"); rp.insert((list(rp).index(ln) + 1) if ln is not None else 0, sf)
def hilite(shape, phrases, base=NAVY):
    pat = re.compile("(" + "|".join(re.escape(x) for x in phrases) + ")")
    for pg in shape._element.iter(A + "p"):
        for r in list(pg.findall(A + "r")):
            t = r.find(A + "t"); parts = [x for x in pat.split(t.text or "") if x]
            rp = r.find(A + "rPr")
            if rp is None:
                rp = etree.Element(A + "rPr"); r.insert(0, rp)
            prev = r
            for i, part in enumerate(parts):
                nr = r if i == 0 else copy.deepcopy(r)
                nr.find(A + "t").text = part; color(nr.find(A + "rPr"), PINK if part in phrases else base)
                if i: prev.addnext(nr)
                prev = nr
def fill(s, hexv, line=None):
    sp = s._element.find(P + "spPr")
    for c in sp.iter(A + "srgbClr"):
        par = c.getparent().getparent()
        if par.tag == P + "spPr": c.set("val", hexv)
        elif par.tag == A + "ln" and line: c.set("val", line)
def rect(x, y, w, h, hexv, kind=1):
    s = sl.shapes.add_shape(kind, int(x * E), int(y * E), int(w * E), int(h * E))
    s.fill.solid(); s.fill.fore_color.rgb = __import__("pptx.dml.color", fromlist=["RGBColor"]).RGBColor.from_string(hexv)
    s.line.fill.background(); s.shadow.inherit = False
    return s
lead = head2 = None
for s in sl.shapes:
    if s.has_text_frame:
        t = s.text_frame.text
        if t.startswith("환경변화와"): lead = s
        elif t.startswith("안정적 유지관리"): head2 = s
    if s.name == "LBL_사업지원방안": fill(s, NAVY, NAVY)
hilite(head2, ["최고 수준 서비스 지속 지원"])
hilite(lead, ["신속하게 반응", "안정적이고 지속적인 운영 지원"])
pill = rect(1.95, 2.94, 5.45, 0.42, "FDEAF1", 5)   # 5 = rounded rectangle
pill.adjustments[0] = 0.5
lead._element.addprevious(pill._element)
rect(1.35, 3.14, 0.45, 0.025, PINK); rect(7.55, 3.14, 0.45, 0.025, PINK)
# thin rule under titles (replace short pink bars)
for s in list(sl.shapes):
    if s.shape_type == 1 and abs(s.height - 0.02 * E) < 0.006 * E and abs(s.width - 0.3 * E) < 0.01 * E:
        cx = s.left + s.width / 2; s.width = int(0.95 * E); s.left = int(cx - s.width / 2); s.height = int(0.012 * E)
        fill(s, "B5C8E6")
# highlighted card (2nd of the right group) + right titles 9pt
def walk(sh):
    for s in sh:
        if s.shape_type == 6: yield from walk(s.shapes)
        else: yield s
for s in walk(sl.shapes):
    if s.has_text_frame and s.text_frame.text.strip() and s.top > 4 * E and s.width < 1.4 * E:
        if abs(s.left - 5.7 * E) < 0.05 * E: fill(s, "C3D6FA")
        if s.left > 4.4 * E and s.height < 1.4 * E and abs(s.width - 1.15 * E) < 0.02 * E:
            for pg in s.text_frame.paragraphs:
                for r in pg.runs: r.font.size = Pt(9)
p.save(sys.argv[2])
