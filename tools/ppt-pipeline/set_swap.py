# -*- coding: utf-8 -*-
"""세트 교체 — 도형 하나씩 지목해 그림을 라이브러리 자산으로 바꾼다(위치·크기 유지).
   원래 그림이 캔버스에서 차지하던 잉크 영역에 새 그림을 비율 유지로 맞춰 넣고,
   다시 칠하기 효과(duotone·clrChange 등)와 자르기(srcRect)는 걷어낸다.
   같은 그림 파트를 여러 도형이 공유해도 도형마다 새 파트를 만들어 서로 안 번진다.
   인자: <src> <dst> <spec.json>   spec = [[슬라이드, "도형이름", "이미지경로", 순번(같은 이름 n번째, 0부터)], ...]"""
import io, json, sys
from PIL import Image
from pptx import Presentation
sys.stdout.reconfigure(encoding="utf-8")
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
src, dst, spec = sys.argv[1:4]
jobs = json.load(open(spec, encoding="utf-8"))
p = Presentation(src)


def walk(shapes):
    for s in shapes:
        if s.shape_type == 6:
            yield from walk(s.shapes)
        else:
            yield s


def drop_img_layer(blip):
    """예술 효과 원본 레이어(a14:imgProps → hdphoto) 제거. 남겨 두면 PowerPoint 가 저장할 때
       그 원본으로 그림을 다시 만들어 교체한 그림이 옛 그림으로 되돌아간다."""
    ext = blip.find(A + "extLst")
    if ext is None:
        return
    for e in list(ext):
        if any(c.tag.endswith("}imgProps") for c in e):
            ext.remove(e)


def ink(im):
    return im.getchannel("A").point(lambda v: 255 if v > 8 else 0).getbbox()


done = 0
for sno, name, img, nth in jobs:
    cands = [s for s in walk(p.slides[sno - 1].shapes) if s.shape_type == 13 and s.name == name]
    if nth >= len(cands):
        print("없음", sno, name, nth); continue
    sh = cands[nth]
    if img == "DELETE":                        # 겹쳐 있던 보조 그림 제거
        sh._element.getparent().remove(sh._element); done += 1; continue
    orig = Image.open(io.BytesIO(sh.image.blob)).convert("RGBA")
    blip = sh._element.find(".//" + A + "blip")
    fill = blip.getparent()
    src_rect = fill.find(A + "srcRect")
    ow, oh = orig.size
    # 자르기가 걸려 있으면 보이는 부분만 기준으로
    if src_rect is not None:
        l, t, r, b = (int(src_rect.get(k, 0)) / 100000 for k in ("l", "t", "r", "b"))
        orig = orig.crop((round(l * ow), round(t * oh), round(ow * (1 - r)), round(oh * (1 - b))))
        fill.remove(src_rect)
        ow, oh = orig.size
    box = ink(orig) or (0, 0, ow, oh)
    # 불투명 사각 배경 그림(인물 사진형)은 잉크 = 전체 → 캔버스 전체에 맞춘다
    tw, th = box[2] - box[0], box[3] - box[1]
    cx, cy = (box[0] + box[2]) / 2, (box[1] + box[3]) / 2
    new = Image.open(img).convert("RGBA")
    nb = ink(new)
    if nb:
        new = new.crop(nb)
    k = min(tw / new.width, th / new.height)
    # 캔버스 해상도가 낮으면 캔버스를 키워 선명도 확보
    up = max(1.0, 1200 / max(ow, oh))   # 2026-09-30: 400 이면 장표에서 키울 때 깨진다(원본 해상도 원칙) → 1200
    W, H = round(ow * up), round(oh * up)
    new = new.resize((max(1, round(new.width * k * up)), max(1, round(new.height * k * up))), Image.LANCZOS)
    canvas = Image.new("RGBA", (W, H), (0, 0, 0, 0))
    canvas.paste(new, (round(cx * up - new.width / 2), round(cy * up - new.height / 2)), new)
    buf = io.BytesIO(); canvas.save(buf, "PNG", optimize=True); buf.seek(0)
    _, rid = sh.part.get_or_add_image_part(buf)
    blip.set(R + "embed", rid)
    for ch in list(blip):                      # 다시 칠하기·밝기 효과 제거(extLst 는 유지)
        if ch.tag != A + "extLst":
            blip.remove(ch)
    drop_img_layer(blip)
    done += 1
p.save(dst)
print("교체", done, "/", len(jobs))
