# -*- coding: utf-8 -*-
"""20쪽(유지보수 관리 체계 4/5) 처리 절차 도식 강조 복원 + 글자 세로 가운데. 인자: <src> <dst>
원본 도식은 '행동' 상자가 짙은 채움 + 흰 글자라 흐름이 한눈에 읽혔다(지금은 옅은 면 + 작은 글자).
- 사용자 열 행동 상자(유지관리요청·요청 답변·요청자 확인) = 남색 143A69 채움, 흰 글자
- 상주 담당자 열 행동 상자(요청 접수·문제 해결·변경 처리·스크립트 작성/확인·자료 처리) = 파랑 2F78E0 채움, 흰 글자
- 판단 마름모 = 흰 면 + 파랑 테두리 + 남색 글자, IT서비스 책임자 승인만 분홍 EC1C68 테두리·글자
- 종료 = 짙은 남색 알약(가운데 세로선 위에), 상자 안 작은 아이콘은 빼고 글자 가운데
- 모든 도형 글자: anchor ctr, 위아래 여백 0, 줄 간격·문단 간격 없음
- 열 머리 글자 12pt, 사람 아이콘 축소
- 오른쪽 표: 열 폭 재배분(조치 1.23 / 세부 0.90 / 내용 1.90), 머리 11pt, 분류 11pt, 세부 10pt, 내용 9.5pt"""
import os, sys
sys.path.insert(0, os.path.dirname(__file__))
from pptx import Presentation
from pptx.util import Pt, Inches as I, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

sys.stdout.reconfigure(encoding="utf-8")
src, dst = sys.argv[1:3]
p = Presentation(src)
s = p.slides[19]
by = {sh.shape_id: sh for sh in s.shapes}

NAVY = RGBColor(0x14, 0x3A, 0x69)
BLUE = RGBColor(0x2F, 0x78, 0xE0)
PINK = RGBColor(0xEC, 0x1C, 0x68)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x1F, 0x4E, 0x79)
TXT = RGBColor(0x33, 0x33, 0x33)
F2, F3 = "a시월구일2", "a시월구일3"
COMP = {"a시월구일3": 0.115, "a시월구일2": 0.07}


def face(r, f):
    rPr = r._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea"):
        e = rPr.find(qn(tag))
        if e is None:
            e = rPr.makeelement(qn(tag), {})
            rPr.append(e)
        e.set("typeface", f)


def text(sh, size, color, font=F3, lr=0.0, tb=(0.0, 0.0), align=None, lnspc=None, comp_on=True):
    """글자 크기·색·글꼴 + 세로 가운데(위아래 여백 0, 문단 간격 0)."""
    tf = sh.text_frame
    bp = tf._txBody.find(qn("a:bodyPr"))
    bp.set("anchor", "ctr")
    bp.set("anchorCtr", "0")
    bp.set("lIns", str(int(I(lr)))); bp.set("rIns", str(int(I(lr))))
    # 이 글꼴(a시월구일)은 글자가 줄 상자 안에서 아래로 앉는다(렌더 실측: 크기의 약 6%).
    # 아래 여백을 그만큼(×2) 줘서 '보이는 글자'가 정확히 세로 가운데에 오게 한다.
    comp = size * (COMP.get(font, 0.0)) if comp_on else 0.0  # 줄 간격 90% 마름모는 실측상 보정 불필요
    bp.set("tIns", str(int(I(tb[0])))); bp.set("bIns", str(int(I(tb[1]) + Pt(comp))))
    for c in list(bp):
        if c.tag in (qn("a:spAutoFit"), qn("a:normAutofit")):
            bp.remove(c)
    for pg in tf.paragraphs:
        pPr = pg._p.get_or_add_pPr()
        for t in ("a:spcBef", "a:spcAft", "a:lnSpc"):
            for e in pPr.findall(qn(t)):
                pPr.remove(e)
        if lnspc:
            pg.line_spacing = lnspc
        if align is not None:
            pg.alignment = align
        for r in pg.runs:
            r.font.size = Pt(size)
            r.font.bold = False
            r.font.color.rgb = color
            face(r, font)
        # 빈 문단의 끝 글자 크기도 맞춘다(빈 줄 높이가 커지지 않게)
        epr = pg._p.find(qn("a:endParaRPr"))
        if epr is not None:
            epr.set("sz", str(int(size * 100)))


def fill(sh, color):
    sh.fill.solid(); sh.fill.fore_color.rgb = color


def line(sh, color=None, w=None):
    if color is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = color
        sh.line.width = Pt(w)


def kill(ids):
    for i in ids:
        if i in by:
            by[i]._element.getparent().remove(by[i]._element)


def hx(i, x0, x1):
    """가로선의 양 끝 x(인치)를 다시 잡는다(화살표 방향은 그대로)."""
    sh = by[i]
    sh.left = I(min(x0, x1)); sh.width = I(abs(x1 - x0))


def widen(i, w):
    sh = by[i]; cx = sh.left + sh.width // 2
    sh.width = I(w); sh.left = cx - sh.width // 2


# ---------- 도식: 상자 안 작은 아이콘 제거 ----------
kill([71, 73, 75, 77, 79, 81])

# 글자가 테두리에 붙지 않게 넓힌다(가운데 유지) + 맞닿는 선 끝을 새 테두리로
for i in (70, 72, 74, 78):              # 사용자 열 0.44~1.77
    by[i].left = I(0.44); by[i].width = I(1.33)
for i in (38, 42, 68):
    hx(i, 1.77, 3.75)
widen(82, 1.36); widen(84, 1.36)        # IT서비스 책임자 승인 1.995~3.355
widen(90, 1.16)                         # 스크립트 지원 여부 3.70~4.86
widen(88, 1.02)                         # 자료처리 지원 4.92~5.94
hx(57, 3.355, 3.70)
hx(63, 3.355, 3.74)
hx(59, 4.86, 5.14)
hx(48, 5.94, 6.04)

# 사용자 열 행동 상자 = 남색
for i in (70, 72, 74, 78):
    sh = by[i]; fill(sh, NAVY); line(sh, None)
    text(sh, 9.5, WHITE, lr=0.02, align=PP_ALIGN.CENTER)
# 상주 담당자 열 행동 상자 = 파랑
for i in (85, 87, 89, 91, 92, 93):
    sh = by[i]; fill(sh, BLUE); line(sh, None)
    text(sh, 9.5, WHITE, lr=0.02, align=PP_ALIGN.CENTER)
# 변경 처리 상자는 글자 폭보다 살짝 좁아 넓힌다(가운데 유지)
b = by[89]; cx = b.left + b.width // 2
b.width = I(0.74); b.left = cx - b.width // 2

# 종료 = 남색 알약(세로선 x=1.10 중심)
for i in (76, 80):
    sh = by[i]; fill(sh, NAVY); line(sh, None)
    cx = I(1.10); sh.width = I(0.66); sh.left = cx - sh.width // 2
    text(sh, 9.5, WHITE, lr=0.0, align=PP_ALIGN.CENTER)
# 기각 → 종료(가운데 열)
sh = by[83]; fill(sh, NAVY); line(sh, None)
text(sh, 9.5, WHITE, align=PP_ALIGN.CENTER)

# 판단 마름모
for i in (86, 88, 90):
    sh = by[i]; fill(sh, WHITE); line(sh, BLUE, 1.25)
    text(sh, 9, NAVY, align=PP_ALIGN.CENTER, lnspc=0.9, comp_on=False)
for i in (82, 84):
    sh = by[i]; fill(sh, WHITE); line(sh, PINK, 1.25)
    text(sh, 9, PINK, align=PP_ALIGN.CENTER, lnspc=0.9, comp_on=False)

# Y/N/기각 표시
for i in (41, 47, 50, 54, 58, 61, 67):
    text(by[i], 8, NAVY, align=PP_ALIGN.CENTER)
text(by[56], 8, PINK, align=PP_ALIGN.CENTER)

# ---------- 열 머리: 글자 12pt, 사람 아이콘 축소 ----------
HEAD_Y, HEAD_H = 2.65, 0.50
heads = [  # (아이콘, 글자, 열 x0, 열 x1, 글자 폭 추정)
    (32, 33, 0.38, 1.82, 0.62),
    (34, 35, 1.86, 3.48, 1.12),
    (36, 37, 3.53, 6.16, 1.62),
]
for pic, tx, x0, x1, tw in heads:
    pc = by[pic]
    ar = pc.width / pc.height
    ih = 0.27 if pic != 36 else 0.24
    pc.height = I(ih); pc.width = int(I(ih) * ar)
    gap = 0.07
    total = pc.width / 914400 + gap + tw
    left = (x0 + x1) / 2 - total / 2
    pc.left = I(left); pc.top = I(HEAD_Y + (HEAD_H - ih) / 2)
    t = by[tx]
    t.left = I(left + pc.width / 914400 + gap - 0.02); t.width = I(tw + 0.04)
    t.top = I(HEAD_Y); t.height = I(HEAD_H)
    text(t, 12 if tx != 35 else 11, NAVY if tx != 33 else BLUE, align=PP_ALIGN.LEFT, lnspc=0.95 if tx == 35 else None)

# ---------- 오른쪽 표 ----------
C1, C2, C3, C4 = 6.43, 7.66, 8.56, 10.46
# 머리
for i, (a, b2) in zip((98, 99, 100), ((C1, C2), (C2, C3), (C3, C4))):
    sh = by[i]; sh.left = I(a); sh.width = I(b2 - a)
    text(sh, 11, WHITE, align=PP_ALIGN.CENTER)
# 분류 묶음 상자(배경)·아이콘·이름
for bg, ic, lab in ((101, 102, 103), (108, 109, 110), (117, 118, 119)):
    g = by[bg]; g.left = I(C1 + 0.01); g.width = I(C2 - C1 - 0.02)
    pc = by[ic]
    pc.width = pc.height = I(0.32)
    pc.left = I(C1 + 0.08); pc.top = g.top + (g.height - pc.height) // 2
    t = by[lab]
    t.left = I(C1 + 0.44); t.width = I(C2 - C1 - 0.46)
    t.top = g.top; t.height = g.height
    col = by[lab].text_frame.paragraphs[0].runs[0].font.color.rgb
    text(t, 11, col, align=PP_ALIGN.LEFT, lnspc=0.95)
# 단어 중간 줄바꿈 방지: 점|검 → 빗금 뒤에서 줄바꿈(문단 글자를 다시 쓰고 아래에서 서식)
by[107].text_frame.paragraphs[0].text = "시스템 기술지원/" + chr(11) + "운영상태 점검"
# 세부유형·내용 칸
rows = [(104, 105), (106, 107), (111, 112), (113, 114), (115, 116),
        (120, 121), (122, 123), (124, 125), (126, 127)]
for sub, con in rows:
    a = by[sub]; b2 = by[con]
    a.left = I(C2); a.width = I(C3 - C2)
    b2.left = I(C3); b2.width = I(C4 - C3)
    col = a.text_frame.paragraphs[0].runs[0].font.color.rgb
    text(a, 10, col, align=PP_ALIGN.CENTER)
    text(b2, 9.5, TXT, font=F2, lr=0.08, align=PP_ALIGN.LEFT, lnspc=0.95)

p.save(dst)
print("저장", dst)
