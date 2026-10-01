# -*- coding: utf-8 -*-
"""18쪽(유지보수 관리 체계 2/5) 원본 문구 복원. 인자: <src> <dst>
- 인물 일러스트 축소, 3대 가치(안정성·지속성·신속성) 원문 설명 2줄씩 복원
- 주요활동 5종: 원본 이름(서비스 요청(SR) 처리 등) + 세부 활동 원문 전부
- S-ISM: 글자가 박힌 3D 블록 그림 4장을 지우고 4개 영역의 프로세스 목록(원문 18항목) 카드로
- 지속성 아이콘(볼록한 3D 순환 화살표 그림)·배포 아이콘(볼록한 화살표 구름) → 납작한 도형/아이콘"""
import os, sys, io
sys.path.insert(0, os.path.dirname(__file__))
from pptx import Presentation
from pptx.util import Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from lib_e import *
from pptx.oxml.ns import qn

sys.stdout.reconfigure(encoding="utf-8")
src, dst = sys.argv[1:3]
p = Presentation(src); s = p.slides[17]
by = by_id(s)

keep_pics = {81: None, 93: None, 104: None, 108: None, 116: None, 120: None, 100: None, 124: None}
flat_cloud = None
for sh in p.slides[18].shapes:            # 19쪽의 납작한 구름 아이콘(배포용)
    if sh.shape_id == 84 and sh.shape_type == 13:
        flat_cloud = sh.image.blob

# 지울 것: 기존 카드·텍스트·3D 블록 그림·작은 화살표
kill(s, [i for i in range(80, 159) if i not in keep_pics])

# 1) 인물 축소
fit_pic(by[160], 0.30 + 1.15 / 2, 1.62 + 1.30 / 2, h=1.30)

# 2) 3대 가치
X0, X1, Y0, H0 = 1.52, 10.53, 1.62, 1.38
GAP = 0.10
cw = (X1 - X0 - 2 * GAP) / 3
vals = [
    ("안정성", 81, [[("프로젝트 착수부터 종료까지\v", F2, TXT), ("안정적인 관리활동 수행", F3, INK)],
                   [("시스템과 서비스의 변화에 사용자 그룹이\v반응하지 않아도 되도록 ", F2, TXT), ("안정적인 서비스 보장", F3, INK)]]),
    ("지속성", None, [[("끊김없이 ", F2, TXT), ("무중단으로 지속적인 서비스 지원", F3, INK)],
                     [("강력한 보안 점검과 철저한 백업·복구 등\v활동의 안정적 지원으로 ", F2, TXT), ("지속성 확보", F3, INK)]]),
    ("신속성", 93, [[("법제도적, 시스템·업무적으로 발생한\v환경변화에 ", F2, TXT), ("신속하게 반응하고 조치", F3, INK)],
                   [("사용자 요구를 능동적으로 파악, ", F2, TXT), ("신속한 응대", F3, INK)]]),
]
for k, (title, pic, items) in enumerate(vals):
    x = X0 + k * (cw + GAP)
    box(s, x, Y0, cw, H0, fill=WHITE, line=LIGHT, lw=1)
    box(s, x, Y0, cw, 0.46, fill=PALE, line=None, adj=0.18)
    icx, icy = x + 0.33, Y0 + 0.23
    if pic:
        fit_pic(by[pic], icx, icy, h=0.40); to_front(s, by[pic])
    else:                                   # 납작한 순환 화살표
        c = box(s, icx - 0.19, icy - 0.19, 0.38, 0.38, fill=BLUE, line=None, shape=MSO_SHAPE.OVAL)
        ca = box(s, icx - 0.13, icy - 0.13, 0.26, 0.26, fill=WHITE, line=None, shape=MSO_SHAPE.CIRCULAR_ARROW)
        av = ca._element.spPr.find(qn("a:prstGeom")).find(qn("a:avLst"))
        for g in list(av):
            av.remove(g)
        for nm, v in (("adj1", 14000), ("adj2", 1142319), ("adj3", 20457681), ("adj4", 2400000), ("adj5", 14000)):
            av.append(av.makeelement(qn("a:gd"), {"name": nm, "fmla": "val %d" % v}))
    textbox(s, x + 0.60, Y0 + 0.05, cw - 0.7, 0.36, [title], 13.5, color=NAVY, font=F4, anchor=MSO_ANCHOR.MIDDLE)
    textbox(s, x + 0.10, Y0 + 0.56, cw - 0.18, 0.78, items, 9, bullet="•", space_after=3, line_spacing=1.0)

# 3) 주요활동 커스터마이징
LY = 3.12
pill = box(s, 0.30, LY, 2.30, 0.30, fill=NAVY, line=None, adj=0.5)
fit_pic(by[100], 0.52, LY + 0.15, h=0.20); to_front(s, by[100])
textbox(s, 0.70, LY, 1.85, 0.30, ["주요활동 커스터마이징"], 10.5, color=WHITE, font=F4, anchor=MSO_ANCHOR.MIDDLE)
hline(s, 2.70, LY + 0.15, 10.53, color=LIGHT, w=1)

AY, AH = 3.48, 1.58
aw = (10.53 - 0.30 - 4 * 0.08) / 5
acts = [
    ("서비스 요청(SR) 처리", 104, ["서비스 요청(SR)의 적시 처리",
                                 "상주 담당자의 단순 답변이나\v원격지원, 기술지원 활동 수행",
                                 "장비 및 솔루션 관련 기술지원\v1차 처리 및 책임 이관"]),
    ("변경 요청(RFC) 처리", 108, ["변경 영향 평가와 적정성\v확보 후 변경 처리",
                                "변경요청 처리 과정 기록관리",
                                "변경 사항이 반영된 산출물\v작성 후 주관기관의 승인 획득"]),
    ("배포 관리", "cloud", ["서비스 배포 결과 기록 관리",
                          "통합테스트가 완료된 변경\v처리의 운영 서버 테스트\v수행 후 주관기관 승인 획득"]),
    ("요구사항 관리", 116, ["고객 요구사항 검증 및 승인,\v변경관리, 추적관리"]),
    ("장애 관리", 120, ["장애 등급에 따른 조치",
                      "장애 근본원인 분석,\v재발방지 및 예방 활동 수행"]),
]
for k, (title, pic, items) in enumerate(acts):
    x = 0.30 + k * (aw + 0.08)
    box(s, x, AY, aw, AH, fill=WHITE, line=LIGHT, lw=1)
    icx, icy = x + 0.24, AY + 0.22
    if pic == "cloud":
        ip = s.shapes.add_picture(io.BytesIO(flat_cloud), I(icx - 0.17), I(icy - 0.15))
        fit_pic(ip, icx, icy, h=0.30)
    else:
        fit_pic(by[pic], icx, icy, h=0.34); to_front(s, by[pic])
    textbox(s, x + 0.45, AY + 0.04, aw - 0.50, 0.36, [title], 10, color=NAVY, font=F4, anchor=MSO_ANCHOR.MIDDLE)
    hline(s, x + 0.10, AY + 0.45, x + aw - 0.10, color=LIGHT, w=0.75)
    textbox(s, x + 0.07, AY + 0.53, aw - 0.12, AH - 0.58, items, 8.5, bullet="•", space_after=2.5, line_spacing=1.0)

# 4) 운영/유지보수 방법론(S-ISM)
PY, PH = 5.16, 1.86
box(s, 0.30, PY, 10.23, PH, fill=NAVY, line=None, adj=0.05)
fit_pic(by[124], 0.55, PY + 0.21, h=0.24); to_front(s, by[124])
textbox(s, 0.75, PY + 0.04, 2.6, 0.34, ["운영/유지보수 방법론(S-ISM)"], 11, color=WHITE, font=F4, anchor=MSO_ANCHOR.MIDDLE)
hline(s, 3.20, PY + 0.21, 10.35, color=GRAY, w=0.75)
groups = [
    ("01", "프로세스 관리", 1.76, [["조직 프로세스 관리", "조직 교육 훈련"]]),
    ("02", "서비스 수행", 3.50, [["서비스 수행 계획 관리", "서비스 수행 감독 및 통제", "서비스 수준 관리", "서비스 연속성 관리"],
                             ["인시던트 관리", "장애(문제) 관리", "공급자 계약 관리", "전략적 서비스 관리"]]),
    ("03", "서비스 개발 및 전달", 2.30, [["변경관리 / 배포관리", "요구사항관리", "전략적 서비스 관리", "테스트/검증 관리"]]),
    ("04", "지원", 2.15, [["구성 관리", "품질 관리", "측정 및 분석 관리", "의사결정 관리"]]),
]
GY, GH = PY + 0.44, PH - 0.56
x = 0.42
for num, title, w, cols in groups:
    box(s, x, GY, w, GH, fill=WHITE, line=None, adj=0.06)
    box(s, x, GY, w, 0.32, fill=BLUE, line=None, adj=0.2)
    textbox(s, x + 0.06, GY, w - 0.12, 0.32,
            [[(num + "  ", F4, LIGHT), (title, F4, WHITE)]], 10.5, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    cwid = (w - 0.16) / len(cols)
    for j, items in enumerate(cols):
        textbox(s, x + 0.10 + j * cwid, GY + 0.40, cwid, GH - 0.46, items, 9, bullet="•", space_after=2.5, line_spacing=1.0)
    x += w + 0.08

p.save(dst)
print("saved", dst)
