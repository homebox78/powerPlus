# -*- coding: utf-8 -*-
"""34쪽: 성과결과 추출 방식 카드 — 아이콘이 머리 띠·글에 걸림 → 본문 칸 안에 작게 세로 가운데, 글자 키움.
'예시' 배지는 머리 띠 높이 안에 들어가게 낮춤. 35쪽: 운영 성과 줄 아이콘 15% 축소. 인자: <src> <dst>"""
import sys
from pptx import Presentation
from pptx.util import Pt
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
sys.stdout.reconfigure(encoding="utf-8")
src, dst = sys.argv[1:3]
IN = 914400
prs = Presentation(src)
sl = prs.slides[33]
by = {s.name: s for s in sl.shapes}

# 1) 카드: 본문 칸 = 머리 띠(시안 55) 아래 ~ 카드(시안 54) 아래
card, head, pic, txt = by["시안 54"], by["시안 55"], by["시안 그림 58"], by["시안 57"]
top = head.top + head.height
bot = card.top + card.height
mid = (top + bot) / 2
k = (bot - top) * 0.78 / pic.height
pic.width, pic.height = int(pic.width * k), int(pic.height * k)
pic.left = int(card.left + 0.14 * IN)
pic.top = int(mid - pic.height / 2)
txt.left = int(pic.left + pic.width + 0.06 * IN)
txt.width = int(card.left + card.width - 0.08 * IN - txt.left)
txt.top, txt.height = int(top), int(bot - top)
tf = txt.text_frame
tf.vertical_anchor = MSO_ANCHOR.MIDDLE
tf.margin_top = tf.margin_bottom = 0
for p in tf.paragraphs:
    p.alignment = PP_ALIGN.CENTER
    for r in p.runs:
        r.font.size = Pt(11)

# 2) 예시 배지: 머리 띠(시안 7) 안에 세로 가운데
bar, badge = by["시안 7"], by["시안 50"]
badge.height = int(bar.height * 0.66)
badge.top = int(bar.top + (bar.height - badge.height) / 2)
badge.left = int(bar.left + bar.width - 0.08 * IN - badge.width)
badge.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
badge.text_frame.margin_top = badge.text_frame.margin_bottom = 0

# 3) 35쪽 아이콘 15% 축소(중심 유지)
s35 = prs.slides[34]
for s in s35.shapes:
    if s.name in ("시안 그림 27", "시안 그림 31", "시안 그림 35", "시안 그림 39"):
        cx, cy = s.left + s.width / 2, s.top + s.height / 2
        s.width, s.height = int(s.width * 0.85), int(s.height * 0.85)
        s.left, s.top = int(cx - s.width / 2), int(cy - s.height / 2)
print("34쪽 카드·배지, 35쪽 아이콘 4")
prs.save(dst)
