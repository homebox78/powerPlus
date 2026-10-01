# -*- coding: utf-8 -*-
"""16쪽(III-1. 수행 및 지원방안 개요) 원본 설명 문장 복원.
- 카드 8장의 짧은 키워드를 원본 설명 문장으로 되돌림(원본 굵게 표시 낱말 = 분홍 강조)
- 번호·제목을 위로 당겨 설명 자리를 넓히고 제목·설명 글자를 키움
- 줄바꿈(^)은 실제 렌더 줄을 읽어 어절 경계로 정해 둔 것(COM 저장 없이 python-pptx 만 사용)
인자: <src> <dst>  (16쪽만 손댄다)
"""
import re, sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

sys.stdout.reconfigure(encoding="utf-8")
src, dst = sys.argv[1:3]
SNO = 16
p = Presentation(src)
s = p.slides[SNO - 1]
by = {sh.shape_id: sh for sh in s.shapes}
I = Inches
TXT, NAVY, PINK = RGBColor(0x33, 0x33, 0x33), RGBColor(0x14, 0x3A, 0x69), RGBColor(0xEC, 0x1C, 0x68)
F2, F3, F4 = "a시월구일2", "a시월구일3", "a시월구일4"
DESC_PT, TITLE_PT = 8.5, 9.5


def face(r, f):
    rPr = r._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea"):
        e = rPr.find(qn(tag))
        if e is None:
            e = rPr.makeelement(qn(tag), {})
            rPr.append(e)
        e.set("typeface", f)


def clear(tf):
    for pg in list(tf.paragraphs)[1:]:
        pg._p.getparent().remove(pg._p)
    for r in list(tf.paragraphs[0].runs):
        r._r.getparent().remove(r._r)
    return tf.paragraphs[0]


def desc_box(x, y, w, h, items):
    """items: 문자열 목록. {…} 로 감싼 부분 = 원본에서 굵게 강조했던 낱말 → 분홍·강조 서체."""
    tb = s.shapes.add_textbox(I(x), I(y), I(w), I(h))
    tb.name = "복원설명"
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = tf.margin_right = I(0.02)
    tf.margin_top = tf.margin_bottom = 0
    p0 = clear(tf)
    ind = int(Pt(DESC_PT) * 0.75)
    for i, it in enumerate(items):
        pg = p0 if i == 0 else tf.add_paragraph()
        pg.alignment = PP_ALIGN.LEFT
        pg.line_spacing = 1.05
        pg.space_after = Pt(2.5)
        pPr = pg._p.get_or_add_pPr()
        pPr.set("marL", str(ind)); pPr.set("indent", str(-ind))
        bc = pPr.makeelement(qn("a:buClr"), {})
        bc.append(bc.makeelement(qn("a:srgbClr"), {"val": "5F7496"}))
        pPr.append(bc)
        pPr.append(pPr.makeelement(qn("a:buFont"), {"typeface": "Arial"}))
        pPr.append(pPr.makeelement(qn("a:buChar"), {"char": "•"}))
        for k, seg in enumerate(re.split(r"[{}]", it.replace("|", "{}"))):
            if not seg:
                continue
            emph = k % 2 == 1
            parts = seg.split("^")
            for q, part in enumerate(parts):
                if q:
                    pg._p.append(pg._p.makeelement(qn("a:br"), {}))
                if part:
                    add_run(pg, part, emph)
    return tb


def add_run(pg, seg, emph):
    r = pg.add_run(); r.text = seg
    small = bool(re.fullmatch(r"\s*\([A-Za-z]+\)[,)]*\s*", seg))   # 영문 괄호 병기는 7.5pt
    r.font.size = Pt(7.5 if small else DESC_PT); r.font.bold = False
    r.font.color.rgb = PINK if emph else TXT
    face(r, F3 if emph else F2)


# (번호, 제목, 지울 키워드칸, 원본 설명)
cards = [
    (68, 69, 70, ["제안사 검증된 ^방법론 기반",
                  "{서비스요청유형}의 특성에 맞는 처리 절차 및 응대 관리 체계"]),
    (73, 74, 75, ["대한민국 정부 ^행정정보시스템 집약 전문성 확보",
                  "{청주시 유지보수 18년 경력}의 수행 책임자"]),
    (78, 79, 80, ["{현 사업자}",
                  "안정된 인계를 위한 공동운영, 비상주 ^운영지원으로 ^{업무연속성 유지 보장}"]),
    (83, 84, 85, ["청주시 시스템 연계 현황 현행화", "연계표준가이드라인 관리", "전사차원의 기술적 지원"]),
    (88, 89, 90, ["전사적 ^기술지원으로 ^확장된 {3단계 ^장애관리 체계}",
                  "백업/복구 체계 및 절차, 수행 기준 ^관리"]),
    (93, 94, 95, ["보안 원칙(기밀성 |(Confidentiality), |무결성 ^|(Integrity), ^|가용성 ^|(Availability))|에 충실한 관리"]),
    (98, 99, 100, ["기능개선 요구 및 법제도·시스템적 환경 변화에 ^신속한 대응이 ^가능토록 지원"]),
    (103, 104, 105, ["교육대상자 ^역할별 차별화된 교육 실시",
                     "실질적 운영능력 배양을 위한 ^체계적 교육지원"]),
]
for num, title, old, items in cards:
    n, t, o = by[num], by[title], by[old]
    n.top = I(3.63)
    t.top = I(4.01); t.height = I(0.40)
    for pg in t.text_frame.paragraphs:
        pg.line_spacing = 0.95
        for r in pg.runs:
            r.font.size = Pt(TITLE_PT); r.font.color.rgb = NAVY; face(r, F4)
    t.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    L, W = o.left / 914400.0, o.width / 914400.0
    o._element.getparent().remove(o._element)
    desc_box(L + 0.05, 4.47, W - 0.08, 1.22, items)

# 사업수행방안·사업지원방안 라벨
for i in (48, 57):
    for pg in by[i].text_frame.paragraphs:
        for r in pg.runs:
            r.font.size = Pt(12)

p.save(dst)
print("저장", dst)
