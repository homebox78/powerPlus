# -*- coding: utf-8 -*-
"""19쪽(유지보수 관리 체계 3/5) 2차: 원본 구조로 재구성. 인자: <src> <dst>
- 큰 3D 일러스트(인물·서버·CI/CD) 삭제 → 원본의 작은 원형 선 아이콘
- 왼쪽: 소제목 → 1·2차 머리 → 1차/2차 두 칸 + MSA 한 줄 → 전환 인력·설계 산출물 두 칸 → 강조 줄
  '직접 전환 경험(분홍) 기반, 가장 빠르고 정확한 시스템 관리'
- 오른쪽: 4행, 왼쪽 색 칸(작은 아이콘 + 제목 11pt) + 원문 두 문장 10pt
- 하단 주석 원문"""
import os, sys
sys.path.insert(0, os.path.dirname(__file__))
from pptx import Presentation
from lib_g import *

sys.stdout.reconfigure(encoding="utf-8")
src, dst = sys.argv[1:3]
p = Presentation(src); s = p.slides[18]
purge(s)
here = os.path.dirname(os.path.abspath(__file__))
o = Presentation(os.path.join(here, "..", "..", "..", "제안서", "청주시_발표자료.원본백업.pptx")).slides[18]
ic = {k: pic_blob(o, v) for k, v in dict(c1=76, c2=90, msa=94, man=100, doc=106, chk=110,
                                         box=114, cicd=119, mon=124, gauge=129).items()}

k3 = lambda t: (t, F3, INK)
k2 = lambda t: (t, F2, TXT)


def disc(cx, cy, d, key, fill=BLUE, ratio=0.6):
    box(s, cx - d / 2, cy - d / 2, d, d, fill=fill, line=None, shape=MSO_SHAPE.OVAL)
    add_pic(s, ic[key], cx, cy, d * ratio)


PW = (W - 0.14) / 2
LX, RX = L, R - PW
PY, PH = 1.62, 0.48
PB = BOT
tb(s, LX, PY, PW, PH, ["클라우드 전환 수행 이력 및 이해도"], 15, fill=BLUE, line=None, color=WHITE, font=F4,
   align=PP_ALIGN.CENTER, margins=(0, 0))
tb(s, RX, PY, PW, PH, ["클라우드 환경 운영관리 방안"], 15, fill=NAVY, line=None, color=WHITE, font=F4,
   align=PP_ALIGN.CENTER, margins=(0, 0))
for x in (LX, RX):
    box(s, x, PY + PH, PW, PB - PY - PH, fill=WHITE, line=LIGHT, lw=1, shape=MSO_SHAPE.RECTANGLE)

# ---------- 왼쪽 ----------
x0, w0 = LX + 0.15, PW - 0.30
y = PY + PH + 0.08
tb(s, x0, y, w0, 0.40, ["전환 사업을 직접 수행한 유지관리 사업자"], 15, color=NAVY, font=F4, align=PP_ALIGN.CENTER,
   margins=(0, 0))
y += 0.46
tb(s, x0, y, w0, 0.32, ["클라우드 전환 사업 1·2차 직접 수행 (2024 ~ 2025)"], 11.5, fill=BLUE, line=None,
   color=WHITE, font=F4, align=PP_ALIGN.CENTER, margins=(0, 0))
y += 0.32
BH1 = 1.56
box(s, x0, y, w0, BH1, fill=PALE, line=None, shape=MSO_SHAPE.RECTANGLE)
cols = [(x0, w0 * 0.52), (x0 + w0 * 0.52, w0 * 0.48)]
rows = [("c1", "1차 (2024) 기반 전환", ["클라우드 아키텍처 구성", "공통서비스 전환\v(인증·파일·연계·결재·기안문)"]),
        ("c2", "2차 (2025) 단위업무 전환", ["단위업무 91종 일괄 전환\v(공통 54·개별 37)", "신·구 시스템 병행운영 완료"])]
for k, (key, head, items) in enumerate(rows):
    cx, cwd = cols[k]
    disc(cx + 0.24, y + 0.30, 0.38, key)
    tb(s, cx + 0.48, y + 0.10, cwd - 0.50, 0.30, [head], 10.5, color=NAVY, font=F4, margins=(0, 0), wrap=False)
    tb(s, cx + 0.48, y + 0.40, cwd - 0.50, 0.66, items, 9.5, bullet="•", margins=(0, 0), space_after=2,
       anchor=MSO_ANCHOR.TOP)
hline(s, x0 + 0.12, y + 1.10, x0 + w0 - 0.12, color=LIGHT, w=1)
disc(x0 + 0.30, y + 1.33, 0.32, "msa")
tb(s, x0 + 0.52, y + 1.18, w0 - 0.56, 0.30,
   [[k2("청주시 전산실 온프레미스 프라이빗 클라우드 "), k3("MSA 구조"), k2("로 재설계·구축")]], 9.5, margins=(0, 0))
y += BH1 + 0.12

BW = (w0 - 0.12) / 2
BH2 = 1.62
lows = [("man", "전환 인력 = 유지관리 인력", ["전환 사업 참여 인력을 본\v유지관리에 투입",
                                            "서비스 경계·API·연계\v구조를 설계 단계부터 보유",
                                            "별도 인수인계 없이\v즉시 안정 운영"]),
        ("doc", "설계 산출물·소스 보유", ["MSA 서비스 분해도,\vAPI 게이트웨이 라우팅 정보",
                                      "컨테이너 구성 정보 및\v배포 스크립트",
                                      "30여 종 연계 인터페이스\v규격서·모듈 소스"])]
for k, (key, head, items) in enumerate(lows):
    bx = x0 + k * (BW + 0.12)
    tb(s, bx, y, BW, 0.32, [head], 11.5, fill=BLUE, line=None, color=WHITE, font=F4,
       align=PP_ALIGN.CENTER, margins=(0, 0))
    box(s, bx, y + 0.32, BW, BH2 - 0.32, fill=PALE, line=None, shape=MSO_SHAPE.RECTANGLE)
    disc(bx + 0.28, y + 0.32 + (BH2 - 0.32) / 2, 0.40, key)
    tb(s, bx + 0.52, y + 0.32, BW - 0.56, BH2 - 0.32, items, 9.5, bullet="•", margins=(0, 0.02), space_after=3)
y += BH2 + 0.12

EH = PB - 0.12 - y
box(s, x0, y, w0, EH, fill=PALE, line=None, shape=MSO_SHAPE.RECTANGLE)
disc(x0 + 0.32, y + EH / 2, 0.38, "chk", fill=NAVY)
tb(s, x0 + 0.58, y, w0 - 0.62, EH,
   [[("직접 전환 경험", F4, PINK), (" 기반, 가장 빠르고 정확한 시스템 관리", F4, NAVY)]], 12, margins=(0, 0))

# ---------- 오른쪽 ----------
x0, w0 = RX + 0.15, PW - 0.30
NH = 0.30
Y0, Y1 = PY + PH + 0.14, PB - 0.08 - NH - 0.06
G = 0.10
rh = (Y1 - Y0 - 3 * G) / 4
LW = 1.62
cards = [("box", "컨테이너·\v오케스트레이션 관리", ["서비스별 자원 기준값·최대치 관리,\v오토스케일링·헬스체크",
                                              "컨테이너 이미지 버전 관리로\v구성 변경 이력 추적"]),
         ("cicd", "CI/CD 기반\v무중단 배포", ["변경된 컨테이너만 빌드·배포하여\v영향 범위 최소화",
                                         "서비스 단위 독립 배포,\v이상 시 이전 버전 즉시 롤백"]),
         ("mon", "통합 모니터링·\v장애 대응", ["API 게이트웨이 측정치·로그 수집/분석으로\v이상 조기 감지",
                                          "Circuit Breaker로 장애 전파 차단,\v서비스별 격리 조치"]),
         ("gauge", "부하·자원 관리", ["WEB/WAS VM Scale-out,\vDB Scale-up 기준 운영",
                                   "성능 기준선(Baseline) 관리 및\v자원 사용률 정기 점검"])]
for k, (key, head, items) in enumerate(cards):
    y = Y0 + k * (rh + G)
    box(s, x0, y, LW, rh, fill=BLUE, line=None, shape=MSO_SHAPE.RECTANGLE)
    add_pic(s, ic[key], x0 + LW / 2, y + 0.24, 0.26)
    tb(s, x0, y + 0.40, LW, rh - 0.44, [head], 11, color=WHITE, font=F4, align=PP_ALIGN.CENTER,
       margins=(0.03, 0), line_spacing=1.0)
    tb(s, x0 + LW, y, w0 - LW, rh, items, 10, fill=PALE, line=None, bullet="•", margins=(0.14, 0.03),
       space_after=4)

tb(s, x0, PB - 0.08 - NH, w0, NH, ["※ 전환 사업(1·2차) 설계 기준 유지, 청주시 클라우드 인프라 운영 기준과 연계 관리"], 9,
   color=GRAY, margins=(0, 0), wrap=False)

p.save(dst)
print("saved", dst)
