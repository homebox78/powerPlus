# -*- coding: utf-8 -*-
"""28쪽: 산출물·보안진단도구·프로세스 설명글 15% 키우고 줄간격 넓힘,
보안점검 단계 아래 설명 4칸은 글자 10% 키우고 줄간격 15% 넓힘. 인자: <src> <dst>"""
import sys
from pptx import Presentation
from pptx.util import Pt
from pptx.enum.text import MSO_ANCHOR
from pptx.oxml.ns import qn
sys.stdout.reconfigure(encoding="utf-8")
src, dst = sys.argv[1:3]
prs = Presentation(src)
sl = prs.slides[27]
by = {s.name: s for s in sl.shapes}


def bump(s, k, sp):
    tf = s.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for p in tf.paragraphs:
        for r in p.runs:
            if r.font.size:
                r.font.size = Pt(round(r.font.size.pt * k * 2) / 2)
        pPr = p._p.get_or_add_pPr()
        ln = pPr.find(qn("a:lnSpc"))
        cur = 100000
        if ln is not None and ln.find(qn("a:spcPct")) is not None:
            cur = int(ln.find(qn("a:spcPct")).get("val"))
            pPr.remove(ln)
        ln = pPr.makeelement(qn("a:lnSpc"), {})
        ln.append(ln.makeelement(qn("a:spcPct"), {"val": str(int(cur * sp))}))
        pPr.insert(0, ln)


for nm in ("시안 25", "시안 28", "시안 31"):
    bump(by[nm], 1.15, 1.2)
for nm in ("시안 43", "시안 45", "시안 47", "시안 49"):
    bump(by[nm], 1.10, 1.15)
# '수 행 역 량 확 보' 알약: 작은 띠 규칙(13pt 이하·a시월구일3·굵게 끔)
sys.path.insert(0, __import__("os").path.dirname(__import__("os").path.dirname(__import__("os").path.abspath(__file__))))
from group_title_lib import walk
for s, *_ in walk(sl.shapes):
    if getattr(s, "has_text_frame", False) and s.text_frame.text.replace(" ", "").strip() == "수행역량확보":
        for p in s.text_frame.paragraphs:
            for r in p.runs:
                r.font.size = Pt(12); r.font.bold = False
                rPr = r._r.get_or_add_rPr()
                for tg in ("a:latin", "a:ea"):
                    e = rPr.find(qn(tg))
                    if e is None:
                        e = rPr.makeelement(qn(tg), {}); rPr.append(e)
                    e.set("typeface", "a시월구일3")
        print("수행역량확보 12pt")
print("28쪽 설명글 7칸")
prs.save(dst)
