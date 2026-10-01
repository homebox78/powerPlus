# -*- coding: utf-8 -*-
"""17~19쪽 복원 스크립트 공용 도우미."""
import copy
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

I = Inches
NAVY = RGBColor(0x14, 0x3A, 0x69)
BLUE = RGBColor(0x2F, 0x78, 0xE0)
LIGHT = RGBColor(0xD0, 0xE6, 0xFA)
PALE = RGBColor(0xE8, 0xF2, 0xFC)
INK = RGBColor(0x1F, 0x4E, 0x79)
TXT = RGBColor(0x33, 0x33, 0x33)
GRAY = RGBColor(0x5F, 0x74, 0x96)
PINK = RGBColor(0xEC, 0x1C, 0x68)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
F2, F3, F4 = "a시월구일2", "a시월구일3", "a시월구일4"


def face(r, f):
    rPr = r._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea"):
        e = rPr.find(qn(tag))
        if e is None:
            e = rPr.makeelement(qn(tag), {})
            rPr.append(e)
        e.set("typeface", f)


def fill_tf(sh, paras, size, color=TXT, font=F2, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
            bullet=None, bullet_color=BLUE, space_after=0, line_spacing=None, wrap=True, margins=(0.02, 0.0)):
    """paras: [str | [(text, font, color), ...]]"""
    tf = sh.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = I(margins[0])
    tf.margin_top = tf.margin_bottom = I(margins[1])
    bp = tf._txBody.find(qn("a:bodyPr"))
    for c in list(bp):
        if c.tag in (qn("a:spAutoFit"), qn("a:normAutofit"), qn("a:noAutofit")):
            bp.remove(c)
    for pg in list(tf.paragraphs)[1:]:
        pg._p.getparent().remove(pg._p)
    p0 = tf.paragraphs[0]
    for r in list(p0._p):
        if r.tag in (qn("a:r"), qn("a:br"), qn("a:fld")):
            p0._p.remove(r)
    for i, para in enumerate(paras):
        pg = p0 if i == 0 else tf.add_paragraph()
        pg.alignment = align
        if line_spacing:
            pg.line_spacing = line_spacing
        if space_after:
            pg.space_after = Pt(space_after)
        segs = [(para, font, color)] if isinstance(para, str) else para
        if bullet:
            pPr = pg._p.get_or_add_pPr()
            ind = int(Pt(size) * 0.95)
            pPr.set("marL", str(ind)); pPr.set("indent", str(-ind))
            for t in ("a:buClr", "a:buFont", "a:buChar", "a:buNone"):
                for e in pPr.findall(qn(t)):
                    pPr.remove(e)
            bc = pPr.makeelement(qn("a:buClr"), {}); sc = bc.makeelement(qn("a:srgbClr"), {"val": str(bullet_color)})
            bc.append(sc); pPr.append(bc)
            pPr.append(pPr.makeelement(qn("a:buFont"), {"typeface": "Arial"}))
            pPr.append(pPr.makeelement(qn("a:buChar"), {"char": bullet}))
        for t, f, c in segs:
            for j, part in enumerate(t.split("")):
                if j:
                    pg.add_line_break()
                if not part:
                    continue
                r = pg.add_run(); r.text = part
                r.font.size = Pt(size); r.font.bold = False; r.font.italic = False
                r.font.color.rgb = c; face(r, f)
    return sh


def textbox(s, x, y, w, h, paras, size, **kw):
    tb = s.shapes.add_textbox(I(x), I(y), I(w), I(h))
    return fill_tf(tb, paras, size, **kw)


def box(s, x, y, w, h, fill=WHITE, line=LIGHT, lw=0.75, shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.08):
    b = s.shapes.add_shape(shape, I(x), I(y), I(w), I(h))
    if fill is None:
        b.fill.background()
    else:
        b.fill.solid(); b.fill.fore_color.rgb = fill
    if line is None:
        b.line.fill.background()
    else:
        b.line.color.rgb = line; b.line.width = Pt(lw)
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        b.adjustments[0] = adj
    b.shadow.inherit = False
    if b.has_text_frame:
        b.text_frame.text = ""
    return b


def hline(s, x1, y, x2, color=LIGHT, w=0.75):
    c = s.shapes.add_connector(1, I(x1), I(y), I(x2), I(y))
    c.line.color.rgb = color; c.line.width = Pt(w)
    return c


def by_id(s):
    return {sh.shape_id: sh for sh in s.shapes}


def kill(s, ids):
    by = by_id(s)
    for i in ids:
        if i in by:
            by[i]._element.getparent().remove(by[i]._element)


def place(sh, x=None, y=None, w=None, h=None):
    if x is not None: sh.left = I(x)
    if y is not None: sh.top = I(y)
    if w is not None: sh.width = I(w)
    if h is not None: sh.height = I(h)


def fit_pic(sh, cx, cy, h=None, w=None):
    """비율 유지 크기 조정(중심 기준)."""
    ar = sh.width / sh.height
    if h is not None:
        nh = I(h); nw = int(nh * ar)
    else:
        nw = I(w); nh = int(nw / ar)
    sh.width, sh.height = nw, nh
    sh.left = I(cx) - nw // 2; sh.top = I(cy) - nh // 2


def to_front(s, sh):
    tree = s.shapes._spTree
    tree.remove(sh._element); tree.append(sh._element)


def nlines(text, width_in, size_pt, kfactor=1.0):
    """대략 줄 수 추정: 한글 1, 영문·숫자·공백 0.55 폭."""
    em = size_pt / 72 * kfactor
    w = sum(em if ord(ch) > 0x2E80 else em * 0.55 for ch in text)
    import math
    return max(1, math.ceil(w / (width_in * 0.97)))
