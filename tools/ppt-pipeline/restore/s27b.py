# -*- coding: utf-8 -*-
"""27쪽: 왼쪽 성능 최적화 절차 아이콘 6개 20% 축소(중심 유지),
'추가 제안' 금메달 스티커 → 장표 톤의 분홍 알약 태그(머리 띠 오른쪽 위 모서리에 걸침). 인자: <src> <dst>"""
import sys
from pptx import Presentation
from pptx.util import Pt, Emu
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
sys.stdout.reconfigure(encoding="utf-8")
src, dst = sys.argv[1:3]
IN = 914400
prs = Presentation(src)
sl = prs.slides[26]
by = {s.name: s for s in sl.shapes}
for nm in ("Picture 231", "Picture 232", "Picture 233", "Picture 234", "Picture 235", "Picture 236"):
    p = by[nm]
    cx, cy = p.left + p.width / 2, p.top + p.height / 2
    p.width, p.height = int(p.width * 0.8), int(p.height * 0.8)
    p.left, p.top = int(cx - p.width / 2), int(cy - p.height / 2)

# 스티커 그룹 찾기·삭제
badge = next(s for s in sl.shapes if s.shape_type == 6 and "추가" in "".join(
    x.text_frame.text for x in s.shapes if getattr(x, "has_text_frame", False)))
bx = badge.left + badge.width
badge._element.getparent().remove(badge._element)
bar = next(s for s in sl.shapes if getattr(s, "has_text_frame", False) and s.text_frame.text.strip() == "변경부서 데이터 이관 기능 구축")
w, h = int(0.92 * IN), int(0.28 * IN)
tag = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, bar.left + bar.width - w - int(0.10 * IN), bar.top - h // 2, w, h)
tag.name = "추가 제안 태그"
tag.adjustments[0] = 0.5
tag.fill.solid(); tag.fill.fore_color.rgb = RGBColor(0xEC, 0x1C, 0x68)
tag.line.color.rgb = RGBColor(0xFF, 0xFF, 0xFF); tag.line.width = Pt(1.5)
tag.shadow.inherit = False
tf = tag.text_frame
tf.margin_left = tf.margin_right = Emu(0); tf.margin_top = tf.margin_bottom = Emu(0)
tf.vertical_anchor = MSO_ANCHOR.MIDDLE
tf.word_wrap = False
p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
r = p.add_run(); r.text = "추가 제안"
r.font.size = Pt(11); r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
rPr = r._r.get_or_add_rPr(); rPr.set("b", "0")
for tg in ("a:latin", "a:ea"):
    e = rPr.makeelement(qn(tg), {"typeface": "a시월구일3"}); rPr.append(e)
print("27쪽 아이콘 6·추가 제안 태그")
prs.save(dst)
