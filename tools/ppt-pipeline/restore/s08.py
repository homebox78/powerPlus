# -*- coding: utf-8 -*-
"""8쪽(사업의 특징 및 고려사항) 원본 문구·위계 복원.
원본: 줄마다 작은 도입문 + 큰 헤드라인("…하는데…") + 강조 낱말이 든 설명 한 줄.
v0.76: 고민/고려사항 지어낸 문구 → 삭제하고 원본 3단 위계로. 일러스트 28% 축소, 행을 왼쪽으로 넓힘.
인자: <src> <dst>"""
import sys
from pptx import Presentation
sys.stdout.reconfigure(encoding="utf-8")
"""restore 스크립트 공용: 글자 넣기(런별 크기·색·서체)."""
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

I = Inches
NAVY, BLUE, INK, GRAY, PINK = (RGBColor.from_string(h) for h in ("143A69", "2F78E0", "1F4E79", "5F7496", "EC1C68"))
F2, F3, F4 = "a시월구일2", "a시월구일3", "a시월구일4"


def face(r, f):
    rPr = r._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea"):
        e = rPr.find(qn(tag))
        if e is None:
            e = rPr.makeelement(qn(tag), {}); rPr.append(e)
        e.set("typeface", f)


def fill(sh, paras, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, margin=0.03, spacing=1.0, gap=0):
    """paras: [[(text, size, color, font), ...], ...]"""
    tf = sh.text_frame; tf.word_wrap = True; tf.auto_size = None
    tf.margin_left = tf.margin_right = I(margin); tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for pg in list(tf.paragraphs)[1:]:
        pg._p.getparent().remove(pg._p)
    p0 = tf.paragraphs[0]
    for r in list(p0._p):
        if r.tag in (qn("a:r"), qn("a:br"), qn("a:fld")):
            p0._p.remove(r)
    for i, runs in enumerate(paras):
        pg = p0 if i == 0 else tf.add_paragraph()
        pg.alignment = align; pg.line_spacing = spacing
        if gap and i: pg.space_before = Pt(gap)
        for text, size, color, f in runs:
            r = pg.add_run(); r.text = text
            r.font.size = Pt(size); r.font.color.rgb = color; r.font.bold = False; face(r, f)
    return sh


def box(sh, x, y, w, h):
    sh.left, sh.top, sh.width, sh.height = I(x), I(y), I(w), I(h)


src, dst = sys.argv[1:3]
p = Presentation(src); s = p.slides[7]
by = {sh.shape_id: sh for sh in s.shapes}


def kill(*ids):
    for i in ids:
        if i in by:
            by[i]._element.getparent().remove(by[i]._element)


ROWS = [
    (47, ["굿모닝 시스템에 장애가 발생하여 업무가 지연되지 않도록", "관문 시스템으로서의 안정성을 확보해야 하는데…"],
     [("시스템 간 연계를 통한 ", 0), ("정보의 공동이용까지 포함", 1)]),
    (61, ["특성이 다른 사용자 층의", "지속적이고 다양한 요구사항을 해결해야 하는데…"],
     [("업무 ", 0), ("효율성", 1), (" 및 사용자 ", 0), ("편리성", 1), (" 중심의 기능개선 필요", 0)]),
    (75, ["법/제도/업무/기술 등 시스템을 둘러싼", "주변 환경의 변화에 적절한 대응이 필요한데…"],
     [("기술 및 시스템적 ", 0), ("개선 방안의 도출 및 적용", 1), (" 필요", 0)]),
    (89, ["장기계속계약사업으로", "사용자와 유기적 관계가 형성되어야 하는데…"],
     [("객관적 성능 검증", 1), (" 및 ", 0), ("완벽한 보안관리체계 수립", 1), (" 필요", 0)]),
]
DX = 1.2
for base, (lead, head), sub in ROWS:
    if base + 5 not in by:          # 이미 적용됨
        continue
    band, strip, hexa, icon, label = (by[base + k] for k in range(5))
    for sh in (hexa, icon, label):
        sh.left = sh.left - I(DX)
    for sh in (band, strip):
        sh.left = sh.left - I(DX); sh.width = sh.width + I(DX)
    kill(*(base + k for k in (5, 6, 7, 9, 10, 11, 12)))
    tx = band.left + I(0.6)
    t = by[base + 8]
    t.left, t.top, t.width, t.height = tx, band.top + I(0.03), I(5.6), strip.top - band.top - I(0.03)
    fill(t, [[(lead, 10, INK, F2)], [(head, 15, NAVY, F4)]], anchor=MSO_ANCHOR.MIDDLE, spacing=0.92, gap=2)
    t2 = by[base + 13]
    t2.left, t2.top, t2.width, t2.height = tx, strip.top, I(5.4), strip.height
    fill(t2, [[(txt, 10.5, PINK if em else INK, F4 if em else F2) for txt, em in sub]], anchor=MSO_ANCHOR.MIDDLE)

# 일러스트 28% 축소 — 내용이 주인공
pic = by.get(104)
if pic is not None and pic.width > I(3.0):
    w, h = int(pic.width * 0.72), int(pic.height * 0.72)
    pic.width, pic.height = w, h
    pic.left = I(0.33); pic.top = I(4.35) - h // 2

p.save(dst)
print("s08 ok")
