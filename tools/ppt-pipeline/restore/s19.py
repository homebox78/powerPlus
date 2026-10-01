# -*- coding: utf-8 -*-
"""19쪽(유지보수 관리 체계 3/5) 원본 문구 복원. 인자: <src> <dst>
- 1·2차 전환: '기반 전환' / '단위업무 전환' 소제목 복원, 글자 확대
- 전환인력·설계 산출물: 원문 문장, 일러스트 축소·잘린 글자 조각 제거
- 오른쪽: '클라우드 환경 운영관리 방안' 머리 복원, 4개 카드 원문 문장 2개씩(키워드 나열 → 원문)
- 볼록한 3D 화살표가 든 그림은 화살표 부분을 잘라냄(구름 화살표·부하 화살표·CI/CD 회전 화살표)
- 하단 주석 '※ 전환 사업(1·2차) 설계 기준 유지, …' 원문"""
import os, sys, copy
sys.path.insert(0, os.path.dirname(__file__))
from pptx import Presentation
from pptx.util import Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from lib_e import *

sys.stdout.reconfigure(encoding="utf-8")
src, dst = sys.argv[1:3]
p = Presentation(src); s = p.slides[18]
by = by_id(s)

# ---------- 왼쪽 ----------
for i, sz in ((85, 12), (86, 11)):
    for r in by[i].text_frame.paragraphs[0].runs:
        r.font.size = Pt(sz)
by[86].width = I(4.0)
# 사업 이력 상자: 그림(화살표 구름 잘라냄) 축소
pic = by[88]
pic.crop_top = 0.40
fit_pic(pic, 5.45, 3.55, h=1.62)
t = by[89]; place(t, w=3.40)
for r in t.text_frame.paragraphs[0].runs:
    r.font.size = Pt(8.5)

kill(s, [93, 94, 95, 96, 99, 100, 101, 102])
rows = [(92, 91, 2.98, "기반 전환", ["클라우드 아키텍처 구성", "공통서비스 전환(인증·파일·연계·결재·기안문)"]),
        (98, 97, 3.62, "단위업무 전환", ["단위업무 91종 일괄 전환(공통 54·개별 37)", "신·구 시스템 병행운영 완료"])]
for badge, dot, y, head, items in rows:
    b = by[badge]; b.top = I(y + 0.05)
    by[dot].top = I(y + 0.22)
    textbox(s, 1.33, y - 0.02, 2.95, 0.24, [head], 9.5, color=NAVY, font=F4)
    textbox(s, 1.33, y + 0.21, 2.95, 0.36, items, 8.5, bullet="•", line_spacing=1.0)
ln = by[90]; place(ln, y=3.24, h=0.63)
place(by[103], y=3.585)
place(by[104], y=4.21)
place(by[105], y=4.26)
t = by[106]; place(t, y=4.25, w=3.40)
for r in t.text_frame.paragraphs[0].runs:
    r.font.size = Pt(8)
place(by[87], y=2.47, h=2.18)

# 아래 두 상자
for i in (115, 124):
    for r in by[i].text_frame.paragraphs[0].runs:
        r.font.size = Pt(9)
kill(s, [111, 112, 117, 119, 121, 126, 127, 128, 130, 132])
p109 = by[109]; p109.crop_left = 0.22
fit_pic(p109, 3.06, 5.90, h=0.60)
p110 = by[110]; p110.crop_left = 0.21
fit_pic(p110, 6.02, 5.85, h=0.66)


def items_col(x, y0, w, pics, items):
    y = y0
    for pc, it in zip(pics, items):
        n = 1 + it.count("\v")
        by[pc].left = I(x); by[pc].top = I(y + 0.035)
        textbox(s, x + 0.16, y - 0.01, w, n * 0.155 + 0.03, [it], 8.5, line_spacing=1.0)
        y += n * 0.155 + 0.08


items_col(0.55, 5.14, 2.05, [116, 118, 120],
          ["전환 사업 참여 인력을\v본 유지관리에 투입",
           "서비스 경계·API·연계 구조를\v설계 단계부터 보유",
           "별도 인수인계 없이 즉시 안정 운영"])
items_col(3.58, 5.14, 2.05, [125, 129, 131],
          ["MSA 서비스 분해도,\vAPI 게이트웨이 라우팅 정보",
           "컨테이너 구성 정보 및\v배포 스크립트",
           "30여 종 연계 인터페이스\v규격서·모듈 소스"])
place(by[107], h=1.52); place(by[108], h=1.52)

# ---------- 오른쪽 ----------
hdr = copy.deepcopy(by[83]._element); s.shapes._spTree.append(hdr)
hb = by_id(s)
newhdr = [sh for sh in s.shapes if sh._element is hdr][0]
place(newhdr, x=6.71, y=1.66, w=3.82, h=0.34)
ic = copy.deepcopy(by[84]._element); s.shapes._spTree.append(ic)
newic = [sh for sh in s.shapes if sh._element is ic][0]
place(newic, x=6.83, y=1.71)
ht = copy.deepcopy(by[85]._element); s.shapes._spTree.append(ht)
newht = [sh for sh in s.shapes if sh._element is ht][0]
place(newht, x=7.22, y=1.66, w=3.2)
newht.text_frame.paragraphs[0].runs[0].text = "클라우드 환경 운영관리 방안"
for r in newht.text_frame.paragraphs[0].runs[1:]:
    r._r.getparent().remove(r._r)

kill(s, [138] + list(range(141, 147)) + list(range(151, 157)) + list(range(161, 167)) + list(range(171, 178)))
cards = [(136, 139, 140, 137, {"crop_left": 0.08},
          ["서비스별 자원 기준값·최대치 관리,\v오토스케일링·헬스체크",
           "컨테이너 이미지 버전 관리로\v구성 변경 이력 추적"]),
         (147, 149, 150, 148, {"crop_right": 0.25},
          ["변경된 컨테이너만 빌드·배포하여\v영향 범위 최소화",
           "서비스 단위 독립 배포,\v이상 시 이전 버전 즉시 롤백"]),
         (157, 159, 160, 158, {},
          ["API 게이트웨이 측정치·로그 수집/분석으로\v이상 조기 감지",
           "Circuit Breaker로 장애 전파 차단,\v서비스별 격리 조치"]),
         (167, 169, 170, 168, {"crop_right": 0.44},
          ["WEB/WAS VM Scale-out,\vDB Scale-up 기준 운영",
           "성능 기준선(Baseline) 관리 및\v자원 사용률 정기 점검"])]
CY0, CY1, G = 2.06, 6.66, 0.08
ch = (CY1 - CY0 - 3 * G) / 4
for k, (card, num, title, pic, crop, items) in enumerate(cards):
    y = CY0 + k * (ch + G)
    place(by[card], y=y, h=ch)
    place(by[num], y=y + 0.07)
    place(by[title], x=7.22, y=y + 0.10, w=2.5)
    for sh in (by[num], by[title]):
        for r in sh.text_frame.paragraphs[0].runs:
            r.font.size = Pt(10.5)
    pc = by[pic]
    for a, v in crop.items():
        setattr(pc, a, v)
    fit_pic(pc, 10.02, y + ch / 2 + 0.05, w=0.86)
    if pc.height > I(0.80):
        fit_pic(pc, 10.02, y + ch / 2 + 0.05, h=0.80)
    textbox(s, 6.86, y + 0.39, 3.05, ch - 0.41, items, 8.5, bullet="•", space_after=2, line_spacing=1.0)

note = by[178]; place(note, x=6.71, y=6.73, w=3.82, h=0.32)
fill_tf(note, ["※ 전환 사업(1·2차) 설계 기준 유지, 청주시 클라우드 인프라\v   운영 기준과 연계 관리"], 8, color=GRAY, font=F2,
        anchor=MSO_ANCHOR.TOP, line_spacing=1.0)

p.save(dst)
print("saved", dst)
