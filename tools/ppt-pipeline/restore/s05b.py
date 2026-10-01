# -*- coding: utf-8 -*-
"""5쪽: '틈새업무 포털'이 왼쪽 묶음 테두리에 닿음 → 같은 줄 '이력 중심'과 함께 줄여 여유 확보.
인자: <src> <dst>"""
import sys
from pptx import Presentation
from pptx.util import Pt
sys.stdout.reconfigure(encoding="utf-8")
src, dst = sys.argv[1:3]
prs = Presentation(src)
sl = prs.slides[4]


def walk(shs):
    for s in shs:
        if s.shape_type == 6:
            yield from walk(s.shapes)
        else:
            yield s


n = 0
for s in walk(sl.shapes):
    if s.has_text_frame and s.text_frame.text.strip() in ("틈새업무 포털", "이력 중심"):
        for p in s.text_frame.paragraphs:
            for r in p.runs:
                if r.font.size and r.font.size.pt > 15.5:
                    r.font.size = Pt(15.5)
        n += 1
print("5쪽 제목", n)
prs.save(dst)
