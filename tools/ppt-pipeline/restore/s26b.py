# -*- coding: utf-8 -*-
"""26쪽: 원본의 모으는 괄호(위 두 상자 → 차이 분석 → 범위 규명 → 요건 도출 → 단계) 3개 복원,
양옆 일러스트 20% 축소·같은 높이·좌우 대칭. 인자: <src> <dst> (cwd D:\\powerPlus)"""
import copy, sys
from pptx import Presentation
from pptx.oxml.ns import qn
sys.stdout.reconfigure(encoding="utf-8")
src, dst = sys.argv[1:3]
IN = 914400
ORIG = "제안서/청주시_발표자료.원본백업.pptx"
prs = Presentation(src)
sl = prs.slides[25]
by = {s.name: s for s in sl.shapes}
osl = Presentation(ORIG).slides[25]
oby = {s.shape_id: s for s in osl.shapes}

# 1) 일러스트
gap = by["시안 12"]
cy = gap.top + gap.height / 2
cx = 10.833 * IN / 2
L, R = by["시안 그림 10"], by["시안 그림 11"]
dl = cx - (L.left + L.width / 2)
dr = (R.left + R.width / 2) - cx
d = (dl + dr) / 2
for p, side in ((L, -1), (R, 1)):
    w, h = int(p.width * 0.8), int(p.height * 0.8)
    p.width, p.height = w, h
    p.left = int(cx + side * d - w / 2)
    p.top = int(cy - h / 2)

# 2) 괄호: 원본 도형(같은 모양·그라데이션·뒤집기)을 복제해 새 자리로
def brace(oid, ccx, ccy, length, thick):
    el = copy.deepcopy(oby[oid]._element)
    x = el.find(qn("p:spPr")).find(qn("a:xfrm"))
    # 90° 회전 도형: ext 의 cy 가 화면 가로 길이
    x.find(qn("a:ext")).set("cx", str(int(thick)))
    x.find(qn("a:ext")).set("cy", str(int(length)))
    x.find(qn("a:off")).set("x", str(int(ccx - thick / 2)))
    x.find(qn("a:off")).set("y", str(int(ccy - length / 2)))
    for c in el.iter(qn("a:srgbClr")):   # 원본 색은 옅은 바탕 위에서 거의 안 보임 → 한 단계 진하게
        if c.get("val") == "D0E3F4":
            c.set("val", "AFC8E8")
    ln = el.find(qn("p:spPr")).find(qn("a:ln"))
    ln.set("w", "9525")
    ln.find(qn("a:solidFill"))[0].set("val", "8FB0DA")
    nv = el.find(qn("p:nvSpPr")).find(qn("p:cNvPr"))
    nv.set("id", str(9000 + oid)); nv.set("name", "괄호 %d" % oid)
    return el

top_bot = by["시안 4"].top + by["시안 4"].height
anchor = by["시안 바탕 2"]._element
els = [
    brace(125, cx, (top_bot + L.top) / 2 + 0.02 * IN, 5.3 * IN, 0.26 * IN),        # 위 두 상자 → 차이 분석
    brace(221, cx, by["시안 15"].top - 0.08 * IN, 5.2 * IN, 0.22 * IN),           # 차이 분석 → 범위 규명(펼침)
    brace(311, cx, (by["시안 21"].top + by["시안 21"].height + by["시안 24"].top) / 2, 5.3 * IN, 0.2 * IN),  # 요건 도출 → 단계
]
for el in reversed(els):
    anchor.addnext(el)
# 원본에 없던 직선·삼각 화살표(괄호와 겹침)는 지움
for nm in ("시안 선 13", "시안 14", "시안 선 22", "시안 23"):
    e = by[nm]._element
    e.getparent().remove(e)
print("26쪽 괄호 3·일러스트 2")
prs.save(dst)
