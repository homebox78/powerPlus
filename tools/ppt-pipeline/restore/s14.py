# -*- coding: utf-8 -*-
"""14쪽(4. 변화 실천) 원본 문구 복원.
- '전략 3' 배지를 ACT.5 알약 왼쪽에 추가(알약·라벨은 오른쪽으로 밀기)
- 사용자 업무 환경 분석 체계: 원본 설명 문장 복원, 제목·설명 글자 확대, 일러스트 축소
- 변경 관리 체계 STEP 1~6: 원본 세부 항목 복원, 글자 확대, 일러스트 축소
인자: <src> <dst>  (14쪽만 손댄다)
"""
import sys
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

sys.stdout.reconfigure(encoding="utf-8")
src, dst = sys.argv[1:3]
p = Presentation(src)
s = p.slides[13]
by = {sh.shape_id: sh for sh in s.shapes}
I = Inches
TXT, INK, NAVY = RGBColor(0x33, 0x33, 0x33), RGBColor(0x1F, 0x4E, 0x79), RGBColor(0x14, 0x3A, 0x69)
WHITE, RED = RGBColor(0xFF, 0xFF, 0xFF), RGBColor(0xC0, 0x00, 0x00)
F2, F3, F4 = "a시월구일2", "a시월구일3", "a시월구일4"
VT = "\x0b"


def face(r, f):
    rPr = r._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea"):
        e = rPr.find(qn(tag))
        if e is None:
            e = rPr.makeelement(qn(tag), {})
            rPr.append(e)
        e.set("typeface", f)


def resize(sh, pt):
    for pg in sh.text_frame.paragraphs:
        for r in pg.runs:
            r.font.size = Pt(pt)


def clear(tf):
    for pg in list(tf.paragraphs)[1:]:
        pg._p.getparent().remove(pg._p)
    for r in list(tf.paragraphs[0].runs):
        r._r.getparent().remove(r._r)
    return tf.paragraphs[0]


def bullets(x, y, w, h, items, size, color=TXT, f=F2, gap=1.5):
    """items: 문자열 목록. 문자열 안의 \\x0b 는 같은 항목 안 줄바꿈(어절 경계)."""
    tb = s.shapes.add_textbox(I(x), I(y), I(w), I(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = tf.margin_right = I(0.02)
    tf.margin_top = tf.margin_bottom = 0
    p0 = clear(tf)
    ind = int(Pt(size) * 0.75)
    for i, it in enumerate(items):
        pg = p0 if i == 0 else tf.add_paragraph()
        pg.alignment = PP_ALIGN.LEFT
        pg.line_spacing = 1.05
        pg.space_after = Pt(gap)
        pPr = pg._p.get_or_add_pPr()
        pPr.set("marL", str(ind))
        pPr.set("indent", str(-ind))
        pPr.append(pPr.makeelement(qn("a:buFont"), {"typeface": "Arial"}))
        pPr.append(pPr.makeelement(qn("a:buChar"), {"char": "•"}))
        for k, seg in enumerate(it.split(VT)):
            if k:
                pg._p.append(pg._p.makeelement(qn("a:br"), {}))
            r = pg.add_run()
            r.text = seg
            r.font.size = Pt(size)
            r.font.color.rgb = color
            r.font.bold = False
            face(r, f)
    return tb


def kill(*ids):
    for i in ids:
        if i in by:
            by[i]._element.getparent().remove(by[i]._element)


def scale_pics(ids, factor, right=None, bottom=None, cx=None, cy=None):
    """여러 그림을 한 덩어리로 같은 비율 축소. right/bottom 또는 중심(cx,cy)으로 자리 잡기."""
    shs = [by[i] for i in ids]
    L = min(x.left for x in shs); T = min(x.top for x in shs)
    R = max(x.left + x.width for x in shs); B = max(x.top + x.height for x in shs)
    W, H = (R - L) * factor, (B - T) * factor
    nl = I(right) - W if right is not None else (I(cx) - W / 2)
    nt = I(bottom) - H if bottom is not None else (I(cy) - H / 2)
    for x in shs:
        x.left = int(nl + (x.left - L) * factor)
        x.top = int(nt + (x.top - T) * factor)
        x.width = int(x.width * factor)
        x.height = int(x.height * factor)


# ── 0) 전략 3 배지 ──────────────────────────────────────────────
act, lab, bar = by[112], by[113], by[111]
bw, gap = I(0.74), I(0.07)
badge = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, act.left, act.top, bw, act.height)
badge.adjustments[0] = 0.18
badge.fill.solid(); badge.fill.fore_color.rgb = RED
badge.line.fill.background()
badge.shadow.inherit = False
tf = badge.text_frame
tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
tf.vertical_anchor = MSO_ANCHOR.MIDDLE
tf.word_wrap = False
pg = clear(tf); pg.alignment = PP_ALIGN.CENTER
for t, sz in (("전략 ", 10), ("3", 15)):
    r = pg.add_run(); r.text = t; r.font.size = Pt(sz); r.font.color.rgb = WHITE; r.font.bold = False; face(r, F4)
badge.name = "전략 배지"
shift = bw + gap
for sh in (act, lab, bar):
    sh.left = sh.left + shift
# 배지를 ACT 알약과 같은 높이 줄 맨 앞에 두되 바탕 막대 위로
act._element.addprevious(badge._element)

# ── 1) 패널 머리 글자 확대 ───────────────────────────────────────
resize(by[117], 13); resize(by[121], 13)

# ── 2) 사용자 업무 환경 분석 체계 ─────────────────────────────────
cards = [  # (카드, 번호, 제목, 지울 설명, 그림들, 설명 항목, 그림 비율)
    (122, 123, 124, (125, 126), (127,), ["청주시청, 구청, 읍면동 등" + VT + "사용자 PC 사양, OS 버전 등 이해"], 0.70),
    (135, 136, 137, (138, 139), (140,), ["시스템 개발/실행 환경 이해", "UI 플랫폼에 대한 이해"], 0.72),
    (150, 151, 152, (153, 154), (155, 156), ["청주시 업무 시스템, 자치단체 표준" + VT + "시스템 등 연계시스템 환경 이해"], 0.70),
]
for card, num, title, olds, pics, items, fac in cards:
    c = by[card]
    ct = c.top / 914400.0; ch = c.height / 914400.0
    kill(*olds)
    t = by[title]
    t.left, t.top, t.width, t.height = I(0.91), I(ct + 0.16), I(2.25), I(0.30)
    t.text_frame.word_wrap = False
    resize(t, 10.5)
    for r in (rr for pg in t.text_frame.paragraphs for rr in pg.runs):
        r.font.color.rgb = NAVY; face(r, F4)
    n = by[num]
    n.top = I(ct + 0.16)
    bullets(0.93, ct + 0.54, 2.18, ch - 0.62, items, 9.5)
    scale_pics(pics, fac, right=3.96, bottom=ct + ch - 0.12)

# 오른쪽 칩 글자 7.5pt (여백 0)
for i in (131, 133, 145, 147, 149, 160, 162):
    sh = by[i]
    tf = sh.text_frame
    tf.margin_left = tf.margin_right = 0
    right = I(4.97)
    sh.width = right - sh.left
    resize(sh, 7.5)

# ── 3) 변경 관리 체계 ────────────────────────────────────────────
for i in (165, 168, 171):
    resize(by[i], 9)
steps = [  # (카드, STEP 알약, 설명칸, 그림, 항목)
    (172, 174, 175, 173, ["기능개선사항 발생 및" + VT + "형상관리 등록", "변경계획 수립"]),
    (176, 178, 179, 177, ["세부사전검토(범위," + VT + "영향분석, 사유, 기능 등)", "결과 보고"]),
    (180, 182, 183, 181, ["관련 업무담당자" + VT + "관리협의체 소집", "상세분석 및 타당성 검토"]),
    (184, 186, 187, 185, ["전체의견 수렴 및" + VT + "적용범위 설정", "적용계획 수립"]),
    (188, 190, 191, 189, ["기능개선작업(소스코드" + VT + "상세한 내역관리)", "통합테스트 및" + VT + "검증결과 반영"]),
    (192, 194, 195, 193, ["운영시스템 적용 및" + VT + "배포관리", "서비스 개시 및 안정화"]),
]
for card, pill, desc, pic, items in steps:
    c = by[card]
    cl = c.left / 914400.0; ct = c.top / 914400.0; cw = c.width / 914400.0; ch = c.height / 914400.0
    resize(by[pill], 8)
    kill(desc)
    bullets(cl + 0.07, ct + 0.25, cw - 0.12, 0.85, items, 9)
    pc = by[pic]
    fac = 0.44 / (pc.height / 914400.0)
    scale_pics((pic,), fac, right=cl + cw - 0.08, bottom=ct + ch - 0.06)

p.save(dst)
print("저장", dst)
