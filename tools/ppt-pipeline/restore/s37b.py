# -*- coding: utf-8 -*-
"""37쪽: 솔리데오 수상 목록 글자색 → 333F50 (글머리표 색 포함). 인자: <src> <dst>"""
import os, sys
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
sys.stdout.reconfigure(encoding="utf-8")
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from group_title_lib import walk
src, dst = sys.argv[1:3]
prs = Presentation(src)
n = 0
for s, *_ in walk(prs.slides[36].shapes):
    if getattr(s, "has_text_frame", False) and "운영지원단 업무혁신" in s.text_frame.text:
        for p in s.text_frame.paragraphs:
            for r in p.runs:
                r.font.color.rgb = RGBColor(0x33, 0x3F, 0x50); n += 1
            pPr = p._p.find(qn("a:pPr"))
            if pPr is not None:
                bc = pPr.find(qn("a:buClr"))
                if bc is not None:
                    for c in list(bc):
                        bc.remove(c)
                    bc.append(bc.makeelement(qn("a:srgbClr"), {"val": "333F50"}))
print("37쪽 수상 목록 글자", n)
prs.save(dst)
