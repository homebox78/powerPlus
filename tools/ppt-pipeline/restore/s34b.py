# -*- coding: utf-8 -*-
"""34쪽: '운영성과 보고 및 보고서' 머리 글이 오른쪽 99.9점 목표·예시 배지에 걸림
→ 글 상자를 배지 왼쪽까지만으로 줄여 남은 폭 가운데에. 인자: <src> <dst>"""
import os, sys
from pptx import Presentation
sys.stdout.reconfigure(encoding="utf-8")
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from group_title_lib import walk
src, dst = sys.argv[1:3]
IN = 914400
prs = Presentation(src)
sl = prs.slides[33]
items = list(walk(sl.shapes))
t = next(s for s, *_ in items if getattr(s, "has_text_frame", False) and s.text_frame.text.strip() == "운영성과 보고 및 보고서")
badge = next(s for s, *_ in items if getattr(s, "has_text_frame", False) and "99.9" in s.text_frame.text)
right = badge.left - int(0.08 * IN)
t.width = int(right - t.left)
print("34쪽 머리 글 폭", round(t.width / IN, 2))
prs.save(dst)
