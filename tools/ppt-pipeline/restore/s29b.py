# -*- coding: utf-8 -*-
"""29쪽: 위 두 카드 일러스트 20% 축소 + 옆 문구·교육 목록 글자 키움,
아래 교육대상/조직/형태·이전 범위/필수요소/방법 설명 15% 키우고 줄간격 10% 넓힘,
기술이전 양옆 일러스트·학사모 25% 축소. 인자: <src> <dst>"""
import sys
from pptx import Presentation
from pptx.util import Pt
from pptx.oxml.ns import qn
sys.stdout.reconfigure(encoding="utf-8")
src, dst = sys.argv[1:3]
prs = Presentation(src)
sl = prs.slides[28]
by = {s.name: s for s in sl.shapes}


def shrink(nm, k, anchor="center"):
    p = by[nm]
    cx, cy, bot = p.left + p.width / 2, p.top + p.height / 2, p.top + p.height
    p.width, p.height = int(p.width * k), int(p.height * k)
    p.left = int(cx - p.width / 2)
    p.top = int(bot - p.height) if anchor == "bottom" else int(cy - p.height / 2)


def font(nm, size=None, k=None, sp=None):
    tf = by[nm].text_frame
    for pg in tf.paragraphs:
        for r in pg.runs:
            if r.font.size:
                r.font.size = Pt(size if size else round(r.font.size.pt * k * 2) / 2)
        if sp:
            pPr = pg._p.get_or_add_pPr()
            ln = pPr.find(qn("a:lnSpc"))
            cur = 100000
            if ln is not None and ln.find(qn("a:spcPct")) is not None:
                cur = int(ln.find(qn("a:spcPct")).get("val"))
            if ln is not None:
                pPr.remove(ln)
            ln = pPr.makeelement(qn("a:lnSpc"), {})
            ln.append(ln.makeelement(qn("a:spcPct"), {"val": str(int(cur * sp))}))
            pPr.insert(0, ln)


# 위 카드
shrink("시안 그림 8", 0.8, "bottom"); shrink("시안 그림 23", 0.8, "bottom")
for nm, pic in (("시안 9", "시안 그림 8"), ("시안 24", "시안 그림 23")):
    t, p = by[nm], by[pic]
    right = t.left + t.width
    t.left = int(p.left + p.width + 0.03 * 914400)
    t.width = int(right - t.left)
    font(nm, size=10)
for nm in ("시안 14", "시안 16", "시안 18", "시안 20"):
    font(nm, size=8)
for nm in ("시안 27", "시안 29", "시안 31", "시안 33"):
    font(nm, size=8)
font("시안 25", size=10)
# 아래 설명
for nm in ("시안 39", "시안 43", "시안 47", "시안 59", "시안 62", "시안 65"):
    font(nm, k=1.15, sp=1.10)
# 기술이전 일러스트
for nm in ("시안 그림 53", "시안 그림 54", "시안 그림 56"):
    shrink(nm, 0.75)
print("29쪽")
prs.save(dst)
