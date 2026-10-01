# -*- coding: utf-8 -*-
"""14쪽(4. 변화 실천) 2차 수정 — v0.77 기준.
- 왼쪽 '사용자 업무 환경 분석 체계': 번호 원 1·2·3 삭제 → 원본처럼 남색 라벨 타일(흰 아이콘 + 사용자/시스템 환경/연계시스템)
  제목 13pt a시월구일4 남색, 설명 10.5pt, PC·모니터·서버 일러스트 축소, 오른쪽 작은 칩 패널 삭제
- 오른쪽 '변경 관리 체계': 3단 흐름 머리 글자 세로 가운데(아이콘 삭제), STEP 카드 글자 10pt·일러스트 축소, 카드 높이 통일
- 원본 문구 유지(원본 낱말 그대로), 어절 경계 줄바꿈
인자: <src> <dst> [원본.pptx]  (14쪽만 손댄다. 라벨 타일 아이콘은 원본 장표에서 가져온다)
"""
import copy, io, os, sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

sys.stdout.reconfigure(encoding="utf-8")
src, dst = sys.argv[1:3]
here = os.path.dirname(os.path.abspath(__file__))
orig_path = sys.argv[3] if len(sys.argv) > 3 else os.path.join(here, "..", "..", "..", "제안서", "청주시_발표자료.원본백업.pptx")
p = Presentation(src)
s = p.slides[13]
by = {sh.shape_id: sh for sh in s.shapes}
I, E = Inches, 914400.0
HEX = lambda h: RGBColor.from_string(h)
NAVY, BLUE, WHITE, TXT = HEX("143A69"), HEX("2F78E0"), HEX("FFFFFF"), HEX("333333")
F2, F3, F4 = "a시월구일2", "a시월구일3", "a시월구일4"


def face(r, f):
    rPr = r._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea"):
        e = rPr.find(qn(tag))
        if e is None:
            e = rPr.makeelement(qn(tag), {})
            rPr.append(e)
        e.set("typeface", f)


def clear(tf):
    for pg in list(tf.paragraphs)[1:]:
        pg._p.getparent().remove(pg._p)
    for r in list(tf.paragraphs[0].runs):
        r._r.getparent().remove(r._r)
    return tf.paragraphs[0]


def put_text(tf, text, size, color, font, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE, m=0.0):
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = I(m)
    tf.margin_top = tf.margin_bottom = 0
    pg = clear(tf)
    pg.alignment = align
    pg.line_spacing = 1.0
    for k, seg in enumerate(text.split("^")):
        if k:
            pg._p.append(pg._p.makeelement(qn("a:br"), {}))
        r = pg.add_run(); r.text = seg
        r.font.size = Pt(size); r.font.bold = False; r.font.color.rgb = color
        face(r, font)


def bullets(x, y, w, h, items, size, gap=2.5, name="복원설명"):
    """items: 문자열 목록. ^ = 같은 항목 안 어절 경계 줄바꿈."""
    tb = s.shapes.add_textbox(I(x), I(y), I(w), I(h))
    tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = tf.margin_right = I(0.01)
    tf.margin_top = tf.margin_bottom = 0
    p0 = clear(tf)
    ind = int(Pt(size) * 0.75)
    for i, it in enumerate(items):
        pg = p0 if i == 0 else tf.add_paragraph()
        pg.alignment = PP_ALIGN.LEFT
        pg.line_spacing = 1.05
        pg.space_after = Pt(gap)
        pPr = pg._p.get_or_add_pPr()
        pPr.set("marL", str(ind)); pPr.set("indent", str(-ind))
        bc = pPr.makeelement(qn("a:buClr"), {})
        bc.append(bc.makeelement(qn("a:srgbClr"), {"val": "5F7496"}))
        pPr.append(bc)
        pPr.append(pPr.makeelement(qn("a:buFont"), {"typeface": "Arial"}))
        pPr.append(pPr.makeelement(qn("a:buChar"), {"char": "•"}))
        for k, seg in enumerate(it.split("^")):
            if k:
                pg._p.append(pg._p.makeelement(qn("a:br"), {}))
            r = pg.add_run(); r.text = seg
            r.font.size = Pt(size); r.font.color.rgb = TXT; r.font.bold = False
            face(r, F2)
    return tb


def kill(*ids):
    for i in ids:
        if i in by:
            by[i]._element.getparent().remove(by[i]._element)


def fit(ids, x0, y0, x1, y1):
    """여러 그림을 한 덩어리로 같은 비율로 축소해 상자(x0,y0)-(x1,y1) 가운데에 둔다."""
    shs = [by[i] for i in ids]
    L = min(a.left for a in shs); T = min(a.top for a in shs)
    R = max(a.left + a.width for a in shs); B = max(a.top + a.height for a in shs)
    k = min(I(x1 - x0) / (R - L), I(y1 - y0) / (B - T))
    W, H = (R - L) * k, (B - T) * k
    nl, nt = I((x0 + x1) / 2) - W / 2, I((y0 + y1) / 2) - H / 2
    for a in shs:
        a.left = int(nl + (a.left - L) * k); a.top = int(nt + (a.top - T) * k)
        a.width = int(a.width * k); a.height = int(a.height * k)


def setbox(sh, x, y, w, h):
    sh.left, sh.top, sh.width, sh.height = I(x), I(y), I(w), I(h)


# ── 원본 장표의 라벨 타일 아이콘(흰색 처리된 그림 2 + 도형 1) ──────────
op = Presentation(orig_path)
os_ = op.slides[13]


def walk(c):
    for sh in c:
        if sh.shape_type == 6:
            yield from walk(sh.shapes)
        else:
            yield sh


oby = {sh.shape_id: sh for sh in walk(os_.shapes)}


def clone_icon(oid, x, y, size):
    o = oby[oid]
    el = copy.deepcopy(o._element)
    if o.shape_type == 13:   # 그림: 이미지 파트를 새로 연결
        _, rId = s.part.get_or_add_image_part(io.BytesIO(o.image.blob))
        el.find(".//" + qn("a:blip")).set(qn("r:embed"), rId)
    s.shapes._spTree.append(el)
    sh = s.shapes[-1]
    k = size / max(o.width, o.height) * E
    sh.width, sh.height = int(o.width * k), int(o.height * k)
    sh.left = int(I(x) - sh.width / 2); sh.top = int(I(y))
    sh.name = "라벨 아이콘"
    return sh


# ── 1) 왼쪽 패널: 사용자 업무 환경 분석 체계 ────────────────────────
kill(123, 136, 151, 124, 137, 152, 198, 199, 200,
     128, 129, 130, 131, 132, 133, 134,
     141, 142, 143, 144, 145, 146, 147, 148, 149,
     157, 158, 159, 160, 161, 162)
resize_head = [117, 121]
for i in resize_head:
    for pg in by[i].text_frame.paragraphs:
        for r in pg.runs:
            r.font.size = Pt(14)
        by[i].text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE

CH, CG, CT0 = 1.34, 0.12, 2.62
left = [  # (카드, 그림들, 타일 라벨, 아이콘(원본 id), 제목, 설명, 줄 수)
    (122, (127,), "사용자", 105, "사용자 PC 및 업무 환경 분석",
     ["청주시청, 구청, 읍면동 등^사용자 PC 사양, OS 버전 등 이해"], 2),
    (135, (140,), "시스템 환경", 101, "업무지원포털 시스템 환경 분석",
     ["시스템 개발/실행 환경 이해", "UI 플랫폼에 대한 이해"], 2),
    (150, (155, 156), "연계시스템", 99, "업무시스템 및 연계 환경 모니터링",
     ["청주시 업무 시스템, 자치단체 표준^시스템 등 연계시스템 환경 이해"], 2),
]
TX, TW, BW = 1.47, 3.50, 2.80     # 글 시작, 제목 폭, 설명 폭
for n, (card, pics, tlab, oid, title, items, nl) in enumerate(left):
    ct = CT0 + n * (CH + CG)
    c = by[card]
    setbox(c, 0.43, ct, 4.61, CH)
    # 라벨 타일
    tw, th = 0.82, 0.92
    tx, ty = 0.55, ct + (CH - th) / 2
    tile = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, I(tx), I(ty), I(tw), I(th))
    tile.adjustments[0] = 0.14
    tile.fill.solid(); tile.fill.fore_color.rgb = NAVY
    tile.line.fill.background()
    sp = tile._element.spPr
    sp.append(sp.makeelement(qn("a:effectLst"), {}))
    tile.name = "라벨 타일"
    put_text(tile.text_frame, tlab, 9, WHITE, F4, PP_ALIGN.CENTER, MSO_ANCHOR.BOTTOM, 0.02)
    tile.text_frame.margin_bottom = I(0.10)
    clone_icon(oid, tx + tw / 2, ty + 0.15, 0.34)
    # 제목 + 설명: 덩어리째 카드 세로 가운데
    bh = nl * 0.19 + (len(items) - 1) * 0.04
    block = 0.30 + 0.07 + bh
    by0 = ct + (CH - block) / 2
    t = s.shapes.add_textbox(I(TX), I(by0), I(TW), I(0.30))
    t.name = "항목 제목"
    put_text(t.text_frame, title, 13, NAVY, F4, PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE, 0.01)
    bullets(TX, by0 + 0.37, BW, bh + 0.05, items, 10.5)
    # 일러스트: 설명 오른쪽 작게
    fit(pics, 4.36, by0 + 0.30, 4.96, by0 + block)

# ── 2) 오른쪽 패널: 변경 관리 체계 ───────────────────────────────
kill(164, 167, 170, 201, 202, 203, 204, 205, 206)
for chev, txt in ((163, 165), (166, 168), (169, 171)):
    cv = by[chev]
    cv.top, cv.height = I(2.64), I(0.50)
    tb = by[txt]
    l = cv.left / E
    setbox(tb, l + 0.06, 2.64, cv.width / E - 0.26, 0.50)
    text = tb.text_frame.text
    put_text(tb.text_frame, text, 11, WHITE, F4, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, 0)

SH, RT = 1.73, (3.28, 5.13)
steps = [  # (카드, 그림, STEP 알약, 행, 항목)
    (172, 173, 174, 0, ["기능개선사항 발생^및 형상관리 등록", "변경계획 수립"]),
    (176, 177, 178, 0, ["세부사전검토^(범위, 영향분석,^사유, 기능 등)", "결과 보고"]),
    (180, 181, 182, 0, ["관련 업무담당자^관리협의체 소집", "상세분석 및^타당성 검토"]),
    (184, 185, 186, 1, ["전체의견 수렴 및^적용범위 설정", "적용계획 수립"]),
    (188, 189, 190, 1, ["기능개선작업^(소스코드 상세한^내역관리)", "통합테스트 및^검증결과 반영"]),
    (192, 193, 194, 1, ["운영시스템 적용^및 배포관리", "서비스 개시 및^안정화"]),
]
for card, pic, pill, row, items in steps:
    c = by[card]
    cl, cw = c.left / E, c.width / E
    ct = RT[row]
    c.top, c.height = I(ct), I(SH)
    pl = by[pill]
    setbox(pl, cl + 0.07, ct + 0.08, 0.68, 0.24)
    put_text(pl.text_frame, pl.text_frame.text, 9.5, WHITE, F4, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, 0)
    bullets(cl + 0.07, ct + 0.42, cw - 0.12, 0.95, items, 10, gap=3)
    fit((pic,), cl + cw - 0.66, ct + SH - 0.42, cl + cw - 0.08, ct + SH - 0.06)

p.save(dst)
print("저장", dst)
