# -*- coding: utf-8 -*-
"""10쪽(수행 전략 개요) 원본 문구·강조 복원.
- 핵심 성공 요소 5: 원본의 작은 도입문 + 큰 강조구(완벽한 이해 / 업무지식 및 기술역량 …)
- 3대 추진 전략: ACT.1~5 원본 설명 전문(두 줄), 글자 확대. 전략 그림은 머리줄로 축소.
- 왼쪽 열은 일러스트·육각 아이콘을 줄여 가운데 열을 넓힘.
인자: <src> <dst>"""
import sys
from pptx import Presentation
sys.stdout.reconfigure(encoding="utf-8")
"""restore 스크립트 공용: 글자 넣기(런별 크기·색·서체)."""
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

I = Inches
NAVY, BLUE, INK, GRAY, PINK = (RGBColor.from_string(h) for h in ("143A69", "2F78E0", "1F4E79", "5F7496", "EC1C68"))
F2, F3, F4 = "a시월구일2", "a시월구일3", "a시월구일4"


def face(r, f):
    rPr = r._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea"):
        e = rPr.find(qn(tag))
        if e is None:
            e = rPr.makeelement(qn(tag), {}); rPr.append(e)
        e.set("typeface", f)


def fill(sh, paras, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, margin=0.03, spacing=1.0, gap=0):
    """paras: [[(text, size, color, font), ...], ...]"""
    tf = sh.text_frame; tf.word_wrap = True; tf.auto_size = None
    tf.margin_left = tf.margin_right = I(margin); tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for pg in list(tf.paragraphs)[1:]:
        pg._p.getparent().remove(pg._p)
    p0 = tf.paragraphs[0]
    for r in list(p0._p):
        if r.tag in (qn("a:r"), qn("a:br"), qn("a:fld")):
            p0._p.remove(r)
    for i, runs in enumerate(paras):
        pg = p0 if i == 0 else tf.add_paragraph()
        pg.alignment = align; pg.line_spacing = spacing
        if gap and i: pg.space_before = Pt(gap)
        for text, size, color, f in runs:
            r = pg.add_run(); r.text = text
            r.font.size = Pt(size); r.font.color.rgb = color; r.font.bold = False; face(r, f)
    return sh


def box(sh, x, y, w, h):
    sh.left, sh.top, sh.width, sh.height = I(x), I(y), I(w), I(h)


src, dst = sys.argv[1:3]
p = Presentation(src); s = p.slides[9]
by = {sh.shape_id: sh for sh in s.shapes}
if 115 not in by:
    print("s10 already applied"); p.save(dst); sys.exit(0)


def kill(*ids):
    for i in ids:
        if i in by:
            by[i]._element.getparent().remove(by[i]._element)


def fit_pic(pic, x, y, w, h):
    """비율 유지하며 (x,y,w,h) 안 가운데."""
    r = min(I(w) / pic.width, I(h) / pic.height)
    pw, ph = int(pic.width * r), int(pic.height * r)
    pic.width, pic.height = pw, ph
    pic.left = I(x) + (I(w) - pw) // 2; pic.top = I(y) + (I(h) - ph) // 2


def vmid(sh, cy, h):
    sh.height = I(h); sh.top = I(cy) - I(h) // 2


# ── 왼쪽 열: 2.42in 로 좁힘, 일러스트 축소 ─────────────────────
LW = 2.42
box(by[83], 0.33, 2.11, LW, 4.78)
box(by[84], 0.33, 1.61, LW, 0.45)
box(by[86], 0.92, 1.61, LW - 0.65, 0.45)
fill(by[86], [[("고객의 고민, 걱정 …", 13, NAVY, F4)]], anchor=MSO_ANCHOR.MIDDLE)
person = by[103]
w, h = int(person.width * 0.70), int(person.height * 0.70)
person.width, person.height = w, h; person.left = I(0.12); person.top = I(6.86) - h
for hexa, lab, ln, dot in ((87, 88, 89, 90), (91, 92, 93, 94), (95, 96, 97, 98), (99, 100, 101, 102)):
    hx = by[hexa]; cy = (hx.top + hx.height // 2) / 914400
    fit_pic(hx, 1.18, cy - 0.32, 0.56, 0.64)
    L = by[lab]; box(L, 1.78, cy - 0.17, 0.98, 0.3)
    L.text_frame.word_wrap = False
    txt = L.text_frame.text
    fill(L, [[(txt, 11.5, NAVY, F4)]], anchor=MSO_ANCHOR.MIDDLE, margin=0.02)
    box(by[ln], 1.8, cy + 0.17, 0.94, 0)
    box(by[dot], 1.78, cy + 0.145, 0.05, 0.05)

# ── 가운데 열: 3.15~6.63 ─────────────────────────────────────
MX, MW = 3.15, 3.48
box(by[107], MX, 1.84, MW, 4.63)
box(by[108], MX - 0.05, 1.61, MW + 0.1, 0.45)
box(by[109], MX + 0.15, 1.68, 0.33, 0.31)
box(by[110], MX + 0.52, 1.61, MW - 0.6, 0.45)
fill(by[110], [[("핵심 성공 요소", 14, RGBColor.from_string("FFFFFF"), F4)]], anchor=MSO_ANCHOR.MIDDLE)
ped = by[106]; ped.left = I(MX + MW / 2) - ped.width // 2
CSF = [(111, 112, 113, 114, 115, "업무지원포털의 목적, 사상에 대한", "완벽한 이해"),
       (116, 117, 118, 119, 120, "구축, 운영사업 수행 경험에 근거한", "업무지식 및 기술역량"),
       (121, 122, 123, 124, 125, "정확한 문제진단을 통한", "최적 해결방안 제시 능력"),
       (126, 127, 128, 129, 130, "선제적 대응 중심의", "능동적 변화관리 체계"),
       (131, 132, 133, 134, 135, "원활한 의사소통 및", "밀접한 신뢰관계 유지")]
CH = 0.64
for i, (card, num, icon, t1, t2, lead, emph) in enumerate(CSF):
    top = 2.30 + i * 0.80; cy = top + CH / 2
    box(by[card], MX + 0.1, top, MW - 0.2, CH)
    nm = by[num]; box(nm, MX + 0.17, cy - 0.21, 0.42, 0.42)
    fit_pic(by[icon], MX + 0.64, cy - 0.17, 0.32, 0.34)
    t = by[t1]; box(t, MX + 1.0, top + 0.03, MW - 1.15, CH - 0.06)
    fill(t, [[(lead, 9, INK, F2)], [(emph, 13, BLUE, F4)]], anchor=MSO_ANCHOR.MIDDLE, spacing=0.95, gap=1)
    kill(t2)

# ── 오른쪽 열: 3대 추진 전략 ─────────────────────────────────
RX, RW = 7.13, 3.32
BLOCKS = [(140, 141, 142, 143, "전략 1", "안정성 확보",
           [(144, 145, 146, "ACT.1", ["시스템 구축과 운영을 모두", "수행한 핵심인력 투입"]),
            (147, 148, 149, "ACT.2", ["철저한 장애관리로", "장애율 Zero유지"])]),
          (150, 151, 152, 153, "전략 2", "관리체계 확립",
           [(154, 155, 156, "ACT.3", ["서비스 지속성을 위한 기능", "개선 체계 구축"]),
            (157, 158, 159, "ACT.4", ["법제도 및 업무 변경", "모니터링 체계 마련"])]),
          (160, 161, 162, 163, "전략 3", "변화 실천",
           [(164, 165, 166, "ACT.5", ["사용자 환경변화에 대한", "능동적 대응체계 구축"])])]
HD, RH, G, PAD = 0.60, 0.42, 0.05, 0.06
y = 2.16
for bg, pic, pill, name, ptxt, ntxt, rows in BLOCKS:
    bh = PAD + HD + len(rows) * (G + RH) + PAD
    box(by[bg], RX, y, RW, bh)
    fit_pic(by[pic], RX + 0.08, y + PAD, 0.66, HD)
    hc = y + PAD + HD / 2
    box(by[pill], RX + 0.82, hc - 0.12, 0.6, 0.24)
    fill(by[pill], [[(ptxt, 8.5, RGBColor.from_string("FFFFFF"), F4)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, margin=0.01)
    box(by[name], RX + 1.5, hc - 0.17, RW - 1.55, 0.34)
    fill(by[name], [[(ntxt, 13.5, NAVY, F4)]], anchor=MSO_ANCHOR.MIDDLE)
    ry = y + PAD + HD + G
    for rbg, rpill, rtxt, act, lines in rows:
        box(by[rbg], RX + 0.09, ry, RW - 0.18, RH)
        box(by[rpill], RX + 0.17, ry + RH / 2 - 0.12, 0.6, 0.24)
        fill(by[rpill], [[(act, 8.5, RGBColor.from_string("FFFFFF"), F4)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, margin=0.01)
        box(by[rtxt], RX + 0.84, ry + 0.02, RW - 0.98, RH - 0.04)
        fill(by[rtxt], [[(ln, 10.5, INK, F3)] for ln in lines], anchor=MSO_ANCHOR.MIDDLE, spacing=1.08, margin=0.02)
        ry += RH + G
    y += bh + 0.07

# ── 열 사이 화살표: 두꺼운 블록 화살표 → 얇은 단색 셰브런 ─────────
from pptx.oxml.ns import qn as _qn
for aid, cx in ((104, 2.95), (105, 6.88)):
    a = by[aid]
    geom = a._element.spPr.find(_qn("a:prstGeom")); geom.set("prst", "chevron")
    av = geom.find(_qn("a:avLst"))
    for g in list(av): av.remove(g)
    av.append(av.makeelement(_qn("a:gd"), {"name": "adj", "fmla": "val 62000"}))
    a.fill.solid(); a.fill.fore_color.rgb = BLUE
    box(a, cx - 0.1, 4.35 - 0.24, 0.2, 0.48)

p.save(dst)
print("s10 ok")
