# -*- coding: utf-8 -*-
"""31쪽(사업관리 방법 — 방법론/공정 구성 전략): 원본에서 강조한 곳만 강조.
- 6개 칸 머리(프로세스 관리·프로젝트 관리(S-PMM)·유지보수 및 운영(S-ISM)·개발(S-SEM)·
  연계확대(S-CEM)·프로젝트 지원)를 같은 모양으로: 파랑 2F78E0 채움 + 흰 글자,
  높이 0.21in, 제목 10pt a시월구일3 / 영문 부제 7pt a시월구일2.
- 원본에서 진한 회색 채움이던 목록 알약(절차서·지침서·양식 / 프로세스 진단 3항목 /
  교육 3항목) → 회청 5F7496 채움 + 흰 글자.
- 원본에서 파랑 채움이던 분석·설계·구현·인도 → 2F78E0 채움 + 흰 글자.
- 오른쪽 전략 본문은 원본 글자 크기(10 / 9.5pt)로 되돌림(문구는 원본과 같음).
인자: <src> <dst>"""
import sys
from pptx import Presentation
from pptx.util import Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.oxml.ns import qn

sys.stdout.reconfigure(encoding="utf-8")
src, dst = sys.argv[1:3]
prs = Presentation(src)
sl = None
for s in prs.slides:
    t = " ".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame)
    if "공정 구성 전략" in t and "S-PMM" in t:
        sl = s
        break
assert sl is not None, "31쪽을 찾지 못함"
shapes = list(sl.shapes)
by = {}
for s in shapes:
    by.setdefault(s.name, s)

BLUE, SLATE, WHITE = RGBColor(0x2F, 0x78, 0xE0), RGBColor(0x5F, 0x74, 0x96), RGBColor(0xFF, 0xFF, 0xFF)
IN = 914400


def set_font(r, face, size=None, color=None):
    rPr = r._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            # latin/ea 는 solidFill 뒤에 와야 함 — 끝에 붙이고 순서 정리
            rPr.append(el)
        el.set("typeface", face)
    if size:
        r.font.size = Pt(size)
    r.font.bold = False
    if color is not None:
        r.font.color.rgb = color
    # 스키마 순서: ln, fill, effect, highlight, uLnTx.., latin, ea, cs, sym, hlink...
    order = ["a:ln", "a:noFill", "a:solidFill", "a:gradFill", "a:blipFill", "a:pattFill", "a:grpFill",
             "a:effectLst", "a:effectDag", "a:highlight", "a:uLnTx", "a:uLn", "a:uFillTx", "a:uFill",
             "a:latin", "a:ea", "a:cs", "a:sym", "a:hlinkClick", "a:hlinkMouseOver", "a:rtl", "a:extLst"]
    kids = list(rPr)
    kids.sort(key=lambda e: next((i for i, t in enumerate(order) if e.tag == qn(t)), 99))
    for k in kids:
        rPr.remove(k)
    for k in kids:
        rPr.append(k)


def fill(sh, rgb):
    sh.fill.solid()
    sh.fill.fore_color.rgb = rgb
    sh.line.fill.background()


def pill(sh):
    g = sh._element.spPr.find(qn("a:prstGeom"))
    g.set("prst", "roundRect")
    av = g.find(qn("a:avLst"))
    for k in list(av):
        av.remove(k)
    av.append(av.makeelement(qn("a:gd"), {"name": "adj", "fmla": "val 50000"}))


def texts_on(box, pool=None):
    """도형 box 와 같은 자리에 겹쳐 있는 글상자들."""
    out = []
    for s in (pool or shapes):
        if s is box or not s.has_text_frame or not s.text_frame.text.strip():
            continue
        if abs(s.left - box.left) < 0.02 * IN and abs(s.top - box.top) < 0.02 * IN:
            out.append(s)
    return out


def center_text(t, box):
    t.left, t.top, t.width, t.height = box.left, box.top, box.width, box.height
    tf = t.text_frame
    tf.word_wrap = False
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for p in tf.paragraphs:
        p.alignment = PP_ALIGN.CENTER


# 1) 칸 머리 6개
H = int(0.21 * IN)
heads = [("한쪽 모서리가 잘린 사각형 26", "직사각형 27"), ("한쪽 모서리가 잘린 사각형 68", "직사각형 69"),
         ("한쪽 모서리가 잘린 사각형 72", "직사각형 73"), ("한쪽 모서리가 잘린 사각형 134", "직사각형 135"),
         ("한쪽 모서리가 잘린 사각형 178", "직사각형 179"), ("한쪽 모서리가 잘린 사각형 208", "직사각형 209")]
cols = {"한쪽 모서리가 잘린 사각형 26": "직사각형 25", "한쪽 모서리가 잘린 사각형 68": "직사각형 67"}
for bn, tn in heads:
    b, t = by[bn], by[tn]
    if bn in cols:  # 좌우 칸 머리도 가운데 칸처럼 칸 폭 가득·칸 윗변에 붙인 띠로
        c = by[cols[bn]]
        b.left, b.top, b.width = c.left, c.top, c.width
    b.adjustments[0] = 0.30
    b.height = H
    fill(b, BLUE)
    center_text(t, b)
    for p in t.text_frame.paragraphs:
        sub, seen = False, ""
        for r in p.runs:
            if "Solideos" in r.text or (r.text.strip() == "" and ")" in seen):
                sub = True
            seen += r.text
            if sub:
                set_font(r, 'a시월구일2', 7, WHITE)
            else:
                set_font(r, 'a시월구일3', 10, WHITE)

# 2) 원본에서 진한 회색이던 목록 알약 → 회청 채움 + 흰 글자
for n in (30, 32, 34, 46, 48, 50, 60, 62, 64):
    b = by[f"모서리가 둥근 직사각형 {n}"]
    fill(b, SLATE)
    pill(b)  # 원본처럼 모두 알약(앞선 정리에서 round2SameRect 로 바뀐 것 되돌림)
    for t in texts_on(b):
        center_text(t, b)
        for p in t.text_frame.paragraphs:
            for r in p.runs:
                set_font(r, "a시월구일3", 7, WHITE)

# 3) 분석·설계·구현·인도 → 파랑 채움 + 흰 글자 (원본 7pt)
for n in (190, 193, 196, 199):
    b = by[f"모서리가 둥근 직사각형 {n}"]
    fill(b, BLUE)
    for t in texts_on(b):
        center_text(t, b)
        for p in t.text_frame.paragraphs:
            for r in p.runs:
                set_font(r, "a시월구일3", 7, WHITE)

# 4) 오른쪽 전략 본문 원본 크기로
body = by["AutoShape 91"]
for p in body.text_frame.paragraphs:
    for r in p.runs:
        r.font.size = Pt(10 if p.level == 0 else 9.5)

print("31쪽 완료")
prs.save(dst)
