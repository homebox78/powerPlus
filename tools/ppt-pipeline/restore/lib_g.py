# -*- coding: utf-8 -*-
"""17~19쪽 2차 재구성(s17b·s18b·s19b) 공용 도우미."""
import io, copy
from pptx.util import Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from lib_e import *

L, R = 0.298, 10.535          # 가이드: 가운데 ±13.0cm
W = R - L
TOP, BOT = 1.033, 7.018       # 가운데 -6.9cm / +8.3cm
MID = RGBColor(0x23, 0x5C, 0xB8)


def purge(s):
    """바탕·제목·키메시지만 남기고 본문 도형을 모두 지운다."""
    for sh in list(s.shapes):
        if sh.name.startswith("시안 바탕") or sh.name == "키메시지" or sh.top < I(0.87):
            continue
        sh._element.getparent().remove(sh._element)


def tb(s, x, y, w, h, paras, size, fill=None, line=None, lw=0.75, shape=MSO_SHAPE.RECTANGLE, adj=0.08,
       color=TXT, font=F2, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE, bullet=None, bullet_color=BLUE,
       margins=(0.06, 0.03), space_after=0, line_spacing=1.05, wrap=True):
    if fill is None and line is None:
        sh = s.shapes.add_textbox(I(x), I(y), I(w), I(h))
    else:
        sh = box(s, x, y, w, h, fill=fill, line=line, lw=lw, shape=shape, adj=adj)
    fill_tf(sh, paras, size, color=color, font=font, align=align, anchor=anchor, bullet=bullet,
            bullet_color=bullet_color, space_after=space_after, line_spacing=line_spacing, wrap=wrap, margins=margins)
    return sh


def pic_blob(slide, sid):
    for sh in slide.shapes:
        if sh.shape_id == sid and sh.shape_type == 13:
            return sh.image.blob
    raise KeyError(sid)


def add_pic(s, blob, cx, cy, h):
    p = s.shapes.add_picture(io.BytesIO(blob), I(0), I(0))
    fit_pic(p, cx, cy, h=h)
    return p


def copy_group(dst_slide, src_slide, sid, x, y, size):
    for sh in src_slide.shapes:
        if sh.shape_id == sid:
            el = copy.deepcopy(sh._element)
            dst_slide.shapes._spTree.append(el)
            g = [g for g in dst_slide.shapes if g._element is el][0]
            g.left, g.top, g.width, g.height = I(x), I(y), I(size), I(size)
            return g
    raise KeyError(sid)


def _walk(shapes):
    for sh in shapes:
        yield sh
        if sh.shape_type == 6:
            yield from _walk(sh.shapes)


def pic_blob(slide, sid):
    for sh in _walk(slide.shapes):
        if sh.shape_id == sid and sh.shape_type == 13:
            return sh.image.blob
    raise KeyError(sid)


def copy_group(dst_slide, src_slide, sid, x, y, size):
    for sh in _walk(src_slide.shapes):
        if sh.shape_id == sid:
            el = copy.deepcopy(sh._element)
            dst_slide.shapes._spTree.append(el)
            g = [g for g in dst_slide.shapes if g._element is el][0]
            g.left, g.top, g.width, g.height = I(x), I(y), I(size), I(size)
            return g
    raise KeyError(sid)
