# -*- coding: utf-8 -*-
"""18쪽(유지보수 관리 체계 2/5) 2차: 원본 구조로 재구성. 인자: <src> <dst>
- 여자·건물 일러스트 삭제, 본문을 가이드 가로 전체로
- 3대 가치(안정성·지속성·신속성): 남색 머리 16pt + 작은 원형 아이콘(원본 아이콘), 원문 설명 10pt
- 왼쪽 행 머리 칸 복원: '주요활동 커스터마이징' / '운영/유지보수 방법론(S-ISM)' 12pt 남색
- 주요활동 5종: 머리 11.5pt, 원문 세부 활동 9.5pt (장식 아이콘 삭제)
- S-ISM 4영역: 머리 12pt, 항목 9.5pt"""
import os, sys
sys.path.insert(0, os.path.dirname(__file__))
from pptx import Presentation
from lib_g import *

sys.stdout.reconfigure(encoding="utf-8")
src, dst = sys.argv[1:3]
p = Presentation(src); s = p.slides[17]
purge(s)
# 원본(같은 파일 안에는 없음) 아이콘: 원본백업 파일에서 가져온다
here = os.path.dirname(os.path.abspath(__file__))
orig_path = os.path.join(here, "..", "..", "..", "제안서", "청주시_발표자료.원본백업.pptx")
o = Presentation(orig_path).slides[17]

k3 = lambda t: (t, F3, INK)
k2 = lambda t: (t, F2, TXT)

# 1) 3대 가치 — 가로 전체
VY, HH, VH = 1.62, 0.52, 1.42
GAP = 0.10
cw = (W - 2 * GAP) / 3
vals = [
    ("안정성", [[k2("프로젝트 착수부터 종료까지\v"), k3("안정적인 관리활동 수행")],
               [k2("시스템과 서비스의 변화에 사용자 그룹이\v반응하지 않아도 되도록 "), k3("안정적인 서비스 보장")]]),
    ("지속성", [[k2("끊김없이 "), k3("무중단으로 지속적인 서비스 지원")],
               [k2("강력한 보안 점검과 철저한 백업·복구 등\v활동의 안정적 지원으로 "), k3("지속성 확보")]]),
    ("신속성", [[k2("법제도적, 시스템·업무적으로 발생한\v환경변화에 "), k3("신속하게 반응하고 조치")],
               [k2("사용자 요구를 능동적으로 파악, "), k3("신속한 응대")]]),
]
for k, (title, items) in enumerate(vals):
    x = L + k * (cw + GAP)
    tb(s, x, VY, cw, HH, [title], 16, fill=NAVY, line=None, color=WHITE, font=F4,
       align=PP_ALIGN.CENTER, margins=(0.0, 0.0))
    icx = x + cw / 2 - 0.62
    if k == 0:
        add_pic(s, pic_blob(o, 142), icx, VY + HH / 2, 0.40)
    else:
        g = copy_group(s, o, 136 if k == 1 else 139, icx - 0.20, VY + HH / 2 - 0.20, 0.40)
    tb(s, x, VY + HH, cw, VH - HH, items, 10, fill=PALE, line=None, bullet="•",
       margins=(0.14, 0.03), space_after=3)

# 2) 왼쪽 행 머리 칸
LBW = 1.22
CX0 = L + LBW + 0.10
CW = R - CX0
R2Y, R2H = VY + VH + 0.14, 2.02
R3Y = R2Y + R2H + 0.12
R3H = BOT - R3Y
for y, h, t in ((R2Y, R2H, "주요활동\v커스터마이징"), (R3Y, R3H, "운영/유지보수\v방법론(S-ISM)")):
    tb(s, L, y, LBW, h, [t], 12, fill=WHITE, line=NAVY, lw=1.25, color=NAVY, font=F4,
       align=PP_ALIGN.CENTER, margins=(0.02, 0.02), line_spacing=1.1)

# 3) 주요활동 5종 (원본 폭 비율)
acts = [
    ("서비스 요청(SR) 처리", 1.86, ["서비스 요청(SR)의 적시 처리",
                                  "상주 담당자의 단순 답변이나\v원격지원, 기술지원\v활동 수행",
                                  "장비 및 솔루션 관련\v기술지원 1차 처리 및\v책임 이관"]),
    ("변경 요청(RFC) 처리", 1.86, ["변경 영향 평가와 적정성\v확보 후 변경 처리",
                                 "변경요청 처리 과정 기록관리",
                                 "변경 사항이 반영된 산출물\v작성 후 주관기관의\v승인 획득"]),
    ("배포 관리", 1.68, ["서비스 배포 결과\v기록 관리",
                       "통합테스트가 완료된\v변경 처리의 운영 서버\v테스트 수행 후\v주관기관 승인 획득"]),
    ("요구사항 관리", 1.38, ["고객 요구사항 검증\v및 승인, 변경관리,\v추적관리"]),
    ("장애 관리", 1.38, ["장애 등급에\v따른 조치",
                       "장애 근본원인\v분석, 재발방지 및\v예방 활동 수행"]),
]
G2 = 0.08
scale = (CW - G2 * 4) / sum(a[1] for a in acts)
x = CX0
AHH = 0.40
for title, w0, items in acts:
    w = w0 * scale
    tb(s, x, R2Y, w, AHH, [title], 11.5, fill=BLUE, line=None, color=WHITE, font=F4,
       align=PP_ALIGN.CENTER, margins=(0.02, 0.0))
    tb(s, x, R2Y + AHH, w, R2H - AHH, items, 9.5, fill=WHITE, line=LIGHT, lw=1, bullet="•",
       margins=(0.07, 0.03), space_after=3)
    x += w + G2

# 4) S-ISM 4영역
groups = [
    ("프로세스 관리", 1.50, [["조직 프로세스 관리", "조직 교육 훈련"]]),
    ("서비스 수행", 3.62, [["서비스 수행 계획 관리", "서비스 수행 감독 및 통제", "서비스 수준 관리", "서비스 연속성 관리"],
                        ["인시던트 관리", "장애(문제) 관리", "공급자 계약 관리", "전략적 서비스 관리"]]),
    ("서비스 개발 및 전달", 1.75, [["변경관리 / 배포관리", "요구사항관리", "전략적 서비스 관리", "테스트/검증 관리"]]),
    ("지원", 1.42, [["구성 관리", "품질 관리", "측정 및 분석 관리", "의사결정 관리"]]),
]
G3 = 0.10
scale = (CW - G3 * 3) / sum(g[1] for g in groups)
x = CX0
for title, w0, cols in groups:
    w = w0 * scale
    tb(s, x, R3Y, w, AHH, [title], 12, fill=NAVY, line=None, color=WHITE, font=F4,
       align=PP_ALIGN.CENTER, margins=(0.02, 0.0))
    box(s, x, R3Y + AHH, w, R3H - AHH, fill=WHITE, line=LIGHT, lw=1, shape=MSO_SHAPE.RECTANGLE)
    cwid = (w - 0.10) / len(cols)
    for j, items in enumerate(cols):
        tb(s, x + 0.05 + j * cwid, R3Y + AHH, cwid, R3H - AHH, items, 9.5, bullet="•",
           margins=(0.07, 0.03), space_after=3)
    x += w + G3

p.save(dst)
print("saved", dst)
