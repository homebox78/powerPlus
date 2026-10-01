# -*- coding: utf-8 -*-
"""12쪽 원본 문구 복원 — 3. 관리체계 확립 (1/2).
   A 풍부한 행정정보화 사업경험(새올행정·세움터·행안부·기타 + 원본 사업명)
   B 공공행정 정보화 사업 특징(원본 4문구)
   C 운영 및 유지관리 (S-ISM)방법론(부처명 회색 / 사업명 파랑, ※ S-ISM 주석)
   하단 5단계 = 원본 글머리 그대로. 인물·단계 그림은 작게 줄여 글이 주인공.
   ACT.3 왼쪽에 '전략 2' 배지.  인자: <src> <dst>"""
import sys, copy
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
sys.stdout.reconfigure(encoding="utf-8")
src, dst = sys.argv[1:3]
p = Presentation(src); s = p.slides[11]
by = {sh.shape_id: sh for sh in s.shapes}
I = Inches
E = 914400
NAVY, BLUE, LIGHT, PALE = RGBColor(0x14, 0x3A, 0x69), RGBColor(0x2F, 0x78, 0xE0), RGBColor(0xD0, 0xE6, 0xFA), RGBColor(0xE8, 0xF2, 0xFC)
INK, TXT, GRAY, WHITE, RED = RGBColor(0x1F, 0x4E, 0x79), RGBColor(0x33, 0x33, 0x33), RGBColor(0x5F, 0x74, 0x96), RGBColor(255, 255, 255), RGBColor(0xC0, 0, 0)
F2, F3, F4 = "a시월구일2", "a시월구일3", "a시월구일4"


def face(r, f):
    rPr = r._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea"):
        e = rPr.find(qn(tag))
        if e is None:
            e = rPr.makeelement(qn(tag), {}); rPr.append(e)
        e.set("typeface", f)


def fill_tf(tf, paras, align=PP_ALIGN.LEFT, spacing=None, bullet=False, after=0, anchor=MSO_ANCHOR.TOP, m=0.04, spc=None):
    """paras: [[(text,size,color,font), ...], ...]"""
    tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = I(m); tf.margin_top = tf.margin_bottom = 0
    for pg in list(tf.paragraphs)[1:]:
        pg._p.getparent().remove(pg._p)
    p0 = tf.paragraphs[0]
    for r in list(p0._p):
        if r.tag in (qn("a:r"), qn("a:br"), qn("a:fld")):
            p0._p.remove(r)
    for i, runs in enumerate(paras):
        pg = p0 if i == 0 else tf.add_paragraph()
        pg.alignment = align
        if spacing: pg.line_spacing = spacing
        pg.space_after = Pt(after)
        pPr = pg._p.get_or_add_pPr()
        for b in pPr.findall(qn("a:buChar")) + pPr.findall(qn("a:buNone")):
            pPr.remove(b)
        if bullet:
            pPr.set("marL", "91440"); pPr.set("indent", "-91440")
            pPr.append(pPr.makeelement(qn("a:buChar"), {"char": "•"}))
        else:
            pPr.set("marL", "0"); pPr.set("indent", "0")
            pPr.append(pPr.makeelement(qn("a:buNone"), {}))
        for t, sz, col, f in runs:
            r = pg.add_run(); r.text = t
            r.font.size = Pt(sz); r.font.color.rgb = col; r.font.bold = False; face(r, f)
            if spc: r._r.get_or_add_rPr().set("spc", str(spc))


def tb(x, y, w, h, paras, **kw):
    t = s.shapes.add_textbox(I(x), I(y), I(w), I(h))
    fill_tf(t.text_frame, paras, **kw)
    return t


def box(x, y, w, h, fill, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=None, lw=0.75):
    b = s.shapes.add_shape(shape, I(x), I(y), I(w), I(h))
    b.shadow.inherit = False
    if fill is None:
        b.fill.background()
    else:
        b.fill.solid(); b.fill.fore_color.rgb = fill
    if line is None:
        b.line.fill.background()
    else:
        b.line.color.rgb = line; b.line.width = Pt(lw)
    if adj is not None:
        b.adjustments[0] = adj
    b.text_frame.text = ""
    return b


def geo(sh, x=None, y=None, w=None, h=None):
    if x is not None: sh.left = I(x)
    if y is not None: sh.top = I(y)
    if w is not None: sh.width = I(w)
    if h is not None: sh.height = I(h)


def fitpic(sh, cx, cy, hmax, wmax=99):
    r = sh.width / sh.height
    h = min(hmax, wmax / r); w = h * r
    geo(sh, cx - w / 2, cy - h / 2, w, h)


def crop(sh, l=0.0, t=0.0):
    w0, h0 = sh.width, sh.height
    sh.crop_left = l; sh.crop_top = t
    sh.width = int(w0 * (1 - l)); sh.height = int(h0 * (1 - t))


def kill(*ids):
    for i in ids:
        if i in by:
            by[i]._element.getparent().remove(by[i]._element)


def front(sh):
    tree = s.shapes._spTree; tree.remove(sh._element); tree.append(sh._element)


# ── 0) 전략 2 배지 + ACT.3 이동 ─────────────────────────────
act, actlab = by[139], by[140]
AY, AH = act.top / E, act.height / E
BX, BW = 0.30, 0.86
badge = box(BX, AY, BW, AH, RED, adj=0.25)
fill_tf(badge.text_frame, [[("전략 ", 10.5, WHITE, F3), ("2", 15, WHITE, F4)]], align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE, m=0)
dx = BX + BW + 0.10 - act.left / E
act.left = act.left + I(dx); actlab.left = actlab.left + I(dx)

# ── 1) 위 3카드 ─────────────────────────────────────────────
TOP, BOT = 2.06, 4.18
L, R = 0.30, 10.535
GAP = 0.20
WA, WB = 3.18, 2.95
WC = R - L - WA - WB - 2 * GAP
XA, XB = L, L + WA + GAP
XC = XB + WB + GAP
cards = [(141, 142, 143, 158, XA, WA, "풍부한 행정정보화 사업경험"),
         (146, 147, 148, 177, XB, WB, "공공행정 정보화 사업 특징"),
         (152, 153, 154, 190, XC, WC, "운영 및 유지관리 (S-ISM)방법론")]
for card, let, ttl, pic, x, w, title in cards:
    geo(by[card], x, TOP, w, BOT - TOP)
    geo(by[let], x + 0.11, TOP + 0.07, 0.36, 0.36)
    for pg in by[let].text_frame.paragraphs:
        for r in pg.runs: r.font.size = Pt(17)
    t = by[ttl]; geo(t, x + 0.52, TOP + 0.07, w - 0.52 - 0.50, 0.36)
    fill_tf(t.text_frame, [[(title, 11, NAVY, F4)]], anchor=MSO_ANCHOR.MIDDLE, m=0.02)
    t.text_frame.word_wrap = False
    if pic in (177, 190):
        crop(by[pic], 0.12, 0.05)       # 원래 그림 왼쪽 위에 붙어 있던 + 기호 조각 제거
    fitpic(by[pic], x + w - 0.27, TOP + 0.26, 0.44, 0.42)   # 인물 = 제목 오른쪽 작은 아이콘
kill(144, 145, 149, 150, 151, 155, 156, 157)
kill(*range(178, 190))            # B 지어낸 4행(법·제도 준수 등)
kill(191, 192, *range(193, 200))  # C 고리 그림·표준화/체계화/지속개선
# + 기호
for c, h1, h2, gx in ((171, 172, 173, XA + WA + GAP / 2), (174, 175, 176, XB + WB + GAP / 2)):
    cy = (TOP + BOT) / 2 + 0.10; d = 0.30
    geo(by[c], gx - d / 2, cy - d / 2, d, d)
    geo(by[h1], gx - 0.07, cy - 0.013, 0.14, 0.026)
    geo(by[h2], gx - 0.013, cy - 0.07, 0.026, 0.14)
    for i in (c, h1, h2): front(by[i])

CY0 = TOP + 0.53   # 내용 시작
# A: 기관 꼬리표 + 원본 사업명
rowsA = [(159, 160, 161, ["지방행정통합정보 시스템 유지보수"]),
         (162, 163, 164, ["건축행정시스템 구축 및 유지보수", "지능형 건축행정 구축 및 유지보수"]),
         (165, 166, 167, ["행정정보공동이용시스템 유지보수", "정부24 유지보수"]),
         (168, 169, 170, ["식·의약품 종합정보 시스템", "복합민원, 교통행정 포털"])]
labels = ["새올행정", "세움터", "행안부", "기타"]
RH, RG = 0.33, 0.05
TW = 0.86
for k, (tile, icon, lab, items) in enumerate(rowsA):
    y = CY0 + k * (RH + RG)
    geo(by[tile], XA + 0.12, y, TW, RH)
    by[tile].fill.solid(); by[tile].fill.fore_color.rgb = WHITE
    by[tile].line.color.rgb = LIGHT
    fitpic(by[icon], XA + 0.12 + 0.17, y + RH / 2, 0.19, 0.22)
    geo(by[lab], XA + 0.12 + 0.30, y, TW - 0.32, RH)
    fill_tf(by[lab].text_frame, [[(labels[k], 9, NAVY, F4)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, m=0)
    by[lab].text_frame.word_wrap = False
    bg = box(XA + 0.12 + TW + 0.05, y, WA - 0.24 - TW - 0.05, RH, PALE, adj=0.12)
    t = tb(XA + 0.12 + TW + 0.08, y, WA - 0.24 - TW - 0.10, RH,
           [[(it, 8.5, TXT, F2)] for it in items], bullet=True, spacing=0.95, anchor=MSO_ANCHOR.MIDDLE, m=0.02)
# B: 원본 4문구
bph = ["장애 및 변경요인 모니터링 중요", "공공성 강화 UI 요구 증대", "SW 품질관리 증가", "절대적 서비스 안정성"]
for k, ph in enumerate(bph):
    y = CY0 + k * (RH + RG)
    b = box(XB + 0.14, y, WB - 0.28, RH, WHITE, LIGHT, adj=0.12)
    bar = box(XB + 0.14, y + 0.05, 0.05, RH - 0.10, BLUE, shape=MSO_SHAPE.RECTANGLE)
    fill_tf(b.text_frame, [[(ph, 10.5, INK, F4)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, m=0.06)
# C: 원본 6행 + 주석
crows = [("행정안전부", "지방행정통합정보시스템 유지보수(새올)"),
         ("국토교통부", "건축행정시스템 유지보수"),
         ("식품의약품안전처", "통합 유지보수(식품/의약품/의료기기)"),
         ("행정안전부", "정부24 유지보수"),
         ("행정안전부", "행정정보공동이용 유지보수"),
         ("서울특별시", "건축주택통계분석시스템 유지보수")]
CRH, CRG = 0.205, 0.035
for k, (mi, pj) in enumerate(crows):
    y = CY0 + k * (CRH + CRG)
    b = box(XC + 0.12, y, WC - 0.24, CRH, WHITE, LIGHT, adj=0.15)
    fill_tf(b.text_frame, [[("(" + mi + ") ", 8.5, GRAY, F3), (pj, 8.5, BLUE, F3)]],
            anchor=MSO_ANCHOR.MIDDLE, m=0.07)
    b.text_frame.word_wrap = False
yn = CY0 + 6 * (CRH + CRG) - 0.01
tb(XC + 0.12, yn, WC - 0.24, 0.18, [[("※ S-ISM : Solideos – IT Service Mate", 8, GRAY, F2)]], align=PP_ALIGN.RIGHT, m=0.02)

# ── 2) 화살표·하단 5단계 ───────────────────────────────────
geo(by[200], 4.82, BOT + 0.045, 1.20, 0.11)
BT = BOT + 0.22
B_BOT = 7.00
geo(by[201], L, BT, R - L, B_BOT - BT)
geo(by[202], L, BT, R - L, 0.30)
fill_tf(by[202].text_frame, [[("S-ISM 방법론기반 청주시에 적합한 기능개선 체계", 11.5, WHITE, F4)]],
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, m=0)
steps = [(203, 204, 205, 206, 207, 208, (209, 210), 211, "서비스요청(SR) 처리",
          ["사용자로부터 접수 받은 서비스요청(SR)의 적시 처리",
           "상주담당자의 단순답변, 원격지원, 기술지원 등의 활동 수행",
           "장비 및 솔루션 관련 기술지원이 필요한 경우 해당 유지보수 사업자에게 서비스 요청 이관"]),
         (212, 213, 214, 215, 216, 217, (218, 219), 220, "변경요청(RFC) 처리",
          ["변경에 대한 영향 평가와 변경의 적정성을 확보한 후 변경처리",
           "변경요청 처리에 대한 일련의 과정을 기록/관리",
           "변경 사항이 반영된 산출물을 주관기관의 승인 획득"]),
         (221, 222, 223, 224, 225, 226, (227, 228), 229, "배포 관리",
          ["서비스 배포결과의 기록을 관리하여 배포 후 발생하는 문제에 대한 즉시 대응 및 배포 결과의 관리",
           "통합테스트가 완료된 변경처리 건에 대하여 운영서버에서 테스트 수행하고 주관기관 승인 획득"]),
         (230, 231, 232, 233, 234, 235, (236, 237), 238, "요구사항 관리",
          ["고객 요구사항의 검증 및 승인, 변경관리, 추적관리를 통하여 요구사항을 일관성 있게 유지 관리"]),
         (239, 240, 241, 242, 243, 244, (245, 246), None, "장애관리",
          ["장애를 정해진 기간 안에 조치할 수 있도록 등록/해결/예방하는 활동을 통해 시스템의 가용성과 연속성을 확보",
           "장애등급에 따른 조치 및 장애의 근본원인 분석을 통한 재발방지 및 예방 활동 수행"])]
CT = BT + 0.38
CB = B_BOT - 0.07
X0, X1 = L + 0.07, R - 0.07
CG = 0.10
CW = (X1 - X0 - 4 * CG) / 5
HH = 0.28
for k, (card, head, num, ttl, pic, pnl, olds, chev, title, buls) in enumerate(steps):
    x = X0 + k * (CW + CG)
    geo(by[card], x, CT, CW, CB - CT)
    geo(by[head], x, CT, CW, HH)
    geo(by[num], x + 0.05, CT + 0.02, 0.24, 0.24)
    for pg in by[num].text_frame.paragraphs:
        for r in pg.runs: r.font.size = Pt(10.5)
    geo(by[ttl], x + 0.30, CT, CW - 0.34, HH)
    fill_tf(by[ttl].text_frame, [[(title, 10, NAVY, F4)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, m=0)
    fitpic(by[pic], x + CW / 2, CT + HH + 0.03 + 0.14, 0.28, CW - 0.3)
    kill(pnl, *olds)
    tb(x + 0.03, CT + HH + 0.36, CW - 0.05, CB - CT - HH - 0.40,
       [[(b, 9.5, TXT, F2)] for b in buls], bullet=True, spacing=0.95, after=3, m=0.02, spc=-10)
    if chev is not None:
        geo(by[chev], x + CW + CG / 2 - 0.035, CT + (CB - CT) / 2 - 0.08, 0.07, 0.16)
p.save(dst); print("ok")
