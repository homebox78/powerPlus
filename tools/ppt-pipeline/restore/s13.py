# -*- coding: utf-8 -*-
"""13쪽 원본 문구 복원 — 3. 관리체계 확립 (2/2).
   AS-IS: '유지관리 담당자' 분홍 강조 배지, 주변 4노드 점선 원+빨간 굵은 글자, 아래 정보통신담당자·업무담당자를
          점선으로 연결, 요청 띠 "공문/서면으로 변경요청" 강조. 지어낸 '수동적 대응' 삭제.
   3D 그림 화살표 → 납작한 셰브런.  TO-BE: 헤드라인 확대(둘째 줄 남색 굵게), 번호 원 삭제,
   작은 아이콘 + 제목(12.5pt) + 원본 설명(10pt).  ACT.4 왼쪽에 '전략 2' 배지.  인자: <src> <dst>"""
import sys, copy
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.dml import MSO_LINE
from pptx.oxml.ns import qn
sys.stdout.reconfigure(encoding="utf-8")
src, dst = sys.argv[1:3]
p = Presentation(src); s = p.slides[12]
by = {sh.shape_id: sh for sh in s.shapes}
I = Inches
E = 914400
NAVY, BLUE, LIGHT, PALE = RGBColor(0x14, 0x3A, 0x69), RGBColor(0x2F, 0x78, 0xE0), RGBColor(0xD0, 0xE6, 0xFA), RGBColor(0xE8, 0xF2, 0xFC)
INK, TXT, GRAY, WHITE, RED = RGBColor(0x1F, 0x4E, 0x79), RGBColor(0x33, 0x33, 0x33), RGBColor(0x5F, 0x74, 0x96), RGBColor(255, 255, 255), RGBColor(0xC0, 0, 0)
PINK = RGBColor(0xEC, 0x1C, 0x68)
F2, F3, F4 = "a시월구일2", "a시월구일3", "a시월구일4"


def face(r, f):
    rPr = r._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea"):
        e = rPr.find(qn(tag))
        if e is None:
            e = rPr.makeelement(qn(tag), {}); rPr.append(e)
        e.set("typeface", f)


def fill_tf(tf, paras, align=PP_ALIGN.LEFT, spacing=None, bullet=False, after=0, anchor=MSO_ANCHOR.TOP, m=0.04, spc=None):
    """paras: [[(text,size,color,font), ...], ...]"""
    tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = I(m); tf.margin_top = tf.margin_bottom = 0
    for pg in list(tf.paragraphs)[1:]:
        pg._p.getparent().remove(pg._p)
    p0 = tf.paragraphs[0]
    for r in list(p0._p):
        if r.tag in (qn("a:r"), qn("a:br"), qn("a:fld")):
            p0._p.remove(r)
    for i, runs in enumerate(paras):
        pg = p0 if i == 0 else tf.add_paragraph()
        pg.alignment = align
        if spacing: pg.line_spacing = spacing
        pg.space_after = Pt(after)
        pPr = pg._p.get_or_add_pPr()
        for b in pPr.findall(qn("a:buChar")) + pPr.findall(qn("a:buNone")):
            pPr.remove(b)
        if bullet:
            pPr.set("marL", "91440"); pPr.set("indent", "-91440")
            pPr.append(pPr.makeelement(qn("a:buChar"), {"char": "•"}))
        else:
            pPr.set("marL", "0"); pPr.set("indent", "0")
            pPr.append(pPr.makeelement(qn("a:buNone"), {}))
        for t, sz, col, f in runs:
            r = pg.add_run(); r.text = t
            r.font.size = Pt(sz); r.font.color.rgb = col; r.font.bold = False; face(r, f)
            if spc: r._r.get_or_add_rPr().set("spc", str(spc))


def tb(x, y, w, h, paras, **kw):
    t = s.shapes.add_textbox(I(x), I(y), I(w), I(h))
    fill_tf(t.text_frame, paras, **kw)
    return t


def box(x, y, w, h, fill, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=None, lw=0.75):
    b = s.shapes.add_shape(shape, I(x), I(y), I(w), I(h))
    b.shadow.inherit = False
    if fill is None:
        b.fill.background()
    else:
        b.fill.solid(); b.fill.fore_color.rgb = fill
    if line is None:
        b.line.fill.background()
    else:
        b.line.color.rgb = line; b.line.width = Pt(lw)
    if adj is not None:
        b.adjustments[0] = adj
    b.text_frame.text = ""
    return b


def geo(sh, x=None, y=None, w=None, h=None):
    if x is not None: sh.left = I(x)
    if y is not None: sh.top = I(y)
    if w is not None: sh.width = I(w)
    if h is not None: sh.height = I(h)


def fitpic(sh, cx, cy, hmax, wmax=99):
    r = sh.width / sh.height
    h = min(hmax, wmax / r); w = h * r
    geo(sh, cx - w / 2, cy - h / 2, w, h)


def crop(sh, l=0.0, t=0.0):
    w0, h0 = sh.width, sh.height
    sh.crop_left = l; sh.crop_top = t
    sh.width = int(w0 * (1 - l)); sh.height = int(h0 * (1 - t))


def kill(*ids):
    for i in ids:
        if i in by:
            by[i]._element.getparent().remove(by[i]._element)


def front(sh):
    tree = s.shapes._spTree; tree.remove(sh._element); tree.append(sh._element)


# ── 0) 전략 2 배지 + ACT.4 이동 ─────────────────────────────
act, actlab = by[88], by[89]
AY, AH = act.top / E, act.height / E
BX, BW = 0.30, 0.86
badge = box(BX, AY, BW, AH, RED, adj=0.25)
fill_tf(badge.text_frame, [[("전략 ", 10.5, WHITE, F3), ("2", 15, WHITE, F4)]], align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE, m=0)
dx = BX + BW + 0.10 - act.left / E
act.left = act.left + I(dx); actlab.left = actlab.left + I(dx)

# ── 1) AS-IS ───────────────────────────────────────────────
kill(92, 113)                                  # 지어낸 '수동적 대응' · 두 사람 사이 양방향 화살표
# 주변 4노드: 점선 원 + 빨간 굵은 글자 (배지와 겹치지 않게 바깥으로 벌림)
ctr = by[109]; ccx = (ctr.left + ctr.width / 2) / E
kill(93, 94, 95, 96)
NODES = ((97, 98, 99, "법제도", 1.06, 3.12), (100, 101, 102, "업무규정", 3.98, 3.12),
         (103, 104, 105, "조례", 1.06, 4.18), (106, 107, 108, "업무시스템", 3.98, 4.18))
D = 0.90
for circ, pic, lab, text, cx, cy in NODES:
    c = by[circ]
    geo(c, cx - D / 2, cy - D / 2, D, D)
    c.fill.solid(); c.fill.fore_color.rgb = WHITE
    c.line.color.rgb = BLUE; c.line.width = Pt(1.25); c.line.dash_style = MSO_LINE.ROUND_DOT
    fitpic(by[pic], cx, cy - 0.13, 0.30, 0.34)
    geo(by[lab], cx - 0.45, cy + 0.06, 0.90, 0.22)
    fill_tf(by[lab].text_frame, [[(text, 9.5, RED, F4)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, m=0)
    by[lab].text_frame.word_wrap = False
    # 원 → 가운데 인물 점선
    import math
    tx, ty = ccx, 3.50
    ang = math.atan2(ty - cy, tx - cx)
    sx, sy = cx + math.cos(ang) * (D / 2 + 0.03), cy + math.sin(ang) * (D / 2 + 0.03)
    ex, ey = tx - math.cos(ang) * 0.46, ty - math.sin(ang) * 0.42
    ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, I(sx), I(sy), I(ex), I(ey))
    ln.line.color.rgb = RGBColor(0x9F, 0xB3, 0xCC); ln.line.width = Pt(1.25); ln.line.dash_style = MSO_LINE.ROUND_DOT
# 가운데 '유지관리 담당자' = 분홍 강조 배지
fitpic(ctr, ccx, 3.50, 0.80)
bd = by[110]; geo(bd, ccx - 0.78, 3.93, 1.56, 0.34)
bd.fill.solid(); bd.fill.fore_color.rgb = PINK
fill_tf(bd.text_frame, [[("유지관리 담당자", 11.5, WHITE, F4)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, m=0)
bd.text_frame.word_wrap = False
front(bd)
# 아래 두 사람(정보통신담당자·업무담당자) + 배지에서 점선 연결
people = [(111, 114, "정보통신담당자", 1.38), (112, 115, "업무담당자", 3.47)]
PY = 5.02
for pic, lab, text, px in people:
    fitpic(by[pic], px, PY + 0.30, 0.62)
    geo(by[lab], px - 0.62, PY + 0.64, 1.24, 0.26)
    fill_tf(by[lab].text_frame, [[(text, 10, WHITE, F4)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, m=0)
    by[lab].text_frame.word_wrap = False
for sx, ex in ((ccx - 0.45, 1.38), (ccx + 0.45, 3.47)):
    ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, I(sx), I(4.27), I(ex), I(PY - 0.02))
    ln.line.color.rgb = BLUE; ln.line.width = Pt(1.25); ln.line.dash_style = MSO_LINE.ROUND_DOT
# 아래 요청 띠 = 원본처럼 강조
rq = by[116]; geo(rq, 0.42, 6.12, 4.03, 0.80)
rq.fill.solid(); rq.fill.fore_color.rgb = NAVY
fill_tf(rq.text_frame, [[("업무처리 담당자 법제도 및 업무변경 요청", 11, WHITE, F3)],
                        [("“공문/서면으로 변경요청”", 14, WHITE, F4)]],
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, spacing=1.0, m=0.05)
geo(by[90], 0.31, 2.09, 4.25, 7.0 - 2.09)

# ── 2) 3D 그림 화살표 → 납작한 셰브런 ─────────────────────
kill(117)
chv = box(4.80, 4.20, 0.36, 0.62, BLUE, shape=MSO_SHAPE.CHEVRON)
chv.adjustments[0] = 0.55

# ── 3) TO-BE ───────────────────────────────────────────────
TB = by[118]; geo(TB, 5.42, 2.09, 10.535 - 5.42, 7.0 - 2.09)
TX, TW = 5.42, 10.535 - 5.42
geo(by[119], TX + 0.13, 2.27, 1.10, 0.38)
hl = by[120]; geo(hl, TX + 1.36, 2.18, TW - 1.46, 0.64)
fill_tf(hl.text_frame, [[("법제도 및 업무 변경 모니터링을 통한", 12, INK, F3)],
                        [("체계적이고 능동적인 관리", 16, NAVY, F4)]],
        anchor=MSO_ANCHOR.MIDDLE, spacing=0.95, m=0.02)
kill(122, 130, 135, 124, 125, 126, 127, 132, 137)
items = [(121, 123, 128, "입법예고 및 요청사항 모니터링",
          [[("주기: ", 10, INK, F4), ("일일 모니터링", 10, TXT, F2)],
           [("대상: ", 10, INK, F4), ("법제처, 자치법규관리시스템", 10, TXT, F2)]]),
         (129, 131, 133, "사용자 및 정보통신담당자 밀착 지원",
          [[("사용자와의 유선, 대면 서비스 요청 접수 및", 10, TXT, F2)],
           [("처리 결과 통보 시 법규정 변경 사항 식별", 10, TXT, F2)]]),
         (134, 136, 138, "전사 공유체계 활용",
          [[("제안사의 업무시스템 유지관리 담당자들 간의", 10, TXT, F2)],
           [("법개정 관련 사항 공유체계 활용", 10, TXT, F2)]])]
Y0, Y1, G = 3.00, 6.90, 0.12
H = (Y1 - Y0 - 2 * G) / 3
for k, (card, ttl, pic, title, desc) in enumerate(items):
    y = Y0 + k * (H + G)
    geo(by[card], TX + 0.15, y, TW - 0.30, H)
    fitpic(by[pic], TX + 0.15 + 0.55, y + H / 2, 0.52, 0.86)
    geo(by[ttl], TX + 1.18, y + 0.20, TW - 1.40, 0.30)
    fill_tf(by[ttl].text_frame, [[(title, 12.5, BLUE, F4)]], anchor=MSO_ANCHOR.MIDDLE, m=0.02)
    by[ttl].text_frame.word_wrap = False
    tb(TX + 1.18, y + 0.56, TW - 1.40, H - 0.66, desc, spacing=1.05, after=2, m=0.02)
p.save(dst); print("ok")
