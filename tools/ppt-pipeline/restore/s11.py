# -*- coding: utf-8 -*-
"""11쪽(2. 안정성 확보) 원본 문구 복원.
- '전략 1' 배지 복원(ACT.1 알약 왼쪽), ACT 문구 원본대로
- ACT.1: 사진 축소(비율 유지), 원본 이력 4줄 그대로, 흐름(현 운영담당자 업무 지속 → 인수위험 Zero →
  안정적인 유지보수 업무 수행)은 아이콘 작게·글자 크게, 납작한 셰브런 화살표
- ACT.2: 시스템 안정화 그래프 폭 축소 + 원본 인용 띠, 4단계 장애율 관리 전략 원본 구성
  (가운데 '장애발생요인 사전제거' 원 + 01~04 사분면, 원본 세부 항목 전부)
인자: <src> <dst>  (11쪽만 수정)
"""
import copy
import sys
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

sys.stdout.reconfigure(encoding="utf-8")
src, dst = sys.argv[1:3]
p = Presentation(src)
s = p.slides[10]
by = {sh.shape_id: sh for sh in s.shapes}
I = lambda v: Inches(v)
NAVY, BLUE, INK, GRAY, PINK, RED, WHITE = (RGBColor.from_string(c) for c in
    ("143A69", "2F78E0", "1F4E79", "5F7496", "EC1C68", "C00000", "FFFFFF"))
LIGHT, PALE = RGBColor.from_string("D0E6FA"), RGBColor.from_string("E8F2FC")
F2, F3, F4 = "a시월구일2", "a시월구일3", "a시월구일4"


def face(r, f):
    rPr = r._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea"):
        e = rPr.find(qn(tag))
        if e is None:
            e = rPr.makeelement(qn(tag), {})
            rPr.append(e)
        e.set("typeface", f)


def run(pg, text, size, color, f):
    r = pg.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.color.rgb = color
    r.font.bold = False
    face(r, f)
    return r


def clear(sh):
    tf = sh.text_frame
    for pg in list(tf.paragraphs)[1:]:
        pg._p.getparent().remove(pg._p)
    p0 = tf.paragraphs[0]
    for r in list(p0.runs):
        r._r.getparent().remove(r._r)
    for br in p0._p.findall(qn("a:br")):
        p0._p.remove(br)
    return tf


def para(tf, first, align, spacing=None, bullet=None, marL=0, after=0):
    pg = tf.paragraphs[0] if first else tf.add_paragraph()
    pg.alignment = align
    if spacing:
        pg.line_spacing = spacing
    pg.space_after = Pt(after)
    if bullet:
        pPr = pg._p.get_or_add_pPr()
        pPr.set("marL", str(int(I(marL))))
        pPr.set("indent", str(-int(I(marL))))
        pPr.append(pPr.makeelement(qn("a:buFont"), {"typeface": "Arial"}))
        pPr.append(pPr.makeelement(qn("a:buChar"), {"char": bullet}))
    return pg


def lines_to(tf, lines, size, color, f, align=PP_ALIGN.LEFT, spacing=None, bullet=None, marL=0.11, after=0):
    """lines: 문자열(\\v = 줄바꿈 a:br) 또는 (문자열, 크기, 색, 글꼴, 글머리) 튜플."""
    for i, ln in enumerate(lines):
        if isinstance(ln, tuple):
            text, sz, col, ff, bu = ln
        else:
            text, sz, col, ff, bu = ln, size, color, f, bullet
        pg = para(tf, i == 0, align, spacing, bu, marL if bu else 0, after)
        parts = text.split("\v")
        for j, part in enumerate(parts):
            if j:
                pg._p.append(pg._p.makeelement(qn("a:br"), {}))
            run(pg, part, sz, col, ff)
    return tf


def textbox(x, y, w, h, lines, size, color, f, anchor=MSO_ANCHOR.TOP, wrap=True, **kw):
    tb = s.shapes.add_textbox(I(x), I(y), I(w), I(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    lines_to(tf, lines, size, color, f, **kw)
    return tb


def shape(kind, x, y, w, h, fill=None, line=None, lw=0.75):
    sh = s.shapes.add_shape(kind, I(x), I(y), I(w), I(h))
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(lw)
    sh.shadow.inherit = False
    tf = sh.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return sh


def place(sh, x=None, y=None, w=None, h=None):
    if x is not None: sh.left = I(x)
    if y is not None: sh.top = I(y)
    if w is not None: sh.width = I(w)
    if h is not None: sh.height = I(h)


def kill(*ids):
    for i in ids:
        if i in by:
            by[i]._element.getparent().remove(by[i]._element)
            del by[i]


def inch(v):
    return Emu(v).inches


def fit_pic(pic, x, y, w=None, h=None):
    """비율 유지: w 또는 h 하나로 맞추고 (x,y)를 가운데 기준점으로."""
    ratio = pic.width / pic.height
    if w is not None:
        nw, nh = I(w), int(I(w) / ratio)
    else:
        nh, nw = I(h), int(I(h) * ratio)
    pic.width, pic.height = nw, nh
    pic.left = I(x) - nw // 2
    pic.top = I(y) - nh // 2


def set_label(sh, text, size, color, f, align=PP_ALIGN.LEFT):
    tf = clear(sh)
    tf.paragraphs[0].alignment = align
    run(tf.paragraphs[0], text, size, color, f)


# ───────────── 0) 전략 1 배지 + ACT 문구 ─────────────
act1, lab1 = by[121], by[122]
badge_el = copy.deepcopy(act1._element)
act1._element.addprevious(badge_el)
badge = [sh for sh in s.shapes if sh._element is badge_el][0]
badge_el.find(qn("p:nvSpPr")).find(qn("p:cNvPr")).set("name", "전략 배지")
badge_el.find(qn("p:nvSpPr")).find(qn("p:cNvPr")).set("id", "990")
badge.fill.solid(); badge.fill.fore_color.rgb = RED
place(badge, x=0.30, w=0.86)
tf = clear(badge)
tf.paragraphs[0].alignment = PP_ALIGN.CENTER
run(tf.paragraphs[0], "전략 ", 10, WHITE, F3)
run(tf.paragraphs[0], "1", 13.5, WHITE, F4)
place(act1, x=1.22)
place(lab1, x=2.26, w=6.0)
set_label(lab1, "시스템 구축과 운영을 모두 수행한 핵심인력 투입", 11.5, NAVY, F3)
set_label(by[165], "철저한 장애관리로 장애율 Zero 유지", 11.5, NAVY, F3)

# ───────────── 1) ACT.1 인물 (사진 축소·비율 유지) ─────────────
place(by[123], x=0.42, y=2.04, w=2.84, h=2.06)          # 인물 패널을 칸 전체로
fit_pic(by[73], x=0.42 + 0.12 + 0.43, y=3.02, w=0.86)   # 1.17 → 0.86 (비율 유지)
place(by[124], x=1.52, y=2.17, w=1.70, h=0.32)
set_label(by[124], "", 16, NAVY, F4)
tf = by[124].text_frame; pg = tf.paragraphs[0]
for r in list(pg.runs): r._r.getparent().remove(r._r)
run(pg, "이도훈 ", 16, NAVY, F4); run(pg, "차장", 10.5, NAVY, F3)
place(by[125], x=1.52, y=2.52, w=1.70, h=0.22)
set_label(by[125], "유지관리 담당자", 10, INK, F3)
# 운영사업 18년 (18년 강조)
t18 = textbox(1.52, 2.76, 1.70, 0.30, [], 9, INK, F2, anchor=MSO_ANCHOR.MIDDLE, wrap=False)
pg = t18.text_frame.paragraphs[0]
run(pg, "운영사업 ", 9.5, INK, F2); run(pg, "18", 15, PINK, F4); run(pg, "년", 10, PINK, F4)
place(by[126], x=1.52, y=3.10, w=1.62, h=0.27)
set_label(by[126], "현재 유지관리 담당자", 9, PINK, F3, align=PP_ALIGN.CENTER)
by[126].fill.solid(); by[126].fill.fore_color.rgb = RGBColor.from_string("FCE4EC")
# 로고: 사진 아래
logo = by[3]
fit_pic(logo, x=2.20, y=3.66, w=0.92)

# ───────────── 2) ACT.1 주요 이력 (원본 문구 그대로) ─────────────
kill(*range(132, 147))   # 타임라인 점·선·짧은 문구
place(by[127], x=3.34)
place(by[129], x=3.42, w=4.02)
place(by[130], x=3.53)
place(by[131], x=3.79, w=2.5)
set_label(by[131], "주요 이력", 10, NAVY, F3)
career = [
    "행정포털지원 시스템 유지관리 경험\v(2006년~2009년, 2012년~현재)",
    "행정정보공유 구축 경험 (2010년)",
    "적용 솔루션(아키하드, 아키뷰어, FMS) 사용 능통",
    "31개 연계시스템의 기관 담당자,\v업체 담당자와의 채널 지속 유지",
]
textbox(3.50, 2.44, 3.92, 1.62, career, 10, RGBColor.from_string("333333"), F2,
        bullet="•", marL=0.14, spacing=1.1, after=6)

# ───────────── 3) ACT.1 흐름 (아이콘↓ 글자↑, 세로 3단 + 납작한 셰브런) ─────────────
place(by[128], x=7.52)
place(by[147], x=7.62, w=2.82)
place(by[148], x=7.70)
place(by[149], x=7.97, w=2.3)
set_label(by[149], "핵심 프로세스", 10, NAVY, F3)
place(by[150], x=7.62, y=2.34, w=2.82, h=1.76)
flow = [(151, 152, "현 운영담당자 업무 지속"), (153, 154, "인수위험 Zero"), (155, 156, "안정적인 유지보수 업무 수행")]
ys = [2.43, 2.99, 3.55]
for (pic, box, text), y in zip(flow, ys):
    fit_pic(by[pic], x=7.95, y=y + 0.20, h=0.36)
    place(by[box], x=8.24, y=y, w=2.12, h=0.40)
    set_label(by[box], text, 10.5, WHITE, F3, align=PP_ALIGN.CENTER)
kill(157, 158)  # 기존 오른쪽 화살표 → 납작한 아래 셰브런
for y in (2.85, 3.41):
    ch = shape(MSO_SHAPE.CHEVRON, 9.24, y, 0.13, 0.12, fill=BLUE)
    ch.rotation = 90

# ───────────── 4) ACT.2 시스템 안정화 그래프 (폭·높이 축소 + 인용 띠) ─────────────
X0, X1, NX1 = 0.44, 5.09, 3.56
Y0, Y1, NY1 = 5.25, 6.70, 6.36
kx, ky = (NX1 - X0) / (X1 - X0), (NY1 - Y0) / (Y1 - Y0)


def sx(v): return X0 + (v - X0) * kx
def sy(v): return Y0 + (v - Y0) * ky


for i in (171, 172, 173, 174, 175, 176, 2, 177, 178, 179, 180, 181):
    sh = by[i]
    l, t, w, h = inch(sh.left), inch(sh.top), inch(sh.width), inch(sh.height)
    sh.left, sh.top = I(sx(l)), I(sy(t))
    sh.width, sh.height = I(w * kx), I(h * ky)
for i in (182, 183, 184, 185, 186, 187):
    sh = by[i]
    cx, cy = inch(sh.left) + inch(sh.width) / 2, inch(sh.top) + inch(sh.height) / 2
    sh.left, sh.top = I(sx(cx) - inch(sh.width) / 2), I(sy(cy) - inch(sh.height) / 2)
z = by[188]
place(z, x=sx(4.83) - 0.40, y=sy(6.30) - 0.30, w=0.44, h=0.22)
set_label(z, "Zero", 9, WHITE, F4, align=PP_ALIGN.CENTER)
place(by[170], x=0.48, y=5.00, w=0.6, h=0.2)
set_label(by[170], "장애율", 8.5, GRAY, F3)
place(by[167], w=NX1 - 0.44 + 0.0)
set_label(by[169], "시스템 안정화", 10, NAVY, F3)
q = by[189]
place(q, x=0.48, y=6.48, w=3.04, h=0.40)
q.fill.solid(); q.fill.fore_color.rgb = NAVY
set_label(q, "“지난 20여년간 운영을 통한 시스템 안정화”", 9.5, WHITE, F3, align=PP_ALIGN.CENTER)
place(by[166], x=3.65)

# ───────────── 5) ACT.2 4단계 장애율 관리 전략 (원본 구성) ─────────────
place(by[190], x=3.74, w=6.64)
place(by[191], x=3.86)
place(by[192], x=4.12, w=3.0)
set_label(by[192], "4단계 장애율 관리 전략", 10, NAVY, F3)
kill(193, 196, 197, 198, 199, 202, 203, 204, 205, 208, 209, 210, 211, 214, 215)

LX, RX, CW = 3.76, 7.62, 2.76
TOP, BH, GAP = 5.05, 0.90, 0.06
CX = (LX + CW + RX) / 2          # 가운데 원 중심 x
CY = TOP + BH + GAP / 2           # 가운데 원 중심 y
quads = [
    # (번호배지, 아이콘, 열, 행, 제목, 세부)
    (195, 194, 0, 0, "예방점검 및 교육 시행",
     [("사전예방활동(1일 3회)", 9, INK, F2, "•"), ("사후관리", 9, INK, F2, "•"),
      ("  - 장애원인분석, 해결방안 마련", 8.5, GRAY, F2, None), ("  - 사례교육 실시(기술,보안)", 8.5, GRAY, F2, None)]),
    (201, 200, 1, 0, "철저한 변경관리 수행",
     [("연계기관 수시점검", 9, INK, F2, "•"), ("업무 팀간 교차테스트 수행", 9, INK, F2, "•"),
      ("운영·테스트서버 동기화", 9, INK, F2, "•")]),
    (207, 206, 1, 1, "연계 모니터링 강화",
     [("오류발생신고", 9, INK, F2, "•"), ("원인파악", 9, INK, F2, "•"),
      ("현황안내 및 조치협조", 9, INK, F2, "•")]),
    (213, 212, 0, 1, "장애이력관리",
     [("조치사항 기록 및 유형분류", 9, INK, F2, "•"), ("유형별 인시던트 등록", 9, INK, F2, "•"),
      ("장애이력DB관리", 9, INK, F2, "•")]),
]
card_ids = []
for num, icon, col, row, title, items in quads:
    x = LX if col == 0 else RX
    y = TOP + row * (BH + GAP)
    card = shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, CW, BH, fill=PALE)
    card.adjustments[0] = 0.06
    card._element.getparent().remove(card._element)
    by[190]._element.addnext(card._element)   # 원·배지보다 뒤
    # 가운데 원 쪽(안쪽) 여백
    inner = 0.42
    tx = x + 0.12 if col == 0 else x + inner
    tw = CW - 0.12 - inner
    # 번호 배지 + 아이콘 + 제목
    nb = by[num]
    place(nb, x=tx, y=y + 0.07, w=0.24, h=0.23)
    set_label(nb, f"0{[1,2,3,4][[195,201,207,213].index(num)]}", 7.5, WHITE, F3, align=PP_ALIGN.CENTER)
    nb.fill.solid(); nb.fill.fore_color.rgb = BLUE
    fit_pic(by[icon], x=tx + 0.43, y=y + 0.185, h=0.24)
    textbox(tx + 0.62, y + 0.06, tw - 0.62, 0.25, [title], 10.5, NAVY, F4, anchor=MSO_ANCHOR.MIDDLE, wrap=False)
    textbox(tx + 0.04, y + 0.31, tw - 0.04, BH - 0.33, items, 9, INK, F2, marL=0.12, spacing=0.96 if len(items) > 3 else 1.1)

# 가운데 원: 장애발생요인 사전제거
ring = shape(MSO_SHAPE.OVAL, CX - 0.52, CY - 0.52, 1.04, 1.04, fill=BLUE)
core = shape(MSO_SHAPE.OVAL, CX - 0.40, CY - 0.40, 0.80, 0.80, fill=WHITE)
ctf = core.text_frame
ctf.word_wrap = False
ctf.vertical_anchor = MSO_ANCHOR.MIDDLE
lines_to(ctf, ["장애", "발생요인", "사전제거"], 9.5, PINK, F4, align=PP_ALIGN.CENTER, spacing=0.92)

p.save(dst)
print("saved", dst)
