# -*- coding: utf-8 -*-
"""서체 통일 2판 — 나눔 Bold 는 역할(흰 글자/크기)로 본문·강조를 가른다.
인자: <src> <dst>   제외: 노란 메모 · 분홍 '디자인' 말풍선
"""
import sys, collections
from pptx import Presentation
NS = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
src, dst = sys.argv[1], sys.argv[2]
TITLE, BODY, EMPH = "G마켓 산스 TTF Bold", "a시월구일2", "a시월구일3"
MAP = {
    "나눔스퀘어 네오 ExtraBold": EMPH,
    "a시월구일1": BODY, "a시월구일4": EMPH,
    "Pretendard Medium": BODY, "Pretendard SemiBold": EMPH, "Pretendard ExtraBold": EMPH,
    "산돌고딕B": EMPH, "다음_SemiBold": EMPH, "a타이틀고딕2": TITLE,
    "굴림": BODY, "환경B": EMPH,
}
ROLE = "나눔스퀘어 네오 Bold"
KEEP = {TITLE, BODY, EMPH, "Arial", "+mj-cs", "+mn-cs"}
p = Presentation(src)
changed = collections.Counter(); skipped = 0
def is_memo(sh):
    try:
        if sh.fill.type == 1 and str(sh.fill.fore_color.rgb) in ("FFFF00","FFFFCC","FFFF99","FFFF87"): return True
    except Exception: pass
    return "말풍선" in sh.name and sh.has_text_frame and "디자인" in sh.text_frame.text
def run_role(r):
    """흰 글자거나 12pt 이상이면 강조."""
    pr = r.find(NS+"rPr")
    sz = None
    if pr is not None and pr.get("sz"): sz = int(pr.get("sz"))/100
    white = False
    if pr is not None:
        fill = pr.find(NS+"solidFill")
        if fill is not None:
            sg = fill.find(NS+"srgbClr"); sc = fill.find(NS+"schemeClr")
            if sg is not None and (sg.get("val") or "").upper() in ("FFFFFF","FEFEFE","FFFFFE"): white = True
            if sc is not None and (sc.get("val") or "") in ("bg1","lt1"): white = True
    return EMPH if (white or (sz is not None and sz >= 12)) else BODY
def fix(txBody):
    for r in txBody.iter(NS+"r"):
        pr = r.find(NS+"rPr")
        if pr is None: continue
        for tag in ("latin","ea","cs"):
            e = pr.find(NS+tag)
            if e is None: continue
            tf = e.get("typeface")
            if not tf or tf in KEEP: continue
            new = run_role(r) if tf == ROLE else MAP.get(tf) or (BODY if tag=="cs" else None)
            if new: e.set("typeface", new); changed[(tf,new)] += 1
def walk(shapes):
    global skipped
    for sh in shapes:
        if is_memo(sh): skipped += 1; continue
        if str(sh.shape_type).startswith("GROUP"): walk(sh.shapes); continue
        if sh.has_text_frame: fix(sh.text_frame._txBody)
        if getattr(sh,"has_table",False) and sh.has_table:
            for row in sh.table.rows:
                for c in row.cells: fix(c.text_frame._txBody)
for s in p.slides: walk(s.shapes)
p.save(dst)
import io
log = io.open("font_log.txt","w",encoding="utf-8")
log.write("제외 도형 %d\n" % skipped)
for (a,b),v in changed.most_common(): log.write("%5d  %s -> %s\n" % (v,a,b))
log.close()
print("saved", dst)
