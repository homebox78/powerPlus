import sys, io
from PIL import Image
from pptx import Presentation
sys.stdout.reconfigure(encoding="utf-8")
p=Presentation(sys.argv[1]); sl=p.slides[10]
def find(container, name):
    for s in container.shapes:
        if s.name==name: return s, container
        if s.shape_type==6:
            r=find(s,name)
            if r: return r
    return None
jobs=[("모서리가 둥근 직사각형 163",1300),("모서리가 둥근 직사각형 174",1374),("모서리가 둥근 직사각형 170",1443),("모서리가 둥근 직사각형 171",1335)]
for tile,icon in jobs:
    t,g=find(sl,tile)
    kids=list(g.shapes)
    L,T,W,H=t.left,t.top,t.width,t.height
    inside=[s for s in kids if s is not t and s.left>=L-5000 and s.top>=T-5000 and s.left+s.width<=L+W+5000 and s.top+s.height<=T+H+5000]
    im=Image.open(f"lib/icon_{icon}.png").convert("RGBA"); im=im.crop(im.getchannel("A").point(lambda v:255 if v>8 else 0).getbbox())
    b=io.BytesIO(); im.save(b,"PNG"); b.seek(0)
    D=int(max(W,H)*1.12); k=D/max(im.size); w,h=int(im.width*k),int(im.height*k)
    g.shapes.add_picture(b,int(L+W/2-w/2),int(T+H/2-h/2),w,h)
    for s in [t]+inside:
        if s._element.getparent() is not None: s._element.getparent().remove(s._element)
    print(tile,"→",icon,"removed",[s.name for s in inside])
p.save(sys.argv[2])
