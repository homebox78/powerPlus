# -*- coding: utf-8 -*-
"""모든 그림을 번호 붙여 시트로. 인자: <pptx> <출력접두> [최소면적in2]"""
import sys, io, os, collections
from pptx import Presentation
from PIL import Image, ImageDraw, ImageFont
src=sys.argv[1]; pre=sys.argv[2]; MIN=float(sys.argv[3]) if len(sys.argv)>3 else 0.6
p=Presentation(src); items=[]
def walk(shapes, sn):
    for sh in shapes:
        st=str(sh.shape_type).split()[0]
        if st=="GROUP": walk(sh.shapes, sn); continue
        if st!="PICTURE": continue
        try:
            w,h=sh.width/914400, sh.height/914400
            if w*h < MIN: continue
            items.append((sn, sh, w, h))
        except Exception: pass
for i,s in enumerate(p.slides,1): walk(s.shapes,i)
os.makedirs("pics", exist_ok=True)
log=io.open(pre+"_list.txt","w",encoding="utf-8")
log.write("그림 %d개 (면적 %.1fin² 이상)\n"%(len(items),MIN))
tiles=[]
for idx,(sn,sh,w,h) in enumerate(items,1):
    try:
        blob=sh.image.blob; ext=sh.image.ext
        fp="pics/p%03d.%s"%(idx,ext)
        open(fp,"wb").write(blob)
        im=Image.open(fp).convert("RGB")
    except Exception as e:
        log.write("%3d s%02d  (열기 실패 %s)\n"%(idx,sn,e)); continue
    log.write("%3d s%02d %-22s %.2f x %.2f in  %dx%d px\n"%(idx,sn,sh.name[:22],w,h,im.width,im.height))
    tiles.append((idx,sn,im))
C=5; CW=380; CH=270
R=(len(tiles)+C-1)//C
sheet=Image.new("RGB",(C*CW,R*(CH+26)),(228,231,238)); dr=ImageDraw.Draw(sheet)
fb=ImageFont.truetype("C:/Windows/Fonts/malgunbd.ttf",20)
for i,(idx,sn,im) in enumerate(tiles):
    im.thumbnail((CW-10,CH-10), Image.LANCZOS)
    x=(i%C)*CW; y=(i//C)*(CH+26)
    sheet.paste(im,(x+(CW-im.width)//2, y+26+(CH-26-im.height)//2))
    dr.rectangle([x+4,y+3,x+108,y+25],fill=(200,30,40))
    dr.text((x+9,y+3),"%d  s%02d"%(idx,sn),fill=(255,255,255),font=fb)
for i in range(0,len(tiles),C*4):
    part=i//(C*4)+1
    sheet.crop((0,(i//C)*(CH+26),C*CW,min(R*(CH+26),(i//C+4)*(CH+26)))).save("%s_%d.png"%(pre,part))
log.close(); print("tiles",len(tiles))
