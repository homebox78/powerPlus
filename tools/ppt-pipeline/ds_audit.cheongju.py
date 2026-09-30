import sys
from collections import Counter,defaultdict
from pptx import Presentation
sys.stdout.reconfigure(encoding="utf-8")
A="{http://schemas.openxmlformats.org/drawingml/2006/main}"
P="{http://schemas.openxmlformats.org/presentationml/2006/main}"
E=914400
p=Presentation(sys.argv[1])
SK=(1,2,3,9,14,23,30,36,40)
TOK={"2F3B6F","0B2E6B","333F50","404040","D23737","1F4E79","FFFFFF","3A46A0","1973D1","0456B6","002060","000000","658EBB"}
FOK={"0B2E6B","3A46A0","0456B6","1973D1","658EBB","C0D2E6","DEEBF7","F2F7FC","FFFFFF","D23737","002060"}
FONT={"a시월구일2","a시월구일3","a시월구일4","G마켓 산스 TTF Bold"}
V=defaultdict(list)
def walk(sh):
    for s in sh:
        yield s
        if s.shape_type==6: yield from walk(s.shapes)
def ri(r):
    rp=r._r.find(A+"rPr")
    if rp is None: return None,None,None
    e=rp.find(A+"ea"); c=rp.find(A+"solidFill")
    return (e.get("typeface") if e is not None else None, int(rp.get("sz"))/100 if rp.get("sz") else None, (c[0].get("val").upper() if c is not None and len(c) else None))
def memo(s):
    x=s._element.xml
    return any(k in x for k in('FFFF00','FFFFCC','FFFF99'))
for i,sl in enumerate(p.slides,1):
    km=0;hd=0
    for s in walk(sl.shapes):
        if memo(s) or s.left/E>=10.8: continue
        x,y,w,h=s.left/E,s.top/E,s.width/E,s.height/E
        if s.shape_type==19:
            for rix,row in enumerate(s.table.rows):
                for c in row.cells:
                    for pg in c.text_frame.paragraphs:
                        for r in pg.runs:
                            if not r.text.strip(): continue
                            f,sz,col=ri(r)
                            if f and f not in FONT: V["서체 밖"].append((i,f))
                            if f is None: V["표 서체 상속"].append(i)
                            if rix==0 and f=="a시월구일2": V["표 머리가 본문 서체"].append((i,s.name))
                            if col and col not in TOK: V["글자색 밖"].append((i,col))
            continue
        if s.shape_type==13: continue
        sp=s._element.find(P+"spPr")
        if sp is not None:
            c=sp.find(A+"solidFill/"+A+"srgbClr")
            if c is not None and c.get("val").upper() not in FOK: V["채움색 밖"].append((i,c.get("val").upper()))
            if sp.find(A+"effectLst") is not None and len(sp.find(A+"effectLst")) and s.has_text_frame and s.text_frame.text.strip() and sp.find(A+"solidFill") is None and sp.find(A+"gradFill") is None:
                V["글상자 도형 효과"].append((i,s.name))
        if not s.has_text_frame or not s.text_frame.text.strip(): continue
        iskm = i not in SK and abs(y-1.03)<0.08 and w>8
        ishd = i not in SK and y<0.9 and x<1 and abs(y-0.35)<0.06
        if iskm:
            km+=1
            if abs(x-0.20)>0.011 or abs(w-10.43)>0.02: V["키메시지 위치"].append((i,round(x,2),round(w,2)))
            for pg in s.text_frame.paragraphs:
                if pg.alignment is None or int(pg.alignment)!=2: V["키메시지 정렬"].append(i)
        for pg in s.text_frame.paragraphs:
            for r in pg.runs:
                if not r.text.strip(): continue
                f,sz,col=ri(r)
                if f and f not in FONT: V["서체 밖"].append((i,f))
                if f is None: V["서체 상속"].append((i,s.name[:14],r.text[:8]))
                if col and col not in TOK: V["글자색 밖"].append((i,col))
                if iskm:
                    if f not in("G마켓 산스 TTF Bold",) : V["키메시지 서체"].append((i,f,r.text[:8]))
                    if sz!=24: V["키메시지 크기"].append((i,sz,r.text[:8]))
                    if col not in("0B2E6B","D23737"): V["키메시지 색"].append((i,col,r.text[:8]))
                if ishd and (f!="a시월구일4" or sz!=17): V["머리 제목"].append((i,f,sz))
        if ishd: hd+=1
    if i not in SK and i!=41 and km==0: V["키메시지 없음(변형)"].append(i)
for k in sorted(V):
    v=V[k]; c=Counter(v)
    print(k,len(v),c.most_common(25))
