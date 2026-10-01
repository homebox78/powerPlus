# -*- coding: utf-8 -*-
"""내 판과 사용자가 손으로 고친 판을 도형 단위로 비교 → 무엇을 고쳤는지 json. 인자: <내판> <사용자판> <out.json> (분류는 hand_cat 로)"""
import sys, json
sys.stdout.reconfigure(encoding="utf-8")
sys.path.insert(0,"tools/ppt-pipeline")
from pptx import Presentation
from group_title_lib import walk
I=914400
def snap(path):
    p=Presentation(path); out=[]
    for n,sl in enumerate(p.slides,1):
        d={}
        for s,x,y,w,h,*_ in walk(sl.shapes):
            key=(s.shape_id,s.name)
            info={"x":x/I,"y":y/I,"w":w/I,"h":h/I,"type":str(s.shape_type)}
            if getattr(s,"has_text_frame",False) and s.text_frame.text.strip():
                tf=s.text_frame; info["text"]=tf.text
                runs=[r for pg in tf.paragraphs for r in pg.runs]
                info["sizes"]=sorted({r.font.size.pt for r in runs if r.font.size})
                fonts=set()
                for r in runs:
                    rp=r._r.find("{http://schemas.openxmlformats.org/drawingml/2006/main}rPr")
                    if rp is not None:
                        e=rp.find("{http://schemas.openxmlformats.org/drawingml/2006/main}ea")
                        if e is not None: fonts.add(e.get("typeface"))
                info["fonts"]=sorted(fonts)
                cols=set()
                for r in runs:
                    try: cols.add(str(r.font.color.rgb))
                    except: pass
                info["cols"]=sorted(cols)
                ls=set()
                for pg in tf.paragraphs:
                    try: ls.add(pg.line_spacing)
                    except: pass
                info["ls"]=sorted(str(v) for v in ls)
            d[str(key)]=info
        out.append(d)
    return out
a=snap(sys.argv[1]); b=snap(sys.argv[2])
res=[]
for n,(da,db) in enumerate(zip(a,b),1):
    for k in set(da)|set(db):
        if k not in db: res.append((n,"삭제",k,da[k].get("text","")[:30])); continue
        if k not in da: res.append((n,"추가",k,db[k].get("text","")[:30],db[k]["type"])); continue
        A,B=da[k],db[k]; ch=[]
        for f in ("x","y","w","h"):
            if abs(A[f]-B[f])>0.02: ch.append(f"{f} {A[f]:.2f}→{B[f]:.2f}")
        for f in ("text","sizes","fonts","cols","ls"):
            if A.get(f)!=B.get(f): ch.append(f"{f} {str(A.get(f))[:40]}→{str(B.get(f))[:40]}")
        if ch: res.append((n,"변경",k,(A.get("text") or "")[:24].replace("\n","/"),"; ".join(ch)))
json.dump(res,open(sys.argv[3],"w",encoding="utf-8"),ensure_ascii=False,indent=0)
from collections import Counter
print(len(res), Counter(r[1] for r in res), Counter(r[0] for r in res).most_common(41))
