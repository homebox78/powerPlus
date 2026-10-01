"""ACT 배지 통일(11~14쪽 머리 배지 5개) + 10쪽 목록 ACT 라벨 색. 인자: <src> <dst>"""
import sys
from lxml import etree
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
sys.stdout.reconfigure(encoding="utf-8")
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
IN = 914400
PINK = "EC1C68"
W, H, LEFT = 0.9, 0.31, 0.30
src, dst = sys.argv[1:3]
p = Presentation(src)


def style(sh, size_pt, pill=True):
    sp = sh._element.spPr
    for tag in ("solidFill", "gradFill", "noFill", "pattFill", "ln", "effectLst", "prstGeom", "custGeom"):
        for e in sp.findall(A + tag):
            sp.remove(e)
    geom = etree.SubElement(sp, A + "prstGeom", prst="roundRect")
    av = etree.SubElement(geom, A + "avLst")
    etree.SubElement(av, A + "gd", name="adj", fmla="val 50000")
    fill = etree.SubElement(sp, A + "solidFill")
    etree.SubElement(fill, A + "srgbClr", val=PINK)
    ln = etree.SubElement(sp, A + "ln")
    etree.SubElement(ln, A + "noFill")
    # 순서: xfrm, prstGeom, solidFill, ln (xfrm 은 맨 앞에 있음)
    st = sh._element.find("{http://schemas.openxmlformats.org/presentationml/2006/main}style")
    if st is not None:
        sh._element.remove(st)
    tf = sh.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, 0)
    txt = tf.text.strip()
    pg = tf.paragraphs[0]
    runs = list(pg.runs)
    runs[0].text = txt
    for r in runs[1:]:
        r._r.getparent().remove(r._r)
    for extra in tf.paragraphs[1:]:
        extra._p.getparent().remove(extra._p)
    pg.alignment = PP_ALIGN.CENTER
    r = pg.runs[0]
    r.font.size = Pt(size_pt)
    r.font.bold = False
    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    r.font.name = "a시월구일3"
    rpr = r._r.get_or_add_rPr()
    for e in rpr.findall(A + "ea"):
        rpr.remove(e)
    rpr.append(rpr.makeelement(A + "ea", {"typeface": "a시월구일3"}))


# 11쪽 깃발 조각 삭제
for s in list(p.slides[10].shapes):
    if s.name in ("시안 10", "시안 11", "시안 12", "시안 53", "시안 54", "시안 55"):
        s._element.getparent().remove(s._element)
        print("11 깃발 조각 삭제", s.name)

HEAD = {11: ["시안 13", "시안 56"], 12: ["시안 6"], 13: ["시안 8"], 14: ["시안 6"]}
for n, names in HEAD.items():
    for s in p.slides[n - 1].shapes:
        if s.name in names and s.has_text_frame and s.text_frame.text.strip().startswith("ACT."):
            cy = s.top + s.height / 2
            s.left, s.width, s.height = Emu(int(LEFT * IN)), Emu(int(W * IN)), Emu(int(H * IN))
            s.top = Emu(int(cy - H * IN / 2))
            style(s, 13)
            print("%d %s 배지 통일" % (n, s.text_frame.text))


def walk(shs):
    for s in shs:
        if s.shape_type == 6:
            yield from walk(s.shapes)
        else:
            yield s


for s in walk(p.slides[9].shapes):
    if s.has_text_frame and s.text_frame.text.strip().startswith("ACT."):
        style(s, 8)
        print("10 %s 색" % s.text_frame.text)
p.save(dst)
