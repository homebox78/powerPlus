"""v0.56 - slide 24 flow pills: text was left-aligned with wide insets and wrapped into an extra line.
   center the paragraphs and narrow the side insets. args: src dst"""
import sys
from pptx import Presentation
from pptx.enum.text import PP_ALIGN
sys.stdout.reconfigure(encoding="utf-8")
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
p = Presentation(sys.argv[1]); n = 0
def walk(shs):
    for s in shs:
        if s.shape_type == 6: yield from walk(s.shapes)
        else: yield s
for s in walk(p.slides[23].shapes):
    if not s.has_text_frame or not s.text_frame.text.strip(): continue
    g = s._element.spPr.find(A + "prstGeom")
    if g is None or g.get("prst") != "roundRect" or s.name not in ("Rectangle 365", "Rectangle 371"): continue
    bp = s.text_frame._txBody.bodyPr
    bp.set("lIns", "18288"); bp.set("rIns", "18288")
    for pg in s.text_frame.paragraphs:
        pg.alignment = PP_ALIGN.CENTER
        pPr = pg._p.pPr
        for k in ("marL", "indent"):
            if pPr.get(k) is not None: pPr.set(k, "0")
    n += 1
print(n); p.save(sys.argv[2])
