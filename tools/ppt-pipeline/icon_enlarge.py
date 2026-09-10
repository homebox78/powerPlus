# -*- coding: utf-8 -*-
"""내가 넣은 아이콘 중 너무 작게 들어간 것을 최소 크기까지 키운다(중심 유지).
   원본 픽토그램이 얇은 선이면 그 bbox 가 작아 아이콘도 작아진다."""
import hashlib, io, sys
from PIL import Image
from pptx import Presentation
from pptx.util import Emu, Inches
sys.path.insert(0, ".")
import map as M
sys.stdout.reconfigure(encoding="utf-8")
src, dst = sys.argv[1], sys.argv[2]
MIN = float(sys.argv[3]) if len(sys.argv) > 3 else 0.30

mine = set()
for p in set(M.KIND.values()):
    im = Image.open(p).convert("RGBA")
    b = im.getchannel("A").point(lambda v: 255 if v > 8 else 0).getbbox()
    if b: im = im.crop(b)
    buf = io.BytesIO(); im.save(buf, "PNG", optimize=True)
    mine.add(hashlib.sha1(buf.getvalue()).hexdigest())

pres = Presentation(src)
n = 0
for sno in range(12, len(pres.slides._sldIdLst) + 1):
    sl = pres.slides[sno - 1]
    def walk(lst):
        global n
        for sh in lst:
            if sh.shape_type == 6:
                walk(sh.shapes); continue
            if sh.shape_type != 13:
                continue
            try:
                h = hashlib.sha1(sh.image.blob).hexdigest()
            except Exception:
                continue
            if h not in mine:
                continue
            w, ht = Emu(sh.width).inches, Emu(sh.height).inches
            m = max(w, ht)
            if m >= MIN:
                continue
            k = min(MIN / m, 2.6)                     # 과확대 방지
            cx = Emu(sh.left).inches + w / 2
            cy = Emu(sh.top).inches + ht / 2
            nw, nh = w * k, ht * k
            sh.width, sh.height = Inches(nw), Inches(nh)
            sh.left, sh.top = Inches(cx - nw / 2), Inches(cy - nh / 2)
            n += 1
    walk(sl.shapes)
pres.save(dst)
print(f"작게 들어간 아이콘 {n}개 확대(최소 {MIN}in) → {dst}")
