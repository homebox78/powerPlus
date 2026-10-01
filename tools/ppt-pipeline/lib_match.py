"""덱 그림 ↔ 파워플러스 라이브러리 썸네일 근접 매칭 → match72.json"""
import io, os, sys, json
import numpy as np
from PIL import Image
from pptx import Presentation
sys.stdout.reconfigure(encoding="utf-8")
E=914400
def vec(im):
    im=im.convert("RGBA"); bg=Image.new("RGBA",im.size,(255,255,255,255)); bg.alpha_composite(im)
    g=bg.convert("RGB")
    # 바깥 흰/연한 여백 자르기
    a=np.asarray(g).astype(int); m=(a.min(axis=2)<235)
    ys,xs=np.where(m)
    if len(xs): g=g.crop((xs.min(),ys.min(),xs.max()+1,ys.max()+1))
    v=np.asarray(g.resize((20,20),Image.BILINEAR)).astype(float).ravel()
    return v
cache="libvec.npz"
idx=json.load(open("lib_index.json",encoding="utf-8"))
if os.path.exists(cache):
    z=np.load(cache,allow_pickle=True); ids=list(z["ids"]); M=z["M"]
else:
    ids=[];V=[]
    for it in idx:
        f="libthumb/%s.png"%it["id"]
        if os.path.exists(f):
            ids.append(it["id"]);V.append(vec(Image.open(f)))
    M=np.array(V);np.savez(cache,ids=np.array(ids),M=M)
p=Presentation(sys.argv[1])
def walk(shs,path=""):
    for s in shs:
        if s.shape_type==6: yield from walk(s.shapes,path+"/"+s.name)
        elif s.shape_type==13: yield s,path
res=[]
for n,sl in enumerate(p.slides,1):
    for s,path in walk(sl.shapes):
        try: im=Image.open(io.BytesIO(s.image.blob))
        except Exception: continue
        v=vec(im); d=np.sqrt(((M-v)**2).mean(axis=1)); j=int(d.argmin())
        res.append(dict(n=n,path=path,name=s.name,size=im.size,w=round(s.width/E,2),h=round(s.height/E,2),best=ids[j],dist=round(float(d[j]),1)))
json.dump(res,open("match72.json","w",encoding="utf-8"),ensure_ascii=False,indent=0)
for r in res:
    if r["dist"]<30: print(r["n"],r["name"],r["size"],r["best"],r["dist"])
print(len(res))
