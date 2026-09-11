# -*- coding: utf-8 -*-
"""오른쪽 화살표: 가로 방향(좌→우) 그라데이션, 좌 투명도 100% → 우 #6890CD 투명도 15%. XML 직접."""
import sys, copy; sys.stdout.reconfigure(encoding='utf-8')
from pptx import Presentation
from lxml import etree
A='http://schemas.openxmlformats.org/drawingml/2006/main'
src,dst=sys.argv[1],sys.argv[2]; p=Presentation(src)
def walk(shs):
    for s in shs:
        if s.shape_type==6: yield from walk(s.shapes)
        else: yield s
NEW=f'''<a:gradFill xmlns:a="{A}" flip="none" rotWithShape="1"><a:gsLst>
<a:gs pos="0"><a:srgbClr val="6890CD"><a:alpha val="0"/></a:srgbClr></a:gs>
<a:gs pos="100000"><a:srgbClr val="6890CD"><a:alpha val="85000"/></a:srgbClr></a:gs>
</a:gsLst><a:lin ang="0" scaled="1"/><a:tileRect/></a:gradFill>'''
n=0
for i,sl in enumerate(p.slides,1):
    for s in walk(sl.shapes):
        x=s._element
        if b'prst="rightArrow"' not in etree.tostring(x): continue
        spPr=x.find('.//{http://schemas.openxmlformats.org/presentationml/2006/main}spPr')
        if spPr is None: continue
        g=spPr.find(f'{{{A}}}gradFill')
        if g is None: continue
        new=etree.fromstring(NEW); g.getparent().replace(g,new); n+=1
p.save(dst); print('rightArrow 재설정',n)
