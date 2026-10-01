# -*- coding: utf-8 -*-
"""작은 도형(높이 0.42in·폭 2.6in 이하) 안 글자는 a시월구일4 대신 a시월구일3(강조).

사용자 규칙(2026-10-01): "작은 도형에 들어가는 폰트는 시월구일4가 사용되면 안 됨, 강조 시 시월구일3".
같이: 23쪽 점검 알약 3개는 '기존 사업 미진 사항 점검' 크기로 통일.
인자: <src> <dst> [--dry]
"""
import sys
from pptx import Presentation
from pptx.util import Pt
from pptx.oxml.ns import qn

sys.stdout.reconfigure(encoding="utf-8")
sys.path.insert(0, __file__.rsplit("\\", 1)[0] if "\\" in __file__ else ".")
from group_title_lib import walk  # noqa: E402

args = [a for a in sys.argv[1:] if not a.startswith("--")]
src, dst = args[0], args[1]
DRY = "--dry" in sys.argv
IN = 914400
prs = Presentation(src)
n = 0
for si, sl in enumerate(prs.slides, 1):
    items = list(walk(sl.shapes))
    for (s, l, t, w, h, par, sx) in items:
        if not getattr(s, "has_text_frame", False) or not s.text_frame.text.strip():
            continue
        if not (h <= 0.42 * IN and w <= 2.6 * IN):
            continue
        hit = False
        for p in s.text_frame.paragraphs:
            for r in p.runs:
                rPr = r._r.find(qn("a:rPr"))
                if rPr is None:
                    continue
                for tg in ("a:latin", "a:ea", "a:cs"):
                    e = rPr.find(qn(tg))
                    if e is not None and e.get("typeface") == "a시월구일4":
                        if not DRY:
                            e.set("typeface", "a시월구일3")
                        hit = True
        if hit:
            n += 1
    if si == 23:
        pills = [x for x in items if getattr(x[0], "has_text_frame", False)
                 and x[0].text_frame.text.strip() in ("기존 사업 미진 사항 점검", "과업 미 종료 사항 점검", "협의예정사항 점검")]
        ref = [r.font.size for x in pills if x[0].text_frame.text.strip().startswith("기존")
               for p in x[0].text_frame.paragraphs for r in p.runs if r.font.size]
        if ref and not DRY:
            for x in pills:
                for p in x[0].text_frame.paragraphs:
                    for r in p.runs:
                        r.font.size = ref[0]
        print("23쪽 점검 알약", len(pills), "→", ref[0].pt if ref else None)
print("작은 도형 시월구일4→3", n)
if not DRY:
    prs.save(dst)
