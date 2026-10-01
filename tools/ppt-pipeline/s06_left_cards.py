# -*- coding: utf-8 -*-
"""6쪽 왼쪽 남색 카드 3장: 작은 설명 삭제, 아이콘·제목 카드 세로 가운데, 제목 15% 키움. 인자: <src> <dst>"""
import sys
from pptx import Presentation
from pptx.util import Pt
sys.stdout.reconfigure(encoding="utf-8")
src, dst = sys.argv[1:3]
p = Presentation(src); s = p.slides[5]
by = {sh.shape_id: sh for sh in s.shapes}
for card, icon, title, desc in ((82, 83, 84, 85), (86, 87, 88, 89), (90, 91, 92, 93)):
    c = by[card]; cy = c.top + c.height // 2
    by[desc]._element.getparent().remove(by[desc]._element)
    ic = by[icon]; ic.top = cy - ic.height // 2
    t = by[title]; tf = t.text_frame; tf.margin_top = tf.margin_bottom = 0
    for pg in tf.paragraphs:
        for r in pg.runs:
            r.font.size = Pt(round(r.font.size.pt * 1.15 * 2) / 2)
    t.top = cy - t.height // 2
p.save(dst); print("ok")
