import sys
from pptx import Presentation
from pptx.util import Pt
sys.stdout.reconfigure(encoding="utf-8")
E=914400
F=r"D:\powerPlus\제안서\청주시_발표자료(제안요약서)_v0.24_목차색·인물방향·s16연결·s17·s19아이콘·지도.pptx"
D=r"D:\powerPlus\제안서\청주시_발표자료(제안요약서)_v0.25_s7하단채움.pptx"
p=Presentation(F); sl=p.slides[6]
g=[s for s in sl.shapes if s.shape_type==6 and len(s.shapes)==26][0]
W=int(8.0*E); g.left=int(5.415*E-W/2); g.width=W
n=0
for c in g.shapes:
    if c.has_text_frame:
        for pg in c.text_frame.paragraphs:
            for r in pg.runs:
                if r.font.size: r.font.size=Pt(r.font.size.pt+1); n+=1
print("group",g.left/E,g.width/E,"runs+1",n); p.save(D)
