import sys
from collections import Counter,defaultdict
from pptx import Presentation
sys.stdout.reconfigure(encoding="utf-8")
A="{http://schemas.openxmlformats.org/drawingml/2006/main}"
P="{http://schemas.openxmlformats.org/presentationml/2006/main}"
E=914400
p=Presentation(r"D:\powerPlus\제안서\청주시_발표자료(제안요약서)_v0.31_디자인시스템검수.pptx")
print("slide",p.slide_width/E,p.slide_height/E)
R=defaultdict(Counter)
def runinfo(r):
    rp=r._r.find(A+"rPr"); 
    if rp is None: return ("-","-","-")
    e=rp.find(A+"ea"); c=rp.find(A+"solidFill"); col="-"
    if c is not None and len(c): col=c[0].get("val")
    return (e.get("typeface") if e is not None else "-", rp.get("sz"), col)
def walk(sh):
    for s in sh:
        yield s
        if s.shape_type==6: yield from walk(s.shapes)
skip=(1,2,3,9,14,23,30,36,40,41)
for i,sl in enumerate(p.slides,1):
    for s in walk(sl.shapes):
        role=None
        if s.shape_type==19:
            t=s.table
            for ri,row in enumerate(t.rows):
                for c in row.cells:
                    tc=c._tc.find(A+"tcPr"); f="-"
                    if tc is not None:
                        sf=tc.find(A+"solidFill")
                        if sf is not None and len(sf): f=sf[0].get("val")
                    rl="표머리" if ri==0 else "표본문"
                    R[rl+" 채움"][f]+=1
                    for pg in c.text_frame.paragraphs:
                        for r in pg.runs:
                            if r.text.strip(): R[rl][runinfo(r)]+=len(r.text)
            continue
        if not s.has_text_frame or not s.text_frame.text.strip() or i in skip: continue
        y=s.top/E; x=s.left/E; w=s.width/E
        if abs(y-1.03)<0.08 and w>8: role="키메시지"
        elif y<0.9 and x<1 : role="머리 제목(y%.2f)"%round(y,1)
        elif y<0.9: role="머리 우측"
        elif y>7.05: role="쪽번호"
        if role:
            R[role+" 위치"][(round(x,2),round(y,2),round(w,2),round(s.height/E,2))]+=1
            for pg in s.text_frame.paragraphs:
                R[role+" 정렬"][str(pg.alignment)]+=1
                for r in pg.runs:
                    if r.text.strip(): R[role][runinfo(r)]+=len(r.text)
        else:
            sp=s._element.find(P+"spPr"); g=sp.find(A+"prstGeom") if sp is not None else None
            pr=g.get("prst") if g is not None else "-"
            sf=sp.find(A+"solidFill") if sp is not None else None
            fc=sf[0].get("val") if sf is not None and len(sf) else "-"
            for pg in s.text_frame.paragraphs:
                for r in pg.runs:
                    if r.text.strip():
                        f,sz,col=runinfo(r)
                        R["도형글 "+pr+" 채움"+fc][(f,sz,col)]+=len(r.text)
for k in sorted(R):
    tot=sum(R[k].values())
    if tot<150 and k.startswith("도형글"): continue
    print(k,tot,R[k].most_common(5))
