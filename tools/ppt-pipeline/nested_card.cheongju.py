"""v0.57 - a card sitting inside a group card of the same face colour gets lost.
   inner card -> white face + dark-blue line at 60% transparency. Largest first, so only one level flips. args: src dst"""
import sys, collections
from lxml import etree
from pptx import Presentation
sys.stdout.reconfigure(encoding="utf-8")
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
E = 914400
p = Presentation(sys.argv[1]); cnt = collections.Counter()
SKIP = {1, 2, 3, 9, 15, 30, 36, 40, 41}
def flat(shapes, tf=(0, 0, 1, 1)):
    ox, oy, sx, sy = tf
    for s in shapes:
        if s.width is None: continue
        if s.shape_type == 6:
            x = s._element.find(P + "grpSpPr").find(A + "xfrm"); co, ce = x.find(A + "chOff"), x.find(A + "chExt")
            kx = s.width / int(ce.get("cx")) if int(ce.get("cx")) else 1; ky = s.height / int(ce.get("cy")) if int(ce.get("cy")) else 1
            gx, gy = ox + s.left * sx, oy + s.top * sy
            yield from flat(s.shapes, (gx - int(co.get("x")) * kx * sx, gy - int(co.get("y")) * ky * sy, sx * kx, sy * ky))
        else: yield s, ox + s.left * sx, oy + s.top * sy, s.width * sx, s.height * sy
def face(s):
    sp = s._element.find(P + "spPr")
    if sp is None: return None
    sf = sp.find(A + "solidFill")
    c = sf.find(A + "srgbClr") if sf is not None else None
    if c is None or c.find(A + "alpha") is not None: return None
    v = c.get("val"); return tuple(int(v[i:i + 2], 16) for i in (0, 2, 4))
for n, sl in enumerate(p.slides, 1):
    if n in SKIP: continue
    it = []
    for z, (s, x, y, w, h) in enumerate(flat(sl.shapes)):
        f = face(s)
        if f is None or w < 0.15 * E or h < 0.15 * E: continue
        it.append([s, x, y, w, h, f, z])
    it.sort(key=lambda t: -t[3] * t[4])
    for i, t in enumerate(it):
        s, x, y, w, h, f, z = t
        if min(f) < 215 or f == (255, 255, 255) or w < 0.6 * E or h < 0.3 * E: continue
        best = None; T = 0.03 * E
        for c in it[:i]:
            if c[6] < z and c[1] - T <= x and c[2] - T <= y and c[1] + c[3] + T >= x + w and c[2] + c[4] + T >= y + h and c[3] * c[4] >= 1.5 * w * h:
                if best is None or c[3] * c[4] < best[3] * best[4]: best = c
        if best is None or max(abs(a - b) for a, b in zip(f, best[5])) > 10: continue
        sp = s._element.find(P + "spPr"); c = sp.find(A + "solidFill").find(A + "srgbClr"); c.set("val", "FFFFFF")
        ln = sp.find(A + "ln")
        if ln is None:
            ln = etree.Element(A + "ln", w="9525"); sp.insert(list(sp).index(sp.find(A + "solidFill")) + 1, ln)
        elif ln.get("w") is None or int(ln.get("w")) > 12700: ln.set("w", "9525")
        for e in list(ln):
            if e.tag in (A + "solidFill", A + "noFill", A + "gradFill", A + "pattFill"): ln.remove(e)
        lf = etree.Element(A + "solidFill"); lc = etree.SubElement(lf, A + "srgbClr", val="143A69"); etree.SubElement(lc, A + "alpha", val="40000")
        ln.insert(0, lf)
        t[5] = (255, 255, 255); cnt[n] += 1
print(sorted(cnt.items()), sum(cnt.values()))
p.save(sys.argv[2])
