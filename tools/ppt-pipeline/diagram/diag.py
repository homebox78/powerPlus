# -*- coding: utf-8 -*-
"""이미지로 박힌 도식을 도형+텍스트로 다시 그리는 빌더.
원본 그림의 픽셀 좌표로 노드·선을 적으면, 그림이 있던 자리에 같은 비율로 도형을 놓는다.
스타일은 청주시 디자인 가이드(팔레트·서체)로 고정 — 도식마다 스타일이 갈리지 않게.
"""
import sys
sys.path.insert(0, ".")
from absbox import abs_box
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from lxml import etree

NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
A = "{%s}" % NS

# 가이드 팔레트
C = dict(navy="002060", teal="255A7B", indigo="3A46A0", blue="0456B6", azure="1973D1",
         cyan="00B0F0", steel="658EBB", mist="C0D2E6", tint="DEEBF7", near="F2F7FC",
         red="EB696D", crimson="C00000", pink="FF3370", orange="F78E3F", green="008400",
         haze="A2AAC2", sky="5FC8F7",
         ink="404040", ink2="4D4D4D", slate="4E5B6F", gray="D3D3D3", silver="808080",
         white="FFFFFF", black="000000")
BODY, EMPH = "a시월구일2", "a시월구일3"
ALIGN = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}
ANCHOR = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}


def rgb(name):
    return RGBColor.from_string(C.get(name, name))


def set_font(run, face, size, color, bold=False, spc=None):
    run.font.size = Pt(size)
    if spc is not None:
        run._r.get_or_add_rPr().set("spc", str(int(spc)))
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    rpr = run._r.get_or_add_rPr()
    for tag in ("latin", "ea", "cs"):
        e = rpr.find(A + tag)
        if e is None:
            e = etree.SubElement(rpr, A + tag)
        e.set("typeface", face)


class Diagram:
    def __init__(self, slide, pic, img_w, img_h, pic_box=None):
        self.slide, self.pic = slide, pic
        x, y, w, h = pic_box or abs_box(pic)
        self.ox, self.oy = x, y
        self.sx, self.sy = w / img_w, h / img_h      # px -> in
        self.grp = slide.shapes.add_group_shape()
        self.sh = self.grp.shapes
        self.extra = []          # 그룹에 못 넣는 것(표) — finish() 에서 같은 z 자리로 옮긴다

    # 좌표 변환 (px -> EMU)
    def X(self, px): return Emu(int((self.ox + px * self.sx) * 914400))
    def Y(self, py): return Emu(int((self.oy + py * self.sy) * 914400))
    def W(self, p):  return Emu(int(p * self.sx * 914400))
    def H(self, p):  return Emu(int(p * self.sy * 914400))

    def _text(self, shp, text, size, color, face, align="c", anchor="m", margin=3, bold=False, spc=None):
        tf = shp.text_frame
        tf.word_wrap = True
        m = self.W(margin)
        tf.margin_left = tf.margin_right = m
        tf.margin_top = tf.margin_bottom = Emu(0)
        tf.vertical_anchor = ANCHOR[anchor]
        for i, ln in enumerate(text.split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = ALIGN[align]
            r = p.add_run()
            r.text = ln
            set_font(r, face, size, color, bold, spc)

    def _style(self, shp, fill, line, lw):
        if fill:
            shp.fill.solid(); shp.fill.fore_color.rgb = rgb(fill)
        else:
            shp.fill.background()
        if line:
            shp.line.color.rgb = rgb(line); shp.line.width = Pt(lw)
        else:
            shp.line.fill.background()
        shp.shadow.inherit = False

    # ── 도형 ──
    def box(self, x, y, w, h, text="", fill="navy", color="white", size=7, face=EMPH,
            line=None, shape=MSO_SHAPE.RECTANGLE, align="c", anchor="m", lw=0.75, radius=None,
            spc=None, margin=3):
        s = self.sh.add_shape(shape, self.X(x), self.Y(y), self.W(w), self.H(h))
        self._style(s, fill, line, lw)
        if radius is not None:
            s.adjustments[0] = radius
        if text:
            self._text(s, text, size, color, face, align, anchor, margin=margin, spc=spc)
        return s

    def rbox(self, x, y, w, h, text="", radius=0.12, **kw):
        return self.box(x, y, w, h, text, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=radius, **kw)

    def diamond(self, x, y, w, h, text, size=7, color="ink", face=BODY, fill="white", line="navy"):
        # 마름모 자체 글상자는 기하상 안쪽이 좁아 글이 접힌다 → 도형 위에 글상자를 따로 얹는다
        s = self.box(x, y, w, h, "", fill=fill, line=line, shape=MSO_SHAPE.DIAMOND)
        self.label(x + w * 0.12, y, w * 0.76, h, text, size=size, color=color, face=face)
        return s

    def end(self, x, y, w, h, text="종료", size=7):
        return self.box(x, y, w, h, text, fill="tint", line="blue", color="navy",
                        shape=MSO_SHAPE.OVAL, size=size, lw=1.0)

    def label(self, x, y, w, h, text, size=7, color="ink", face=BODY, align="c", anchor="m", spc=None):
        s = self.sh.add_textbox(self.X(x), self.Y(y), self.W(w), self.H(h))
        self._text(s, text, size, color, face, align, anchor, margin=0, spc=spc)
        return s

    def line(self, pts, color="steel", lw=1.0, arrow=True, dash=False, head=False):
        """pts = [(x,y),...] 픽셀. 기본은 끝에 삼각 화살촉."""
        x0, y0 = pts[0]
        ff = self.sh.build_freeform(self.X(x0), self.Y(y0), scale=1.0)
        ff.add_line_segments([(self.X(x), self.Y(y)) for x, y in pts[1:]], close=False)
        s = ff.convert_to_shape()
        s.fill.background()
        s.line.color.rgb = rgb(color); s.line.width = Pt(lw)
        ln = s._element.spPr.find(A + "ln")
        if dash:
            d = etree.SubElement(ln, A + "prstDash"); d.set("val", "dash")
        if head:
            t = etree.SubElement(ln, A + "headEnd"); t.set("type", "triangle"); t.set("w", "med"); t.set("len", "med")
        if arrow:
            t = etree.SubElement(ln, A + "tailEnd"); t.set("type", "triangle"); t.set("w", "med"); t.set("len", "med")
        s.shadow.inherit = False
        return s

    def vline(self, x, y0, y1, color="gray", lw=0.75, dash=False):
        return self.line([(x, y0), (x, y1)], color=color, lw=lw, arrow=False, dash=dash)

    def hline(self, y, x0, x1, color="gray", lw=0.75, dash=False):
        return self.line([(x0, y), (x1, y)], color=color, lw=lw, arrow=False, dash=dash)

    def table(self, x, y, w, h, rows, colw=None, size=7, head_fill="navy", head_color="white",
              body_color="ink", first_col_fill="tint", align_cols=None, row_h=None):
        n, m = len(rows), max(len(r) for r in rows)
        # python-pptx 는 그룹 안에 표를 못 만든다 → 슬라이드에 두고 finish() 에서 자리만 맞춘다
        gf = self.slide.shapes.add_table(n, m, self.X(x), self.Y(y), self.W(w), self.H(h))
        self.extra.append(gf._element)
        tbl = gf.table
        tblPr = tbl._tbl.tblPr
        tblPr.set("bandRow", "0"); tblPr.set("firstRow", "0")
        # 표 기본 스타일 참조 제거(줄무늬·테두리 기본값이 덮어쓰지 않게)
        for st in list(tblPr):
            if st.tag.endswith("tableStyleId"):
                tblPr.remove(st)
        tw = sum(colw) if colw else m
        for j in range(m):
            tbl.columns[j].width = self.W(w * (colw[j] if colw else 1) / tw)
        rh = row_h or [1] * n
        for i in range(n):
            tbl.rows[i].height = self.H(h * rh[i] / sum(rh))
            for j in range(m):
                cell = tbl.cell(i, j)
                txt = rows[i][j] if j < len(rows[i]) else ""
                head = i == 0
                firstc = j == 0 and not head and bool(first_col_fill)
                cell.fill.solid()
                cell.fill.fore_color.rgb = rgb(head_fill if head else (first_col_fill if firstc else "white"))
                cell.margin_left = cell.margin_right = Emu(int(0.04 * 914400))
                cell.margin_top = cell.margin_bottom = Emu(int(0.008 * 914400))
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                tf = cell.text_frame; tf.word_wrap = True
                al = "c" if (head or j == 0) else (align_cols or {}).get(j, "l")
                for k, ln in enumerate(txt.split("\n")):
                    p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
                    p.alignment = ALIGN[al]; p.line_spacing = 0.9
                    r = p.add_run(); r.text = ln
                    set_font(r, EMPH if (head or firstc) else BODY, size,
                             head_color if head else ("navy" if firstc else body_color))
                self._cell_border(cell)
        return gf

    def _cell_border(self, cell, color="D3D3D3", w=6350):
        tcPr = cell._tc.get_or_add_tcPr()
        for tag in ("lnL", "lnR", "lnT", "lnB"):
            old = tcPr.find(A + tag)
            if old is not None:
                tcPr.remove(old)
        # 스키마 순서: lnL, lnR, lnT, lnB 가 solidFill 등보다 앞에 와야 한다
        for idx, tag in enumerate(("lnL", "lnR", "lnT", "lnB")):
            ln = etree.Element(A + tag); ln.set("w", str(w))
            sf = etree.SubElement(ln, A + "solidFill")
            c = etree.SubElement(sf, A + "srgbClr"); c.set("val", color)
            tcPr.insert(idx, ln)

    def image(self, x, y, w, h, path):
        """비율 유지로 (x,y,w,h) 칸 안 중앙에 그림을 넣는다(알파 여백은 잘라낸다)."""
        from PIL import Image
        import io as _io
        im = Image.open(path).convert("RGBA")
        b = im.getchannel("A").point(lambda v: 255 if v > 8 else 0).getbbox()
        if b: im = im.crop(b)
        if max(im.size) > 500: im.thumbnail((500, 500))
        buf = _io.BytesIO(); im.save(buf, "PNG"); buf.seek(0)
        aw, ah = im.size
        sc = min(w / aw, h / ah)
        nw, nh = aw * sc, ah * sc
        return self.sh.add_picture(buf, self.X(x + (w - nw) / 2), self.Y(y + (h - nh) / 2),
                                   self.W(nw), self.H(nh))

    def finish(self, remove_pic=True):
        """그림이 있던 z 자리로 그룹을 옮기고 원본 그림을 지운다."""
        pe = self.pic._element
        parent = pe.getparent()
        ge = self.grp._element
        if parent.tag.endswith("spTree"):
            for el in [ge] + self.extra:
                el.getparent().remove(el)
                parent.insert(parent.index(pe), el)
        if remove_pic:
            parent.remove(pe)
        return self.grp


def find_pic(slide, shape_id):
    def walk(shapes):
        for sh in shapes:
            if str(sh.shape_type).startswith("GROUP"):
                r = walk(sh.shapes)
                if r is not None:
                    return r
            elif sh.shape_id == shape_id:
                return sh
        return None
    return walk(slide.shapes)
