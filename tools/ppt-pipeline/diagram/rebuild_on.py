# -*- coding: utf-8 -*-
"""이미 도식을 다시 그린 파일 위에서 도식만 새로 그린다.
인자: <원본그림이 있는 pptx> <대상 pptx> <저장 pptx> 모듈명...

원본 그림의 위치는 첫 파일에서 재고, 대상 파일에서는 그 자리에 있던 이전 도식
(python-pptx 가 만든 'Group N'·'Table N', 빈 그룹)을 지운 뒤 같은 z 자리에 새로 그린다.
다른 단계(서체·줄바꿈·색)를 다시 돌리지 않고 도식만 고칠 때 쓴다.
"""
import sys, re, importlib
sys.path.insert(0, ".")
from pptx import Presentation
import diag
from absbox import abs_box

src, tgt, dst = sys.argv[1:4]
P0, P1 = Presentation(src), Presentation(tgt)
_orig_find = diag.find_pic


class Stand:
    """원본 그림 대신 쓰는 자리표시 — 박스는 원본 그림, _element 는 대상 파일의 이전 도식."""
    def __init__(self, box, el):
        self.box, self._element = box, el


def find_on_target(slide, sid):
    i = [s for s in P1.slides].index(slide)
    pic = _orig_find(P0.slides[i], sid)
    x, y, w, h = abs_box(pic)
    old = []
    for sh in list(slide.shapes):
        if not re.fullmatch(r"(Group|Table) \d+", sh.name):
            continue
        if sh.shape_type == 6 and len(sh.shapes) == 0:
            old.append(sh._element); continue
        cx = (sh.left + sh.width / 2) / 914400
        cy = (sh.top + sh.height / 2) / 914400
        if x <= cx <= x + w and y <= cy <= y + h:
            old.append(sh._element)
    assert old, f"slide {i+1}: 이전 도식을 못 찾음 (sid {sid})"
    keep = old[0]
    for el in old[1:]:
        el.getparent().remove(el)
    return Stand((x, y, w, h), keep)


diag.find_pic = find_on_target
_orig_init = diag.Diagram.__init__


def _init(self, slide, pic, img_w, img_h, pic_box=None):
    _orig_init(self, slide, pic, img_w, img_h, pic_box=pic.box if isinstance(pic, Stand) else pic_box)


diag.Diagram.__init__ = _init
for m in sys.argv[4:]:
    mod = importlib.import_module(m)
    mod.find_pic = find_on_target
    mod.build(P1)
    print("rebuilt", m)
P1.save(dst)
