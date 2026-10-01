# -*- coding: utf-8 -*-
"""장표 글자 크기·줄간격·서체 분포(글자 수 비중). 기준 판과 나란히 비교. 인자: <pptx> [<pptx> ...]"""
import sys
from collections import Counter
sys.stdout.reconfigure(encoding="utf-8")
sys.path.insert(0,"tools/ppt-pipeline")
from pptx import Presentation
from pptx.oxml.ns import qn
from group_title_lib import walk
I=914400
def stats(path):
    p=Presentation(path)
    size=Counter(); lsr=Counter(); fonts=Counter(); vbreak=0; paras=0; small=0
    for n,sl in enumerate(p.slides,1):
        if n in (1,2,3,9,15,30,36,40,41): continue
        for s,x,y,w,h,*_ in walk(sl.shapes):
            if not getattr(s,"has_text_frame",False): continue
            for pg in s.text_frame.paragraphs:
                t="".join(r.text for r in pg.runs)
                if not t.strip(): continue
                paras+=1; vbreak+=("\x0b" in t) or ("\v" in t)
                rs=[r for r in pg.runs if r.font.size]
                if not rs: continue
                sz=max(r.font.size.pt for r in rs); size[sz]+=len(t)
                if sz<7: small+=len(t)
                rp=rs[0]._r.find(qn("a:rPr")); e=rp.find(qn("a:ea")) if rp is not None else None
                fonts[e.get("typeface") if e is not None else "상속"]+=len(t)
                pPr=pg._p.find(qn("a:pPr"))
                ln=pPr.find(qn("a:lnSpc")) if pPr is not None else None
                if ln is None: lsr["기본"]+=len(t)
                elif ln.find(qn("a:spcPts")) is not None: lsr["pt %.2f×" % (int(ln.find(qn("a:spcPts")).get("val"))/100/sz)]+=len(t)
                else: lsr["%% %d" % (int(ln.find(qn("a:spcPct")).get("val"))//1000)]+=len(t)
    tot=sum(size.values())
    print(path.split("/")[-1][:40])
    print(" 크기(글자 비중):", ", ".join("%g:%.0f%%" % (k,v*100/tot) for k,v in sorted(size.items()) if v*100/tot>=1.5))
    print(" 7pt 미만 %.1f%% / 강제줄바꿈 문단 %.0f%%" % (small*100/tot, vbreak*100/paras))
    print(" 줄간격:", ", ".join("%s:%.0f%%" % (k,v*100/tot) for k,v in lsr.most_common(8)))
    print(" 서체:", ", ".join("%s:%.0f%%" % (k,v*100/tot) for k,v in fonts.most_common(6)))
for a in sys.argv[1:]: stats(a)
