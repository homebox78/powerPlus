# -*- coding: utf-8 -*-
"""찌그러진 그림 비율 복원 — 화면에 보이는 가로세로 비(그룹 배율 포함)를 그림 원래 비로.
   긴 쪽을 줄여 원래 칸 안에 들어가게 하고 가운데는 유지. 자르기(srcRect)가 있으면 보이는 부분 기준.
   인자: <src> <dst> <targets.json>  targets = [[슬라이드, "도형이름", 그림순번, ...], ...]  (없으면 전체 그림 중 8% 넘게 어긋난 것)"""
import io, json, sys
from PIL import Image
from pptx import Presentation
sys.stdout.reconfigure(encoding="utf-8")
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
src, dst = sys.argv[1:3]
T = json.load(open(sys.argv[3], encoding="utf-8")) if len(sys.argv) > 3 else None
p = Presentation(src)
def gsc(g):
    x = g._element.find(".//" + A + "xfrm"); e, c = x.find(A + "ext"), x.find(A + "chExt")
    return (int(e.get("cx")) / max(1, int(c.get("cx"))), int(e.get("cy")) / max(1, int(c.get("cy")))) if c is not None else (1, 1)
def walk(sh, sx=1, sy=1):
    for s in sh:
        if s.shape_type == 6:
            gx, gy = gsc(s); yield from walk(s.shapes, sx * gx, sy * gy)
        else: yield s, sx, sy
n = 0
for i, sl in enumerate(p.slides, 1):
    cnt = {}
    for s, sx, sy in walk(sl.shapes):
        if s.shape_type != 13: continue
        k = cnt.get(s.name, 0); cnt[s.name] = k + 1
        if T is not None and not any(t[0] == i and t[1] == s.name and t[2] == k for t in T): continue
        im = Image.open(io.BytesIO(s.image.blob))
        ia = im.width * (1 - s.crop_left - s.crop_right) / max(1e-6, im.height * (1 - s.crop_top - s.crop_bottom))
        w, h = s.width * sx, s.height * sy
        d = (w / h) / ia
        if abs(d - 1) <= (0.03 if T is not None else 0.08): continue
        cx, cy = s.left + s.width / 2, s.top + s.height / 2
        if d > 1:  # 너무 넓다 → 가로 줄임
            nw = h * ia / sx; s.width = int(nw)
        else:
            nh = w / ia / sy; s.height = int(nh)
        s.left, s.top = int(cx - s.width / 2), int(cy - s.height / 2)
        n += 1; print(i, s.name, k, round(d, 2))
p.save(dst); print("복원", n)
