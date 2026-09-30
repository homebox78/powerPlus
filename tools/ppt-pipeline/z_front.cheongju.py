import sys
from pptx import Presentation
sys.stdout.reconfigure(encoding="utf-8")
A="{http://schemas.openxmlformats.org/drawingml/2006/main}"
P="{http://schemas.openxmlformats.org/presentationml/2006/main}"
E=914400
APPLY=len(sys.argv)>1
D=r"D:\powerPlus\제안서\청주시_발표자료(제안요약서)_v0.35_가려진도형앞으로.pptx"
p=Presentation("l34.pptx")
def opaque(s):
    if s.shape_type==6: return any(opaque(c) for c in s.shapes)
    if s.shape_type in(13,19): return True
    sp=s._element.find(P+"spPr")
    if sp is None: return False
    return sp.find(A+"solidFill") is not None or sp.find(A+"gradFill") is not None
for i,sl in enumerate(p.slides,1):
    sh=list(sl.shapes); mv=[]
    for k,s in enumerate(sh):
        w,h=s.width/E,s.height/E
        if not(0.12<=w<=0.85 and 0.12<=h<=0.85) or s.left/E>10.8: continue
        a=w*h
        for t in sh[k+1:]:
            if t.width*t.height/E/E<a*4 or not opaque(t) or t.left/E>10.8: continue
            ix=min(s.left+s.width,t.left+t.width)-max(s.left,t.left); iy=min(s.top+s.height,t.top+t.height)-max(s.top,t.top)
            if ix>0 and iy>0 and ix*iy/E/E>a*0.1 and ix*iy/E/E<a*0.98 if True else 0:
                mv.append(s); print(i,s.name,int(s.shape_type or 0),f"x{s.left/E:.2f} y{s.top/E:.2f} {w:.2f}x{h:.2f}","<-",t.name,int(t.shape_type or 0)); break
    if APPLY and i!=8:
        tree=sl.shapes._spTree
        for s in mv: tree.remove(s._element); tree.append(s._element)
if APPLY: p.save(D)
