"""v0.58 - slide 16: drop the top-right illustration, centre headline block, gap above the labels,
   card titles 3 lines -> 2 lines, left titles a시월구일4 / right titles a시월구일3. args: src dst"""
import sys
from lxml import etree
from pptx import Presentation
from pptx.util import Emu
sys.stdout.reconfigure(encoding="utf-8")
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
E = 914400
p = Presentation(sys.argv[1]); sl = p.slides[15]
CX = p.slide_width / 2
TITLES = {"꼼꼼한": ["꼼꼼한", "유지관리 체계"], "전문적": ["전문적", "수행 조직 지원"], "흔들림": ["흔들림 없는", "인수인계 유지"],
          "시스템": ["시스템 연계(31개)", "유지관리"], "안정적": ["안정적인", "비상상황 대응관리"], "철저한": ["철저한", "보안관리"],
          "신속한": ["신속한", "환경변화 대응"], "실질적": ["실질적인", "사용자 교육 관리"]}
def font(s, name):
    for r in s._element.iter(A + "r"):
        rp = r.find(A + "rPr")
        if rp is None:
            rp = etree.Element(A + "rPr"); r.insert(0, rp)
        for t in ("latin", "ea", "cs"):
            for e in rp.findall(A + t): rp.remove(e)
        rp.append(etree.Element(A + "latin", typeface=name)); rp.append(etree.Element(A + "ea", typeface=name))
        if rp.get("b"): rp.set("b", "0")
def retitle(s, left):
    ps = s.text_frame.paragraphs
    key = ps[0].text.strip()[:3]
    lines = TITLES[key]
    for pg, txt in zip(ps, lines):
        rs = pg.runs
        rs[0].text = txt
        for r in rs[1:]: pg._p.remove(r._r)
    for pg in ps[2:]: pg._p.getparent().remove(pg._p)
    bp = s.text_frame._txBody.bodyPr
    bp.set("lIns", "0"); bp.set("rIns", "0")
    font(s, "a시월구일4" if left else "a시월구일3")
for s in list(sl.shapes):
    if s.name == "SW13": s._element.getparent().remove(s._element); continue
    if s.name == "Text Box 38" and s.top < 2 * E:                 # sub headline
        s.left = int(CX - s.width / 2); s.top = int(1.62 * E)
    elif s.name == "Rounded Rectangle 171":
        dx = int(CX - s.width / 2) - s.left; s.left += dx; s.top = int(2.6 * E)
    elif s.shape_type == 6:
        left = s.name == "그룹 2"
        for c in s.shapes:
            if c.name.startswith("양쪽 모서리"): retitle(c, left)
DX = int(CX - (1.6 + 6.12 / 2) * E); DY = int((2.6 - 2.94) * E)
for s in sl.shapes:
    if (s.name == "Text Box 38" and s.top > 2 * E) or s.name in ("Rectangle 172", "Rectangle 173"):
        s.left += DX; s.top += DY
p.save(sys.argv[2]); print("ok")
