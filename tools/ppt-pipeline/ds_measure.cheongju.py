import sys,hashlib,json
from collections import Counter,defaultdict
from pptx import Presentation
from pptx.util import Pt
sys.stdout.reconfigure(encoding="utf-8")
A="{http://schemas.openxmlformats.org/drawingml/2006/main}"
E=914400
p=Presentation(r"D:\powerPlus\제안서\청주시_발표자료(제안요약서)_v0.30_컨소시엄표.pptx")
def walk(sh):
    for s in sh:
        yield s
        if s.shape_type==6: yield from walk(s.shapes)
def memo(s):
    x=s._element.xml
    return any(k in x for k in('FFFF00','FFFFCC','FFFF99')) or s.left/E>10.9
font=Counter();fsz=defaultdict(Counter);size=Counter();tcol=Counter();fill=Counter();lw=Counter()
V=defaultdict(list)
PAL={"002060","3A46A0","0456B6","1973D1","658EBB","C0D2E6","DEEBF7","F2F7FC","C00000","F71148","222A35","FFFFFF","000000","1F4E79"}
for i,sl in enumerate(p.slides,1):
    hs=Counter()
    for s in walk(sl.shapes):
        if memo(s): continue
        tfs=[]
        if s.shape_type==19: tfs=[c.text_frame for c in s.table.iter_cells()]
        elif s.has_text_frame: tfs=[s.text_frame]
        for tf in tfs:
            for pg in tf.paragraphs:
                for r in pg.runs:
                    if not r.text.strip(): continue
                    rp=r._r.find(A+"rPr"); ea=lat=None;sz=None
                    if rp is not None:
                        e=rp.find(A+"ea"); l=rp.find(A+"latin")
                        ea=e.get("typeface") if e is not None else None
                        lat=l.get("typeface") if l is not None else None
                        sz=rp.get("sz")
                        c=rp.find(A+"solidFill/"+A+"srgbClr")
                        if c is not None: tcol[c.get("val").upper()]+=len(r.text)
                        if rp.get("b")=="1": V["합성 굵게"].append(i)
                    f=ea or lat or "(상속)"; font[f]+=len(r.text)
                    if sz:
                        v=int(sz)/100; size[v]+=len(r.text); fsz[f][v]+=len(r.text)
                        if v!=int(v): V["소수 크기"].append((i,v,r.text[:10]))
                        if v<7: V["7pt 미만"].append((i,v,r.text[:10]))
        if s.shape_type!=13 and s.shape_type!=19:
            sp=s._element.find("{http://schemas.openxmlformats.org/presentationml/2006/main}spPr")
            if sp is not None:
                c=sp.find(A+"solidFill/"+A+"srgbClr")
                if c is not None: fill[c.get("val").upper()]+=1
                ln=sp.find(A+"ln")
                if ln is not None and ln.get("w") and ln.find(A+"noFill") is None: lw[round(int(ln.get("w"))/12700,2)]+=1
        if s.shape_type==13:
            try:
                im=s.image; h=hashlib.md5(im.blob).hexdigest(); hs[h]+=1
                vis=1-s.crop_left-s.crop_right
                ppi=im.size[0]*vis/(s.width/E) if s.width else 0
                if ppi<150 and s.width/E>0.5: V["저해상 그림(<150ppi)"].append((i,s.name,round(ppi),round(s.width/E,2)))
            except Exception as e: pass
        # 가이드 밖
        if i not in(1,2,3,9,14,23,30,36,40,41) and s.shape_type!=6 and s.top/E>1.0:
            if s.left/E<0.27 or (s.left+s.width)/E>10.56: V["좌우 가이드 밖"].append((i,s.name,round(s.left/E,2),round((s.left+s.width)/E,2)))
    for h,n in hs.items():
        if n>1: V["같은 장 그림 중복"].append((i,n))
tot=sum(font.values())
print("FONT",[(k,round(v/tot*100,1)) for k,v in font.most_common(8)])
for f in list(font)[:6]: print(" ",f,sorted(fsz[f].items(),key=lambda x:-x[1])[:8])
print("SIZE",sorted(size.items()))
print("TCOL",tcol.most_common(14))
print("FILL",fill.most_common(16))
print("LW",lw.most_common(8))
for k,v in V.items():
    print("V",k,len(v), (Counter(v).most_common(12) if k=="합성 굵게" else v[:14]))
