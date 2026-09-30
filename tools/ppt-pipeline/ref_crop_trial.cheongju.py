import sys,glob,os
from PIL import Image
import numpy as np
from scipy import ndimage
from pptx import Presentation
from pptx.util import Emu
sys.stdout.reconfigure(encoding="utf-8")
E=914400
src=sorted(glob.glob(r"D:\powerPlus\제안서\톤 및 구성 디자인 시안\*.png"))[5]
im=Image.open(src).convert("RGB"); k=im.width/1500
def cut(box,name):
    c=im.crop(tuple(int(v*k) for v in box)); a=np.asarray(c).astype(int)
    bg=(a.min(2)>236)&((a.max(2)-a.min(2))<14)
    lab,_=ndimage.label(bg)
    edge=set(lab[0,:])|set(lab[-1,:])|set(lab[:,0])|set(lab[:,-1]); edge.discard(0)
    m=np.isin(lab,list(edge))
    al=np.where(m,0,255).astype(np.uint8)
    al=np.asarray(Image.fromarray(al).filter(__import__("PIL.ImageFilter",fromlist=["x"]).GaussianBlur(1.2)))
    al=np.where(m,np.minimum(al,90),al).astype(np.uint8)
    o=Image.fromarray(np.dstack([a.astype(np.uint8),al]),"RGBA"); o=o.crop(o.getbbox()); o.save(name); return o.size
boxes=[(55,545,215,700),(255,545,445,700),(470,550,645,700),(690,550,835,700),(865,560,1000,690),(1030,550,1150,700),(1185,555,1315,695),(1340,555,1465,700)]
sz=[cut(b,f"ti{i}.png") for i,b in enumerate(boxes)]
ps=cut((515,92,935,322),"tpair.png"); print(sz,ps)
p=Presentation("t37.pptx"); sl=p.slides[15]
for s in list(sl.shapes):
    if s.shape_type==13 and s.left>8.3*E and s.top<3.2*E:
        print("del",s.name); s._element.getparent().remove(s._element)
w=2.9; h=w*ps[1]/ps[0]
sl.shapes.add_picture("tpair.png",Emu(int((10.5-w)*E)),Emu(int(1.5*E)),Emu(int(w*E)))
X0=0.30;W1=1.265;G=0.05;W2=1.15;X1=X0+3*W1+2*G+0.30
cx=[X0+i*(W1+G)+W1/2 for i in range(3)]+[X1+i*(W2+G)+W2/2 for i in range(5)]
H=0.58
for i,c in enumerate(cx):
    wi=H*sz[i][0]/sz[i][1]
    sl.shapes.add_picture(f"ti{i}.png",Emu(int((c-wi/2)*E)),Emu(int((6.93-H)*E)),height=Emu(int(H*E)))
p.save("t37b.pptx")
