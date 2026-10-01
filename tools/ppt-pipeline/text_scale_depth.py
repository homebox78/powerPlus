"""v0.72 글자: 37쪽 목록 정렬 → 본문 글자 10% 축소(28·29쪽 -2pt) → 깊이별 글자색 → 흰 글자 세로 가운데 보정. 인자: <src> <dst>"""
import sys
from pptx import Presentation
from pptx.util import Pt
sys.stdout.reconfigure(encoding="utf-8")
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
I = 914400
src, dst = sys.argv[1:3]
p = Presentation(src)
SKIP_SLIDES = {1, 2, 3, 9, 15, 30}
MINUS2 = {28, 29}
BODY = "46566E"   # 2단계: 일반 본문
SUB = "6E7D93"    # 3단계: 하위 항목·괄호 보충


def walk(shs):
    for s in shs:
        if s.shape_type == 6:
            yield from walk(s.shapes)
        else:
            yield s


def frames(s):
    if s.has_text_frame:
        yield s.text_frame
    if s.shape_type == 19 and s.has_table:
        for row in s.table.rows:
            for c in row.cells:
                yield c.text_frame


# ---- 37쪽: 오른쪽 목록을 왼쪽과 같은 시작(위)·같은 크기 ----
s37 = p.slides[36]
L = next(s for s in s37.shapes if s.name == "시안 7")
R = next(s for s in s37.shapes if s.name == "시안 12")
for s in (L, R):
    s.text_frame._txBody.find(A + "bodyPr").set("anchor", "t")
R.top = L.top
for pg in R.text_frame.paragraphs:
    for r in pg.runs:
        r.font.size = Pt(8.8)
# 문단 간격도 왼쪽과 같게
lp = L.text_frame.paragraphs[0]._p.find(A + "pPr")
for pg in R.text_frame.paragraphs:
    rp = pg._p.find(A + "pPr")
    for tag in ("lnSpc", "spcBef", "spcAft"):
        for e in rp.findall(A + tag):
            rp.remove(e)
        le = lp.find(A + tag)
        if le is not None:
            import copy
            rp.insert(0, copy.deepcopy(le))
print("37 목록 위 맞춤")


def rgb(v):
    return int(v[:2], 16), int(v[2:4], 16), int(v[4:], 16)


def is_dark(v):
    r, g, b = rgb(v)
    lum = 0.3 * r + 0.59 * g + 0.11 * b
    return lum < 110 and (b - r) < 140 and (r - b) < 60  # 짙은 남색·검정 계열(선명한 파랑·빨강 제외)


def run_color(rpr):
    if rpr is None:
        return None
    f = rpr.find(A + "solidFill")
    if f is None or not len(f):
        return None
    c = f[0]
    return c.get("val") if c.tag == A + "srgbClr" else ("scheme:" + c.get("val", ""))


def set_color(rpr, v):
    f = rpr.find(A + "solidFill")
    for e in list(f):
        f.remove(e)
    f.append(f.makeelement(A + "srgbClr", {"val": v}))


nsz = ncol = ncen = 0
for n, sl in enumerate(p.slides, 1):
    if n in SKIP_SLIDES:
        continue
    for s in walk(sl.shapes):
        if s.top is None:
            continue
        if s.name == "키메시지" or s.top < 0.95 * I or s.top > 7.05 * I:
            continue
        for tf in frames(s):
            for pg in tf.paragraphs:
                txt = pg.text.strip()
                ppr = pg._p.find(A + "pPr")
                lvl = int(ppr.get("lvl", "0")) if ppr is not None else 0
                sub = lvl >= 1 or txt[:1] in ("-", "–", "·", "※") or (txt.startswith("(") and txt.endswith(")"))
                for el in list(pg._p.findall(A + "r")) + list(pg._p.findall(A + "endParaRPr")):
                    rpr = el.find(A + "rPr") if el.tag == A + "r" else el
                    if rpr is None:
                        continue
                    sz = rpr.get("sz")
                    if sz:
                        v = int(sz)
                        if v < 1600:
                            nv = v - 200 if n in MINUS2 else int(round(v * 0.9 / 50.0) * 50)
                            nv = max(nv, 600)
                            if nv != v:
                                rpr.set("sz", str(nv)); nsz += 1
                    if el.tag != A + "r":
                        continue
                    c = run_color(rpr)
                    if c and not c.startswith("scheme") and is_dark(c) and rpr.get("b") != "1":
                        bold_font = any((e.get("typeface") or "").endswith(("3", "4", "Bold", "Bd")) for e in rpr if e.tag in (A + "latin", A + "ea"))
                        if bold_font and not sub:
                            continue
                        set_color(rpr, SUB if sub else BODY); ncol += 1
        # 흰 글자 한 줄(띠·배지): 글꼴이 아래로 앉는 만큼 위로
        if s.has_text_frame and s.height < 0.8 * I:
            tf = s.text_frame
            runs = [r for pg in tf.paragraphs for r in pg.runs if r.text.strip()]
            paras = [pg for pg in tf.paragraphs if pg.text.strip()]
            if runs and len(paras) == 1 and "\n" not in tf.text.strip() and "\v" not in tf.text:
                cols = [run_color(r._r.find(A + "rPr")) for r in runs]
                if all(c in ("FFFFFF", "scheme:bg1", "scheme:lt1") for c in cols):
                    bp = tf._txBody.find(A + "bodyPr")
                    if bp.get("anchor", "t") == "ctr":
                        sz = max(int(r._r.find(A + "rPr").get("sz", "1200")) for r in runs) / 100.0
                        add = int(0.36 * sz * 12700)
                        bp.set("bIns", str(int(bp.get("bIns", "45720")) + add)); ncen += 1
print("글자 크기", nsz, "· 색", ncol, "· 흰 글자 가운데", ncen)
p.save(dst)
