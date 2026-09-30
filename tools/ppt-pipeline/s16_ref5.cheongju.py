"""v0.60 - slide 16: card icons 15% smaller (centre kept), card titles 15px (0.156in) up. args: src dst"""
import sys
from pptx import Presentation
p = Presentation(sys.argv[1]); sl = p.slides[15]; UP = 142875; n = 0
for s in sl.shapes:
    if s.shape_type == 13 and s.name.startswith("Picture 1") and s.top > 4 * 914400:
        w, h = s.width, s.height
        s.width, s.height = int(w * 0.85), int(h * 0.85)
        s.left += (w - s.width) // 2; s.top += (h - s.height) // 2; n += 1
    elif s.shape_type == 6:
        for c in s.shapes:
            if c.name.startswith("양쪽 모서리"):
                bp = c.text_frame._txBody.bodyPr
                bp.set("bIns", str(int(bp.get("bIns", "45720")) + UP)); n += 1
print(n); p.save(sys.argv[2])
