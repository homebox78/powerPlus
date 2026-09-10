# -*- coding: utf-8 -*-
"""groups2.json 좌표 기준으로 픽토그램 도형만 교체(배경 카드는 유지)."""
import io, json, sys
from PIL import Image
from pptx import Presentation
from pptx.util import Emu, Inches
sys.path.insert(0, ".")
import map as M, map2 as M2  # map.sample.py 를 복사해 작성

sys.stdout.reconfigure(encoding="utf-8")
src, dst = sys.argv[1], sys.argv[2]
G = json.load(open("groups2.json", encoding="utf-8"))
pres = Presentation(src)
cache = {}
def load(p):
    if p not in cache:
        im = Image.open(p).convert("RGBA")
        b = im.getchannel("A").point(lambda v: 255 if v > 8 else 0).getbbox()
        if b: im = im.crop(b)
        cache[p] = im
    return cache[p]

done = 0
for sno, plan in sorted(M2.ASSIGN2.items()):
    slide = pres.slides[sno - 1]
    for no, kind in sorted(plan.items()):
        g = next((x for x in G if x["slide"] == sno and x["no"] == no), None)
        if not g:
            print(f"  ! s{sno}-{no} 없음"); continue
        x0, y0, x1, y1 = g["x"], g["y"], g["x"] + g["w"], g["y"] + g["h"]
        # bbox 안에 완전히 들어가는 글자 없는 도형만 제거(배경 카드는 더 커서 안 걸린다)
        rm = []
        def walk(lst):
            for sh in lst:
                if sh.shape_type == 6:
                    walk(sh.shapes); continue
                if sh.shape_type == 13:
                    continue
                try:
                    sx, sy = Emu(sh.left).inches, Emu(sh.top).inches
                    sw, sh_ = Emu(sh.width).inches, Emu(sh.height).inches
                except Exception:
                    continue
                if sh.has_text_frame and sh.text_frame.text.strip():
                    continue
                if sx >= x0 - .01 and sy >= y0 - .01 and sx + sw <= x1 + .01 and sy + sh_ <= y1 + .01:
                    rm.append(sh)
        walk(slide.shapes)
        if not rm:
            print(f"  ! s{sno}-{no} 대상 도형 없음"); continue
        im = load(M.KIND[kind])
        buf = io.BytesIO(); im.save(buf, "PNG", optimize=True); buf.seek(0)
        W, H = Inches(g["w"]), Inches(g["h"])
        s = min(W / im.width, H / im.height)
        nw, nh = int(im.width * s), int(im.height * s)
        for sh in rm:
            el = sh._element; el.getparent().remove(el)
        slide.shapes.add_picture(buf, int(Inches(x0) + (W - nw) / 2), int(Inches(y0) + (H - nh) / 2), nw, nh)
        done += 1
    print(f"s{sno}: {len(plan)}개")
pres.save(dst)
print(f"\n2차 교체 {done}개 → {dst}")
