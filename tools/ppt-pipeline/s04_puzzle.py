# -*- coding: utf-8 -*-
"""4쪽: 분야·업무 카드를 원본 같은 퍼즐 모양으로, 조직 하위 설명 글자 키우기. 인자: <src> <dst>"""
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
sys.stdout.reconfigure(encoding="utf-8")
src, dst = sys.argv[1:3]
p = Presentation(src); s = p.slides[3]
by = {sh.shape_id: sh for sh in s.shapes}
NAVY, BLUE, TAB, LAB, LINE = RGBColor(0x14, 0x3A, 0x69), RGBColor(0x2F, 0x78, 0xE0), RGBColor(0xE8, 0xF2, 0xFC), RGBColor(0x14, 0x3A, 0x69), RGBColor(0xBF, 0xD3, 0xEE)
box = by[87]; X, Y, W, H = [v / 914400 for v in (box.left, box.top, box.width, box.height)]
anchor = box._element
for i in range(87, 95):
    if i in by and i != 87:
        by[i]._element.getparent().remove(by[i]._element)


def font(run, size, color, face):
    run.font.size = Pt(size); run.font.color.rgb = color; run.font.bold = False
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea"):
        e = rPr.find(qn(tag))
        if e is None:
            e = rPr.makeelement(qn(tag), {}); rPr.append(e)
        e.set("typeface", face)


def shape(kind, x, y, w, h, fill, rot=0, name=""):
    sh = s.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if rot in (90, 270):   # 돌린 도형: 보이는 상자 (x,y,w,h) 가 되도록 가로세로 바꿔 가운데 고정
        cx, cy = x + w / 2, y + h / 2
        sh.width, sh.height = Inches(h), Inches(w); sh.left, sh.top = Inches(cx - h / 2), Inches(cy - w / 2)
    sh.rotation = rot
    sh.fill.solid(); sh.fill.fore_color.rgb = fill; sh.line.fill.background(); sh.shadow.inherit = False
    sh.name = name
    return sh


def text(x, y, w, h, runs, align=PP_ALIGN.CENTER, name=""):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tb.name = name
    tf = tb.text_frame; tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE; tf.word_wrap = False
    pg = tf.paragraphs[0]; pg.alignment = align
    for t, sz, c, f in runs:
        r = pg.add_run(); r.text = t; font(r, sz, c, f)
    return tb


tw = 0.40; px = X + tw - 0.06; pw = X + W - px; rh = H / 2
mid = Y + rh
new = []
new.append(shape(MSO_SHAPE.ROUND_2_SAME_RECTANGLE, X, Y, tw, H, TAB, rot=270, name="퍼즐 탭"))
new.append(shape(MSO_SHAPE.ROUND_2_SAME_RECTANGLE, px, Y, pw, rh + 0.02, NAVY, name="퍼즐 분야"))
new.append(shape(MSO_SHAPE.ROUND_2_SAME_RECTANGLE, px, mid, pw, rh, BLUE, rot=180, name="퍼즐 업무"))
d = 0.20
new.append(shape(MSO_SHAPE.OVAL, px + 0.05, mid - d / 2 + 0.01, d, d, BLUE, name="퍼즐 홈 왼"))
new.append(shape(MSO_SHAPE.OVAL, X + W - 0.05 - d, mid - d / 2 + 0.01, d, d, NAVY, name="퍼즐 홈 오른"))
ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(X + 0.06), Inches(mid), Inches(px - 0.04), Inches(mid))
ln.line.color.rgb = LINE; ln.line.width = Pt(0.75); ln.name = "퍼즐 구분선"; new.append(ln)
F2, F3 = "a시월구일2", "a시월구일3"
new.append(text(X, Y, tw - 0.04, rh, [("분야", 10, LAB, F3)], name="퍼즐 분야 라벨"))
new.append(text(X, mid, tw - 0.04, rh, [("업무", 10, LAB, F3)], name="퍼즐 업무 라벨"))
W1 = RGBColor(255, 255, 255)
new.append(text(px, Y, pw, rh, [("5", 17, W1, F3), ("개 분야", 10, W1, F3)], name="퍼즐 분야 값"))
new.append(text(px, mid, pw, rh, [("227", 17, W1, F3), ("종", 10, W1, F3)], name="퍼즐 업무 값"))
# 카드 자리(87) 바로 뒤로 순서 옮기고 카드 삭제
for sh in new:
    anchor.addnext(sh._element); anchor = sh._element
box._element.getparent().remove(box._element)
# 조직 하위 설명 글자 키우기 6.5 → 8.5pt
for i in (118, 119, 120):
    sh = by[i]; c = sh.left + sh.width // 2
    sh.width = Inches(0.90); sh.left = c - sh.width // 2; sh.height = Inches(0.60)
    for pg in sh.text_frame.paragraphs:
        for r in pg.runs:
            r.font.size = Pt(8.5)
p.save(dst); print("ok", round(X, 2), round(Y, 2), round(W, 2), round(H, 2))
