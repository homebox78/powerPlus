# -*- coding: utf-8 -*-
"""카드 상단 라벨 밴드 → 윗모서리만 둥근 사각(round2SameRect, 10px). python-pptx/XML."""
import sys
from pptx import Presentation
from pptx.util import Emu
from pptx.oxml.ns import qn
from lxml import etree
sys.stdout.reconfigure(encoding="utf-8")
src, dst = sys.argv[1:3]; IN = 914400; R_PT = 7.5
pres = Presentation(src); hit = {}
def walk(shs):
    for s in shs:
        if s.shape_type == 6: yield from walk(s.shapes)
        else: yield s
for sno, sl in enumerate(pres.slides, 1):
    shapes = [s for s in walk(sl.shapes) if s.shape_type in (1, 17)]
    for card in shapes:
        for band in shapes:
            if band is card or band.height >= card.height * 0.6 or band.height < 0.2 * IN: continue
            if abs(band.left - card.left) > 0.03 * IN or abs(band.width - card.width) > 0.03 * IN or abs(band.top - card.top) > 0.03 * IN: continue
            if band.has_text_frame and len(band.text_frame.text.splitlines()) > 2: continue
            spPr = band._element.spPr; pg = spPr.find(qn("a:prstGeom"))
            if pg is None: continue
            pg.set("prst", "round2SameRect")
            av = pg.find(qn("a:avLst"))
            if av is None: av = etree.SubElement(pg, qn("a:avLst"))
            for g in list(av): av.remove(g)
            m = min(Emu(band.width).pt, Emu(band.height).pt)
            for name, val in (("adj1", int(min(50000, R_PT / m * 100000))), ("adj2", 0)):
                gd = etree.SubElement(av, qn("a:gd")); gd.set("name", name); gd.set("fmla", f"val {val}")
            hit[sno] = hit.get(sno, 0) + 1
pres.save(dst); print("밴드 라운드:", sum(hit.values()), hit)
