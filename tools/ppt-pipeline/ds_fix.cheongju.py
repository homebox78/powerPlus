import sys,math
from collections import Counter
from pptx import Presentation
sys.stdout.reconfigure(encoding="utf-8")
A="{http://schemas.openxmlformats.org/drawingml/2006/main}"
E=914400
S=r"D:\powerPlus\제안서\청주시_발표자료(제안요약서)_v0.30_컨소시엄표.pptx"
D=r"D:\powerPlus\제안서\청주시_발표자료(제안요약서)_v0.31_디자인시스템검수.pptx"
p=Presentation(S)
TC={"4D4D4D":"404040","4E5B6F":"333F50","103574":"0B2E6B","C53030":"D23737"}
FC={"0071C5":"0456B6","0D47A1":"0456B6","0D5DB8":"0456B6"}
n=Counter(); out=[]
def walk(sh):
    for s in sh:
        yield s
        if s.shape_type==6: yield from walk(s.shapes)
for i,sl in enumerate(p.slides,1):
    for el in sl._element.iter(A+"rPr",A+"endParaRPr",A+"defRPr"):
        sz=el.get("sz")
        if sz and int(sz)%100 and int(sz)>=700 and i!=25:
            el.set("sz",str(int(sz)//100*100)); n["size"]+=1
        c=el.find(A+"solidFill/"+A+"srgbClr")
        if c is not None and c.get("val").upper() in TC: c.set("val",TC[c.get("val").upper()]); n["tcol"]+=1
    for c in sl._element.iter(A+"srgbClr"):
        if c.get("val").upper() in FC: c.set("val",FC[c.get("val").upper()]); n["fill"]+=1
    if i in(1,2,3,9,14,23,30,36,40,41): continue
    for s in walk(sl.shapes):
        if s.shape_type==6 or s.top/E<1.0: continue
        l=s.left/E; r=(s.left+s.width)/E
        if l>=10.8: continue
        if (l<0.27 or r>10.56) and not (abs(l-0.2)<0.03 and abs(r-10.63)<0.03) and not (l<0.02 and r>10.8):
            out.append((i,s.name,round(l,2),round(r,2),round(s.top/E,2)))
print(dict(n)); print(len(out)); [print(o) for o in out]
p.save(D)
