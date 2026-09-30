# 교체 그림을 원본 해상도로 재삽입 — 보이는 잉크 영역은 그 자리 그대로.
import io,json,sys,os,win32com.client as win32
from PIL import Image
from pptx import Presentation
sys.stdout.reconfigure(encoding="utf-8")
LIB=r"C:\Users\hbox7\AppData\Local\Temp\claude\d--powerPlus\818b2ce8-2260-4251-abfd-8c6a90609354\scratchpad\v16"
picks=json.load(open(LIB+r"\picks.json",encoding="utf-8"))
amap={f"SW{k:02d}":p[2] for k,p in enumerate(picks)}
app=win32.GetActiveObject("PowerPoint.Application")
for p in app.Presentations:
    if "v0.20" in p.FullName: p.SaveCopyAs(os.path.abspath("l20.pptx"))
A="{http://schemas.openxmlformats.org/drawingml/2006/main}";P="{http://schemas.openxmlformats.org/presentationml/2006/main}";R="{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
p=Presentation("l20.pptx")
def walk(shs):
    for s in shs:
        if s.shape_type==6: yield from walk(s.shapes)
        else: yield s
def chain(el):
    out=[];g=el.getparent()
    while g is not None and g.tag==P+"grpSp": out.append(g);g=g.getparent()
    return out[::-1]
def gx(g):
    x=g.find(P+"grpSpPr").find(A+"xfrm");off,ext,co,ce=[x.find(A+k) for k in("off","ext","chOff","chExt")]
    return int(off.get("x")),int(off.get("y")),int(co.get("x")),int(co.get("y")),int(ext.get("cx"))/max(1,int(ce.get("cx"))),int(ext.get("cy"))/max(1,int(ce.get("cy")))
def to_slide(ch,x,y,w,h):
    for g in reversed(ch):
        ox,oy,cx,cy,sx,sy=gx(g);x,y,w,h=ox+(x-cx)*sx,oy+(y-cy)*sy,w*sx,h*sy
    return x,y,w,h
def to_child(ch,x,y,w,h):
    for g in ch:
        ox,oy,cx,cy,sx,sy=gx(g);x,y,w,h=cx+(x-ox)/sx,cy+(y-oy)/sy,w/sx,h/sy
    return x,y,w,h
def ink(im): return im.getchannel("A").point(lambda v:255 if v>8 else 0).getbbox()
n=0
for sno,sl in enumerate(p.slides,1):
    for s in walk(sl.shapes):
        if s.shape_type!=13: continue
        nm=s.name
        aid=amap.get(nm) or (int(nm.split("_")[1]) if nm.startswith("ADD_") else None)
        if not aid: continue
        cur=Image.open(io.BytesIO(s.image.blob)).convert("RGBA"); cw,chh=cur.size
        if max(cw,chh)>=700: continue
        el=s._element;ch=chain(el);xf=el.find(P+"spPr").find(A+"xfrm");off,ext=xf.find(A+"off"),xf.find(A+"ext")
        X,Y,W,H=to_slide(ch,int(off.get("x")),int(off.get("y")),int(ext.get("cx")),int(ext.get("cy")))
        blip=el.find(".//"+A+"blip");fill=blip.getparent();sr=fill.find(A+"srcRect")
        l=t=r=b=0
        if sr is not None: l,t,r,b=[int(sr.get(k,0))/100000 for k in("l","t","r","b")]
        vis=cur.crop((round(l*cw),round(t*chh),round(cw*(1-r)),round(chh*(1-b))));vw,vh=vis.size
        bb=ink(vis) or (0,0,vw,vh)
        flip=xf.get("flipH")=="1"
        fx0,fx1=bb[0]/vw,bb[2]/vw
        if flip: fx0,fx1=1-fx1,1-fx0
        nx=X+fx0*W; nw=(fx1-fx0)*W; ny=Y+bb[1]/vh*H; nh=(bb[3]-bb[1])/vh*H
        orig=Image.open(os.path.join(LIB,"lib",f"illust_{aid}.png")).convert("RGBA");ob=ink(orig);orig=orig.crop(ob)
        # 잘려 있던 그림(srcRect b)이면 원본도 같은 비율로 아래를 자른다 → 종횡비 유지
        ar_cur=nw/nh; ar_o=orig.width/orig.height
        if abs(ar_cur-ar_o)/ar_o>0.03:
            # 높이를 잘라 맞춤(아래쪽), 폭 맞춤이 필요하면 좌우 가운데 자르기
            if ar_o<ar_cur: newh=round(orig.width/ar_cur); orig=orig.crop((0,0,orig.width,newh))
            else: neww=round(orig.height*ar_cur); x0=(orig.width-neww)//2; orig=orig.crop((x0,0,x0+neww,orig.height))
        cx,cy,cww,chh2=to_child(ch,nx,ny,nw,nh)
        off.set("x",str(round(cx)));off.set("y",str(round(cy)));ext.set("cx",str(round(cww)));ext.set("cy",str(round(chh2)))
        if sr is not None: fill.remove(sr)
        buf=io.BytesIO();orig.save(buf,"PNG",optimize=True);buf.seek(0)
        _,rid=s.part.get_or_add_image_part(buf);blip.set(R+"embed",rid)
        for c in list(blip):
            if c.tag!=A+"extLst": blip.remove(c)
        ex=blip.find(A+"extLst")
        if ex is not None:
            for e in list(ex):
                if any(c.tag.endswith("}imgProps") for c in e): ex.remove(e)
        print(f"{sno:>2} {nm:<9} {cw}x{chh} → {orig.width}x{orig.height}");n+=1
p.save(r"D:\powerPlus\제안서\청주시_발표자료(제안요약서)_v0.21_원본해상도.pptx");print("재삽입",n)
