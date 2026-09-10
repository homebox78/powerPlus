# -*- coding: utf-8 -*-
"""내가 넣은 아이콘의 '덩어리감'을 슬라이드별로 통일 — 크기가 비슷한 것끼리 묶어 √(w·h) 를 중앙값으로.
사용: python massnorm.py <src> <dst> [허용비 1.6]"""
import glob, hashlib, io, statistics, sys
from PIL import Image
from pptx import Presentation
from pptx.util import Emu
sys.stdout.reconfigure(encoding="utf-8")
src, dst = sys.argv[1:3]; TOL = float(sys.argv[3]) if len(sys.argv) > 3 else 1.6
def prep(p):
    im = Image.open(p).convert("RGBA")
    bb = im.getchannel("A").point(lambda v: 255 if v > 8 else 0).getbbox()
    if bb: im = im.crop(bb)
    buf = io.BytesIO(); im.save(buf, "PNG", optimize=True); return hashlib.sha1(buf.getvalue()).hexdigest()
MINE = {prep(p) for p in glob.glob("k*/*.png")}
pres = Presentation(src); total = 0
for sno, sl in enumerate(pres.slides, 1):
    pics = []
    for sh in sl.shapes:
        if sh.shape_type == 13 and hashlib.sha1(sh.image.blob).hexdigest() in MINE:
            g = (Emu(sh.width).inches * Emu(sh.height).inches) ** 0.5
            pics.append((g, sh))
    if len(pics) < 2: continue
    pics.sort(key=lambda t: t[0])
    # 인접 크기 비가 TOL 안이면 같은 무리
    clusters, cur = [], [pics[0]]
    for a, b in zip(pics, pics[1:]):
        if b[0] / a[0] <= TOL: cur.append(b)
        else: clusters.append(cur); cur = [b]
    clusters.append(cur)
    for cl in clusters:
        if len(cl) < 2: continue
        tgt = statistics.median(g for g, _ in cl); changed = 0
        for g, sh in cl:
            if abs(g / tgt - 1) < 0.04: continue
            k = tgt / g
            cx, cy = sh.left + sh.width / 2, sh.top + sh.height / 2
            nw, nh = int(sh.width * k), int(sh.height * k)
            sh.left, sh.top, sh.width, sh.height = int(cx - nw / 2), int(cy - nh / 2), nw, nh
            changed += 1
        total += changed
        print(f"s{sno}: {len(cl)}개 무리 {min(g for g,_ in cl):.2f}~{max(g for g,_ in cl):.2f}in → {tgt:.2f}in ({changed}개 조정)")
pres.save(dst); print(f"총 {total}개 → {dst}")
