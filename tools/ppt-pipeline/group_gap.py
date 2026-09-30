"""7쪽 위·아래 묶음 사이 간격 줄이기 — 두 묶음을 세로로 늘려 간격 GAP 로. 인자: <pptx>(제자리 저장)"""
import sys
from pptx import Presentation
from pptx.util import Emu
sys.stdout.reconfigure(encoding="utf-8")
IN = 914400
f = sys.argv[1]
p = Presentation(f)
s = p.slides[6]
GAP = 0.25
sh = [x for x in s.shapes if x.name != "시안 바탕" and x.top is not None]
top = [x for x in sh if 1.0 * IN <= x.top < 3.9 * IN]
low = [x for x in sh if x.top >= 4.4 * IN and x.top + x.height <= 7.1 * IN]
t0 = min(x.top for x in top) / IN
t1 = max(x.top + x.height for x in top) / IN
l0 = min(x.top for x in low) / IN
l1 = max(x.top + x.height for x in low) / IN
r = (l1 - t0 - GAP) / ((t1 - t0) + (l1 - l0))
B = t0 + (t1 - t0) * r
L0 = B + GAP
print("위 %.2f~%.2f 아래 %.2f~%.2f 비율 %.3f → 위 끝 %.2f 아래 시작 %.2f" % (t0, t1, l0, l1, r, B, L0))


def scale(group, a_old, a_new):
    for x in group:
        top, h = x.top / IN, x.height / IN
        if x.shape_type == 13:          # 그림: 크기 유지, 중심만 이동
            c = a_new + (top + h / 2 - a_old) * r
            x.top = Emu(int((c - h / 2) * IN))
        else:
            x.top = Emu(int((a_new + (top - a_old) * r) * IN))
            x.height = Emu(int(h * r * IN))


scale(top, t0, t0)
scale(low, l0, L0)
p.save(f)
print("위 %d · 아래 %d 도형" % (len(top), len(low)))
