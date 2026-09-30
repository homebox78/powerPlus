"""시안 구성 요소 이식: 하이라이트 핑크 · 카드 둥근 모서리 + 테두리 없는 면 · 알약 라벨 · 표 합계 핑크.
   인자: 원본 결과"""
import sys, collections
from pptx import Presentation
from lxml import etree
sys.stdout.reconfigure(encoding="utf-8")
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"; P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"; E = 914400
PINK = "EC1C68"; PALE = {"D0E6FA", "E8F2FC"}; HEAD = {"0B2E6B", "2F78E0", "78A8F0", "7A90B2", PINK}
LINE = {"C2DCF5", "78A8F0", "2F78E0", "0B2E6B", "808080", "D3D3D3"}
p = Presentation(sys.argv[1]); cnt = collections.Counter()
def spPr(s): return s._element.find(P + "spPr")
def fillc(sp):
    sf = sp.find(A + "solidFill")
    if sf is None: return None
    c = sf.find(A + "srgbClr")
    if c is not None: return c.get("val").upper()
    sc = sf.find(A + "schemeClr")
    return "FFFFFF" if sc is not None and sc.get("val") == "bg1" and len(sc) == 0 else "?"
def linec(sp):
    ln = sp.find(A + "ln")
    if ln is None or ln.find(A + "noFill") is not None: return None, False
    sf = ln.find(A + "solidFill"); c = sf.find(A + "srgbClr") if sf is not None else None
    return (c.get("val").upper() if c is not None else None), ln.find(A + "prstDash") is not None
def flat(shapes, tf=(0, 0, 1, 1)):
    ox, oy, sx, sy = tf
    for s in shapes:
        if s.width is None: continue
        if s.shape_type == 6:
            x = s._element.find(P + "grpSpPr").find(A + "xfrm"); co, ce = x.find(A + "chOff"), x.find(A + "chExt")
            kx = s.width / int(ce.get("cx")) if int(ce.get("cx")) else 1; ky = s.height / int(ce.get("cy")) if int(ce.get("cy")) else 1
            gx, gy = ox + s.left * sx, oy + s.top * sy
            yield from flat(s.shapes, (gx - int(co.get("x")) * kx * sx, gy - int(co.get("y")) * ky * sy, sx * kx, sy * ky))
        else: yield s, ox + s.left * sx, oy + s.top * sy, s.width * sx, s.height * sy
def setgeom(sp, prst, **adj):
    g = sp.find(A + "prstGeom"); g.set("prst", prst)
    av = g.find(A + "avLst")
    if av is None: av = etree.SubElement(g, A + "avLst")
    for c in list(av): av.remove(c)
    for k, v in adj.items(): etree.SubElement(av, A + "gd", name=k, fmla="val %d" % v)
def noline(sp):
    ln = sp.find(A + "ln")
    for c in list(ln):
        if c.tag in (A + "solidFill", A + "prstDash"): ln.remove(c)
    ln.insert(0, etree.Element(A + "noFill"))
def setfill(sp, v):
    sf = sp.find(A + "solidFill")
    for c in list(sf): sf.remove(c)
    etree.SubElement(sf, A + "srgbClr", val=v)
for n, sl in enumerate(p.slides, 1):
    # a) 키메시지 하이라이트
    for c in sl._element.iter(A + "srgbClr"):
        if c.get("val").upper() == "2FAC73": c.set("val", PINK); cnt["hl"] += 1
    # d) 표 합계 행
    for s in sl.shapes:
        if getattr(s, "has_table", False) and s.has_table:
            for r in s.table.rows:
                if r.cells[0].text_frame.text.strip() in ("계", "합계", "합 계"):
                    for cell in list(r.cells)[1:]:
                        for pg in cell.text_frame.paragraphs:
                            for run in pg.runs:
                                rp = run._r.get_or_add_rPr()
                                for o in rp.findall(A + "solidFill"): rp.remove(o)
                                sf = etree.Element(A + "solidFill"); etree.SubElement(sf, A + "srgbClr", val=PINK)
                                rp.insert(0, sf); rp.set("b", "1"); cnt["total"] += 1
    fs = []
    for s, x, y, w, h in flat(sl.shapes):
        sp = spPr(s)
        if sp is None or s.shape_type == 13 or sp.find(A + "prstGeom") is None: continue
        xf = sp.find(A + "xfrm")
        fs.append(dict(s=s, sp=sp, x=x, y=y, w=w, h=h, g=sp.find(A + "prstGeom").get("prst"), f=fillc(sp), l=linec(sp),
                       rot=bool(xf is not None and xf.get("rot"))))
    for a in fs:
        if a["g"] not in ("rect", "roundRect") or a["rot"]: continue
        x, y, w, h, f = a["x"], a["y"], a["w"], a["h"], a["f"]
        lc, dash = a["l"]
        # b) 카드
        if w > 1.0 * E and h > 0.7 * E and not (w > 10.4 * E) and f in PALE | {"FFFFFF"} and not dash:
            label = any(b is not a and b["f"] in HEAD and b["h"] < 0.6 * E and abs(b["y"] - y) < 0.07 * E
                        and abs(b["x"] - x) < 0.09 * E and abs(b["w"] - w) < 0.18 * E for b in fs)
            inside = any(b is not a and b["f"] in PALE and b["x"] <= x + 0.02 * E and b["y"] <= y + 0.02 * E
                         and b["x"] + b["w"] >= x + w - 0.02 * E and b["y"] + b["h"] >= y + h - 0.02 * E for b in fs)
            if lc in LINE:
                if f == "FFFFFF" and not inside: setfill(a["sp"], "E8F2FC"); cnt["fill"] += 1
                noline(a["sp"]); cnt["noline"] += 1
            r = (5 if label else 10) / 96 * E
            setgeom(a["sp"], "roundRect", adj=min(50000, r / min(w, h) * 100000)); cnt["card"] += 1
        # c) 알약 라벨
        elif f in HEAD and 0.15 * E < h < 0.4 * E and 0.5 * E < w < 3.0 * E:
            stacked = any(b is not a and abs(b["x"] - x) < 0.1 * E and abs(b["w"] - w) < 0.2 * E
                          and (abs(b["y"] - (y + h)) < 0.1 * E or abs(b["y"] - y) < 0.05 * E and b["h"] > h * 1.5) for b in fs)
            row = sum(1 for b in fs if b is not a and b["f"] in HEAD and abs(b["y"] - y) < 0.03 * E and abs(b["h"] - h) < 0.03 * E
                      and (abs(b["x"] - (x + w)) < 0.03 * E or abs(b["x"] + b["w"] - x) < 0.03 * E))
            if not stacked and not row:
                setgeom(a["sp"], "roundRect", adj=50000); cnt["pill"] += 1
print(dict(cnt)); p.save(sys.argv[2])
