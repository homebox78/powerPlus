# -*- coding: utf-8 -*-
"""7쪽: 위(유지보수 대상·요구사항)는 세로 압축, 아래 시스템 구성도는 v0.63 의 상세 구성도로 교체.
인자: <src> <ref v0.63> <dst>"""
import sys, copy
from pptx import Presentation
from pptx.oxml.ns import qn
sys.stdout.reconfigure(encoding="utf-8")
src, ref, dst = sys.argv[1:4]
E = 914400
p = Presentation(src); s = p.slides[6]
r = Presentation(ref); rs = r.slides[6]
by = {sh.shape_id: sh for sh in s.shapes}
# 1) 아래 구성도 삭제
for sh in list(s.shapes):
    if sh.shape_id == 2 or 299 <= sh.shape_id <= 364:
        sh._element.getparent().remove(sh._element)
# 2) 위 구역 세로 압축: 1.04 ~ 4.04 → 1.04 ~ 3.38
T0, B0, B1 = 1.04 * E, 4.04 * E, 3.38 * E
k = (B1 - T0) / (B0 - T0)
ymap = lambda y: T0 + (y - T0) * k
for sh in s.shapes:
    if not (225 <= sh.shape_id <= 298):
        continue
    cy = sh.top + sh.height / 2
    if sh.shape_type == 13:                       # 그림: 비율 유지 축소
        cx = sh.left + sh.width / 2
        sh.width = int(sh.width * k); sh.height = int(sh.height * k)
        sh.left = int(cx - sh.width / 2); sh.top = int(ymap(cy) - sh.height / 2)
    elif sh.has_text_frame and sh.text_frame.text.strip():   # 글: 크기 유지, 가운데만 옮김
        sh.top = int(ymap(cy) - sh.height / 2)
    else:                                          # 상자·선: 세로로 줄임
        top = ymap(sh.top); bot = ymap(sh.top + sh.height)
        sh.top = int(top); sh.height = max(0, int(bot - top))
# 3) v0.63 아래 구성도 복사 (그림 관계 다시 연결)
RE = qn("r:embed"); RL = qn("r:link")
anchor = s.shapes._spTree
for sh in rs.shapes:
    if sh.top < 3.45 * E or sh.shape_id == 222:
        continue
    el = copy.deepcopy(sh._element)
    for node in el.iter():
        for attr in (RE, RL):
            rid = node.get(attr)
            if rid:
                part = rs.part.related_part(rid)
                img_part = s.part.package.get_or_add_image_part(__import__("io").BytesIO(part.blob))
                node.set(attr, s.part.relate_to(img_part, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"))
    anchor.append(el)
# 쪽 번호·바닥 띠는 맨 위로
for sh in list(s.shapes):
    if sh.name.startswith("시안 바탕 2"):
        anchor.append(sh._element)
p.save(dst); print("ok k=%.3f" % k)
