import sys
from pptx import Presentation
E=914400
p=Presentation(sys.argv[1]); sl=p.slides[15]
for s in sl.shapes:
    if s.shape_type==1 and abs(s.top-2.94*E)<0.02*E and abs(s.height-0.42*E)<0.02*E:
        s.left=int(1.6*E); s.width=int(6.12*E)
    if s.shape_type==1 and abs(s.height-0.025*E)<0.004*E and abs(s.width-0.45*E)<0.01*E:
        s.left=int(1.05*E) if s.left<4*E else int(7.82*E)
    if s.name=="LBL_사업지원방안": s.adjustments[0]=0.5
p.save(sys.argv[1])
