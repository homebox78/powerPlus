# -*- coding: utf-8 -*-
"""map.ASSIGN 대로 도형 아이콘을 powerPlus 자산으로 일괄 교체."""
import io, os, sys
from PIL import Image
from pptx import Presentation
from pptx.util import Emu
sys.path.insert(0, ".")
import map as M

sys.stdout.reconfigure(encoding="utf-8")

cache = {}
def load(path):
    if path not in cache:
        im = Image.open(path).convert("RGBA")
        b = im.getchannel("A").point(lambda v: 255 if v > 8 else 0).getbbox()
        if b:
            im = im.crop(b)                    # 자산 자체 여백 제거 → 시각 크기 확보
        cache[path] = im
    return cache[path]

def groups_of(slide):
    shapes = []
    def walk(lst):
        for sh in lst:
            if sh.shape_type == 6:
                walk(sh.shapes); continue
            try:
                x, y = Emu(sh.left).inches, Emu(sh.top).inches
                w, h = Emu(sh.width).inches, Emu(sh.height).inches
            except Exception:
                continue
            t = sh.text_frame.text.strip() if sh.has_text_frame else ""
            shapes.append({"sh": sh, "x": x, "y": y, "w": w, "h": h, "text": t})
    walk(slide.shapes)
    anchors = [s for s in shapes if not s["text"] and 0.25 <= max(s["w"], s["h"]) <= 1.1
               and 0.8 <= (s["w"] / s["h"] if s["h"] else 0) <= 1.25]
    anchors.sort(key=lambda s: -(s["w"] * s["h"]))
    used, groups = set(), []
    for a in anchors:
        if id(a["sh"]) in used:
            continue
        parts = []
        for s in shapes:
            if id(s["sh"]) in used or s["text"]:
                continue
            cx, cy = s["x"] + s["w"] / 2, s["y"] + s["h"] / 2
            if a["x"] - .02 <= cx <= a["x"] + a["w"] + .02 and a["y"] - .02 <= cy <= a["y"] + a["h"] + .02 \
               and s["w"] <= a["w"] * 1.05 and s["h"] <= a["h"] * 1.05:
                parts.append(s)
        for s in parts:
            used.add(id(s["sh"]))
        if parts:
            groups.append({"a": a, "parts": parts})
    return groups

