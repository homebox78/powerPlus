"""원본 vs 최종 장별 글 추출 → 문장 단위 대조. 인자: <원본> <최종> <출력폴더>"""
import os, re, sys, difflib
from pptx import Presentation
sys.stdout.reconfigure(encoding="utf-8")
a, b, out = sys.argv[1:4]
os.makedirs(out, exist_ok=True)


def walk(shs):
    for sh in shs:
        if sh.shape_type == 6:
            yield from walk(sh.shapes)
        else:
            yield sh


def lines(slide):
    res = []
    for sh in walk(slide.shapes):
        if sh.has_text_frame:
            for pg in sh.text_frame.paragraphs:
                t = "".join(r.text for r in pg.runs).strip()
                if t:
                    res.append(t)
        elif getattr(sh, "has_table", False) and sh.has_table:
            for row in sh.table.rows:
                for c in row.cells:
                    for pg in c.text_frame.paragraphs:
                        t = "".join(r.text for r in pg.runs).strip()
                        if t:
                            res.append(t)
    return res


def norm(t):
    return re.sub(r"[\s·•\-–—∙※▶►■□◆◇○●]+", "", t)


pa, pb = Presentation(a), Presentation(b)
print("장 수", len(pa.slides), len(pb.slides))
summary = []
for i in range(max(len(pa.slides), len(pb.slides))):
    la = lines(pa.slides[i]) if i < len(pa.slides) else []
    lb = lines(pb.slides[i]) if i < len(pb.slides) else []
    blob_b = norm("".join(lb))
    blob_a = norm("".join(la))
    missing = [t for t in la if len(norm(t)) >= 2 and norm(t) not in blob_b]
    added = [t for t in lb if len(norm(t)) >= 2 and norm(t) not in blob_a]
    with open(os.path.join(out, "s%02d.txt" % (i + 1)), "w", encoding="utf-8") as f:
        f.write("### 원본\n" + "\n".join(la) + "\n\n### 최종\n" + "\n".join(lb) +
                "\n\n### 원본에만(빠짐 후보)\n" + "\n".join(missing) + "\n\n### 최종에만(새·바뀐 문구)\n" + "\n".join(added) + "\n")
    summary.append((i + 1, len(missing), len(added)))
for s, m, ad in summary:
    print("%02d 빠짐후보 %d  새문구 %d" % (s, m, ad))
