"""v0.55 - heads that carry the emphasis go blue face + white text on slides that lacked it. args: src dst
   17: card heads x3 (y~2.4in), 33: heads x3 (y~2.9in), 38: light head bands (6.1x0.5in)."""
import sys, collections
from lxml import etree
from pptx import Presentation
sys.stdout.reconfigure(encoding="utf-8")
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
E = 914400
HL, WHITE = "143A69", "FFFFFF"
LIGHT = ("CCE3FD", "D0E6FA", "E1EEFD")
p = Presentation(sys.argv[1]); cnt = collections.Counter()
def spPr(s): return s._element.find(P + "spPr")
def flat(shapes, tf=(0, 0, 1, 1)):
    ox, oy, sx, sy = tf
    for s in shapes:
        if s.width is None: continue
        if s.shape_type == 6:
            x = s._element.find(P + "grpSpPr").find(A + "xfrm"); co, ce = x.find(A + "chOff"), x.find(A + "chExt")
            kx = s.width / int(ce.get("cx")) if int(ce.get("cx")) else 1; ky = s.height / int(ce.get("cy")) if int(ce.get("cy")) else 1
            gx, gy = ox + s.left * sx, oy + s.top * sy
            yield from flat(s.shapes, (gx - int(co.get("x")) * kx * sx, gy - int(co.get("y")) * ky * sy, sx * kx, sy * ky))
        else: yield s, ox + s.left * sx, oy + s.top * sy, s.width * sx, s.height * sy
def fillc(s):
    sp = spPr(s)
    sf = sp.find(A + "solidFill") if sp is not None else None
    return sf.find(A + "srgbClr") if sf is not None else None
def set_runs(el, hexv):
    for tag in ("r", "endParaRPr"):
        for r in el.iter(A + tag):
            rp = r if tag == "endParaRPr" else r.find(A + "rPr")
            if rp is None:
                rp = etree.Element(A + "rPr"); r.insert(0, rp)
            for t in ("solidFill", "gradFill", "noFill", "pattFill"):
                for e in rp.findall(A + t): rp.remove(e)
            sf = etree.Element(A + "solidFill"); etree.SubElement(sf, A + "srgbClr", val=hexv)
            ln = rp.find(A + "ln"); rp.insert((list(rp).index(ln) + 1) if ln is not None else 0, sf)
T = {17: lambda x, y, w, h: abs(y - 2.4 * E) < 0.2 * E and 2.8 * E < w < 3.3 * E and h < 0.9 * E,
     33: lambda x, y, w, h: abs(y - 2.9 * E) < 0.2 * E and 3.1 * E < w < 3.6 * E and h < 0.6 * E,
     38: lambda x, y, w, h: w > 5.8 * E and h < 0.7 * E}
for n, sl in enumerate(p.slides, 1):
    if n not in T: continue
    fs = list(flat(sl.shapes)); done = []
    for s, x, y, w, h in fs:
        c = fillc(s)
        if c is None or c.get("val").upper() not in LIGHT or not T[n](x, y, w, h): continue
        for a in list(c): c.remove(a)
        c.set("val", HL)
        ln = spPr(s).find(A + "ln")
        if ln is not None and ln.find(A + "solidFill") is not None:
            lc = ln.find(A + "solidFill")
            for a in list(lc): lc.remove(a)
            etree.SubElement(lc, A + "srgbClr", val=HL)
        if s.has_text_frame: set_runs(s._element, WHITE)
        done.append((x, y, w, h)); cnt[n] += 1
    for s, x, y, w, h in fs:
        if not getattr(s, "has_text_frame", False) or not s.text_frame.text.strip() or fillc(s) is not None: continue
        cx, cy = x + w / 2, y + h / 2
        if any(a <= cx <= a + c_ and b <= cy <= b + d and h < d * 1.6 for a, b, c_, d in done):
            set_runs(s._element, WHITE); cnt["txt"] += 1
print(dict(cnt))
p.save(sys.argv[2])
