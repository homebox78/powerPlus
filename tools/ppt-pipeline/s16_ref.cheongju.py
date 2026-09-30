"""v0.53 slide 16 - icons cut from the reference capture, labels bigger, left label emphasized. args: src dst ref_png"""
import sys, io
import numpy as np
from PIL import Image
from scipy import ndimage
from lxml import etree
from pptx import Presentation
from pptx.util import Pt
sys.stdout.reconfigure(encoding="utf-8")
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
E = 914400
p = Presentation(sys.argv[1]); sl = p.slides[15]
ref = Image.open(sys.argv[3]).convert("RGB")
CX = [120, 357, 590, 800, 995, 1185, 1367, 1555]
def cut(cx):
    im = ref.crop((cx - 88, 328, cx + 88, 458)); a = np.asarray(im).astype(int)
    bg = a[2, 2]
    near = (np.abs(a - bg).sum(axis=2) < 40) | ((a.min(axis=2) > 205) & (a.max(axis=2) - a.min(axis=2) < 45) & (a[:, :, 2] >= a[:, :, 0]))
    lab, _ = ndimage.label(near)
    edge = set(lab[0, :]) | set(lab[-1, :]) | set(lab[:, 0]) | set(lab[:, -1]); edge.discard(0)
    mask = np.isin(lab, list(edge))
    alpha = np.where(mask, 0, 255).astype(np.uint8)
    alpha = np.asarray(Image.fromarray(alpha).filter(__import__("PIL.ImageFilter", fromlist=["x"]).GaussianBlur(0.7)))
    out = Image.fromarray(np.dstack([a.astype(np.uint8), alpha]), "RGBA")
    out = out.crop(out.getbbox()); b = io.BytesIO(); out.save(b, "PNG"); b.seek(0); return b, out.size
old = sorted([s for s in sl.shapes if s.shape_type == 13 and s.top > 4 * E and abs(s.width - 0.52 * E) < 0.02 * E], key=lambda s: s.left)
assert len(old) == 8, len(old)
dump = Image.new("RGB", (8 * 180, 140), "#E1EEFD")
for i, s in enumerate(old):
    b, (w, h) = cut(CX[i])
    dump.paste(Image.open(b), (i * 180 + 5, 5), Image.open(b)); b.seek(0)
    H = 0.6 * E; W = H * w / h
    if W > 0.95 * E: W = 0.95 * E; H = W * h / w
    cx = s.left + s.width / 2; top = s.top - 0.02 * E
    s._element.getparent().remove(s._element)
    sl.shapes.add_picture(b, int(cx - W / 2), int(top + (0.6 * E - H) / 2), int(W), int(H))
dump.save(sys.argv[2] + ".icons.png")
SPEC = {"LBL_사업수행방안": (2.5, 0.54, 20, "143A69"), "LBL_사업지원방안": (2.1, 0.44, 16, "2F78E0")}
bottoms = {}
for s in sl.shapes:
    if s.name in SPEC:
        w, h, pt, col = SPEC[s.name]
        cx = s.left + s.width / 2; cy = 3.62 * E
        s.width = int(w * E); s.height = int(h * E); s.left = int(cx - s.width / 2); s.top = int(cy - s.height / 2)
        sp = s._element.find(P + "spPr")
        for c in sp.iter(A + "srgbClr"):
            if c.getparent().getparent().tag in (P + "spPr", A + "ln"): c.set("val", col)
        for r in s.text_frame.paragraphs[0].runs: r.font.size = Pt(pt)
        bottoms[s.name] = s.top + s.height
for s in sl.shapes:
    for k, b in bottoms.items():
        if s.name == "FAN_" + k:
            end = s.top + s.height; s.top = int(b); s.height = int(end - b)
p.save(sys.argv[2])
