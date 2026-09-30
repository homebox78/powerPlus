"""카드 위에 얹힌 머리 라벨(채움 사각형)의 위쪽 두 모서리를 둥글게. 인자: 원본 결과 [반지름px=5]"""
import sys, collections
from pptx import Presentation
from lxml import etree
sys.stdout.reconfigure(encoding="utf-8")
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"; P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"; E = 914400
HEAD = {"0B2E6B", "0456B6", "658EBB", "7A90B2"}
R = (float(sys.argv[3]) if len(sys.argv) > 3 else 5) / 96 * E
p = Presentation(sys.argv[1])
def spPr(s): return s._element.find(P + "spPr")
def col(s):
    sp = spPr(s); sf = sp.find(A + "solidFill") if sp is not None else None
    c = sf.find(A + "srgbClr") if sf is not None else None
    return c.get("val").upper() if c is not None else None
def flat(shapes, tf=(0, 0, 1, 1)):
    ox, oy, sx, sy = tf
    for s in shapes:
        if s.width is None: continue
        if s.shape_type == 6:
            x = s._element.find(P + "grpSpPr").find(A + "xfrm"); co, ce = x.find(A + "chOff"), x.find(A + "chExt")
            kx = s.width / int(ce.get("cx")) if int(ce.get("cx")) else 1; ky = s.height / int(ce.get("cy")) if int(ce.get("cy")) else 1
            gx, gy = ox + s.left * sx, oy + s.top * sy
            yield from flat(s.shapes, (gx - int(co.get("x")) * kx * sx, gy - int(co.get("y")) * ky * sy, sx * kx, sy * ky))
        else: yield s, ox + s.left * sx, oy + s.top * sy, s.width * sx, s.height * sy
per = collections.Counter()
for n, sl in enumerate(p.slides, 1):
    fs = list(flat(sl.shapes))
    for s, x, y, w, h in fs:
        sp = spPr(s)
        if sp is None or col(s) not in HEAD: continue
        g = sp.find(A + "prstGeom")
        if g is None or g.get("prst") not in ("rect", "roundRect") or h > 0.6 * E or w < 0.8 * E or h < 0.12 * E: continue
        xf = sp.find(A + "xfrm")
        if xf is not None and (xf.get("rot") or xf.get("flipV")): continue
        body = [b for b, X, Y, W, H in fs if b is not s and (abs(Y - (y + h)) < 0.07 * E or (abs(Y - y) < 0.05 * E and H > h * 1.5)) and abs(X - x) < 0.09 * E
                and abs(W - w) < 0.18 * E and H > h * 0.8 and b.shape_type != 13]
        if not body: continue
        g.set("prst", "round2SameRect")
        av = g.find(A + "avLst")
        if av is None: av = etree.SubElement(g, A + "avLst")
        for c in list(av): av.remove(c)
        etree.SubElement(av, A + "gd", name="adj1", fmla="val %d" % min(50000, R / min(w, h) * 100000))
        etree.SubElement(av, A + "gd", name="adj2", fmla="val 0")
        per[n] += 1
print(sum(per.values()), dict(sorted(per.items())))
p.save(sys.argv[2])
