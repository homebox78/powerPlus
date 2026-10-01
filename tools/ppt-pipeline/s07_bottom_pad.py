import sys
from pptx import Presentation
sys.stdout.reconfigure(encoding="utf-8")
src,dst=sys.argv[1:3]
E=914400; CM=360000
p=Presentation(src); s=p.slides[6]; H=p.slide_height
B=int(H/2+8.3*CM)
by={sh.shape_id:sh for sh in s.shapes}
I=lambda v:int(v*E)
sec=[2]+list(range(299,365))
for i in sec:
    if i in by: by[i].top-=I(0.06)
outer=by[299]; outer.height=B-outer.top
pad=by[303].left-outer.left
cb=B-pad
for card,inner in ((303,304),(316,317),(338,339)):
    by[card].height=cb-by[card].top
    by[inner].height=cb-I(0.02)-by[inner].top
for ids,d in ((range(305,313),0.06),(range(318,337),0.037)):
    for i in ids:
        if i in (305,306,318,319,320): continue  # 머리 아이콘·제목은 그대로
        by[i].top+=I(d)
# 화살표 세로 가운데
mid=(by[303].top+by[303].height//2)
for i in (313,314): by[i].top=mid-by[i].height//2
p.save(dst)
print("guide",B/E,"card bottom",cb/E)
