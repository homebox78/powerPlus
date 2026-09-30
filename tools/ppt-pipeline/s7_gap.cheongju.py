from pptx import Presentation
from pptx.util import Emu
E=914400
A="{http://schemas.openxmlformats.org/drawingml/2006/main}"
D=r"D:\powerPlus\제안서\청주시_발표자료(제안요약서)_v0.34_s7간격.pptx"
p=Presentation("l33.pptx"); sl=p.slides[6]
DY=0.38; K=0.895
LEFT={"그룹 22","그룹 28","그룹 29","그룹 2","직사각형 51","직사각형 60","직사각형 233","직사각형 255"}
T0=[s.top/E for s in sl.shapes if s.name=="그룹 29"][0]
f=lambda y: T0+DY+(y-T0)*K
for s in sl.shapes:
    y=s.top/E; h=s.height/E
    if y<3.1 or s.name.startswith("CARD"): continue
    if s.name in LEFT:
        s.top=Emu(int(f(y)*E))
        if s.name!="직사각형 255": r=s.left+s.width; s.height=Emu(int(h*K*E)); s.width=Emu(int(s.width*K)); s.left=r-s.width
        continue
    if s.shape_type==9 and s.left/E<2.0:
        xf=s._element.find(".//"+A+"xfrm"); fv=xf.get("flipV")=="1"
        yL,yR=(y+h,y) if fv else (y,y+h)
        yL=f(yL); yR=yR+DY
        s.top=Emu(int(min(yL,yR)*E)); s.height=Emu(int(abs(yL-yR)*E))
        if yL>yR: xf.set("flipV","1")
        elif "flipV" in xf.attrib: del xf.attrib["flipV"]
        continue
    s.top=Emu(int((y+DY)*E))
p.save(D)
