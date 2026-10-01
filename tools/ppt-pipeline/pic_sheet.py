import io, sys, json
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
sys.stdout.reconfigure(encoding="utf-8")
E=914400
p=Presentation(sys.argv[1])
def walk(shs,path=""):
    for s in shs:
        if s.shape_type==6: yield from walk(s.shapes,path+"/"+s.name)
        elif s.shape_type==13: yield s,path
L=[]
for n,sl in enumerate(p.slides,1):
    if n<4: continue
    for s,path in walk(sl.shapes):
        if s.width>0.9*E or s.height>0.9*E: continue
        try: im=Image.open(io.BytesIO(s.image.blob)).convert("RGBA")
        except Exception: continue
        L.append((n,path,s.name,im))
f=ImageFont.truetype("C:/Windows/Fonts/malgun.ttf",13)
C=14;W=110
for part in range(0,len(L),196):
    chunk=L[part:part+196]
    R=(len(chunk)+C-1)//C
    S=Image.new("RGB",(C*W,R*(W+18)),(225,232,242));d=ImageDraw.Draw(S)
    for i,(n,path,name,im) in enumerate(chunk):
        t=im.copy();t.thumbnail((W-10,W-10))
        x,y=(i%C)*W,(i//C)*(W+18)
        S.paste(t,(x+5,y+5),t)
        d.text((x+3,y+W),"%d %s"%(n,name.replace("시안 그림 ","g")),fill="black",font=f)
    S.save("picsheet_%d.png"%(part//196))
print(len(L))
