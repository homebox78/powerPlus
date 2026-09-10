# -*- coding: utf-8 -*-
"""같은 슬라이드 안 동일 아이콘 자산 중복 제거 — 2번째 이후 등장은 같은 종류의 다른 자산으로."""
import glob, hashlib, io, json, os, sys, urllib.request
from PIL import Image
from pptx import Presentation
from pptx.util import Emu
import map as M
sys.stdout.reconfigure(encoding="utf-8")
src, dst = sys.argv[1:3]
ALT = {"체크": ["icon_247", "icon_974", "icon_172"], "상승": ["icon_1099", "icon_432", "icon_268"],
       "분기": ["icon_1061", "icon_1024"], "타겟": ["icon_517", "icon_232"], "필터": ["icon_334", "icon_651"],
       "돋보기": ["icon_1189", "icon_1050"], "문서": ["icon_1044", "icon_398"], "데이터": ["icon_1180", "icon_1202"],
       "격자": ["icon_1241"], "링크": ["icon_1178"], "방패": ["icon_1229"], "이중화": ["icon_1187", "icon_1152"],
       "사람": ["icon_984", "icon_1215"], "지구본": ["icon_1239", "icon_1029"]}
API = "https://hom2box.com/powerPlus/api/public/assets"
os.makedirs("kalt", exist_ok=True)
def prep_img(im):
    bb = im.getchannel("A").point(lambda v: 255 if v > 8 else 0).getbbox()
    if bb: im = im.crop(bb)
    buf = io.BytesIO(); im.save(buf, "PNG", optimize=True); return im, buf.getvalue()
def fetch(aid):
    p = f"kalt/{aid}.png"
    if not os.path.exists(p):
        d = json.loads(urllib.request.urlopen(f"{API}?ids={aid}").read())
        url = d["data"][0]["image_url"]; open(p, "wb").write(urllib.request.urlopen(url).read())
    return prep_img(Image.open(p).convert("RGBA"))
kind_of = {}
for k, p in M.KIND.items():
    _, b = prep_img(Image.open(p).convert("RGBA")); kind_of[hashlib.sha1(b).hexdigest()] = k
pres = Presentation(src); n = 0
for sno, sl in enumerate(pres.slides, 1):
    pics = [s for s in sl.shapes if s.name == "PPICON"]
    hashes = {hashlib.sha1(s.image.blob).hexdigest() for s in pics}
    by = {}
    for s in sorted(pics, key=lambda s: (s.top, s.left)):
        by.setdefault(hashlib.sha1(s.image.blob).hexdigest(), []).append(s)
    for h, lst in by.items():
        k = kind_of.get(h)
        if len(lst) < 2 or not k: continue
        alts = [a for a in ALT.get(k, [])]
        for s in lst[1:]:
            while alts:
                aid = alts.pop(0); im, b = fetch(aid); ah = hashlib.sha1(b).hexdigest()
                if ah in hashes: continue
                break
            else:
                print(f"s{sno} {k}: 대체 부족"); break
            L, T, W, H = s.left, s.top, s.width, s.height; cx, cy = L + W / 2, T + H / 2
            sc = min(W / im.width, H / im.height); nw, nh = int(im.width * sc), int(im.height * sc)
            s._element.getparent().remove(s._element)
            sl.shapes.add_picture(io.BytesIO(b), int(cx - nw / 2), int(cy - nh / 2), nw, nh).name = "PPICON"
            hashes.add(ah); n += 1; print(f"s{sno} {k} → {aid}")
pres.save(dst); print("교체", n)
