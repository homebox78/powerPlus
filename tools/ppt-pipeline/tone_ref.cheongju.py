# 시안 톤 반영: 키메시지 하이라이트 빨강→밝은 파랑, 작은 강조 글자→핑크, 큰 카드 짙은 테두리→옅은 선
import sys,collections
from pptx import Presentation
sys.stdout.reconfigure(encoding="utf-8")
A="{http://schemas.openxmlformats.org/drawingml/2006/main}"
P="{http://schemas.openxmlformats.org/presentationml/2006/main}"
SRC,DST=sys.argv[1],sys.argv[2]
HI="2070E8"; PINK="E84078"; SOFT="C0D2E6"; E=914400
p=Presentation(SRC); n=collections.Counter()
def walk(sh):
    for s in sh:
        if s.shape_type==6: yield from walk(s.shapes)
        else: yield s
for sl in p.slides:
    for s in walk(sl.shapes):
        if s.has_text_frame:
            for r in s._element.iter(A+"r"):
                rp=r.find(A+"rPr")
                if rp is None: continue
                c=rp.find(A+"solidFill/"+A+"srgbClr")
                if c is not None and c.get("val").upper()=="D23737":
                    big=int(rp.get("sz") or 0)>=2200
                    c.set("val",HI if big else PINK); n["키메시지" if big else "작은강조"]+=1
        sp=s._element.find(P+"spPr")
        if sp is None or s.shape_type==13: continue
        l=sp.find(A+"ln")
        if l is None: continue
        c=l.find(A+"solidFill/"+A+"srgbClr")
        if c is not None and c.get("val").upper() in("0B2E6B","0456B6","3A46A0") and int(l.get("w") or 9525)<=12700 \
           and s.width>1.5*E and s.height>0.8*E:
            c.set("val",SOFT); l.set("w","9525"); n["카드선"]+=1
p.save(DST); print(dict(n))
