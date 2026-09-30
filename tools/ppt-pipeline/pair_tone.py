"""좌우로 나란한 제목 띠 두 개가 같은 색이면 왼쪽을 한 단계 밝게. 인자: 원본 결과
   그룹 안 도형은 그룹 변환을 풀어 슬라이드 좌표로 비교한다."""
import sys
from pptx import Presentation
sys.stdout.reconfigure(encoding="utf-8")
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"; P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"; E = 914400
UP = {"0B2E6B": "0456B6", "0456B6": "658EBB"}
p = Presentation(sys.argv[1])
def col(s):
    sp = s._element.find(P + "spPr")
    sf = sp.find(A + "solidFill") if sp is not None else None
    return sf.find(A + "srgbClr") if sf is not None else None
def flat(shapes, tf=(0, 0, 1, 1)):
    ox, oy, sx, sy = tf
    for s in shapes:
        if s.width is None: continue
        if s.shape_type == 6:
            x = s._element.find(P + "grpSpPr").find(A + "xfrm")
            co, ce = x.find(A + "chOff"), x.find(A + "chExt")
            kx = s.width / int(ce.get("cx")) if int(ce.get("cx")) else 1
            ky = s.height / int(ce.get("cy")) if int(ce.get("cy")) else 1
            gx, gy = ox + s.left * sx, oy + s.top * sy
            yield from flat(s.shapes, (gx - int(co.get("x")) * kx * sx, gy - int(co.get("y")) * ky * sy, sx * kx, sy * ky))
        else:
            yield s, ox + s.left * sx, oy + s.top * sy, s.width * sx, s.height * sy
for n, sl in enumerate(p.slides, 1):
    c = [t for t in flat(sl.shapes) if col(t[0]) is not None and col(t[0]).get("val").upper() in UP
         and t[3] > 2.5 * E and t[4] < 0.6 * E and t[2] < 2.6 * E]
    for a in c:
        m = [b for b in c if abs(b[2] - a[2]) < 0.1 * E and abs(b[4] - a[4]) < 0.1 * E
             and col(b[0]).get("val") == col(a[0]).get("val")]
        if len(m) == 2:
            l, r = sorted(m, key=lambda t: t[1])
            if l is a and r[1] >= l[1] + l[3] - 0.05 * E:
                v = col(l[0]).get("val").upper(); col(l[0]).set("val", UP[v]); print(n, v, "→", UP[v], round(l[2] / E, 2))
p.save(sys.argv[2])
